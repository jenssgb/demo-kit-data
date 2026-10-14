"""Build four small Cowork starters and their JSON presentation views.

Run this script, then tools\\Update-Manifest.ps1 -All.
No tenant changes. Expected answers stay outside the installed demo folder.
"""
import argparse
import importlib.util
import json
from copy import deepcopy
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
PACKAGE = ROOT / "cowork-starters"
DOMAIN = "M365CPI98544940.OnMicrosoft.com"
PEOPLE = {"Teresa Sac": "TeresaS", "Vance DeLeon": "VanceD", "Sonia Rees": "SoniaR"}
DATE = "2026-10-14"
USE = "https://learn.microsoft.com/microsoft-365/copilot/cowork/use-cowork"
OVERVIEW = "https://learn.microsoft.com/microsoft-365/copilot/cowork/"
GET_STARTED = "https://learn.microsoft.com/microsoft-365/copilot/cowork/get-started"
FAQ = "https://learn.microsoft.com/microsoft-365/copilot/cowork/cowork-faq"
spec = importlib.util.spec_from_file_location("cowork_base", ROOT / "tools" / "build-cowork-usecases.py")
BASE = importlib.util.module_from_spec(spec)
spec.loader.exec_module(BASE)
BASE.FIXED = datetime(2026, 10, 14, 8, 0, 0)


def bi(en, de):
    return {"en": en, "de": de}


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8", newline="\n")


def write_json(path, value):
    write(path, json.dumps(value, ensure_ascii=False, indent=2))


MAILS = [
    {
        "file": "1_Teresa_Carrier_Confirmation.eml", "sender": "Teresa Sac",
        "subject": bi("[CO-19] Carrier confirmation needed", "[CO-19] Carrier-Bestätigung benötigt"),
        "body": bi(
            "Hi,\n\nFor Contoso's Fargo plan, please confirm that we still need the Northwind exception for the 05:30 arrival. I need your response before the next operations review; without it I cannot finalize the carrier note. Please do not promise a slot that has not been approved.\n\nThanks,\nTeresa",
            "Hallo,\n\nBitte bestätige für Contosos Fargo-Plan, ob wir die Northwind-Ausnahme für die Ankunft um 05:30 weiterhin benötigen. Ich brauche deine Antwort vor dem nächsten Operations-Review; ohne sie kann ich die Carrier-Notiz nicht abschließen. Bitte keinen ungeprüften Slot zusagen.\n\nDanke,\nTeresa"),
    },
    {
        "file": "2_Vance_Workshop_Room.eml", "sender": "Vance DeLeon",
        "subject": bi("[CO-19] Workshop room decision", "[CO-19] Raumentscheidung für den Workshop"),
        "body": bi(
            "Hi,\n\nFor the Contoso operations workshop, do you prefer the training room or a Teams meeting? I can prepare either option after your decision. There is no booking or deadline yet.\n\nThanks,\nVance",
            "Hallo,\n\nMöchtest du für den Contoso-Operations-Workshop den Schulungsraum oder ein Teams-Meeting nutzen? Ich kann nach deiner Entscheidung beide Varianten vorbereiten. Es gibt noch keine Buchung und keine Frist.\n\nDanke,\nVance"),
    },
    {
        "file": "3_Sonia_Training_Update.eml", "sender": "Sonia Rees",
        "subject": bi("[CO-19] Training update - FYI", "[CO-19] Trainingsupdate - zur Info"),
        "body": bi(
            "Hi,\n\nFor information: the updated training outline for Contoso's new Fargo wing is ready for discussion. This is a planning update, not a sign-off of the training itself. No reply or action is needed from you.\n\nThanks,\nSonia",
            "Hallo,\n\nZur Information: Die überarbeitete Trainingsübersicht für Contosos neuen Fargo-Flügel ist bereit zur Besprechung. Das ist ein Planungsupdate, keine Freigabe des Trainings selbst. Du musst nicht antworten und nichts veranlassen.\n\nDanke,\nSonia"),
    },
]
MEETING = {
    "subject": bi("[CO-21] Contoso operations review", "[CO-21] Contoso-Operations-Review"),
    "weekday": 4, "start": "10:00", "minutes": 30,
    "attendees": ["Teresa Sac", "Vance DeLeon", "Sonia Rees"],
    "body": bi(
        "Planning review for Contoso's Fargo wing. Sources are in OneDrive > Demo-Kit > Cowork > cowork-starters: inputs/en/Meeting_Brief.md and Contoso_Fargo_GoLive_Readiness_Tracker.xlsx. Discuss R02 carrier confirmation, R04 electrical work and R08 shift-lead training. The source tracker is a fictional April-June 2027 planning snapshot, not today's operational status. No transcript is provided.",
        "Planungsreview für Contosos Fargo-Flügel. Quellen liegen in OneDrive > Demo-Kit > Cowork > cowork-starters: inputs/de/Meeting_Brief.md und Contoso_Fargo_GoLive_Readiness_Tracker.xlsx. R02 Carrier-Bestätigung, R04 Elektroarbeiten und R08 Schichtleitertraining besprechen. Der Quelltracker ist ein fiktiver Planungsstand für April bis Juni 2027, kein aktueller Betriebsstatus. Es gibt kein Transkript."),
}
EVENT = bi(
    "# Contoso operations workshop\n\nFictional internal workshop; not a booked event.\n\nGoal: agree responsibilities for the Fargo wing planning work. Proposed date: the next Friday after the current conversation, 14:00-15:30 Europe/Berlin. Resolve and show the actual date before creating an event. MOD Administrator is the organizer; invite Teresa Sac, Vance DeLeon and Sonia Rees. Use a Teams meeting, not a booked physical room.\n\nBudget ceiling: EUR 500 for preparation and materials, not travel. No purchases or bookings are authorized. Use an agenda with carrier readiness, construction dependencies, safety training and agreed next steps. Prepare an owner/task list; unconfirmed commitments remain open.\n\nCalendar subject: [CO-20] Contoso operations workshop. Check for an existing event with that subject before creating one. Show attendees, time zone and date for approval. Never send invitations without explicit approval.",
    "# Contoso-Operations-Workshop\n\nFiktiver interner Workshop; noch kein gebuchter Termin.\n\nZiel: Verantwortlichkeiten für die Fargo-Flügelplanung vereinbaren. Terminvorschlag: der nächste Freitag nach dem aktuellen Verlauf, 14:00-15:30 Europe/Berlin. Das konkrete Datum vor der Anlage auflösen und zeigen. MOD Administrator organisiert; Teresa Sac, Vance DeLeon und Sonia Rees einladen. Ein Teams-Meeting nutzen, keinen physischen Raum buchen.\n\nBudgetobergrenze: 500 EUR für Vorbereitung und Material, nicht für Reisen. Keine Käufe oder Buchungen sind freigegeben. Agenda mit Carrier-Vorbereitung, Bauabhängigkeiten, Sicherheitstraining und vereinbarten nächsten Schritten erstellen. Aufgaben mit Verantwortlichen vorbereiten; unbestätigte Zusagen bleiben offen.\n\nKalenderbetreff: [CO-20] Contoso operations workshop. Vor der Anlage nach einem vorhandenen Termin mit diesem Betreff suchen. Teilnehmer, Zeitzone und Datum zur Freigabe zeigen. Einladungen nie ohne ausdrückliche Freigabe versenden.")
MEETING_BRIEF = bi(
    "# Contoso operations review\n\nPrepare for the actual calendar event [CO-21] Contoso operations review. Read its date, time zone and attendees from the calendar; the installer schedules Thursday 10:00 in the VM's local time zone.\n\nUse the shared Contoso_Fargo_GoLive_Readiness_Tracker.xlsx in this demo folder. It is a fictional April-June 2027 planning snapshot. Focus on R02 carrier confirmations (Teresa), R04 electrical work (Vance) and R08 shift-lead training (Sonia). Respect the tracker dates; do not rewrite them as current delivery promises.\n\nWe want the open decisions, dependencies and questions for the owners. Distinguish recorded status from proposed actions. There is no previous meeting transcript or approval record. Produce a one-page briefing with source links; do not send it or modify the tracker.",
    "# Contoso-Operations-Review\n\nDen tatsächlichen Kalendertermin [CO-21] Contoso-Operations-Review vorbereiten. Datum, Zeitzone und Teilnehmer aus dem Kalender lesen; der Installer legt Donnerstag 10:00 in der lokalen Zeitzone der VM an.\n\nDen geteilten Contoso_Fargo_GoLive_Readiness_Tracker.xlsx in diesem Demo-Ordner nutzen. Das ist ein fiktiver Planungsstand für April bis Juni 2027. Fokus auf R02 Carrier-Bestätigungen (Teresa), R04 Elektroarbeiten (Vance) und R08 Schichtleitertraining (Sonia). Tracker-Termine respektieren und nicht als aktuelle Lieferzusagen umdeuten.\n\nWir brauchen offene Entscheidungen, Abhängigkeiten und Fragen an die Verantwortlichen. Dokumentierten Status von Vorschlägen trennen. Es gibt kein früheres Meetingtranskript und keinen Freigabebeleg. Ein einseitiges Briefing mit Quelllinks erstellen; weder versenden noch den Tracker ändern.")
REVENUE = [
    ("2026-07", "North", 100000, 110000), ("2026-07", "South", 80000, 85000),
    ("2026-08", "North", 120000, 115000), ("2026-08", "South", 90000, 90000),
    ("2026-09", "North", 140000, 130000), ("2026-09", "South", 100000, 95000),
]
SCENARIOS = [
    {
        "code": "CO-19", "title": bi("Triage three emails", "Drei Mails priorisieren"),
        "chatTitle": bi("Inbox triage", "Posteingang sortieren"),
        "story": bi("Three Contoso emails need different responses. You want a clear order and useful drafts without changing your whole inbox.",
                    "Drei Contoso-Mails brauchen unterschiedliche Reaktionen. Du möchtest eine klare Reihenfolge und passende Entwürfe, ohne deinen ganzen Posteingang umzubauen."),
        "preparation": bi(
            ["The three **[CO-19]** emails from Teresa, Vance and Sonia are in the inbox. Read their exact subjects in the email reference blocks.",
             "Start a new **Cowork** conversation. No custom skill or plugin is required. If a message is ambiguous, select that original email through **Add attachments › Add work context**."],
            ["Die drei **[CO-19]**-Mails von Teresa, Vance und Sonia liegen im Posteingang. Die genauen Betreffzeilen stehen in den Mail-Referenzblöcken.",
             "Neuen **Cowork**-Verlauf starten. Kein eigener Skill und kein Plugin nötig. Bei mehrdeutigen Treffern die Originalmail über **Add attachments › Add work context** auswählen."]),
        "prompt": bi(
            "Please look at the three emails in my inbox whose subjects start with [CO-19]. Tell me what I should deal with first and why, and save suitable replies to Teresa and Vance as drafts. Keep anything we still need to confirm as an open question. Don't send anything or move or delete messages.",
            "Schau dir bitte die drei Mails in meinem Posteingang an, deren Betreff mit [CO-19] beginnt. Sag mir, was ich zuerst angehen sollte und warum, und speichere passende Antworten an Teresa und Vance als Entwürfe. Was wir noch bestätigen müssen, soll als offene Frage stehen bleiben. Bitte nichts senden und keine Mails verschieben oder löschen."),
        "follow": bi(
            "Please show me the saved draft for Teresa. Send only that reviewed reply to Teresa Sac after asking me to approve the recipient and message. Don't promise the 05:30 exception is approved, and don't send the other draft.",
            "Zeig mir bitte den gespeicherten Entwurf für Teresa. Sende nur diese geprüfte Antwort an Teresa Sac, nachdem du mich Empfänger und Inhalt freigeben lässt. Sag die 05:30-Ausnahme nicht als genehmigt zu und versende den anderen Entwurf nicht."),
        "followPrep": bi(
            ["Stay in this conversation and open Teresa's actual saved draft. Review the internal recipient and body before selecting **Send**; use a one-time approval."],
            ["In diesem Verlauf bleiben und Teresas tatsächlich gespeicherten Entwurf öffnen. Internen Empfänger und Inhalt vor **Send** prüfen; nur einmalig freigeben."]),
        "expect": bi("Three original messages, justified priorities and two actual unsent Outlook drafts; no invented carrier approval.",
                     "Drei Originalmails, begründete Prioritäten und zwei tatsächlich ungesendete Outlook-Entwürfe; keine erfundene Carrier-Freigabe."),
        "followExpect": bi("Only Teresa's reviewed reply is sent after approval; Vance's draft remains unsent.",
                           "Nur Teresas geprüfte Antwort wird nach Freigabe gesendet; Vances Entwurf bleibt ungesendet."),
        "checks": bi(
            ["Teresa is blocking the carrier note; Vance needs a decision but has no deadline; Sonia explicitly needs no action.",
             "Check actual Outlook drafts and, after approval, the actual sent reply. Reject any fabricated slot confirmation."],
            ["Teresa blockiert die Carrier-Notiz; Vance braucht eine Entscheidung ohne Frist; Sonia braucht ausdrücklich keine Aktion.",
             "Tatsächliche Outlook-Entwürfe und nach Freigabe die gesendete Antwort prüfen. Erfundenen Slot-Zusagen widersprechen."]),
    },
    {
        "code": "CO-20", "title": bi("Organize a small workshop", "Einen kleinen Workshop organisieren"),
        "chatTitle": bi("Operations workshop", "Operations-Workshop"),
        "story": bi("You are coordinating a Contoso workshop. You want an agenda and owner list before committing people or spending money.",
                    "Du koordinierst einen Contoso-Workshop. Du möchtest Agenda und Verantwortlichkeiten klären, bevor du Teilnehmer verbindlich einlädst oder Geld ausgibst."),
        "preparation": bi(
            ["Select **inputs/en/Event_Brief.md** in **OneDrive › Demo-Kit › Cowork › cowork-starters** through **Add attachments › Attach cloud files and folders**.",
             "Start a new Cowork conversation. The workshop is not already scheduled; do not reuse the separate **[CO-21]** review meeting."],
            ["**inputs/de/Event_Brief.md** unter **OneDrive › Demo-Kit › Cowork › cowork-starters** über **Add attachments › Attach cloud files and folders** auswählen.",
             "Neuen Cowork-Verlauf starten. Der Workshop ist noch nicht angelegt; nicht den getrennten **[CO-21]**-Reviewtermin verwenden."]),
        "prompt": bi(
            "I'm organizing the Contoso operations workshop in the selected Event_Brief.md. Please prepare an agenda, a short owner/task list and a calendar invitation preview. Work out the actual date for next Friday and show the time zone and attendees. Stay within the brief's budget and leave unconfirmed arrangements open. Don't create the event, send invitations or book anything yet.",
            "Ich organisiere den Contoso-Operations-Workshop aus der ausgewählten Event_Brief.md. Bereite bitte Agenda, eine kurze Aufgabenliste mit Verantwortlichen und eine Vorschau der Kalendereinladung vor. Löse nächsten Freitag als konkretes Datum auf und zeig Zeitzone und Teilnehmer. Bleib innerhalb des Budgets im Brief und lass unbestätigte Absprachen offen. Noch keinen Termin anlegen, keine Einladungen senden und nichts buchen."),
        "follow": bi(
            "Please create the workshop event from the reviewed preview, with a Teams link. Check whether [CO-20] Contoso operations workshop already exists first. Ask me to approve the date, time zone and internal attendees before sending invitations, then show me the actual calendar event.",
            "Lege bitte den Workshop aus der geprüften Vorschau mit einem Teams-Link an. Prüfe zuerst, ob [CO-20] Contoso operations workshop schon existiert. Lass mich Datum, Zeitzone und interne Teilnehmer vor dem Einladungsversand freigeben und zeig mir danach den tatsächlichen Kalendertermin."),
        "followPrep": bi(
            ["Stay in this conversation. Review the resolved date, **14:00–15:30 Europe/Berlin**, internal attendees and Teams link. Approve only this event."],
            ["In diesem Verlauf bleiben. Konkretes Datum, **14:00–15:30 Europe/Berlin**, interne Teilnehmer und Teams-Link prüfen. Nur diesen Termin freigeben."]),
        "expect": bi("Agenda, task owners and an invitation preview; EUR 500 budget respected; nothing booked or sent.",
                     "Agenda, Aufgabenverantwortliche und Einladungsvorschau; Budget von 500 EUR beachtet; nichts gebucht oder versendet."),
        "followExpect": bi("One verified calendar event after approval, with correct date, attendees, time zone and Teams link.",
                           "Nach Freigabe ein überprüfter Kalendertermin mit richtigem Datum, Teilnehmern, Zeitzone und Teams-Link."),
        "checks": bi(
            ["Agenda covers carriers, construction, safety and next steps; no invented purchase or confirmed venue.",
             "Open the actual Outlook event. A preview is not an invitation; repeat requests must not duplicate it."],
            ["Agenda enthält Carrier, Bau, Sicherheit und nächste Schritte; keine erfundenen Käufe oder gebuchten Räume.",
             "Tatsächlichen Outlook-Termin öffnen. Eine Vorschau ist keine Einladung; bei Wiederholung keinen Doppeltermin anlegen."]),
    },
    {
        "code": "CO-21", "title": bi("Prepare for a real calendar meeting", "Einen echten Kalendertermin vorbereiten"),
        "chatTitle": bi("Operations meeting prep", "Operations-Meetingvorbereitung"),
        "story": bi("Your Contoso operations review is in the calendar. You want a focused briefing with decisions and questions grounded in the planning tracker.",
                    "Dein Contoso-Operations-Review steht im Kalender. Du möchtest ein kompaktes Briefing mit Entscheidungen und Fragen, die auf dem Planungstracker beruhen."),
        "preparation": bi(
            ["The **[CO-21] Contoso operations review** event is in Outlook, Thursday **10:00**, in the VM's local time zone. Check its actual date and attendees.",
             "In a new Cowork conversation, select that event via **Add attachments › Add work context** and **inputs/en/Meeting_Brief.md** plus **Contoso_Fargo_GoLive_Readiness_Tracker.xlsx** from this demo's OneDrive folder."],
            ["**[CO-21] Contoso-Operations-Review** steht Donnerstag **10:00** in Outlook, in der lokalen Zeitzone der VM. Tatsächliches Datum und Teilnehmer prüfen.",
             "Im neuen Cowork-Verlauf den Termin über **Add attachments › Add work context** sowie **inputs/de/Meeting_Brief.md** und **Contoso_Fargo_GoLive_Readiness_Tracker.xlsx** aus dem OneDrive-Ordner dieser Demo auswählen."]),
        "prompt": bi(
            "Please prepare me for the selected [CO-21] Contoso operations review. Use Meeting_Brief.md and the selected readiness tracker to create a one-page briefing with the open decisions and questions for Teresa, Vance and Sonia. Link each important point to its source row. Keep the tracker's 2027 planning dates as they are and flag what we don't know. Don't invent a past meeting transcript or change the tracker.",
            "Bereite mich bitte auf den ausgewählten [CO-21] Contoso-Operations-Review vor. Erstelle aus Meeting_Brief.md und dem ausgewählten Readiness-Tracker ein einseitiges Briefing mit offenen Entscheidungen und Fragen an Teresa, Vance und Sonia. Verknüpfe jeden wichtigen Punkt mit seiner Quellzeile. Behalte die Planungsdaten für 2027 bei und markiere, was wir nicht wissen. Erfinde kein früheres Meetingtranskript und ändere den Tracker nicht."),
        "follow": bi(
            "Please turn this briefing into a short agenda and draft an email to Teresa with our three most important open questions. Save it as a draft, don't send it, and show me the source links and the saved draft.",
            "Mach bitte aus diesem Briefing eine kurze Agenda und formuliere eine Mail an Teresa mit unseren drei wichtigsten offenen Fragen. Speichere sie als Entwurf, versende sie nicht und zeig mir die Quelllinks und den gespeicherten Entwurf."),
        "followPrep": bi(["Use this same briefing and its verified sources; do not add a different meeting or a supposed transcript."],
                              ["Dasselbe Briefing mit seinen geprüften Quellen nutzen; keinen anderen Termin und kein vermeintliches Transkript ergänzen."]),
        "expect": bi("A real briefing with links to R02, R04 and R08, actual calendar details and clearly open decisions.",
                     "Ein tatsächliches Briefing mit Links zu R02, R04 und R08, echten Kalenderdetails und klar offenen Entscheidungen."),
        "followExpect": bi("A short agenda and a saved, unsent Outlook draft to Teresa; original planning sources unchanged.",
                           "Kurze Agenda und gespeicherter, ungesendeter Outlook-Entwurf an Teresa; ursprüngliche Planungsquellen unverändert."),
        "checks": bi(
            ["R02 carrier confirmation, R04 electrical delay and R08 training dates are open; none is reported as approved.",
             "Read actual calendar details. Tracker dates are April–June 2027, not current operational promises. No transcript exists."],
            ["R02 Carrier-Bestätigung, R04 Elektroverzug und R08 Trainingstermine sind offen; nichts davon als freigegeben darstellen.",
             "Echte Kalenderdetails lesen. Tracker-Daten betreffen April–Juni 2027, keine aktuellen Betriebszusagen. Es gibt kein Transkript."]),
    },
    {
        "code": "CO-22", "title": bi("Turn revenue figures into slides", "Umsatzzahlen in Folien verwandeln"),
        "chatTitle": bi("Q3 revenue slides", "Q3-Umsatzfolien"),
        "story": bi("You need a small Contoso Q3 revenue update. You want readable slides whose figures still match the selected Excel, not a generic success story.",
                    "Du brauchst ein kleines Contoso-Q3-Umsatzupdate. Du möchtest verständliche Folien, deren Zahlen zur ausgewählten Excel passen, statt einer allgemeinen Erfolgsgeschichte."),
        "preparation": bi(
            ["Select one fresh **runs/en/01/Revenue_Q3.xlsx** through **Add attachments › Attach cloud files and folders**; use **02–04** for repeats.",
             "The workbook contains six synthetic monthly region records in **EUR**, July–September **2026**. Read **Readme** for the current run's output folder. No custom PowerPoint skill is required."],
            ["Eine frische **runs/de/01/Revenue_Q3.xlsx** über **Add attachments › Attach cloud files and folders** auswählen; für Wiederholungen **02–04** nutzen.",
             "Die Excel enthält sechs synthetische Monats-/Regionswerte in **EUR**, Juli–September **2026**. Im Blatt **Readme** steht der Ergebnisordner dieses Laufs. Kein eigener PowerPoint-Skill nötig."]),
        "prompt": bi(
            "Please turn the selected Contoso Revenue_Q3.xlsx into a three-slide Q3 2026 update: an overview against budget, a monthly revenue chart and the key points by region. Use EUR throughout and base every figure on the Revenue sheet. Separate observations from possible explanations; don't invent causes. Save the PowerPoint in the output folder named on Readme and show me its actual link. Leave the source Excel unchanged.",
            "Mach bitte aus der ausgewählten Contoso Revenue_Q3.xlsx ein Q3-2026-Update mit drei Folien: Überblick gegen Budget, ein Diagramm mit dem Monatsumsatz und die wichtigsten Punkte je Region. Nutze durchgehend EUR und leite jede Zahl aus dem Blatt Revenue ab. Trenne Beobachtungen von möglichen Erklärungen und erfinde keine Ursachen. Speichere die PowerPoint im Ergebnisordner aus Readme und zeig mir den tatsächlichen Link. Lass die Quell-Excel unverändert."),
        "follow": bi(
            "Please add one slide about the North region, with its Q3 revenue and budget variance. Keep the existing three slides and update this same presentation. Show me the revised file, and don't change the Excel.",
            "Ergänze bitte eine Folie zur Region North mit ihrem Q3-Umsatz und der Budgetabweichung. Behalte die drei vorhandenen Folien und aktualisiere dieselbe Präsentation. Zeig mir die überarbeitete Datei und ändere die Excel nicht."),
        "followPrep": bi(["Keep this conversation, the same cloud Excel and the actual generated presentation. Approve editing that presentation only."],
                              ["Diesen Verlauf, dieselbe Cloud-Excel und die tatsächlich erzeugte Präsentation behalten. Nur deren Bearbeitung freigeben."]),
        "expect": bi("An actual three-slide PowerPoint with correct monthly, quarterly and region figures and the requested output location.",
                     "Eine tatsächliche PowerPoint mit drei Folien, richtigen Monats-, Quartals- und Regionszahlen am gewünschten Speicherort."),
        "followExpect": bi("The same presentation now has four slides; the original input and the other run folders remain unchanged.",
                           "Dieselbe Präsentation hat nun vier Folien; Originaleingabe und andere Laufordner bleiben unverändert."),
        "checks": bi(
            ["Open PowerPoint and compare all chart values, totals and budget variances with the presenter reference; no unit/currency changes.",
             "Show the actual file URL and output folder. No source for business causes is provided; explanations remain hypotheses."],
            ["PowerPoint öffnen und Diagrammwerte, Summen und Budgetabweichungen mit der Presenter-Referenz abgleichen; keine Einheiten-/Währungswechsel.",
             "Echte Datei-URL und Ergebnisordner zeigen. Es gibt keine Quelle für Geschäftstreiber; Erklärungen bleiben Hypothesen."]),
    },
]


def revenue_workbook(lang, run):
    wb = Workbook()
    ws = wb.active
    ws.title = "Revenue"
    BASE.write_sheet(ws, ["Month", "Region", "Revenue (EUR)", "Budget (EUR)"], REVENUE,
                     widths={1: 18, 2: 18, 3: 22, 4: 22})
    table = Table(displayName="ContosoRevenue", ref="A1:D7")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.sheet_view.showGridLines = False
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.font = Font(name="Arial", size=11, color="0000FF" if c.column >= 3 else "000000")
            c.alignment = Alignment(vertical="center")
            if c.column >= 3:
                c.number_format = '#,##0;(#,##0);"-"'
    for c in ws[1]:
        c.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    readme = wb.create_sheet("Readme")
    BASE.write_sheet(readme, ["Field", "Value"], [
        ("Company", "Contoso - synthetic demo data"),
        ("Period / currency", "Q3 2026 / EUR"),
        ("Source", "Synthetic Contoso dataset, 2026-10-14, six monthly region records; no real company data."),
        ("Run", f"{lang.upper()}-{run:02}"),
        ("Output folder", f"Demo-Kit/Cowork/cowork-starters/runs/{lang}/{run:02}/results"),
        ("Input contract", "Revenue contains raw inputs only. Keep it unchanged; compute totals in the output."),
    ], widths={1: 24, 2: 105})
    for row in readme:
        for c in row:
            c.font = Font(name="Arial", size=11, bold=c.row == 1, color="FFFFFF" if c.row == 1 else "000000")
            c.alignment = Alignment(wrap_text=True, vertical="top")
    path = PACKAGE / "runs" / lang / f"{run:02}" / "Revenue_Q3.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    BASE.save_wb(wb, path)
    wb.close()
    write(path.parent / "results" / "README.md", bi(
        "# Output folder\n\nSave this run's PowerPoint here. Empty of results until Cowork creates the real output.",
        "# Ergebnisordner\n\nHier die PowerPoint dieses Laufs ablegen. Noch keine Ergebnisse vorhanden; Cowork erstellt die echte Ausgabe.")[lang])


def build_data():
    for lang in ("en", "de"):
        write(PACKAGE / "inputs" / lang / "Event_Brief.md", EVENT[lang])
        write(PACKAGE / "inputs" / lang / "Meeting_Brief.md", MEETING_BRIEF[lang])
        for m in MAILS:
            msg = EmailMessage(policy=SMTP)
            msg["From"] = f'{m["sender"]} <{PEOPLE[m["sender"]]}@{DOMAIN}>'
            msg["To"] = f"MOD Administrator <admin@{DOMAIN}>"
            msg["Subject"] = m["subject"][lang]
            msg["Date"] = "Wed, 14 Oct 2026 08:00:00 +0200"
            msg["Message-ID"] = f'<cowork-starters-{lang}-{m["file"]}@{DOMAIN}>'
            msg.set_content(m["body"][lang])
            path = PACKAGE / f"mails-{lang}" / m["file"]
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(msg.as_bytes())
        for n in range(1, 5):
            revenue_workbook(lang, n)
        write(PACKAGE / f"Preparation-{lang}.md", bi(
            "# Cowork starters preparation\n\nFiles: OneDrive > Demo-Kit > Cowork > cowork-starters.\n\n1. Review the global -Tenant -WhatIf command before the presenter runs -Tenant. Preserve changed cloud work files first; the installer is not a rehearsal reset.\n2. Verify the three [CO-19] inbox messages from Teresa, Vance and Sonia, and the [CO-21] Contoso operations review in Outlook (Thursday 10:00, VM local time zone). Read the event's actual date/time zone. The installer verifies existing senders/mailboxes; no extra sign-in is introduced by this package.\n3. CO-19 uses the original messages, not uploaded .eml copies. CO-20 uses inputs/en/Event_Brief.md; its workshop is NOT pre-created. CO-21 uses inputs/en/Meeting_Brief.md and the shared readiness tracker; 2027 dates are a planning snapshot.\n4. CO-22: reference just runs/en/01/Revenue_Q3.xlsx, and use 02-04 for fresh repeats. Read Readme for the output location. Keep source records unchanged.\n5. Start a separate Cowork conversation for each exercise. No custom skills/plugins are required. Select existing cloud references through Add attachments, never replace them with local uploads.\n6. Initial actions produce drafts/previews. Review recipient/date/time zone and approve each requested send/create/edit individually. No Always allow, purchases, deleted mails or broad inbox rules.\n7. Check actual Outlook items, file links and PowerPoint contents. An output in a different folder is placement incomplete, not success. The installer and package checker cannot prove a live Cowork result.\n8. For repeats, keep calendar subjects unique or reuse the verified existing event; do not send replies twice. No global installer rerun as reset.",
            "# Cowork-Einstiege vorbereiten\n\nDateien: OneDrive > Demo-Kit > Cowork > cowork-starters.\n\n1. Vor dem Presenter-Lauf mit -Tenant den globalen -Tenant -WhatIf-Befehl prüfen. Geänderte Cloud-Arbeitsdateien vorher sichern; der Installer ist kein Probelauf-Reset.\n2. Die drei [CO-19]-Posteingangsmails von Teresa, Vance und Sonia und [CO-21] Contoso-Operations-Review in Outlook prüfen (Donnerstag 10:00, lokale Zeitzone der VM). Tatsächliches Datum und Zeitzone lesen. Der Installer prüft vorhandene Absender/Mailboxen; dieses Paket benötigt keine zusätzliche Anmeldung.\n3. CO-19 nutzt die Originalmails, keine hochgeladenen .eml-Kopien. CO-20 nutzt inputs/de/Event_Brief.md; sein Workshop ist NICHT vorab angelegt. CO-21 nutzt inputs/de/Meeting_Brief.md und den geteilten Readiness-Tracker; 2027-Daten sind ein Planungsstand.\n4. CO-22: nur runs/de/01/Revenue_Q3.xlsx referenzieren, für frische Wiederholungen 02-04 nutzen. In Readme den Ergebnisordner lesen. Quelldaten unverändert lassen.\n5. Pro Exercise einen eigenen Cowork-Verlauf starten. Keine eigenen Skills/Plugins nötig. Bestehende Cloud-Referenzen über Add attachments auswählen und nie durch lokale Uploads ersetzen.\n6. Anfangsaufträge erstellen Entwürfe/Vorschauen. Empfänger/Datum/Zeitzone prüfen und jeden Versand/jede Anlage/Bearbeitung einzeln freigeben. Kein Always allow, keine Käufe, gelöschten Mails oder breiten Posteingangsregeln.\n7. Tatsächliche Outlook-Objekte, Dateilinks und PowerPoint-Inhalt prüfen. Ausgabe in einem anderen Ordner bedeutet unvollständige Ablage, keinen Erfolg. Installer und Paketprüfer beweisen keinen Live-Cowork-Erfolg.\n8. Für Wiederholungen vorhandenen Kalendertermin geprüft wiederverwenden; Antworten nicht doppelt versenden. Kein globaler Installer-Rerun als Reset.")[lang])
    write_json(PACKAGE / "shared.json", ["cowork-copilot/Contoso_Fargo_GoLive_Readiness_Tracker.xlsx"])
    manual = bi(
        ["Verify the three [CO-19] inbox mails and the [CO-21] calendar event; check the event's actual date and VM time zone.",
         "Open Cowork with the licensed demo account; verify availability, usage billing and cloud-source access. No custom skills required.",
         "The [CO-20] workshop is created only in the approved follow-up. Keep sends/creates individually approved.",
         "Select a fresh CO-22 run (01-04); preserve edited files before any installer update. Do not use installer reruns as a reset."],
        ["Drei [CO-19]-Posteingangsmails und den [CO-21]-Kalendertermin prüfen; tatsächliches Datum und VM-Zeitzone beachten.",
         "Cowork mit dem lizenzierten Demokonto öffnen; Verfügbarkeit, Nutzungsabrechnung und Cloud-Zugriff prüfen. Keine eigenen Skills nötig.",
         "Der [CO-20]-Workshop entsteht erst im freigegebenen Folgeauftrag. Versand/Anlage jeweils einzeln freigeben.",
         "Frischen CO-22-Lauf (01-04) wählen; bearbeitete Dateien vor Installer-Updates sichern. Installer-Reruns sind kein Reset."])
    write_json(PACKAGE / "tenant.json", {
        "mails": [f'mails-{{lang}}/{m["file"]}' for m in MAILS], "events": [MEETING],
        "links": [{"name": bi("Microsoft Copilot Cowork", "Microsoft Copilot Cowork"), "url": "https://copilot.cloud.microsoft/"}],
        "manual": manual,
    })
    write_json(PACKAGE / "manifest.json", {"title": bi("Simple Cowork starters", "Einfache Cowork-Einstiege"),
                                         "next": manual, "files": []})
    write_json(ROOT / "tools" / "expected" / "cowork-starters.json", {
        "revenue": {
            "period": "Q3 2026", "currency": "EUR", "records": 6,
            "total": 630000, "budget": 625000, "variance": 5000, "variance_percent": 0.8,
            "monthly": {"2026-07": 180000, "2026-08": 210000, "2026-09": 240000},
            "regions": {"North": {"revenue": 360000, "budget": 355000, "variance": 5000},
                        "South": {"revenue": 270000, "budget": 270000, "variance": 0}},
            "initial_slides": 3, "followup_slides": 4, "explanations": "Hypotheses only; no driver evidence provided.",
        },
        "triage": {"messages": 3, "drafts": 2, "first": "Teresa Sac", "no_action": "Sonia Rees",
                   "followup_send": "Teresa Sac only after approval; no approved 05:30 slot is evidenced."},
        "event": {"subject": "[CO-20] Contoso operations workshop", "budget_eur": 500,
                  "start": "14:00", "end": "15:30", "timezone": "Europe/Berlin",
                  "date": "Next Friday after current conversation; review actual date before creating.",
                  "organizer": "MOD Administrator", "attendees": list(PEOPLE),
                  "initial": "Preview only", "followup": "One approved event with Teams link"},
        "meeting": {"subject": MEETING["subject"], "sources": ["R02", "R04", "R08"],
                    "planning_horizon": "April-June 2027", "transcript": False, "followup": "Unsent draft to Teresa"},
    })


def demo_page():
    steps = [{
        "kind": "callout", "tone": "info",
        "title": bi("Start with familiar work", "Mit vertrauter Arbeit beginnen"),
        "text": bi(
            "Four small tasks using built-in Cowork capabilities. Files: **OneDrive › Demo-Kit › Cowork › cowork-starters**. Each section is a separate conversation; follow-ups stay in that section's conversation.",
            "Vier kleine Aufgaben mit integrierten Cowork-Fähigkeiten. Dateien: **OneDrive › Demo-Kit › Cowork › cowork-starters**. Jeden Abschnitt als eigenen Verlauf starten; Folgeaufträge bleiben im jeweiligen Verlauf."),
    }, {
        "kind": "screen", "title": bi("Where do I start?", "Wo starte ich?"),
        "screen": {
            "layout": "chat", "app": "Microsoft Copilot", "url": "copilot.cloud.microsoft", "mode": "cowork",
            "chatText": bi("Describe the small task here", "Hier den kleinen Auftrag eingeben"),
            "markers": [
                {"target": "mode-toggle", "text": bi("Select **Cowork** at the top.", "Oben **Cowork** auswählen.")},
                {"target": "input", "text": bi("Use **Add attachments** for the exact existing sources, then paste the prompt.",
                                              "Über **Add attachments** die genauen vorhandenen Quellen wählen und den Prompt einfügen.")},
            ],
        },
    }, {
        "kind": "callout", "tone": "warning",
        "title": bi("Draft first, approve individually", "Erst Entwürfe, dann einzeln freigeben"),
        "text": bi(
            "Review every recipient, date, time zone and requested file edit. Use one-time approval; do not select **Always allow**. New Office files appear under **Output folder**; open the actual link and verify content and location. A preview or filename alone is not completion.",
            "Jeden Empfänger, Termin, jede Zeitzone und angeforderte Dateibearbeitung prüfen. Einmalig freigeben; nicht **Always allow** wählen. Neue Office-Dateien stehen unter **Output folder**; tatsächlichen Link öffnen und Inhalt sowie Ablage prüfen. Vorschau oder Dateiname allein bedeuten keinen Abschluss."),
    }]
    for s in SCENARIOS:
        steps.append({"kind": "section", "title": s["title"], "exercise": {
            "code": s["code"], "chatTitle": s["chatTitle"], "story": s["story"]}})
        if s["code"] == "CO-19":
            for m in MAILS:
                steps.append({"kind": "content", "title": bi(
                    f'Email from {m["sender"]} (in the inbox)', f'Mail von {m["sender"]} (im Posteingang)'),
                    "text": {lang: f'Subject: {m["subject"][lang]}\n\n{m["body"][lang]}' for lang in ("en", "de")}})
        if s["code"] == "CO-21":
            steps.append({"kind": "content", "title": bi("Meeting reference (in Outlook)", "Terminreferenz (in Outlook)"),
                          "text": {lang: f'{MEETING["subject"][lang]}\n\n{MEETING["body"][lang]}' for lang in ("en", "de")}})
        steps.extend([
            {"kind": "prompt", "title": bi("Start with this request", "Mit diesem Auftrag starten"),
             "preparation": s["preparation"], "text": s["prompt"], "expect": s["expect"],
             "where": bi("A new Cowork conversation.", "Neuer Cowork-Verlauf.")},
            {"kind": "prompt", "title": bi("Short follow-up", "Kurzer Folgeauftrag"),
             "preparation": s["followPrep"], "text": s["follow"], "expect": s["followExpect"],
             "where": bi("The same conversation, after reviewing the first result.", "Derselbe Verlauf, nach Prüfung des ersten Ergebnisses.")},
            {"kind": "checklist", "title": bi("Check the actual result", "Das tatsächliche Ergebnis prüfen"),
             "checks": [{"status": "info", "label": {lang: s["checks"][lang][i] for lang in ("en", "de")}} for i in range(2)]},
        ])
    steps.append({
        "kind": "tiles", "title": bi("More ideas for later", "Weitere Ideen für später"),
        "tiles": [
            {"title": bi("Daily briefing", "Tagesbriefing"), "text": bi("Upcoming meetings, relevant emails and next actions.", "Anstehende Termine, relevante Mails und nächste Aktionen.")},
            {"title": bi("Prepare for a 1:1", "1:1 vorbereiten"), "text": bi("A focused agenda from existing work and open questions.", "Gezielte Agenda aus vorhandener Arbeit und offenen Fragen.")},
            {"title": bi("Organize the week", "Woche organisieren"), "text": bi("Review calendar conflicts and propose work blocks before changing anything.", "Kalenderkonflikte prüfen und Arbeitsblöcke vor Änderungen vorschlagen.")},
            {"title": bi("Stakeholder update", "Stakeholder-Update"), "text": bi("A source-backed project update and an unsent communication draft.", "Belegtes Projektupdate und ungesendeter Kommunikationsentwurf.")},
        ],
    })
    return {
        "id": "cowork-starters", "band": "usecases", "accent": "cowork", "data": "cowork-starters",
        "title": bi("Simple Cowork starters", "Einfache Cowork-Einstiege"),
        "kicker": bi("Email, workshop, meeting and revenue slides", "Mails, Workshop, Meeting und Umsatzfolien"),
        "why": bi("Start small and show a concrete result, not just a chat response.", "Klein anfangen und ein konkretes Ergebnis zeigen statt nur einer Chatantwort."),
        "steps": steps,
        "sources": [
            {"label": bi("Use Cowork: sources, skills and action approval", "Cowork nutzen: Quellen, Skills und Freigaben"), "url": USE},
            {"label": bi("Cowork capabilities and starter suggestions", "Cowork-Fähigkeiten und Einstiegsvorschläge"), "url": OVERVIEW},
            {"label": bi("Cowork prerequisites", "Cowork-Voraussetzungen"), "url": GET_STARTED},
            {"label": bi("Cowork common questions and limitations", "Cowork-Fragen und Grenzen"), "url": FAQ},
            {"label": bi("Official role-based scenario material", "Offizielle rollenbezogene Szenariomaterialien"),
             "url": "https://learn.microsoft.com/training/paths/empower-workforce-copilot-use-cases/"},
        ],
    }


def build_decks(decks):
    if not decks.is_dir():
        raise FileNotFoundError(f"Deck folder not found: {decks}")
    cowork = json.loads((decks / "cowork.json").read_text(encoding="utf-8"))
    bonn = json.loads((decks / "cowork-bonn.json").read_text(encoding="utf-8"))
    for f in decks.glob("*.json"):
        existing = json.loads(f.read_text(encoding="utf-8"))
        for d in existing["demos"]:
            if d["id"] == "cowork-starters":
                continue
            exercises = [d.get("exercise"), *(s.get("exercise") for s in d["steps"])]
            if any(e and e["code"] in {s["code"] for s in SCENARIOS} for e in exercises):
                raise ValueError(f"Starter exercise code already used in {f.name}/{d['id']}")
    page = demo_page()
    for deck, version, band, canonical in [(cowork, "1.7", "usecases", True), (bonn, "1.6", "run", False)]:
        view = deepcopy(page)
        view["band"] = band
        if not canonical:
            del view["data"]
        deck["demos"] = [d for d in deck["demos"] if d["id"] != page["id"]]
        deck["demos"].insert(0, view)
        title = page["title"]
        detail = bi("Four small exercises with built-in capabilities and reviewed follow-ups.",
                    "Vier kleine Exercises mit integrierten Fähigkeiten und geprüften Folgeaufträgen.")
        deck["agenda"]["items"] = [a for a in deck["agenda"]["items"] if a["title"] != title]
        deck["agenda"]["items"].insert(0, {"title": title, "detail": detail})
        deck["meta"].update(version=version, date=DATE)
        note = bi("Also includes four simple starters on a separate page. RFI / RFP has its own deck switch.",
                  "Zusätzlich vier einfache Einstiege auf einer eigenen Seite. RFI / RFP hat einen eigenen Deck-Schalter.")
        for lang in ("en", "de"):
            if note[lang] not in deck["meta"]["scope"][lang]:
                deck["meta"]["scope"][lang] += " " + note[lang]
    write_json(decks / "cowork.json", cowork)
    write_json(decks / "cowork-bonn.json", bonn)
    rfi = deepcopy(next(d for d in cowork["demos"] if d["id"] == "cowork-rfi-rfp"))
    del rfi["data"]
    rfi["band"] = "rfi"
    write_json(decks / "rfi-rfp.json", {
        "meta": {
            "title": bi("RFI / RFP", "RFI / RFP"), "version": "1.0", "date": DATE, "order": 2,
            "scope": bi("A standalone switch for CO-17: answer a complex tender and verify the actual draft outputs.",
                        "Eigener Schalter für CO-17: komplexe Ausschreibung beantworten und tatsächliche Entwurfsdateien prüfen."),
            "dataNote": bi("Uses the existing Cowork RFI package; no second installation.", "Nutzt das vorhandene Cowork-RFI-Paket; keine zweite Installation."),
        },
        "agenda": {"title": bi("Agenda & setup", "Agenda & Setup"), "items": [
            {"title": rfi["title"], "detail": rfi["kicker"]},
        ]},
        "files": {
            "demoData": {},
            "note": bi("The same global setup supplies the RFI files. Then select one fresh cloud workbook.",
                       "Dasselbe globale Setup liefert die RFI-Dateien. Danach eine frische Cloud-Excel auswählen."),
            "lines": bi(
                ["Files: **OneDrive › Demo-Kit › Cowork › cowork-rfi-rfp**. Read **Preparation-en.md**.",
                 "Select the existing **contoso-rfi-rfp-response** skill and a fresh compact/full run. Enable browser access and sign in to Atlas before presenting.",
                 "Use unsaved Atlas previews only. Verify edits to the same Excel and the two actual Word outputs in the run's results folder.",
                 "Keep changed cloud workbooks before updating tenant data. Preview **-Tenant -WhatIf** first; an installer rerun is not a reset.",
                 "This deck has its own switch. The existing CO-17 page under Cowork uses the same scenario and data."],
                ["Dateien: **OneDrive › Demo-Kit › Cowork › cowork-rfi-rfp**. **Preparation-de.md** lesen.",
                 "Vorhandenen Skill **contoso-rfi-rfp-response** und frischen kompakten/großen Lauf wählen. Browserzugriff aktivieren und Atlas vorher anmelden.",
                 "Atlas nur als ungespeicherte Vorschau nutzen. Bearbeitung derselben Excel und zwei tatsächliche Word-Ausgaben im Ergebnisordner des Laufs prüfen.",
                 "Geänderte Cloud-Excel vor Tenant-Updates sichern. Erst **-Tenant -WhatIf** prüfen; ein Installer-Rerun ist kein Reset.",
                 "Dieses Deck hat einen eigenen Schalter. Die bestehende CO-17-Seite unter Cowork nutzt dasselbe Szenario und dieselben Daten."]),
        },
        "bands": [{"id": "rfi", "label": bi("RFI / RFP", "RFI / RFP")}], "demos": [rfi],
    })


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--decks", type=Path, default=ROOT.parent / "demo-kit-new2026" / "decks")
    args = parser.parse_args()
    build_data()
    build_decks(args.decks)
    print("Built Cowork starters (4 exercises, 8 revenue runs) and the standalone RFI deck. Run Update-Manifest.ps1 -All.")
