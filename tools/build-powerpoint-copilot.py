"""Builds the fictional Contoso brand demo files for the powerpoint-copilot demo.

Run:  python tools/build-powerpoint-copilot.py
Needs: pip install python-docx python-pptx pillow
All content is fictional (Contoso is Microsoft's demo company).
"""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor, Inches, Cm
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.dml.color import RGBColor as PRGB
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as PI, Pt as PPt

OUT = Path(__file__).resolve().parent.parent / "powerpoint-copilot"
OUT.mkdir(exist_ok=True)

NAVY = (0x0B, 0x3D, 0x91)
TEAL = (0x00, 0xA3, 0xA1)
AMBER = (0xFF, 0xB0, 0x00)
CHARCOAL = (0x2B, 0x2B, 0x2B)
MIST = (0xF4, 0xF6, 0xF8)
PALETTE = [
    ("Contoso Navy", NAVY, "Primary. Headlines, title slides, logo."),
    ("Harbor Teal", TEAL, "Secondary. Charts, highlights, icons."),
    ("Signal Amber", AMBER, "Accent only. Max. one element per slide (call to action, warning)."),
    ("Charcoal", CHARCOAL, "Body text."),
    ("Mist", MIST, "Backgrounds and cards."),
]


def font(size, bold=False):
    for name in (["segoeuib.ttf"] if bold else []) + ["segoeui.ttf", "arial.ttf"]:
        try:
            return ImageFont.truetype(name, size)
        except OSError:
            continue
    return ImageFont.load_default()


def hexagon(draw, cx, cy, r, fill):
    import math
    pts = [(cx + r * math.cos(math.radians(60 * i - 30)), cy + r * math.sin(math.radians(60 * i - 30))) for i in range(6)]
    draw.polygon(pts, fill=fill)


def build_logos():
    # horizontal logo
    img = Image.new("RGBA", (1200, 360), (255, 255, 255, 0))
    d = ImageDraw.Draw(img)
    hexagon(d, 180, 180, 130, NAVY)
    hexagon(d, 180, 180, 62, TEAL)
    d.text((350, 95), "Contoso", font=font(150, True), fill=NAVY)
    img.save(OUT / "Contoso_Logo.png")
    # white version for dark backgrounds
    w = Image.new("RGBA", (1200, 360), (255, 255, 255, 0))
    d = ImageDraw.Draw(w)
    hexagon(d, 180, 180, 130, (255, 255, 255))
    hexagon(d, 180, 180, 62, TEAL)
    d.text((350, 95), "Contoso", font=font(150, True), fill=(255, 255, 255))
    w.save(OUT / "Contoso_Logo_White.png")
    # icon
    i = Image.new("RGBA", (512, 512), (255, 255, 255, 0))
    d = ImageDraw.Draw(i)
    hexagon(d, 256, 256, 230, NAVY)
    hexagon(d, 256, 256, 110, TEAL)
    i.save(OUT / "Contoso_Icon.png")
    # brand pattern / background image
    p = Image.new("RGB", (1920, 1080), NAVY)
    d = ImageDraw.Draw(p)
    for row in range(-1, 9):
        for col in range(-1, 15):
            cx = col * 150 + (75 if row % 2 else 0)
            cy = row * 130
            hexagon(d, cx, cy, 70, (0x0E, 0x47, 0xA3))
    hexagon(d, 1500, 760, 260, TEAL)
    hexagon(d, 1500, 760, 120, NAVY)
    p.save(OUT / "Contoso_Brand_Background.png")


def set_cell_bg(cell, rgb):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), "%02X%02X%02X" % rgb)
    tcPr.append(shd)


def build_guidelines():
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = "Segoe UI"
    st.font.size = Pt(10.5)
    st.font.color.rgb = RGBColor(*CHARCOAL)
    for s in ("Heading 1", "Heading 2", "Title"):
        doc.styles[s].font.name = "Segoe UI Semibold"
        doc.styles[s].font.color.rgb = RGBColor(*NAVY)

    doc.add_picture(str(OUT / "Contoso_Logo.png"), width=Inches(2.4))
    doc.add_paragraph("Contoso Brand Guidelines", style="Title")
    doc.add_paragraph("Version 3.2 · Corporate Communications · For internal and partner use")

    doc.add_heading("1. Our brand in one sentence", 1)
    doc.add_paragraph(
        "Contoso keeps goods moving. Our brand is calm, precise and dependable: we show real operations, "
        "clear numbers and the people who make it work. We never shout."
    )

    doc.add_heading("2. Logo", 1)
    for t in [
        "Use the full-color logo on white or Mist backgrounds and the white logo on Contoso Navy or on photos.",
        "Keep clear space around the logo equal to the height of the hexagon mark.",
        "Minimum width: 30 mm in print, 120 px on screen.",
        "Never stretch, recolor, rotate or add effects (shadows, glow, outlines) to the logo.",
        "On slides, the logo sits bottom right on content slides and top left on title slides.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("3. Colors", 1)
    tbl = doc.add_table(rows=1, cols=4)
    tbl.style = "Table Grid"
    for c, h in zip(tbl.rows[0].cells, ["", "Name", "HEX / RGB", "Use"]):
        c.text = h
    for name, rgb, use in PALETTE:
        r = tbl.add_row().cells
        set_cell_bg(r[0], rgb)
        r[1].text = name
        r[2].text = "#%02X%02X%02X  ·  %d / %d / %d" % (rgb + rgb)
        r[3].text = use
    for row in tbl.rows:
        row.cells[0].width = Cm(1.5)
    doc.add_paragraph()
    doc.add_paragraph("Ratio rule: about 60 % white/Mist, 30 % Navy, 10 % Teal. Amber is a signal, not a decoration.")

    doc.add_heading("4. Typography", 1)
    for t in [
        "Headlines: Segoe UI Semibold, sentence case, max. 8 words.",
        "Body text: Segoe UI Regular, at least 18 pt on slides, 10.5 pt in documents.",
        "Numbers: big and bold in Contoso Navy, with the unit in Charcoal next to it (e.g. 24 weeks).",
        "No underlines, no all caps except abbreviations such as KPI or ETA.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("5. Imagery", 1)
    for t in [
        "Real operations: warehouses, docks, forklifts, people at work. Bright, natural light.",
        "Safety first: everyone in a photo of an operations area wears PPE (helmet, high-visibility vest, gloves).",
        "No stock-photo handshakes, no abstract light trails, no cartoon illustrations.",
        "Icons: simple line icons in Navy or Teal. The hexagon may be used as a background pattern.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("6. Presentations", 1)
    for t in [
        "Format 16:9. Title slide on Contoso Navy with the white logo and the brand background.",
        "One message per slide; the headline states the message, not the topic (\u201cDock 3 closes on Apr 27\u201d, not \u201cDock 3\u201d).",
        "Max. 5 bullets per slide, max. 2 lines per bullet.",
        "Charts: Navy for actuals, Teal for plan/target, Amber only for the one value that needs attention.",
        "Last slide: next steps with owners and dates. No \u201cThank you\u201d or \u201cQuestions?\u201d slides.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    doc.add_heading("7. Voice and tone", 1)
    for t in [
        "Clear and friendly, like a good shift lead: short sentences, active voice, concrete numbers.",
        "Say \u201cwe\u201d for Contoso and \u201cyou\u201d for the reader.",
        "Avoid jargon and buzzwords (\u201csynergies\u201d, \u201cbest-in-class\u201d, \u201cleverage\u201d).",
        "German texts use the informal \u201cdu\u201d internally and \u201cSie\u201d for customers and partners.",
    ]:
        doc.add_paragraph(t, style="List Bullet")

    p = doc.add_paragraph("Contoso is a fictional company used for Microsoft demos. All names and data are invented.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].font.size = Pt(8)
    doc.save(OUT / "Contoso_Brand_Guidelines.docx")


def rgb(t):
    return PRGB(*t)


def build_deck():
    prs = Presentation()
    prs.slide_width, prs.slide_height = PI(13.333), PI(7.5)
    blank = prs.slide_layouts[6]
    W, H = prs.slide_width, prs.slide_height

    def text(slide, x, y, w, h, s, size=20, bold=False, color=CHARCOAL, fontname="Segoe UI"):
        tb = slide.shapes.add_textbox(x, y, w, h)
        tf = tb.text_frame
        tf.word_wrap = True
        lines = s if isinstance(s, list) else [s]
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            r = p.add_run()
            r.text = line
            r.font.size = PPt(size)
            r.font.bold = bold
            r.font.name = fontname
            r.font.color.rgb = rgb(color)
            p.space_after = PPt(8)
        return tb

    def content_slide(title):
        s = prs.slides.add_slide(blank)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, PI(0.18), H)
        bar.fill.solid(); bar.fill.fore_color.rgb = rgb(NAVY); bar.line.fill.background()
        text(s, PI(0.7), PI(0.5), PI(11.5), PI(1.0), title, 32, True, NAVY, "Segoe UI Semibold")
        s.shapes.add_picture(str(OUT / "Contoso_Logo.png"), W - PI(1.9), H - PI(0.75), width=PI(1.5))
        return s

    # 1 title
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(str(OUT / "Contoso_Brand_Background.png"), 0, 0, W, H)
    s.shapes.add_picture(str(OUT / "Contoso_Logo_White.png"), PI(0.7), PI(0.6), width=PI(2.6))
    text(s, PI(0.7), PI(2.3), PI(8), PI(1.8), "Fargo expansion: where we stand", 44, True, (255, 255, 255), "Segoe UI Semibold")
    text(s, PI(0.7), PI(4.4), PI(8), PI(0.8), "Operations update · Week 14", 22, False, (255, 255, 255))

    # 2 key numbers
    s = content_slide("We are on track for go-live on Jun 22")
    for i, (num, unit, label, col) in enumerate([
        ("24", "weeks", "total project plan", NAVY),
        ("14", "of 24", "weeks completed", TEAL),
        ("1", "week", "possible delay (electrical)", AMBER),
    ]):
        x = PI(0.7 + i * 4.1)
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, PI(2.0), PI(3.7), PI(3.0))
        card.fill.solid(); card.fill.fore_color.rgb = rgb(MIST); card.line.fill.background()
        card.adjustments[0] = 0.06
        text(s, x + PI(0.35), PI(2.3), PI(3.2), PI(1.3), num, 66, True, col, "Segoe UI Semibold")
        text(s, x + PI(0.35), PI(3.6), PI(3.2), PI(0.5), unit, 20, False, CHARCOAL)
        text(s, x + PI(0.35), PI(4.15), PI(3.2), PI(0.6), label, 18, False, CHARCOAL)

    # 3 dock 3
    s = content_slide("Dock 3 closes on Apr 27 – inbound moves to Dock 2")
    text(s, PI(0.7), PI(1.8), PI(11.5), PI(4), [
        "All inbound trucks use Gate 4 and Dock 2 from Apr 27.",
        "Carriers are informed; two still need confirmation (Megan).",
        "Extra Saturday shifts at Dock 2 for wave 1 (A items).",
        "Yellow forklift floor markings in Hall 2 ordered (Riley).",
    ], 22)

    # 4 safety
    s = content_slide("East corridor escape route is closed")
    text(s, PI(0.7), PI(1.8), PI(7), PI(4), [
        "Detour via mezzanine A and the north stairs.",
        "New assembly point: C.",
        "PPE required in all zones next to the construction site.",
        "Updated evacuation maps go up in every hall this week (Anne).",
    ], 22)
    warn = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, PI(8.4), PI(1.9), PI(4.2), PI(2.0))
    warn.fill.solid(); warn.fill.fore_color.rgb = rgb(AMBER); warn.line.fill.background()
    warn.text_frame.text = "Not all new maps are up yet"
    for p in warn.text_frame.paragraphs:
        for r in p.runs:
            r.font.size = PPt(22); r.font.bold = True; r.font.color.rgb = rgb(CHARCOAL); r.font.name = "Segoe UI Semibold"

    # 5 next steps
    s = content_slide("Next steps")
    rows = [("What", "Who", "By"),
            ("Confirm remaining carriers for Dock 2", "Megan", "Apr 20"),
            ("Approve sprinkler test", "Facility + insurer", "Apr 10"),
            ("Train shift leads on new safety rules", "Anne", "Apr 24"),
            ("Floor markings Hall 2", "Riley", "Apr 17")]
    tbl = s.shapes.add_table(len(rows), 3, PI(0.7), PI(1.8), PI(11.5), PI(3.2)).table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(NAVY if r == 0 else (MIST if r % 2 else (255, 255, 255)))
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    run.font.size = PPt(18); run.font.name = "Segoe UI"
                    run.font.bold = r == 0
                    run.font.color.rgb = rgb((255, 255, 255) if r == 0 else CHARCOAL)
    prs.save(OUT / "Contoso_Fargo_Update_Branded.pptx")

    # plain, off-brand deck to apply the style to
    prs2 = Presentation()
    prs2.slide_width, prs2.slide_height = PI(13.333), PI(7.5)
    t = prs2.slides.add_slide(prs2.slide_layouts[0])
    t.shapes.title.text = "Vendor access during construction"
    t.placeholders[1].text = "Draft – please make this look like Contoso"
    for title, bullets in [
        ("Hours", ["Vendors only between 6 am and 8 pm", "Saturday work needs an escort"]),
        ("Parking and check-in", ["Park in Lot C", "Check in at the north entrance", "Badge required at all times"]),
        ("Safety", ["Helmet, high-visibility vest and gloves", "Stay on marked walkways", "Assembly point C"]),
    ]:
        sl = prs2.slides.add_slide(prs2.slide_layouts[1])
        sl.shapes.title.text = title
        tf = sl.placeholders[1].text_frame
        tf.text = bullets[0]
        for b in bullets[1:]:
            tf.add_paragraph().text = b
    prs2.save(OUT / "Contoso_Vendor_Access_Draft.pptx")


def export_pdf():
    """Brand kits ingest a guidelines PDF; export the .docx with Word (Windows only)."""
    import subprocess
    src, dst = OUT / "Contoso_Brand_Guidelines.docx", OUT / "Contoso_Brand_Guidelines.pdf"
    ps = (
        "$w=New-Object -ComObject Word.Application;$w.Visible=$false;"
        f"try{{$d=$w.Documents.Open('{src}',$false,$true);$d.ExportAsFixedFormat('{dst}',17);$d.Close($false)}}"
        "finally{$w.Quit()}"
    )
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


if __name__ == "__main__":
    build_logos()
    build_guidelines()
    export_pdf()
    build_deck()
    print("built:", sorted(p.name for p in OUT.iterdir()))
