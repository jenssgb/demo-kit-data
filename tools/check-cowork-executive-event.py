"""Offline fixture checks, not a substitute for real tenant/Cowork rehearsals."""
import hashlib
import importlib.util
import json
import re
from datetime import datetime, timedelta
from email import policy
from email.parser import BytesParser
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
PACKAGE = ROOT / "cowork-executive-event"
spec = importlib.util.spec_from_file_location("event_builder", ROOT / "tools" / "build-cowork-executive-event.py")
BUILDER = importlib.util.module_from_spec(spec)
spec.loader.exec_module(BUILDER)


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def text(path):
    doc = Document(path)
    return "\n".join(p.text for p in doc.paragraphs)


expected = read(ROOT / "tools" / "expected" / "cowork-executive-event.json")
tenant = read(PACKAGE / "tenant.json")
config = tenant["executiveEvent"]
assert len(tenant["mails"]) == 6 and len(config["mailReplies"]) == 2
assert len(config["events"]) == 4 and len(config["threads"]) == 2
assert config["account"] == f"admin@{BUILDER.DOMAIN}"
assert config["people"] == {name: f"{alias}@{BUILDER.DOMAIN}" for name, alias in BUILDER.CAST.items()}
assert all(reply["parentSender"] in config["people"] for reply in config["mailReplies"])
assert {e["key"] for e in config["events"]} == {"forum", "prep", "prep-clash", "forum-clash"}
for event in config["events"]:
    assert event["timeZone"] == expected["facts"]["timeZone"]
    assert event["date"] in {"2026-10-20", "2026-10-22"}
    assert all(event["subject"][lang].startswith("[LEF26]") for lang in ("en", "de"))
    assert all(name in BUILDER.CAST for name in event["attendees"])
    assert "weekday" not in event

for first, second in [("prep", "prep-clash"), ("forum", "forum-clash")]:
    events = {e["key"]: e for e in config["events"]}
    a, b = events[first], events[second]
    start_a = datetime.fromisoformat(f"{a['date']}T{a['start']}")
    start_b = datetime.fromisoformat(f"{b['date']}T{b['start']}")
    assert max(start_a, start_b) < min(start_a + timedelta(minutes=a["minutes"]), start_b + timedelta(minutes=b["minutes"]))

deck_path = ROOT.parent / "demo-kit-new2026" / "decks" / "cowork.json"
deck = read(deck_path)
page = next(d for d in deck["demos"] if d["id"] == PACKAGE.name)
assert page == BUILDER.page(), "Generated deck and canonical builder differ"
sections = [s for s in page["steps"] if s["kind"] == "section"]
assert [s["exercise"]["code"] for s in sections if "exercise" in s] == expected["codes"]
assert len(sections) == expected["stageDemonstrations"] == 8
assert len([s for s in page["steps"] if s["kind"] == "prompt"]) == 8
assert all(s["kind"] in {"section", "prompt"} and "screen" not in s for s in page["steps"])
assert expected["codes"] == ["CO-26", "CO-27", "CO-28", "CO-29", "CO-23", "CO-24", "CO-25"]
assert len(set(expected["codes"])) == 7, "The existing controlled change stays in CO-24; never reuse or renumber codes"
for lang in ("en", "de"):
    prompt_steps = [s for s in page["steps"] if s["kind"] == "prompt"]
    briefing_prompt = prompt_steps[4]
    briefing_instructions = BUILDER.ACTS[0]["instructions"][lang]
    assert "[LEF26]" in briefing_instructions, "Calendar access must stay inside the synthetic scenario"
    assert any("[LEF26]" in item for item in briefing_prompt["preparation"][lang])
    assert {
        "en": "only entries whose subjects start with [LEF26]",
        "de": "ausschließlich Einträge, deren Betreff mit [LEF26] beginnt",
    }[lang] in briefing_instructions
    assert {"en": "not news reports", "de": "nicht aus Presseberichten"}[lang] in briefing_instructions
    assert {"en": "two readable pages", "de": "zwei gut lesbaren Seiten"}[lang] in briefing_instructions
    assert {"en": "not just source labels", "de": "nicht nur Quellennamen"}[lang] in briefing_instructions
    assert {"en": "rather than a narrow five-column table", "de": "statt in einer schmalen fünfspaltigen Tabelle"}[lang] in briefing_instructions
    for url in (
        "https://www.marriott.com/de/hotels/framd-le-meridien-frankfurt/events/",
        "https://hrewards.com/de/steigenberger-icon-frankfurter-hof/meeting-event",
    ):
        assert url in briefing_instructions, "Use the tested primary operator pages"
    assert {"en": "all file contents in English", "de": "sämtliche Datei-Inhalte auf Deutsch"}[lang] in briefing_instructions
    followup_instructions = BUILDER.ACTS[2]["instructions"][lang]
    assert {
        "en": "outputs subfolder next to the selected Action_Register.xlsx",
        "de": "Unterordner outputs neben der ausgewählten Action_Register.xlsx",
    }[lang] in followup_instructions, "Standalone follow-up must name its selected workbook's output folder"
    assert "Arial 11" in followup_instructions
    assert {
        "en": "Arial 11 applies to the worksheet cells",
        "de": "Arial 11 gilt für die Zellen der Aufgabenmappe",
    }[lang] in followup_instructions
    assert {
        "en": "not completed, unless the selected AFTER evidence explicitly confirms completion",
        "de": "noch nicht abgeschlossen, solange die ausgewählten AFTER-Belege den Abschluss nicht ausdrücklich bestätigen",
    }[lang] in followup_instructions
    assert all(value in followup_instructions for value in ["Actions!A2:F2", "2026-10-20", "General"])
    assert {"en": "Append A01-A05 starting at row 3", "de": "Hänge A01-A05 ab Zeile 3"}[lang] in followup_instructions
    assert {"en": "wrap long source references", "de": "brich lange Quellennachweise um"}[lang] in followup_instructions
    assert {"en": "A01-A05", "de": "A01 bis A05"}[lang] in followup_instructions
    for step in page["steps"]:
        if step["kind"] == "prompt":
            assert step["text"][lang] and step["preparation"][lang]
            assert all(f["kind"] and f["origin"] for f in step["files"][lang])
            assert "Act as" not in step["text"][lang]
            assert len(step["text"][lang].split()) <= 60, "Live requests must remain short; procedures belong in skills"
            assert "\n" not in step["text"][lang] and "https://" not in step["text"][lang]
    change_prompt = prompt_steps[6]
    change_instructions = BUILDER.ACTS[1]["changeInstructions"][lang]
    assert all(name in " ".join(change_prompt["preparation"][lang]) for name in ["Leadership_Forum.pptx", "Agenda.docx", "outputs"])
    assert {"en": "default Cowork working copies", "de": "Cowork-Arbeitskopien"}[lang] in " ".join(change_prompt["preparation"][lang])
    assert {"en": "existing agenda-decision draft", "de": "vorhandenen Entwurf zur Agendaentscheidung"}[lang] in change_instructions
    assert {"en": "whole forum", "de": "ganze Forum"}[lang] in change_instructions
    assert {"en": "replace the full timeline labels", "de": "Beschriftungen der Zeitachse vollständig"}[lang] in change_instructions
    pack_instructions = BUILDER.ACTS[1]["instructions"][lang]
    assert {"en": "one paragraph without manual line breaks", "de": "einem Absatz ohne händische Zeilenumbrüche"}[lang] in pack_instructions
    conversation = next(a for a in BUILDER.STARTERS if a["code"] == "CO-28")
    assert "LinkedIn" in conversation["instructions"][lang] and "Conversation_Brief.docx" in conversation["instructions"][lang]
    travel_instructions = next(a for a in BUILDER.STARTERS if a["code"] == "CO-29")["instructions"][lang]
    assert {"en": "not a dated fallback connection", "de": "nicht als datierte Rückfallverbindung"}[lang] in travel_instructions
    travel = (PACKAGE / "inputs" / lang / "before" / "Travel_Request.md").read_text(encoding="utf-8")
    assert all(value in travel for value in ["BEFORE", "Billie Vester", "Berlin Hbf", "18:00", "09:15", "15:00", "650"])
    assert "A01:" not in travel and "D01:" not in travel and "11:30-12:15" not in travel
    assert expected["facts"]["travel"]["insideEventBudget"] is True
    brief = text(PACKAGE / "inputs" / lang / "before" / "Event_Brief.docx")
    assert "42" in brief and ("9,000" in brief or "9.000" in brief)
    assert "A01:" not in brief and "D01:" not in brief and "11:30-12:15" not in brief
    after = text(PACKAGE / "inputs" / lang / "after" / "Post_Event_Notes.docx")
    assert all(code in after for code in ["D01", "D02", "A01", "A02", "A03", "A04", "A05"])
    assert "2026" in after
    for relative in tenant["mails"]:
        path = PACKAGE / relative.replace("{lang}", lang)
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        sender = message["From"].addresses[0]
        assert sender.addr_spec == config["people"][sender.display_name]
        assert message["To"].addresses[0].addr_spec.lower() == config["account"].lower()
        assert str(message["Subject"]).startswith("[LEF26-BEFORE-")
        assert any(str(message["Subject"]) in f["name"] for s in page["steps"] if s["kind"] == "prompt"
                   for f in s["files"][lang] if f["kind"] == "email")
    for reply in config["mailReplies"]:
        assert any(reply["parentSubject"][lang] in f["name"] for s in page["steps"] if s["kind"] == "prompt"
                   for f in s["files"][lang] if f["kind"] == "email")
    for path in [PACKAGE / "baseline" / lang / "Action_Register.xlsx",
                 *(PACKAGE / "runs" / lang / f"{i:02}" / "Action_Register.xlsx" for i in range(1, 5))]:
        workbook = load_workbook(path)
        try:
            assert workbook.sheetnames == ["Actions", "Readme"]
            assert workbook["Actions"].max_row == 2
            assert workbook["Actions"]["A2"].value == "P01"
            assert workbook["Actions"]["D2"].value == "2026-10-20"
            assert workbook["Actions"]["D2"].data_type == "s"
            assert workbook["Actions"]["D2"].number_format == "General"
            assert not workbook._external_links
            assert "EventActions" in workbook["Actions"].tables
            for sheet in workbook:
                for row in sheet:
                    for cell in row:
                        assert cell.data_type not in {"e", "f"}
                        if cell.value is not None:
                            assert cell.font.name == "Arial"
        finally:
            workbook.close()
    runbook = ROOT / "tools" / "expected" / f"cowork-executive-event-runbook-{lang}.md"
    runbook_text = runbook.read_text(encoding="utf-8")
    assert runbook_text.rstrip() == BUILDER.runbook(lang).rstrip(), "Runbook and canonical generator differ"
    assert all(code in runbook_text for code in expected["codes"])
    assert runbook_text.count(" | OPEN |" if lang == "en" else " | OFFEN |") == 8
    assert {"en": "The presenter owns the live rehearsal", "de": "Die eigene Generalprobe"}[lang] in runbook_text
    assert "Require two consecutive successful" not in runbook_text
    assert "Zwei erfolgreiche Proben mit frischen Läufen hintereinander verlangen" not in runbook_text
    assert runbook_text.count("### Prompt") == 8, "The controlled change also needs a copy-ready prompt"

assert "Preparation package, not live acceptance" in expected["readiness"]

assert len(BUILDER.SKILLS) == 3
assert set(expected["skills"]) == set(BUILDER.SKILLS)
for key, skill in BUILDER.SKILLS.items():
    assert re.fullmatch(r"[A-Za-z0-9._-]{1,64}", skill["name"]), "Enforce the actual tenant's supported SKILL.md name syntax"
    source = PACKAGE / "skills" / key / "SKILL.md"
    content = source.read_text(encoding="utf-8")
    assert content == BUILDER.skill_text(key)
    assert content.startswith("---\nname: ") and "\ndescription: " in content
    assert "derive the actual run name" in content.lower()
    assert "no external contact" in content.lower()
    if key == "contoso-event-pack":
        assert "one readable A4 page" in content and "18 pt" in content
        assert "gut lesbaren A4-Seite" in content and "18 Punkt" in content
    for act in BUILDER.WORKFLOWS:
        if act["code"] in skill["codes"]:
            assert act["skill"] == key
            for lang in ("en", "de"):
                assert act["instructions"][lang] in content
                assert skill["name"] in " ".join(act["prep"][lang])
                if "changeInstructions" in act:
                    assert act["changeInstructions"][lang] in content
    with ZipFile(PACKAGE / "skills" / f"{key}.zip") as archive:
        assert archive.namelist() == ["SKILL.md"], "Only the root skill is needed; no stale companion-file snapshots"
        assert archive.read("SKILL.md") == source.read_bytes()

manifest = read(PACKAGE / "manifest.json")
hashes = read(PACKAGE / "baseline" / "hashes.json")
assert set(hashes) == {"en", "de"}
for lang, sha in hashes.items():
    assert hashlib.sha256((PACKAGE / "baseline" / lang / "Action_Register.xlsx").read_bytes()).hexdigest() == sha
paths = [f["path"] for f in manifest["files"]]
assert len(paths) == len(set(paths))
assert all("expected" not in p and "runbook" not in p and p != "tenant.json" and "settings" not in p for p in paths)
assert not any(p.endswith((".pptx", ".mp4", ".png")) for p in paths), "Do not fabricate result backups"
for entry in manifest["files"]:
    file = PACKAGE / entry["path"]
    assert file.is_file() and hashlib.sha256(file.read_bytes()).hexdigest() == entry["sha256"], entry["path"]
    if file.suffix in {".docx", ".xlsx"}:
        with ZipFile(file) as archive:
            assert archive.testzip() is None
for name in ["ExecutiveEvent.ps1", "Reset-CoworkExecutiveEvent.ps1"]:
    assert (PACKAGE / name).read_bytes() == (ROOT / "tools" / name).read_bytes()
owners = [app["id"] for app in read(ROOT / "catalog.json")["apps"] if PACKAGE.name in app["demos"]]
assert owners == ["cowork"]
assert set(BUILDER.CAST) <= set(read(ROOT / "people.json"))
print("LEF26 offline checks passed: eight demonstrations, seven stable exercise codes, bilingual prompts/travel inputs, explicit original-file change preparation, phase isolation, calendar clashes, cast, 10 pristine workbooks and manifest/catalog/runtime parity. Real Cowork execution is NOT verified.")
