"""Builds fictional Contoso Fargo Agent Builder demo files.

Run: python tools/build-agent-builder-copilot.py
Needs: python-docx, openpyxl
All content is fictional Contoso demo data.
"""
from datetime import date
from pathlib import Path
import csv
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "agent-builder-copilot"
OUT.mkdir(exist_ok=True)
YEAR = 2027
NAVY = RGBColor(12, 53, 106)
BLUE = "DDEBFF"
GREEN = "E3FCEF"
AMBER = "FFF4CE"
RED = "FDE7E9"

RULES = [
    ("Emergency", "Sprinkler alarm in new wing", "Stop all forklift movement. Shift lead calls 555-0100 and uses radio channel 4. Evacuate through East Corridor B unless blocked. Assemble at Muster Point C. Do not reset the panel; wait for Anne Weiler or Facility."),
    ("Emergency", "Fire or smoke", "Pull nearest manual alarm, call 555-0100, evacuate by the posted route, close dock doors if safe, and account for visitors and carriers."),
    ("Dock", "Dock 3 closure", "Dock 3 is closed from Apr 27. Inbound freight moves to Dock 2 / Gate 4. Dock 2 priority window is 7:00-9:00 for the transition week."),
    ("Dock", "Vendor hours", "Standard carrier window is 6:00-20:00. Northwind may use a temporary 5:30 arrival exception on Apr 29 only if Safety signs off before Apr 26."),
    ("PPE", "New wing PPE", "High-visibility vest, safety shoes, and hard hat are mandatory beyond the blue line until Construction clears the wing."),
    ("Forklift", "Forklift routes", "Use the one-way loop from Dock 2 to Aisle F. No forklift traffic through East Corridor B during evacuation drills, sprinkler tests, or visitor tours."),
    ("Training", "Shift lead briefing", "Every shift lead must complete the 90-minute new-wing safety briefing before Jun 12 and sign the roster."),
]

SCHEDULE = [
    ["2027-04-27", "06:30", "08:30", "Dock 2", "Gate 4", "Fabrikam Logistics", "Inbound", "Pallet racks", "Transition slot while Dock 3 closes", "Megan Bowen", "Confirmed"],
    ["2027-04-28", "07:00", "09:00", "Dock 2", "Gate 4", "Fabrikam Logistics", "Inbound", "Packaging supplies", "Use Gate 4 detour signage", "Megan Bowen", "Confirmed"],
    ["2027-04-29", "05:30", "06:15", "Dock 2", "Gate 4", "Northwind", "Inbound", "Line-side totes", "Safety exception requested; early arrival before vendor window", "Anne Weiler", "Exception pending"],
    ["2027-04-29", "07:15", "08:30", "Dock 2", "Gate 4", "Fabrikam Logistics", "Inbound", "Conveyor spares", "Northwind must clear Gate 4 first", "Megan Bowen", "Confirmed"],
    ["2027-04-30", "08:00", "10:00", "Dock 2", "Gate 4", "Fabrikam Logistics", "Inbound", "A-item cartons", "Gate 4 standard route starts", "Megan Bowen", "Confirmed"],
    ["2027-05-02", "09:00", "12:00", "Dock 2", "Gate 4", "Northwind", "Outbound", "Tailspin Toys partial shipment", "Hold until sprinkler sign-off is recorded", "Megan Bowen", "Tentative"],
    ["2027-05-09", "07:00", "13:00", "Dock 2", "Gate 4", "Internal move", "Inventory", "Wave 1 A-items", "Needs four more forklift drivers", "Megan Bowen", "At risk"],
    ["2027-06-15", "10:00", "11:30", "New wing", "East Corridor B", "Contoso", "Safety", "Evacuation drill", "Must complete before go/no-go", "Anne Weiler", "Planned"],
    ["2027-06-22", "06:00", "14:00", "Dock 2", "Gate 4", "All carriers", "Go-live", "New wing cutover", "Use new safety routes", "Ops lead", "Planned"],
]

INSTRUCTIONS = """Fargo Site Assistant – instructions for Agent Builder

Name: Fargo Site Assistant
Short description: Answers shift lead questions for the Fargo new-wing go-live.

Use a calm, practical, human style. Answer like an experienced Contoso operations lead who is helping shift leads during a busy site transition.

Stay grounded in the Fargo safety handbook, the readiness tracker, and the Fargo dock schedule. If the answer is not in those sources, say that you cannot find it in the Fargo material and suggest who to ask: Megan Bowen for logistics and dock slots, Anne Weiler for safety, Riley Johnson for construction, Alex Wilber for IT.

When a question is about a carrier, dock, gate, alarm, evacuation route, training, or go/no-go readiness, give the direct answer first, then the reason and the source you used.

For safety incidents, be brief and decisive. Include the first action, who to contact, and what not to do. Never invent emergency procedures.

If dates are mentioned without a year, assume the Fargo go-live year in the demo data. Keep April 27 Dock 3 closure, June 15 go/no-go, and June 22 go-live in mind.

Use English by default. If the user asks in German, answer in German and keep UI or source names unchanged.

Starter prompts:
- Which gate does Northwind use on April 29?
- What do I do if the sprinkler alarm goes off in the new wing?
- Are we ready for the June 15 go/no-go?
"""


def patch_docx_for_validation(path):
    with zipfile.ZipFile(path, "r") as zin:
        entries = {name: zin.read(name) for name in zin.namelist()}
    settings = entries.get("word/settings.xml", b"").decode("utf-8")
    settings = settings.replace('<w:zoom w:val="bestFit"/>', '<w:zoom w:val="bestFit" w:percent="100"/>')
    entries["word/settings.xml"] = settings.encode("utf-8")
    entries["word/fontTable.xml"] = b'<?xml version="1.0" encoding="utf-8"?>\n<w:fonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">\n  <w:font w:name="Aptos">\n    <w:charset w:val="00"/>\n    <w:family w:val="swiss"/>\n    <w:pitch w:val="variable"/>\n  </w:font>\n  <w:font w:name="Aptos Display">\n    <w:charset w:val="00"/>\n    <w:family w:val="swiss"/>\n    <w:pitch w:val="variable"/>\n  </w:font>\n</w:fonts>\n'
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for name, data in entries.items():
            zout.writestr(name, data)


def set_styles(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    for name in ["Heading 1", "Heading 2"]:
        styles[name].font.name = "Aptos Display"
        styles[name].font.color.rgb = NAVY


def build_docx():
    doc = Document()
    set_styles(doc)
    sec = doc.sections[0]
    sec.top_margin = Pt(54)
    sec.bottom_margin = Pt(54)
    sec.left_margin = Pt(54)
    sec.right_margin = Pt(54)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Contoso Fargo Safety Handbook")
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = NAVY
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("New warehouse wing | Shift lead quick guide | Demo data").italic = True

    doc.add_heading("1. What changes for the new wing", level=1)
    doc.add_paragraph("Dock 3 closes on April 27. During the transition, inbound freight moves to Dock 2 / Gate 4. The go/no-go decision is June 15 and the new wing go-live is June 22.")
    doc.add_paragraph("Shift leads own the floor conversation: keep carriers moving, keep forklifts out of blocked routes, and escalate safety questions early.")

    doc.add_heading("2. Emergency rules", level=1)
    table = doc.add_table(rows=1, cols=3)
    table.style = "Light Shading Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(["Situation", "First action", "Do not"]):
        hdr[i].text = h
    rows = [
        ("Sprinkler alarm in new wing", "Stop forklift movement. Call 555-0100 and use radio channel 4. Evacuate through East Corridor B unless blocked. Muster Point C.", "Do not reset the panel. Wait for Anne Weiler or Facility."),
        ("Smoke or fire", "Pull the nearest manual alarm, call 555-0100, close dock doors if safe, and account for visitors.", "Do not re-enter until Facility clears the area."),
        ("Injury", "Call the onsite first responder, secure the area, and keep one person with the injured colleague.", "Do not move the person unless the area is unsafe."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, v in enumerate(row):
            cells[i].text = v

    doc.add_heading("3. Dock and carrier rules", level=1)
    for topic, title, text in RULES[2:4]:
        doc.add_paragraph(f"{title}: {text}", style=None)
    doc.add_paragraph("Northwind's April 29 slot is the only early-arrival exception in the demo schedule. If Safety does not approve it, ask Megan Bowen to move the arrival to the regular carrier window.")

    doc.add_heading("4. PPE, forklift routes, and training", level=1)
    for topic, title, text in RULES[4:]:
        doc.add_paragraph(f"{title}: {text}")
    doc.add_paragraph("All 12 shift leads must complete training before June 12. Anne Weiler owns safety content; Megan Bowen owns carrier and dock slot questions.")

    doc.add_page_break()
    doc.add_heading("Deutsche Kurzfassung", level=1)
    doc.add_paragraph("Diese Kurzfassung ist für deutschsprachige Multiplikatoren gedacht. Die Arbeitsdaten bleiben absichtlich auf Englisch, damit die Microsoft 365 Demo-Umgebung konsistent bleibt.")
    doc.add_heading("Notfallregeln", level=2)
    doc.add_paragraph("Sprinkleralarm im neuen Flügel: Staplerverkehr sofort stoppen, 555-0100 anrufen, Funkkanal 4 nutzen, über East Corridor B evakuieren, sofern der Weg frei ist, und am Sammelpunkt C zählen. Das Panel nicht zurücksetzen.")
    doc.add_paragraph("Feuer oder Rauch: nächsten Handmelder auslösen, 555-0100 anrufen, Docktore schließen, wenn es sicher ist, und Besucher sowie Speditionen mitzählen.")
    doc.add_heading("Dock- und Gate-Regeln", level=2)
    doc.add_paragraph("Dock 3 ist ab dem 27. April geschlossen. Eingehende Lieferungen laufen über Dock 2 / Gate 4. Das Standardfenster für Speditionen ist 6:00 bis 20:00 Uhr.")
    doc.add_paragraph("Northwind darf am 29. April nur mit Safety-Ausnahme um 5:30 Uhr kommen. Ohne Ausnahme muss Megan Bowen den Slot verlegen.")
    doc.add_heading("Kontakte", level=2)
    doc.add_paragraph("Megan Bowen: Logistik, Dock-Slots und Speditionen. Anne Weiler: Sicherheit, Training, Evakuierung. Riley Johnson: Bau und Sperrungen. Alex Wilber: IT, Scanner und WLAN.")

    path = OUT / "Contoso_Fargo_Safety_Handbook.docx"
    doc.save(path)
    patch_docx_for_validation(path)


def build_schedule():
    headers = ["Date", "Start", "End", "Dock", "Gate", "Carrier", "Type", "Load", "Notes", "Owner", "Status"]
    csv_path = OUT / "Fargo_Dock_Schedule.csv"
    with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(headers)
        w.writerows(SCHEDULE)

    wb = Workbook()
    ws = wb.active
    ws.title = "Fargo dock schedule"
    ws.append(headers)
    for row in SCHEDULE:
        ws.append(row)
    tab = Table(displayName="FargoDockSchedule", ref=f"A1:K{len(SCHEDULE)+1}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    widths = [13, 10, 10, 13, 10, 20, 13, 24, 44, 16, 18]
    for i, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E79")
        cell.alignment = Alignment(horizontal="center")
    fill_by_status = {"Exception pending": RED, "At risk": AMBER, "Tentative": AMBER, "Planned": GREEN, "Confirmed": GREEN}
    for row in range(2, len(SCHEDULE)+2):
        ws[f"K{row}"].fill = PatternFill("solid", fgColor=fill_by_status.get(ws[f"K{row}"].value, "FFFFFF"))
        for col in range(1, 12):
            ws.cell(row, col).alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    wb.save(OUT / "Fargo_Dock_Schedule.xlsx")


def build_instructions():
    (OUT / "Fargo_Site_Assistant_Instructions.txt").write_text(INSTRUCTIONS, encoding="utf-8")


def build_manifest_stub():
    manifest = '''{
  "title": {
    "en": "Agent Builder – Contoso Fargo Site Assistant",
    "de": "Agent Builder – Contoso Fargo Site Assistant"
  },
  "next": {
    "en": [
      "Upload Contoso_Fargo_Safety_Handbook.docx and Contoso_Fargo_GoLive_Readiness_Tracker.xlsx to the same SharePoint site or OneDrive folder.",
      "Import Fargo_Dock_Schedule.csv into SharePoint as a list named Fargo dock schedule (or keep Fargo_Dock_Schedule.xlsx as the fallback knowledge file).",
      "Copy the text from Fargo_Site_Assistant_Instructions.txt into Agent Builder instructions.",
      "Continue with the preparation steps in the Demo Kit."
    ],
    "de": [
      "Contoso_Fargo_Safety_Handbook.docx und Contoso_Fargo_GoLive_Readiness_Tracker.xlsx in dieselbe SharePoint-Site oder denselben OneDrive-Ordner hochladen.",
      "Fargo_Dock_Schedule.csv als SharePoint-Liste Fargo dock schedule importieren (oder Fargo_Dock_Schedule.xlsx als Fallback-Wissensdatei behalten).",
      "Den Text aus Fargo_Site_Assistant_Instructions.txt in die Agent-Builder-Anweisungen kopieren.",
      "Mit den Vorbereitungsschritten im Demo Kit weitermachen."
    ]
  }
}
'''
    (OUT / "manifest.json").write_text(manifest, encoding="utf-8")


if __name__ == "__main__":
    build_docx()
    build_schedule()
    build_instructions()
    build_manifest_stub()
    print(f"Built {OUT}")

