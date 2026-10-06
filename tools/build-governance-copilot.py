
"""Builds fictional Contoso governance demo files for Microsoft Copilot trust demos.

Run: python tools/build-governance-copilot.py
Needs: python-docx openpyxl
All content is fictional (Contoso). Scenario: Contoso Fargo distribution center expansion.
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter

OUT = Path(__file__).resolve().parent.parent / "governance-copilot"
OUT.mkdir(parents=True, exist_ok=True)
NAVY = "0B3D91"
BLUE = "EAF2FF"
AMBER = "FFF4CE"
RED = "FDE7E9"
GREEN = "DFF6DD"


def style_doc(doc):
    styles = doc.styles
    styles["Normal"].font.name = "Segoe UI"
    styles["Normal"].font.size = Pt(10.5)
    for name in ("Heading 1", "Heading 2"):
        styles[name].font.name = "Segoe UI"
        styles[name].font.color.rgb = RGBColor.from_string(NAVY)


def build_public_fact_sheet():
    doc = Document()
    style_doc(doc)
    doc.add_heading("Contoso Fargo Distribution Center – public site fact sheet", 0)
    doc.add_paragraph("Classification to apply in the demo tenant: Public")
    doc.add_heading("What is public", level=1)
    for text in [
        "The Fargo distribution center serves customers across North Dakota, Minnesota, and South Dakota.",
        "A new warehouse wing is scheduled to go live on June 22.",
        "Inbound traffic temporarily moves from Dock 3 to Dock 2 / Gate 4 starting April 27.",
        "The customer-facing promise is unchanged: standard orders ship within the existing SLA.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Public talking points", level=1)
    doc.add_paragraph(
        "The site expansion increases resilience and separates construction traffic from everyday inbound operations. "
        "Customers such as Tailspin Toys should see no service interruption."
    )
    doc.save(OUT / "Contoso_Fargo_Public_Site_Fact_Sheet.docx")


def build_ma_note():
    doc = Document()
    style_doc(doc)
    doc.add_heading("Highly Confidential – Fargo M&A note", 0)
    doc.add_paragraph("Classification to apply in the demo tenant: Highly Confidential")
    p = doc.add_paragraph()
    run = p.add_run("Demo-only fictional content. Do not use with real company data.")
    run.bold = True
    doc.add_heading("Board-prep note", level=1)
    for text in [
        "Contoso is evaluating a fictional warehouse services acquisition in the Upper Midwest.",
        "Early diligence estimate: integration costs could reach USD 2.8M over 18 months.",
        "No employee communication before the board decides whether to proceed.",
        "This note must not be summarized by Copilot in the live demo once the DLP policy is enabled.",
    ]:
        doc.add_paragraph(text, style="List Bullet")
    doc.add_heading("Why it is in the demo", level=1)
    doc.add_paragraph(
        "The presenter references this file in Copilot after labeling it Highly Confidential. "
        "The expected result is that Copilot excludes or refuses to process it because of the Purview DLP policy."
    )
    doc.save(OUT / "Contoso_Fargo_Highly_Confidential_MA_Note.docx")


def build_labor_costs():
    wb = Workbook()
    ws = wb.active
    ws.title = "Labor costs"
    ws["A1"] = "Contoso Fargo labor cost forecast – confidential"
    ws["A1"].font = Font(bold=True, size=14, color=NAVY)
    ws["A2"] = "Classification to apply in the demo tenant: Confidential"
    ws["A2"].fill = PatternFill("solid", fgColor=AMBER)
    rows = [
        ["Workstream", "Owner", "Apr", "May", "Jun", "Risk", "Note"],
        ["Dock 2 weekend receiving", "Teresa Sac", 18400, 27500, 12600, "Medium", "Temporary inbound move while Dock 3 closes"],
        ["Forklift driver backfill", "Teresa Sac", 7200, 19800, 6400, "High", "Four drivers missing for May 9 wave 1"],
        ["Safety shift-lead training", "Sonia Rees", 0, 4800, 9600, "Medium", "12 shift leads before June 15 drill"],
        ["Construction escort overtime", "Vance DeLeon", 11200, 15400, 6400, "Medium", "Electrical slips to May 22"],
        ["IT floor support", "Billie Vester", 0, 6200, 13400, "Low", "Scanners and Wi-Fi commissioning"],
    ]
    start = 4
    for r_idx, row in enumerate(rows, start):
        for c_idx, val in enumerate(row, 1):
            cell = ws.cell(r_idx, c_idx, val)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if r_idx == start:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=NAVY)
            elif c_idx == 6:
                fill = {"High": RED, "Medium": AMBER, "Low": GREEN}[val]
                cell.fill = PatternFill("solid", fgColor=fill)
    end = start + len(rows) - 1
    tab = Table(displayName="LaborCosts", ref=f"A{start}:G{end}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tab)
    for col, width in zip(range(1, 8), [28, 18, 12, 12, 12, 12, 48]):
        ws.column_dimensions[get_column_letter(col)].width = width
    for row in ws.iter_rows(min_row=start+1, min_col=3, max_col=5, max_row=end):
        for cell in row:
            cell.number_format = '$#,##0'
    ws["A12"] = "Demo prompt idea"
    ws["A12"].font = Font(bold=True, color=NAVY)
    ws["A13"] = "Summarize the top cost drivers for the Fargo go-live call. Copilot-generated output should inherit the Confidential label."
    ws["A13"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A13:G13")
    wb.save(OUT / "Fargo_Labor_Costs_Confidential.xlsx")


def build_readme():
    text = """# Governance Copilot demo setup / Demo-Setup Governance Copilot

All files are fictional Contoso data for the Fargo distribution center expansion. Do not replace them with real customer data.

## English setup

1. Install the demo data with the Demo Kit one-liner. The folder should appear as **OneDrive > Demo-governance-copilot**.
2. In Excel, open **Fargo_Labor_Costs_Confidential.xlsx**. On **Home > Sensitivity**, apply your tenant's **Confidential** label. Save the file to OneDrive or SharePoint.
3. In Word, open **Contoso_Fargo_Public_Site_Fact_Sheet.docx**. On **Home > Sensitivity**, apply **Public**. Save it.
4. In Word, open **Contoso_Fargo_Highly_Confidential_MA_Note.docx**. On **Home > Sensitivity**, apply **Highly Confidential**. Save it.
5. Ask your Purview admin to create or confirm a DLP policy for the **Microsoft 365 Copilot and Copilot Chat** location that prevents Copilot from processing content with the **Highly Confidential** label.
6. For the admin-view part, prepare accounts/permissions for Microsoft 365 admin center, SharePoint admin center, Microsoft Purview portal, and Viva Insights/Copilot Dashboard as needed.

## Deutsches Setup

1. Demodaten mit dem One-Liner aus dem Demo Kit installieren. Der Ordner sollte als **OneDrive > Demo-governance-copilot** erscheinen.
2. In Excel **Fargo_Labor_Costs_Confidential.xlsx** öffnen. Unter **Start > Vertraulichkeit** (oder **Sensitivity**, je nach UI-Sprache) das Tenant-Label **Confidential/Vertraulich** anwenden. In OneDrive oder SharePoint speichern.
3. In Word **Contoso_Fargo_Public_Site_Fact_Sheet.docx** öffnen. Unter **Start > Vertraulichkeit/Sensitivity** das Label **Public/Öffentlich** anwenden. Speichern.
4. In Word **Contoso_Fargo_Highly_Confidential_MA_Note.docx** öffnen. Unter **Start > Vertraulichkeit/Sensitivity** das Label **Highly Confidential/Hoch vertraulich** anwenden. Speichern.
5. Die Purview-Administration soll eine DLP-Richtlinie für den Ort **Microsoft 365 Copilot and Copilot Chat** erstellen oder bestätigen, die Copilot das Verarbeiten von Inhalten mit dem Label **Highly Confidential/Hoch vertraulich** verbietet.
6. Für den Admin-Teil Konten/Berechtigungen für Microsoft 365 Admin Center, SharePoint Admin Center, Microsoft Purview Portal und Viva Insights/Copilot Dashboard vorbereiten.

## Files / Dateien

- **Fargo_Labor_Costs_Confidential.xlsx** — cost forecast, apply Confidential.
- **Contoso_Fargo_Public_Site_Fact_Sheet.docx** — harmless public grounding file, apply Public.
- **Contoso_Fargo_Highly_Confidential_MA_Note.docx** — blocked-file demo, apply Highly Confidential.
"""
    (OUT / "README_setup.md").write_text(text, encoding="utf-8")


def build_manifest_seed():
    manifest = {
        "title": {
            "en": "Governance & trust for Microsoft Copilot",
            "de": "Governance & Vertrauen für Microsoft Copilot",
        },
        "next": {
            "en": [
                "Apply the labels named in README_setup.md before the session.",
                "Confirm the Purview DLP policy blocks Highly Confidential content for Microsoft 365 Copilot and Copilot Chat.",
                "Use the demo page prompts to show label inheritance first, then the DLP block, then the admin dashboards.",
            ],
            "de": [
                "Die Labels aus README_setup.md vor der Session anwenden.",
                "Prüfen, dass die Purview-DLP-Richtlinie Highly-Confidential-Inhalte für Microsoft 365 Copilot and Copilot Chat blockiert.",
                "Mit den Prompts der Demoseite zuerst Label-Vererbung, dann den DLP-Block und danach die Admin-Dashboards zeigen.",
            ],
        },
    }
    (OUT / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build_public_fact_sheet()
    build_ma_note()
    build_labor_costs()
    build_readme()
    build_manifest_seed()
    print("built", sorted(str(p.relative_to(OUT)) for p in OUT.rglob("*") if p.is_file()))
