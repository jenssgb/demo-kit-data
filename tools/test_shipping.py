"""Deterministic corpus checks; no network, sign-in or tenant changes."""
import csv
import json
import unittest
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree

from docx import Document
from openpyxl import load_workbook
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / "word-shipping-brief"


def rows(name):
    with (BASE / name).open(encoding="utf-8", newline="") as stream:
        return list(csv.DictReader(stream))


class ShippingCorpus(unittest.TestCase):
    def test_volume_table(self):
        wb = load_workbook(BASE / "Volumes_2026.xlsx")
        ws = wb["Volumes"]
        self.assertEqual(ws.max_row - 1, 216)
        self.assertEqual(len(ws.tables), 1)
        self.assertEqual(len({row[0] for row in ws.iter_rows(min_row=2, values_only=True)}), 12)
        self.assertEqual(wb.calculation.calcMode, "auto")
        self.assertEqual(sum(row[3] for row in ws.iter_rows(min_row=2, values_only=True)), 30780)
        self.assertEqual(sum(row[4] for row in ws.iter_rows(min_row=2, values_only=True)), 34560)

    def test_exact_variances_and_rules(self):
        wb = load_workbook(BASE / "DnD_reconciliation.xlsx")
        charges = list(wb["Charges"].iter_rows(min_row=2, values_only=True))
        self.assertEqual(charges, [
            ("Fabrikam", 12400, 13950), ("Northwind", 8200, 8350),
            ("Tailspin", 15600, 17100), ("Wide World", 6900, 6950),
            ("Adatum", 10300, 10920), ("Contoso Retail", 9800, 9900),
        ])
        self.assertEqual([c for c, a, b in charges if (b-a)/a > .05], ["Fabrikam", "Tailspin", "Adatum"])
        self.assertEqual([c for c, a, b in charges if (b-a)/a > .10], ["Fabrikam"])
        self.assertEqual(wb[".Rules"].sheet_state, "visible")
        self.assertEqual(wb[".Rules"].max_column, 1)
        self.assertTrue(all(row[0].value for row in wb[".Rules"]))
        self.assertEqual(len(rows("DnD_expected.csv")), 6)

    def test_five_reconciliation_exceptions(self):
        wb = load_workbook(BASE / "Carrier_invoices_Q3.xlsx")
        sources = [{str(r[0]): r[2] for r in wb[name].iter_rows(min_row=2, values_only=True)}
                   for name in ["Invoices", "Accruals"]]
        invoices, accruals = sources
        self.assertEqual([len(x) for x in sources], [60, 60])
        exceptions = {k for k in invoices.keys() | accruals.keys() if invoices.get(k) != accruals.get(k)}
        self.assertEqual(exceptions, {"5008", "5019", "5040", "5055", "5999"})
        self.assertEqual(exceptions, {r["Booking"] for r in rows("Carrier_reconciliation_expected.csv")})
        self.assertEqual(len(invoices.keys() & accruals.keys()), 59)

    def test_analyst_join_and_ranking(self):
        invoices, disputes = rows("invoices_Q3.csv"), rows("invoice_disputes_Q3.csv")
        invoice_ids = {r["invoice_id"] for r in invoices}
        self.assertEqual(len(invoices), 300)
        self.assertEqual(len(disputes), 300)
        self.assertEqual(len(invoice_ids), 300)
        self.assertEqual(invoice_ids, {r["invoice_id"] for r in disputes})
        totals = {}
        for row in disputes:
            reason = row["dispute_reason"]
            totals[reason] = totals.get(reason, 0) + int(row["corrected_amount_eur"])
        self.assertEqual(sorted(totals.items(), key=lambda r: -r[1])[:3], [
            ("Duplicate charge", 12000), ("Wrong rate", 7200), ("Free-time dispute", 4800)])
        self.assertEqual(totals, {r["Reason"]: int(r["Corrections EUR"]) for r in rows("Analyst_expected.csv")})

    def test_office_structure_and_document_content(self):
        for pattern in ["*.docx", "*.xlsx"]:
            for path in BASE.glob(pattern):
                with self.subTest(path=path.name), ZipFile(path) as archive:
                    self.assertIsNone(archive.testzip())
                    for name in archive.namelist():
                        if name.endswith(".xml"):
                            ElementTree.fromstring(archive.read(name))
        draft = Document(BASE / "Customer_brief_draft.docx")
        self.assertEqual(sum(p.style.name == "Heading 1" for p in draft.paragraphs), 4)
        xml = draft.element.xml
        self.assertEqual(xml.count('w:type="page"'), 3)
        contract = "\n".join(p.text for p in Document(BASE / "Supplier_Agreement_Draft.docx").paragraphs)
        playbook = "\n".join(p.text for p in Document(BASE / "Procurement_Playbook.docx").paragraphs)
        for value in ["7.1", "30 days", "9.2", "unlimited", "12.4", "14 days"]:
            self.assertIn(value, contract)
        for value in ["4.1", "60 days", "6.3", "12 months", "90 days"]:
            self.assertIn(value, playbook)
        thread = "\n".join(p.text for p in Document(BASE / "Channel_thread_seed.docx").paragraphs)
        self.assertGreater(len(thread), 1000)
        self.assertEqual(thread.count("Reply "), 12)

    def test_all_sources_resolve_and_mail_options_are_unique(self):
        from importlib.util import module_from_spec, spec_from_file_location
        spec = spec_from_file_location("builder", Path(__file__).with_name("build-shipping-playbook.py"))
        module = module_from_spec(spec)
        spec.loader.exec_module(module)
        for folder in [module.BASE, *module.FILES]:
            manifest = json.loads((ROOT / folder / "manifest.json").read_text(encoding="utf-8-sig"))
            for entry in [*manifest["files"], *manifest.get("shared", [])]:
                source = ROOT / entry.get("from", folder) / entry["path"]
                self.assertTrue(source.is_file(), source)
            tenant = json.loads((ROOT / folder / "tenant.json").read_text(encoding="utf-8"))
            subjects = [item["subject"] for item in tenant.get("mailOptions", [])]
            self.assertEqual(len(set(subjects)), len(subjects))
            for file in tenant.get("mails", []):
                self.assertTrue((ROOT / folder / file).is_file())
        with ZipFile(BASE / "dg-pre-check.zip") as archive:
            self.assertEqual(archive.read("SKILL.md"), (BASE / "skills/dg-pre-check/SKILL.md").read_bytes())

    def test_authored_presentation_is_editable_and_honestly_labeled(self):
        deck = Presentation(BASE / "QBR_authored_reference.pptx")
        self.assertEqual(len(deck.slides), 6)
        for slide in deck.slides:
            text = "\n".join(s.text for s in slide.shapes if s.has_text_frame)
            self.assertIn("Not a Copilot result", text)
            self.assertTrue(slide.has_notes_slide)
            self.assertTrue(slide.notes_slide.notes_text_frame.text)
        charts = [s.chart for slide in deck.slides for s in slide.shapes if s.has_chart]
        self.assertEqual(len(charts), 1)
        self.assertEqual(len(charts[0].series[0].values), 6)


if __name__ == "__main__":
    unittest.main()
