"""Offline acceptance for the original Contoso sources; no tenant writes."""
import hashlib
import importlib.util
import json
import re
import subprocess
import xml.etree.ElementTree as ET
from email import policy
from email.parser import BytesParser
from pathlib import Path
from zipfile import ZipFile

import imageio_ffmpeg
from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("port_sources", Path(__file__).with_name("build-port-operations.py"))
sources = importlib.util.module_from_spec(spec)
spec.loader.exec_module(sources)
folders = [sources.BASE, *sources.APPS]
catalog = json.loads((ROOT / "catalog.json").read_text(encoding="utf-8"))
installed = [demo for app in catalog["apps"] for demo in app["demos"]]
assert len(folders) == 15
for demo in folders:
    assert installed.count(demo) == 1, demo
    folder = ROOT / demo
    manifest = json.loads((folder / "manifest.json").read_text(encoding="utf-8"))
    assert manifest["title"]["en"] and manifest["title"]["de"]
    for entry in [*manifest["files"], *manifest.get("shared", [])]:
        path = ROOT / entry.get("from", demo) / entry["path"]
        assert path.is_file(), path
        assert hashlib.sha256(path.read_bytes()).hexdigest() == entry["sha256"], path
    actual = {str(p.relative_to(folder)).replace("\\", "/") for p in folder.rglob("*") if p.is_file()
              and p.name not in ("manifest.json", "shared.json", "tenant.json")
              and not p.relative_to(folder).parts[0] == "tenant"}
    assert actual == {entry["path"] for entry in manifest["files"]}, demo
    cfg = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    assert cfg["links"], demo
    for template in cfg.get("mails", []):
        for lang in ("en", "de"):
            path = folder / template.replace("{lang}", lang)
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            assert message["From"] and message["Subject"] and message.get_body().get_content().strip()
            assert any(upn.lower() in str(message["From"]).lower() for upn in cfg["requiredUsers"])
    if cfg.get("manual"):
        assert cfg["manual"]["en"] and cfg["manual"]["de"]

base = ROOT / sources.BASE
book = load_workbook(base / "port_exceptions.xlsx", data_only=False)
sheet = book["Shipments"]
rows = list(sheet.values)
assert rows[0] == tuple(sources.HEADERS)
assert len(rows) == 13 and len(sheet.tables) == 1 and "PortShipments" in sheet.tables
data = rows[1:]
assert len({row[0] for row in data}) == 12
assert sum(row[3] for row in data) == 115
held = [row for row in data if row[5] != "Ready"]
assert len(held) == 4 and sum(row[3] for row in held) == 43
assert sum(row[3] for row in data if row[1] == "Rotterdam") == 60
assert sum(row[3] for row in data if row[1] == "Hamburg") == 55
assert sum(row[3] for row in data if row[5] == "Ready") == 72
assert all(cell.data_type not in ("e", "f") for row in sheet for cell in row)
for name in ("port_operations_brief.docx", "port_status_v1.docx", "port_status_v2.docx"):
    with ZipFile(base / name) as package:
        assert package.testzip() is None
        ET.fromstring(package.read("word/document.xml"))
    doc = Document(base / name)
    assert doc.sections[0].page_width.inches == 8.5
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "43" in text and "115" in text and "pending" in text
    assert "CT-1005" in text
assert "has arrived" in "\n".join(p.text for p in Document(base / "port_status_v2.docx").paragraphs)

meeting = json.loads((ROOT / "teams-port-operations" / "tenant.json").read_text(encoding="utf-8"))["events"][0]
assert meeting["onlineMeeting"] is True and meeting["weekday"] == "Monday" and not meeting.get("fixedDate")
assert set(meeting["attendees"]) == {"Teresa Sac", "Sonia Rees", "Billie Vester"}
assert len((ROOT / "planner-port-operations" / "task_seeds.txt").read_text(encoding="utf-8").splitlines()[3:]) == 10
notebook = ROOT / "onenote-port-operations"
cfg = json.loads((notebook / "tenant.json").read_text(encoding="utf-8"))["onenote"]
for lang in ("en", "de"):
    path = notebook / cfg["pages"][0]["file"].replace("{lang}", lang)
    assert "43 containers" in path.read_text(encoding="utf-8")

video = ROOT / "clipchamp-port-operations"
vtt = (video / "port_handover_briefing.en.vtt").read_text(encoding="utf-8")
assert vtt.startswith("WEBVTT\n")
cues = re.findall(r"(\d{2}:\d{2}:\d{2}\.\d{3}) --> (\d{2}:\d{2}:\d{2}\.\d{3})\n(.+?)(?=\n\n|\Z)", vtt, re.S)
assert len(cues) >= 20


def seconds(stamp):
    h, m, s = stamp.split(":")
    return int(h) * 3600 + int(m) * 60 + float(s)


previous = 0
transcript = []
for start, end, text in cues:
    assert seconds(start) >= previous - 0.001
    assert seconds(end) > seconds(start)
    assert len(text.splitlines()) <= 3
    previous = seconds(end)
    transcript.extend(text.split())
assert 60 < previous < 600
assert len(transcript) > 100
expected = " ".join(text for _, text in sources.VIDEO).split()
assert transcript == expected, "The attached transcript must match the actual narration."
subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-v", "error", "-i", str(video / "port_handover_briefing.mp4"),
                "-map", "0:v:0", "-map", "0:a:0", "-f", "null", "-"], check=True)
print(f"PASS: 15 catalog scenarios, manifest/shared hashes, bilingual inbox/notes, exact workbook totals, "
      f"Office packages, Teams payload, 10 Planner seeds, decodable narrated video ({previous:.1f}s / {len(transcript)} words).")
