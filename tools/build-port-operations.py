"""Build original, fictional Contoso port-operation sources. No tenant access.

Uses the same python-docx/openpyxl conventions as the other demo builders.
Video additionally uses Pillow, imageio-ffmpeg and Windows' built-in English
speech synthesizer. Run Update-Manifest.ps1 -All after building.
"""
import csv
import json
import subprocess
import tempfile
import textwrap
import wave
from datetime import datetime, timezone
from email.message import EmailMessage
from pathlib import Path

import imageio_ffmpeg
from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parent.parent
BASE = "word-excel-port-operations"
DOMAIN = "M365CPI98544940.OnMicrosoft.com"
TRAINING = "https://learn.microsoft.com/training/modules/empower-workforce-copilot-operations/"
HEADERS = ["Shipment", "Port", "Consignee", "Containers", "Delay hours",
           "Status", "Owner", "Next action", "Due"]
SHIPMENTS = [
    ["CT-1001", "Rotterdam", "Fabrikam", 18, 0, "Ready", "Teresa Sac", "Confirm gate slot", "Monday 15:00"],
    ["CT-1002", "Rotterdam", "Tailspin Toys", 12, 18, "Document hold", "Teresa Sac", "Obtain corrected manifest", "Monday 15:00"],
    ["CT-1003", "Hamburg", "Fabrikam", 6, 24, "Inspection hold", "Sonia Rees", "Obtain authorized inspection outcome", "Monday 16:00"],
    ["CT-1004", "Hamburg", "Contoso", 10, 0, "Ready", "Teresa Sac", "Confirm gate slot", "Monday 15:00"],
    ["CT-1005", "Rotterdam", "Contoso", 14, 12, "Document hold", "Teresa Sac", "Obtain missing packing list", "Monday 15:00"],
    ["CT-1006", "Hamburg", "Tailspin Toys", 9, 0, "Ready", "Teresa Sac", "Confirm gate slot", "Monday 15:00"],
    ["CT-1007", "Rotterdam", "Fabrikam", 7, 0, "Ready", "Billie Vester", "Verify portal status", "Monday 15:00"],
    ["CT-1008", "Hamburg", "Contoso", 11, 8, "Carrier hold", "Billie Vester", "Obtain written cut-off confirmation", "Monday 15:00"],
    ["CT-1009", "Rotterdam", "Tailspin Toys", 5, 0, "Ready", "Teresa Sac", "Confirm gate slot", "Monday 15:00"],
    ["CT-1010", "Hamburg", "Fabrikam", 13, 0, "Ready", "Teresa Sac", "Confirm gate slot", "Monday 15:00"],
    ["CT-1011", "Rotterdam", "Contoso", 4, 0, "Ready", "Billie Vester", "Verify portal status", "Monday 15:00"],
    ["CT-1012", "Hamburg", "Tailspin Toys", 6, 0, "Ready", "Teresa Sac", "Confirm gate slot", "Monday 15:00"],
]
BRIEF = [
    ("Situation", "Contoso coordinates a fictional twelve-shipment handover across Rotterdam and Hamburg. "
     "There are 115 containers. Four shipments, covering 43 containers, have an unresolved hold. "
     "The other eight shipments cover 72 containers. All quantities, events and company activities are synthetic."),
    ("Decision", "The operations lead needs a verified exception list before Monday 15:00. "
     "Teresa Sac owns CT-1002 and CT-1005 document follow-up. Sonia Rees owns CT-1003 inspection follow-up. "
     "Billie Vester owns CT-1008 carrier confirmation. A confirmed delivery date is not available."),
    ("Control", "A document received is not a release authorization. Never describe a held shipment as released. "
     "Only the authorized operational process can clear a hold. The training material is not customs, "
     "legal, safety or routing advice. Do not send a customer commitment without the operations lead's review."),
    ("Source precedence", "port_exceptions.xlsx is the original snapshot. port_status_v1.docx is the 14:00 update. "
     "port_status_v2.docx is the 15:00 update: CT-1005's packing list has arrived, but release approval remains pending. "
     "CT-1008's cut-off is still unconfirmed. Do not silently merge a newer update into the original workbook."),
    ("Reusable clock", "Monday is the exercise's fictional operating day, not a claim about today's date. "
     "Use the actual calendar date shown in the installed Contoso port handover meeting for a live rehearsal. "
     "Keep all output drafts in the demo account's OneDrive."),
]
MAILS = [
    ("1_Teresa_Port_Exception", "Teresa Sac", "TeresaS",
     "Contoso port handover - document exceptions",
     "Contoso Hafenuebergabe - fehlende Dokumente",
     "Hi,\n\nPlease prepare the port handover for Contoso. CT-1002 has 12 containers and is waiting for a corrected "
     "manifest. CT-1005 has 14 containers and was waiting for a packing list at 14:00. "
     "The 15:00 update is important: the packing list has now arrived, but release approval has NOT been received. "
     "Both shipments remain on hold until the authorized process clears them.\n\n"
     "Sonia is following up the inspection outcome for CT-1003, with 6 containers. Billie is following up "
     "the carrier cut-off for CT-1008, with 11 containers. The four held shipments cover 43 containers. "
     "There are 115 containers in the twelve-shipment source workbook.\n\n"
     "Please distinguish confirmed facts from open questions in the update. Do not promise a delivery date or "
     "release a shipment. I need a short list of owners and the next evidence we need before Monday 15:00 "
     "in the fictional exercise. The actual rehearsal date is in the installed calendar invitation.\n\n"
     "This is synthetic Contoso training data, not a real shipping instruction.\n\nThanks,\nTeresa",
     "Hallo,\n\nbitte bereite die Hafenuebergabe fuer Contoso vor. Bei CT-1002 mit 12 Containern fehlt das korrigierte "
     "Manifest. Bei CT-1005 mit 14 Containern fehlte um 14 Uhr die Packliste. Das Update von 15 Uhr ist wichtig: "
     "Die Packliste ist eingegangen, die Freigabe liegt aber NICHT vor. Beide Sendungen bleiben gesperrt, "
     "bis das zustaendige Verfahren sie freigibt.\n\nSonia klaert das Pruefergebnis fuer CT-1003 mit 6 Containern. "
     "Billie klaert den Annahmeschluss fuer CT-1008 mit 11 Containern. Die vier gesperrten Sendungen umfassen "
     "43 Container; die zwoelf Sendungen in der Arbeitsmappe umfassen insgesamt 115 Container.\n\n"
     "Trenne im Update bestaetigte Fakten von offenen Fragen. Versprich keinen Liefertermin und gib keine Sendung frei. "
     "Ich brauche die Verantwortlichen und die jeweils noch fehlende Bestaetigung vor Montag 15 Uhr im fiktiven "
     "Szenario. Der wirkliche Probetermin steht in der installierten Kalendereinladung.\n\n"
     "Das sind synthetische Contoso-Uebungsdaten, keine echte Transportanweisung.\n\nDanke,\nTeresa"),
    ("2_Billie_Carrier_Confirmation", "Billie Vester", "BillieV",
     "Contoso port handover - carrier confirmation pending",
     "Contoso Hafenuebergabe - Speditionsbestaetigung offen",
     "Hi,\n\nThe carrier has reserved a provisional slot for CT-1008, but has NOT confirmed the cut-off in writing. "
     "The shipment covers 11 containers. Please keep its status as Carrier hold. "
     "A provisional slot is not a delivery commitment. I will obtain written confirmation before we update "
     "the operational status. No live portal or carrier system has been changed by this exercise.\n\nBillie",
     "Hallo,\n\nfuer CT-1008 gibt es einen vorlaeufig reservierten Slot, aber noch KEINE schriftliche Bestaetigung "
     "des Annahmeschlusses. Die Sendung umfasst 11 Container und bleibt im Status Carrier hold. "
     "Ein vorlaeufiger Slot ist keine Lieferzusage. Ich hole die schriftliche Bestaetigung ein, bevor wir den "
     "operativen Status aendern. Diese Uebung aendert kein echtes Portal oder Speditionssystem.\n\nBillie"),
]
VIDEO = [
    ("The situation", "Welcome to the Contoso port handover rehearsal. This is a fictional operations exercise, "
     "not a real shipping instruction. Our source workbook contains twelve shipments and one hundred and fifteen "
     "containers across Rotterdam and Hamburg. Eight shipments are ready for the next coordination step. "
     "Four shipments, covering forty-three containers, still have an unresolved hold."),
    ("Documents", "Teresa Sac owns the document follow-up. Shipment CT-1002 has twelve containers and needs a corrected "
     "manifest. Shipment CT-1005 has fourteen containers. Its packing list arrived in the fifteen hundred hours update, "
     "but release approval is still pending. Receiving a document does not authorize release. "
     "The older fourteen hundred hours status report must not override the newer evidence."),
    ("Inspection", "Sonia Rees is following up shipment CT-1003, which has six containers and an inspection hold. "
     "The team needs an outcome from the authorized inspection process. This rehearsal provides no customs or safety "
     "advice. Do not infer an approval, bypass a check, or describe the shipment as released just to make the update sound positive."),
    ("Carrier", "Billie Vester owns shipment CT-1008, with eleven containers. A provisional carrier slot exists, but "
     "the cut-off has not been confirmed in writing. The carrier hold stays open. No confirmed delivery date is available. "
     "The customer-facing draft must separate what is confirmed from what is still an open question."),
    ("Handover", "Before the fictional Monday fifteen hundred hours handover, ask each owner for the missing evidence. "
     "Teresa needs the corrected manifest and release decision. Sonia needs the inspection outcome. Billie needs "
     "written carrier confirmation. Use the actual calendar date in the rehearsal invitation rather than treating "
     "the fictional Monday as today's operating date."),
    ("Review", "The operations lead reviews the exception list and any stakeholder draft before it is sent. "
     "Keep the original workbook unchanged while comparing updates. A good handover names the shipment, its owner, "
     "the missing evidence, and the next action. It does not invent release approvals, confirmed arrival times, "
     "or messages that have not actually been sent. All content in this recording is synthetic Contoso training material."),
]
APPS = {
    "powerpoint-port-operations": ("PowerPoint", "https://powerpoint.cloud.microsoft/"),
    "onenote-port-operations": ("OneNote", "https://onenote.cloud.microsoft/"),
    "cowork-port-operations": ("Cowork", "https://copilot.cloud.microsoft/"),
    "chat-port-operations": ("Copilot Chat", "https://copilot.cloud.microsoft/chat"),
    "outlook-port-operations": ("Outlook", "https://outlook.cloud.microsoft/"),
    "teams-port-operations": ("Teams", "https://teams.microsoft.com/"),
    "forms-port-operations": ("Forms", "https://forms.office.com/"),
    "whiteboard-port-operations": ("Whiteboard", "https://whiteboard.microsoft.com/"),
    "loop-port-operations": ("Loop", "https://loop.cloud.microsoft/"),
    "onedrive-port-operations": ("OneDrive", "https://onedrive.cloud.microsoft/"),
    "sharepoint-port-operations": ("SharePoint", "https://www.microsoft365.com/launch/sharepoint"),
    "planner-port-operations": ("Planner", "https://planner.cloud.microsoft/"),
    "engage-port-operations": ("Viva Engage", "https://engage.cloud.microsoft/"),
    "clipchamp-port-operations": ("Clipchamp", "https://m365.cloud.microsoft/launch/Clipchamp/"),
}


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def write_doc(path, title, sections):
    doc = Document()
    doc.settings.element.find(qn("w:zoom")).set(qn("w:percent"), "100")
    doc.sections[0].page_width = Inches(8.5)
    doc.sections[0].page_height = Inches(11)
    for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(doc.sections[0], edge, Inches(1))
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, 0)
    doc.add_paragraph("CONTOSO | SYNTHETIC TRAINING DATA")
    for heading, text in sections:
        doc.add_heading(heading, 1)
        doc.add_paragraph(text)
    doc.core_properties.author = "Contoso"
    doc.core_properties.title = title
    doc.save(path)


def build_sources():
    folder = ROOT / BASE
    folder.mkdir(exist_ok=True)
    write_doc(folder / "port_operations_brief.docx", "Contoso port operations briefing", BRIEF)
    for version, detail in [
        (1, "14:00 snapshot: CT-1005's packing list is missing. Release approval is pending."),
        (2, "15:00 update: CT-1005's packing list has arrived. Release approval is still pending. "
         "CT-1008's carrier cut-off remains unconfirmed."),
    ]:
        write_doc(folder / f"port_status_v{version}.docx", f"Contoso port status - version {version}", [
            ("Scope", "Synthetic handover update. Twelve shipments, 115 containers; four holds covering 43 containers."),
            ("Document evidence", detail),
            ("Unchanged controls", "CT-1002 awaits the corrected manifest; CT-1003 awaits the authorized inspection outcome. "
             "No hold is cleared. No delivery date is confirmed. Do not mistake a document receipt for release authorization."),
        ])
    wb = Workbook()
    ws = wb.active
    ws.title = "Shipments"
    ws.append(HEADERS)
    for row in SHIPMENTS:
        ws.append(row)
    table = Table(displayName="PortShipments", ref=f"A1:I{ws.max_row}")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "D2"
    for col, width in zip("ABCDEFGHI", [16, 16, 20, 14, 14, 20, 20, 48, 21]):
        ws.column_dimensions[col].width = width
    for row in ws:
        for cell in row:
            cell.font = Font(name="Arial", size=11, bold=cell.row == 1)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.print_options.horizontalCentered = True
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    notes = wb.create_sheet("Source notes")
    for line in [
        "All rows are original synthetic Contoso training data; no real shipment identifiers.",
        "Monday is a fictional exercise day, not today's date.",
        "Ready is an administrative coordination status, not a delivery or release guarantee.",
        "Read port_status_v2.docx separately for the newer CT-1005 document receipt; approval remains pending.",
        "No calculated totals are embedded: Copilot builds the summary from the source table during the exercise.",
    ]:
        notes.append([line])
    notes.column_dimensions["A"].width = 110
    wb.save(folder / "port_exceptions.xlsx")
    with (folder / "port_exceptions.csv").open("w", newline="", encoding="utf-8") as out:
        csv.writer(out).writerows([HEADERS, *SHIPMENTS])
    (folder / "handover_notes.txt").write_text(
        "\n\n".join(f"{title}\n{text}" for title, text in BRIEF) + "\n", encoding="utf-8")


def manifest(folder, title):
    path = folder / "manifest.json"
    value = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {
        "next": {
            "en": ["Open the matching app page in the Demo Kit."],
            "de": ["Die passende App-Seite im Demo Kit öffnen."]}, "files": []}
    value["title"] = {"en": title, "de": title.replace("Contoso port operations", "Contoso-Hafenbetrieb")}
    write_json(path, value)


def build_tenant_sources():
    manifest(ROOT / BASE, "Contoso port operations - Word & Excel")
    write_json(ROOT / BASE / "tenant.json", {"links": [
        {"name": {"en": "Word", "de": "Word"}, "url": "https://www.microsoft365.com/launch/word"},
        {"name": {"en": "Excel", "de": "Excel"}, "url": "https://www.microsoft365.com/launch/excel"},
    ]})
    for demo, (app, url) in APPS.items():
        folder = ROOT / demo
        folder.mkdir(exist_ok=True)
        manifest(folder, f"Contoso port operations - {app}")
        write_json(folder / "shared.json", [
            f"{BASE}/port_operations_brief.docx", f"{BASE}/port_exceptions.xlsx",
            f"{BASE}/port_status_v1.docx", f"{BASE}/port_status_v2.docx", f"{BASE}/handover_notes.txt",
        ])
        path = folder / "tenant.json"
        config = json.loads(path.read_text(encoding="utf-8")) if path.exists() else {}
        config["links"] = [{"name": {"en": app, "de": app}, "url": url}]
        write_json(path, config)
    folder = ROOT / "outlook-port-operations"
    for stem, name, alias, subject_en, subject_de, body_en, body_de in MAILS:
        for lang, subject, body in [("en", subject_en, body_en), ("de", subject_de, body_de)]:
            msg = EmailMessage()
            msg["From"] = f"{name} <{alias}@{DOMAIN}>"
            msg["Subject"] = subject
            msg["Date"] = datetime(2026, 10, 23, 12, tzinfo=timezone.utc).strftime("%a, %d %b %Y %H:%M:%S %z")
            msg["Message-ID"] = f"<contoso-port-{stem}-{lang}@demo.invalid>"
            msg.set_content(body, charset="utf-8")
            target = folder / f"mails-{lang}" / f"{stem}.eml"
            target.parent.mkdir(exist_ok=True)
            target.write_bytes(msg.as_bytes())
    config = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    config["mails"] = [f"mails-{{lang}}/{mail[0]}.eml" for mail in MAILS]
    config["requiredUsers"] = [f"{mail[2]}@{DOMAIN}" for mail in MAILS]
    write_json(folder / "tenant.json", config)
    folder = ROOT / "teams-port-operations"
    config = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    config["events"] = [{
        "subject": {"en": "Contoso port handover rehearsal", "de": "Contoso Hafenübergabe - Generalprobe"},
        "weekday": "Monday", "start": "10:00", "minutes": 30, "onlineMeeting": True,
        "attendees": ["Teresa Sac", "Sonia Rees", "Billie Vester"],
        "body": {
            "en": "Synthetic Contoso rehearsal. Review the 12-shipment source table, the 15:00 document update "
                  "and the 4 unresolved holds. No real customer commitments. Rehearse with transcription "
                  "only after participant consent and in accordance with tenant policy.",
            "de": "Synthetische Contoso-Generalprobe. Die 12 Sendungen, das Dokumenten-Update von 15 Uhr und "
                  "die 4 offenen Sperren besprechen. Keine echten Kundenzusagen. Transkription nur nach "
                  "Zustimmung der Teilnehmenden und gemäß Tenant-Richtlinie verwenden.",
        },
    }]
    config["requiredUsers"] = [f"{alias}@{DOMAIN}" for alias in ("TeresaS", "SoniaR", "BillieV")]
    write_json(folder / "tenant.json", config)
    folder = ROOT / "onenote-port-operations"
    for lang, title in [("en", "Port handover notes"), ("de", "Notizen zur Hafenübergabe")]:
        from html import escape
        html = f"<html><head><title>{title}</title></head><body>"
        html += "".join(f"<h1>{escape(h)}</h1><p>{escape(t)}</p>" for h, t in BRIEF)
        html += "</body></html>"
        path = folder / "tenant" / "pages" / f"handover-{lang}.html"
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(html, encoding="utf-8")
    config = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    config["onenote"] = {"notebook": "Contoso Port Operations", "section": "Handover",
                        "pages": [{"title": {"en": "Port handover notes", "de": "Notizen zur Hafenübergabe"},
                                   "file": "tenant/pages/handover-{lang}.html"}]}
    write_json(folder / "tenant.json", config)
    folder = ROOT / "planner-port-operations"
    tasks = ["Verify CT-1002 manifest", "Review CT-1005 packing list", "Obtain CT-1005 release decision",
             "Obtain CT-1003 inspection outcome", "Confirm CT-1008 carrier cut-off",
             "Compare the 14:00 and 15:00 updates", "Draft stakeholder update",
             "Review held-container totals", "Prepare handover agenda", "Review draft before sharing"]
    (folder / "task_seeds.txt").write_text(
        "CONTOSO SYNTHETIC REHEARSAL TASKS\nPaste these ten tasks into an approved group-based demo plan.\n\n"
        + "\n".join(tasks) + "\n", encoding="utf-8")


def stamp(seconds):
    ms = round(seconds * 1000)
    return f"{ms // 3600000:02}:{ms // 60000 % 60:02}:{ms // 1000 % 60:02}.{ms % 1000:03}"


def build_video():
    folder = ROOT / "clipchamp-port-operations"
    with tempfile.TemporaryDirectory(prefix="contoso-port-video-") as tmp:
        temp = Path(tmp)
        write_json(temp / "speech.json", [text for _, text in VIDEO])
        ps = r"""
param([string]$Folder)
$ErrorActionPreference = 'Stop'
Add-Type -AssemblyName System.Speech
$speaker = New-Object System.Speech.Synthesis.SpeechSynthesizer
$voice = $speaker.GetInstalledVoices() | Where-Object { $_.Enabled -and $_.VoiceInfo.Culture.Name -like 'en-*' } | Select-Object -First 1
if (-not $voice) { throw 'An installed English Windows speech voice is required.' }
$speaker.SelectVoice($voice.VoiceInfo.Name)
$speaker.Rate = -1
$paragraphs = Get-Content -LiteralPath (Join-Path $Folder 'speech.json') -Raw -Encoding UTF8 | ConvertFrom-Json
try {
    for ($i = 0; $i -lt $paragraphs.Count; $i++) {
        $speaker.SetOutputToWaveFile((Join-Path $Folder "$i.wav"))
        $speaker.Speak($paragraphs[$i])
        $speaker.SetOutputToNull()
    }
} finally { $speaker.Dispose() }
"""
        (temp / "speak.ps1").write_text(ps, encoding="utf-8")
        subprocess.run(["powershell.exe", "-NoProfile", "-ExecutionPolicy", "Bypass", "-File",
                        str(temp / "speak.ps1"), "-Folder", str(temp)], check=True)
        elapsed = 0
        cues = ["WEBVTT", ""]
        images = []
        for i, (heading, text) in enumerate(VIDEO):
            with wave.open(str(temp / f"{i}.wav"), "rb") as audio:
                duration = audio.getnframes() / audio.getframerate()
            image = Image.new("RGB", (1280, 720), "white")
            draw = ImageDraw.Draw(image)
            title_font = ImageFont.truetype(r"C:\Windows\Fonts\arialbd.ttf", 42)
            body_font = ImageFont.truetype(r"C:\Windows\Fonts\arial.ttf", 29)
            draw.text((70, 55), "CONTOSO | PORT OPERATIONS", font=body_font, fill="#115EA3")
            draw.text((70, 145), heading, font=title_font, fill="black")
            draw.multiline_text((70, 230), "\n".join(textwrap.wrap(text, 66)),
                                font=body_font, fill="black", spacing=14)
            draw.text((70, 650), "Synthetic training data | Not an actual meeting recording", font=body_font, fill="black")
            image.save(temp / f"{i}.png")
            images.append((temp / f"{i}.png", duration))
            words = text.split()
            for offset in range(0, len(words), 12):
                chunk = words[offset:offset + 12]
                start = elapsed + duration * offset / len(words)
                end = elapsed + duration * (offset + len(chunk)) / len(words)
                cues.extend([f"{stamp(start)} --> {stamp(end)}",
                             "\n".join(textwrap.wrap(" ".join(chunk), 48)), ""])
            elapsed += duration
        with wave.open(str(temp / "narration.wav"), "wb") as out:
            for i in range(len(VIDEO)):
                with wave.open(str(temp / f"{i}.wav"), "rb") as part:
                    if i == 0:
                        out.setparams(part.getparams())
                    out.writeframes(part.readframes(part.getnframes()))
        concat = []
        for image, duration in images:
            concat.extend([f"file '{image.as_posix()}'", f"duration {duration}"])
        concat.append(f"file '{images[-1][0].as_posix()}'")
        (temp / "frames.txt").write_text("\n".join(concat), encoding="utf-8")
        subprocess.run([imageio_ffmpeg.get_ffmpeg_exe(), "-hide_banner", "-loglevel", "error", "-y",
                        "-f", "concat", "-safe", "0", "-i", str(temp / "frames.txt"),
                        "-i", str(temp / "narration.wav"), "-c:v", "libx264", "-pix_fmt", "yuv420p",
                        "-r", "24", "-c:a", "aac", "-b:a", "128k", "-shortest", "-movflags", "+faststart",
                        str(folder / "port_handover_briefing.mp4")], check=True)
        (folder / "port_handover_briefing.en.vtt").write_text("\n".join(cues), encoding="utf-8")
        (folder / "video_narration.txt").write_text(
            "\n\n".join(text for _, text in VIDEO) + "\n", encoding="utf-8")


if __name__ == "__main__":
    build_sources()
    build_tenant_sources()
    build_video()
    print("Built Contoso port sources, tenant inputs and narrated MP4. Run tools\\Update-Manifest.ps1 -All.")
