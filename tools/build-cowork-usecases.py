"""Builds the fictional Contoso Fargo logistics demo files for six Copilot Cowork use cases.

Run:  python tools/build-cowork-usecases.py
Needs: pip install python-docx openpyxl reportlab
All content is fictional (Contoso, Tailspin Toys, Woodgrove Bank, Northwind, Fabrikam, Adatum, Litware and
Proseware are Microsoft demo companies). People are real users of the CDX demo tenant.
Scenario: Contoso distribution center Fargo (Dock 3 closes on Apr 27, go-live of the new wing on Jun 22, 2027).

Writes one folder per demo next to this tools folder and the presenter-only expected results to tools/expected/.
The expected results must NOT live in a demo folder: Cowork could find them in OneDrive.
"""
import csv
import random
import re
import zipfile
from datetime import date, datetime
from pathlib import Path

from docx import Document
from docx.shared import Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import simpleSplit
from reportlab.pdfgen import canvas

import cowork_usecases_text as T

ROOT = Path(__file__).resolve().parent.parent
EXPECTED = Path(__file__).resolve().parent / "expected"
EXPECTED.mkdir(exist_ok=True)
NAVY = "0B3D91"


def folder(name):
    p = ROOT / name
    p.mkdir(exist_ok=True)
    return p


FIXED = datetime(2027, 3, 1, 8, 0, 0)


def freeze(path):
    """Make Office files byte-identical across runs: fixed core properties and zip timestamps."""
    path = Path(path)
    src = zipfile.ZipFile(path)
    items = [(i.filename, src.read(i.filename)) for i in src.infolist()]
    src.close()
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            if name == "docProps/core.xml":
                stamp = FIXED.strftime("%Y-%m-%dT%H:%M:%SZ").encode()
                data = re.sub(rb"(<dcterms:(?:created|modified)[^>]*>)[^<]*", rb"\g<1>" + stamp, data)
            zi = zipfile.ZipInfo(name, FIXED.timetuple()[:6])
            zi.compress_type = zipfile.ZIP_DEFLATED
            z.writestr(zi, data)


def save_wb(wb, path):
    wb.properties.created = wb.properties.modified = FIXED
    wb.properties.creator = wb.properties.lastModifiedBy = "Contoso Demo"
    wb.save(path)
    freeze(path)


def save_doc(d, path):
    cp = d.core_properties
    cp.created = cp.modified = cp.last_printed = FIXED
    cp.author = cp.last_modified_by = "Contoso Demo"
    cp.revision = 1
    d.save(path)
    freeze(path)


# ---------------------------------------------------------------- helpers
def write_sheet(ws, headers, rows, widths=None, money_cols=(), pct_cols=()):
    ws.append(headers)
    for r in rows:
        ws.append(list(r))
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=NAVY)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, h in enumerate(headers, 1):
        w = (widths or {}).get(i) or max(12, min(40, len(str(h)) + 4))
        ws.column_dimensions[get_column_letter(i)].width = w
    for col in money_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = "#,##0.00"
    for col in pct_cols:
        for cell in ws[get_column_letter(col)][1:]:
            cell.number_format = "0.0%"
    ws.freeze_panes = "A2"


def doc_with(title, subtitle, note, sections):
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(title, 0)
    d.add_paragraph(subtitle)
    p = d.add_paragraph(note)
    p.runs[0].italic = True
    for head, paras in sections:
        d.add_heading(head, 1)
        for t in paras:
            d.add_paragraph(t)
    return d


def pdf_lines(path, lines, title):
    c = canvas.Canvas(str(path), pagesize=A4, invariant=1)
    c.setTitle(title)
    w, h = A4
    y = h - 60
    for kind, text in lines:
        if kind == "h1":
            c.setFont("Helvetica-Bold", 16)
            size, lead = 16, 22
        elif kind == "h2":
            c.setFont("Helvetica-Bold", 11)
            size, lead = 11, 16
        elif kind == "mono":
            c.setFont("Courier", 9.5)
            size, lead = 9.5, 13
        else:
            c.setFont("Helvetica", 10)
            size, lead = 10, 14
        font = {"h1": "Helvetica-Bold", "h2": "Helvetica-Bold", "mono": "Courier"}.get(kind, "Helvetica")
        for part in simpleSplit(text, font, size, w - 120) or [""]:
            if y < 70:
                c.showPage()
                y = h - 60
                c.setFont(font, size)
            c.drawString(60, y, part)
            y -= lead
        y -= 3 if kind in ("h1", "h2") else 2
    c.save()


# ================================================================ 1) freight invoice check
CARRIERS = [
    # id, legal name, status, currency, country, tax id, terms, AP owner
    ("N100", "Northwind Freight Inc.", "Active", "USD", "United States", "US-MN-554201", "Net 60", "Teresa Sac"),
    ("F200", "Fabrikam Logistics LLC", "Active", "USD", "United States", "US-IL-608811", "Net 30", "Teresa Sac"),
    ("A300", "Adatum Cargo B.V.", "Active", "EUR", "Netherlands", "NL-8264.19.771.B01", "Net 45", "Sydney Mattos"),
    ("L400", "Litware Transport GmbH", "Blocked", "EUR", "Germany", "DE-811-204-993", "Net 30", "Sydney Mattos"),
    ("P500", "Proseware Courier Inc.", "Active", "USD", "United States", "US-ND-580432", "Net 30", "Teresa Sac"),
]
CARRIER = {c[0]: c for c in CARRIERS}

# TO number, line, carrier, lane, service, chargeable kg, rate per 100 kg, currency, fuel %, tax code, tax rate
TO_LINES = [
    ("TO-27-0101", 1, "N100", "FAR-MSP", "LTL", 12400, 18.50, "USD", 0.08, "US-FRT-0", 0.0),
    ("TO-27-0101", 2, "N100", "FAR-DEN", "LTL", 6000, 21.00, "USD", 0.08, "US-FRT-0", 0.0),
    ("TO-27-0102", 1, "F200", "FAR-CHI", "LTL", 15000, 24.80, "USD", 0.09, "US-FRT-0", 0.0),
    ("TO-27-0103", 1, "A300", "FAR-ROT", "Ocean-ready", 9000, 31.00, "EUR", 0.11, "NL-STD-21", 0.21),
    ("TO-27-0103", 2, "A300", "FAR-AMS", "Ocean-ready", 8500, 32.50, "EUR", 0.11, "NL-STD-21", 0.21),
    ("TO-27-0104", 1, "L400", "FAR-OMA", "LTL", 7200, 17.20, "EUR", 0.07, "DE-STD-19", 0.19),
    ("TO-27-0105", 1, "P500", "FAR-SEA", "Express", 4800, 38.00, "USD", 0.10, "US-FRT-0", 0.0),
    ("TO-27-0106", 1, "N100", "FAR-MSP", "LTL", 10000, 18.50, "USD", 0.08, "US-FRT-0", 0.0),
]
TO = {(t[0], t[1]): t for t in TO_LINES}

# TO, line, delivered kg, delivery date, status
PODS = [
    ("TO-27-0101", 1, 12400, "2027-04-02", "Confirmed"),
    ("TO-27-0101", 2, 6000, "2027-04-03", "Confirmed"),
    ("TO-27-0102", 1, 15000, "2027-04-02", "Confirmed"),
    ("TO-27-0103", 1, 9000, "2027-04-04", "Confirmed"),
    ("TO-27-0103", 2, 7960, "2027-04-04", "Confirmed"),
    ("TO-27-0104", 1, 7200, "2027-04-03", "Confirmed"),
    ("TO-27-0105", 1, 0, "", "Pending"),
    ("TO-27-0106", 1, 10000, "2027-04-05", "Confirmed"),
]

# invoice id, carrier, invoice date, currency, lines [(TO or "", line, lane, kg, rate, fuel%, tax rate)], expected codes
INVOICES = [
    ("NW-1042", "N100", "2027-04-06", "USD",
     [("TO-27-0101", 1, "FAR-MSP", 12400, 18.50, 0.08, 0.0), ("TO-27-0101", 2, "FAR-DEN", 6000, 21.00, 0.08, 0.0)], []),
    ("NW-1043", "N100", "2027-04-07", "USD",
     [("TO-27-0106", 1, "FAR-MSP", 10000, 19.60, 0.08, 0.0)], ["E03"]),
    ("FL-7710", "F200", "2027-04-06", "USD",
     [("TO-27-0102", 1, "FAR-CHI", 15000, 24.80, 0.09, 0.0)], []),
    ("FL-7711", "F200", "2027-04-09", "USD",
     [("TO-27-0102", 1, "FAR-CHI", 15000, 24.80, 0.09, 0.0)], ["E02"]),
    ("AC-8821", "A300", "2027-04-07", "EUR",
     [("TO-27-0103", 1, "FAR-ROT", 9000, 31.00, 0.11, 0.21)], []),
    ("AC-8822", "A300", "2027-04-07", "EUR",
     [("TO-27-0103", 2, "FAR-AMS", 8500, 32.50, 0.11, 0.21)], ["E07"]),
    ("LW-3901", "L400", "2027-04-08", "EUR",
     [("TO-27-0104", 1, "FAR-OMA", 7200, 17.20, 0.07, 0.19)], ["E05"]),
    ("PC-5520-A", "P500", "2027-04-08", "USD",
     [("TO-27-0105", 1, "FAR-SEA", 4800, 38.00, 0.10, 0.0)], ["E04"]),
    ("PC-5520-B", "P500", "2027-04-09", "USD",
     [("", 0, "FAR-SEA", 4800, 38.00, 0.10, 0.0)], ["E01"]),
]


def line_amounts(kg, rate, fuel, tax):
    net = round(kg / 100 * rate, 2)
    fsc = round(net * fuel, 2)
    tx = round((net + fsc) * tax, 2)
    return net, fsc, tx, round(net + fsc + tx, 2)


def build_freight():
    out = folder("cowork-freight-invoice-check")
    d = doc_with(T.POLICY_TITLE, T.POLICY_SUBTITLE, T.POLICY_NOTE, T.POLICY)
    save_doc(d, out / "freight_audit_policy.docx")

    wb = Workbook()
    ws = wb.active
    ws.title = "Transport Order Lines"
    rows = []
    for t in TO_LINES:
        net, fsc, tx, gross = line_amounts(t[5], t[6], t[8], t[10])
        rows.append((*t, net, gross))
    write_sheet(ws, ["TO Number", "Line", "Carrier ID", "Lane Code", "Service", "Chargeable Weight (kg)",
                     "Rate per 100 kg", "Currency", "Fuel Surcharge %", "Tax Code", "Tax Rate",
                     "Authorized Net Freight", "Authorized Gross"], rows,
                money_cols=(7, 12, 13), pct_cols=(9, 11), widths={1: 14, 4: 12})
    ws2 = wb.create_sheet("TO Summary")
    summ = {}
    for r in rows:
        k = (r[0], r[2], r[7])
        summ[k] = summ.get(k, 0) + r[12]
    write_sheet(ws2, ["TO Number", "Carrier ID", "Currency", "Authorized Gross"],
                [(k[0], k[1], k[2], round(v, 2)) for k, v in summ.items()], money_cols=(4,), widths={1: 14})
    save_wb(wb, out / "freight_transport_orders.xlsx")

    with open(out / "freight_pod_receipts.csv", "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["TO Number", "Line", "Delivered Weight (kg)", "Delivery Date", "POD Status"])
        w.writerows(PODS)

    wb = Workbook()
    ws = wb.active
    ws.title = "Carrier Master"
    write_sheet(ws, ["Carrier ID", "Legal Name", "Status", "Default Currency", "Country", "Tax Registration ID",
                     "Payment Terms", "AP Owner"], CARRIERS, widths={2: 28, 6: 22})
    save_wb(wb, out / "freight_carrier_master.xlsx")

    for inv_id, cid, inv_date, cur, lines, _ in INVOICES:
        c = CARRIER[cid]
        pdf = [("h1", f"INVOICE {inv_id}"), ("h2", c[1]), ("p", f"Carrier ID: {cid}   Tax ID: {c[5]}"),
               ("p", f"Invoice date: {inv_date}   Currency: {cur}"), ("p", "Bill to: Contoso Distribution Center Fargo, Accounts Payable"),
               ("p", "")]
        tot_net = tot_fsc = tot_tax = tot_gross = 0.0
        for i, (to, ln, lane, kg, rate, fuel, tax) in enumerate(lines, 1):
            net, fsc, tx, gross = line_amounts(kg, rate, fuel, tax)
            tot_net += net; tot_fsc += fsc; tot_tax += tx; tot_gross += gross
            pdf += [("h2", f"Line {i}"),
                    ("mono", f"Transport order : {to if to else '(not stated)'}"),
                    ("mono", f"TO line         : {ln if to else '(not stated)'}"),
                    ("mono", f"Lane code       : {lane}"),
                    ("mono", f"Chargeable wt   : {kg:,} kg"),
                    ("mono", f"Rate per 100 kg : {rate:.2f} {cur}"),
                    ("mono", f"Net freight     : {net:,.2f} {cur}"),
                    ("mono", f"Fuel surcharge  : {fuel*100:.1f}%  {fsc:,.2f} {cur}"),
                    ("mono", f"Tax {tax*100:.0f}%        : {tx:,.2f} {cur}"),
                    ("mono", f"Line total      : {gross:,.2f} {cur}"), ("p", "")]
        pdf += [("h2", f"Invoice total: {tot_gross:,.2f} {cur}"), ("p", f"Payable within {c[6].replace('Net ', '')} days.")]
        pdf_lines(out / f"freight_invoice_{inv_id}.pdf", pdf, f"Invoice {inv_id}")
    return out


def freight_expected():
    pod = {(p[0], p[1]): p for p in PODS}
    seen = set()
    rows = []
    for inv_id, cid, inv_date, cur, lines, codes in INVOICES:
        for to, ln, lane, kg, rate, fuel, tax in lines:
            net, fsc, tx, gross = line_amounts(kg, rate, fuel, tax)
            rows.append((inv_id, cid, to or "-", ln or "-", cur, gross, codes))
    # sanity check of the seeded flaws
    for inv_id, cid, inv_date, cur, lines, codes in INVOICES:
        for to, ln, lane, kg, rate, fuel, tax in lines:
            found = []
            if not to:
                found.append("E01")
            else:
                t = TO[(to, ln)]
                if (to, ln) in seen:
                    found.append("E02")
                seen.add((to, ln))
                if abs(rate / t[6] - 1) > 0.02 or fuel != t[8]:
                    found.append("E03")
                p = pod[(to, ln)]
                if p[4] != "Confirmed":
                    found.append("E04")
                elif kg > p[2] * 1.03:
                    found.append("E07")
            if CARRIER[cid][2] == "Blocked":
                found.append("E05")
            assert sorted(found) == sorted(codes), (inv_id, found, codes)
    totals = {}
    md = ["# Expected result: cowork-freight-invoice-check", "",
          "| Invoice | Carrier | TO | Line | Status | Codes | Amount |", "|---|---|---|---|---|---|---|"]
    for inv_id, cid, to, ln, cur, gross, codes in rows:
        status = "BLOCKED" if codes else "RELEASED"
        totals.setdefault(cur, {"RELEASED": 0.0, "BLOCKED": 0.0})[status] += gross
        md.append(f"| {inv_id} | {cid} | {to} | {ln} | {status} | {', '.join(codes) or '-'} | {gross:,.2f} {cur} |")
    md += ["", "| Currency | Released | Blocked | Total |", "|---|---|---|---|"]
    for cur, v in totals.items():
        md.append(f"| {cur} | {v['RELEASED']:,.2f} | {v['BLOCKED']:,.2f} | {v['RELEASED'] + v['BLOCKED']:,.2f} |")
    md += ["", "Decision cases for the user: NW-1043 (rate +5.9%), FL-7711 (duplicate), AC-8822 (weight +6.8% over POD), "
           "LW-3901 (carrier blocked), PC-5520-A (POD pending), PC-5520-B (no TO number, do not guess)."]
    (EXPECTED / "cowork-freight-invoice-check.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    return totals


# ================================================================ 2) carrier contracts
def build_contracts():
    out = folder("cowork-carrier-contracts")
    for key, c in T.CONTRACTS.items():
        lines = [("h1", "CARRIER SERVICES AGREEMENT"),
                 ("p", "Between Contoso Distribution Center Fargo (the customer) and " + c["address"] + " (the carrier)."),
                 ("p", "Subject: less-than-truckload transport from the Fargo distribution center. Synthetic demo contract."), ("p", "")]
        for head, text in c["clauses"]:
            lines += [("h2", head), ("p", text)]
        pdf_lines(out / c["file"], lines, c["carrier"] + " agreement")
    exp = ["# Expected result: cowork-carrier-contracts", "",
           "| Topic | Northwind Freight | Fabrikam Logistics | Adatum Cargo |", "|---|---|---|---|",
           "| Rate Fargo-Minneapolis (per 100 kg) | 18.50 | 19.90 | 19.20 |",
           "| Fuel surcharge | uncapped | capped at 12% | uncapped |",
           "| Payment | 60 days | 30 days | 45 days |",
           "| Price adjustment | fixed 2% per year | CPI, max 3% | CPI + 2 points, no cap |",
           "| On-time target / credit | 94% / 2% below 92% | 98% / 5% below 96% | 96% / 3% below 94% |",
           "| Liability | USD 2/kg | invoice value up to USD 50,000 | USD 6/kg |",
           "| Data | US only, subprocessors without notice, breach 7 days | US/EU, subprocessor list, 72 hours | any country, no notice, 14 days |",
           "| Exit | after month 24, 6 months + 25% fee, 90% volume commitment | 90 days, no fee, no minimum | 180 days + 10% fee |",
           "| Renewal | auto 36 months, 12 months notice | 12 months, 90 days notice | auto 24 months, 6 months notice |", "",
           "Recommendation: Fabrikam Logistics (best service, liability, data and exit), despite the highest rate. "
           "Biggest risks: Northwind 36-month auto-renewal plus 25% exit fee and USD 2/kg liability; Adatum uncapped price adjustment "
           "and open data processing; Fabrikam rate 7.6% above Northwind and 3% annual increase cap."]
    (EXPECTED / "cowork-carrier-contracts.md").write_text("\n".join(exp) + "\n", encoding="utf-8")
    return out


# ================================================================ 3) lane margin
LANES = [("FAR-MSP", 118), ("FAR-CHI", 164), ("FAR-DEN", 210), ("FAR-SEA", 255), ("FAR-OMA", 142), ("FAR-STL", 188),
         ("FAR-KC", 176), ("FAR-DAL", 232), ("FAR-WIN", 128), ("FAR-BIS", 62), ("FAR-SLC", 238), ("FAR-ATL", 268)]
SERVICES = {"Express": (1.8, 8.5), "Standard": (1.0, 5.0), "LTL Freight": (0.85, 6.5), "Contract Logistics": (1.3, 7.5)}
CHANNELS = {"Direct": 0.02, "Broker": 0.06, "Marketplace": 0.12}
CUSTOMERS = ["Tailspin Toys", "Woodgrove Bank", "Fourth Coffee", "Lucerne Publishing", "Wide World Importers",
             "Margie's Travel", "Adventure Works", "Proseware", "Coho Vineyard", "Trey Research"]


def margin_rows():
    rnd = random.Random(2027)
    rows = []
    for i in range(600):
        lane, base = rnd.choice(LANES)
        svc = rnd.choice(list(SERVICES))
        ch = rnd.choices(list(CHANNELS), weights=[5, 3, 3])[0]
        mult, handling = SERVICES[svc]
        price = round(base * mult * rnd.uniform(0.92, 1.08) * 0.55, 2)
        ratio = rnd.uniform(0.50, 0.62)
        if lane == "FAR-SEA" and svc == "Standard":
            ratio = rnd.uniform(0.86, 0.95)
        carrier = round(price * ratio, 2)
        fuel = round(price * 0.07, 2)
        hand = round(handling * rnd.uniform(0.9, 1.15), 2)
        fee = round(price * CHANNELS[ch], 2)
        if ch == "Broker" and svc == "Express":
            disc = round(rnd.uniform(0.18, 0.28), 2)
        elif ch == "Broker":
            disc = round(rnd.uniform(0.05, 0.12), 2)
        elif ch == "Marketplace":
            disc = round(rnd.uniform(0.03, 0.10), 2)
        else:
            disc = round(rnd.uniform(0.0, 0.08), 2)
        units = rnd.randint(150, 3000)
        rows.append((f"LN-{10000 + i}", lane, svc, ch, rnd.choice(CUSTOMERS), price, carrier, fuel, hand, fee, disc, units))
    return rows


def build_margin():
    out = folder("cowork-lane-margin")
    rows = margin_rows()
    wb = Workbook()
    ws = wb.active
    ws.title = "lane_margin_raw"
    write_sheet(ws, ["Booking ID", "Lane", "Service", "Sales Channel", "Customer", "Rate per Shipment ($)",
                     "Carrier Cost ($)", "Fuel Cost ($)", "Handling ($)", "Channel Fee ($)", "Discount %",
                     "Shipments (Qtr)"], rows, widths={1: 12, 3: 20, 5: 22}, money_cols=(6, 7, 8, 9, 10))
    for cell in ws["K"][1:]:
        cell.number_format = "0%"
    save_wb(wb, out / "lane_margin_raw.xlsx")

    def profit(r, disc=None):
        d = r[10] if disc is None else disc
        return (r[5] * (1 - d) - r[6] - r[7] - r[8] - r[9]) * r[11]

    total = sum(profit(r) for r in rows)
    by = {}
    for r in rows:
        by.setdefault(("svc", r[2], r[3]), 0.0)
        by[("svc", r[2], r[3])] += profit(r)
    lane = {}
    for r in rows:
        lane.setdefault((r[1], r[2]), 0.0)
        lane[(r[1], r[2])] += profit(r)
    loss_rows = [r for r in rows if profit(r) < 0]
    leak_rows = [r for r in rows if r[3] == "Broker" and r[2] == "Express"]
    uplift_disc = sum(profit(r, 0.12) - profit(r) for r in leak_rows if r[10] > 0.12)
    sea = [r for r in rows if r[1] == "FAR-SEA" and r[2] == "Standard"]
    uplift_sea = sum(r[5] * 0.15 * r[11] for r in sea)
    mk = [r for r in rows if r[3] == "Marketplace"]
    uplift_fee = sum(r[9] * 0.25 * r[11] for r in mk)
    actions = sorted([(uplift_disc, "Cap the broker discount on Express at 12%"),
                      (uplift_sea, "Reprice FAR-SEA Standard by 15% (carrier cost is 86-95% of the rate)"),
                      (uplift_fee, "Renegotiate Marketplace fees by 25%")], reverse=True)
    md = ["# Expected result: cowork-lane-margin", "",
          f"Rows: {len(rows)}. Quarterly profit across all rows: {total:,.0f} USD. Loss-making rows: {len(loss_rows)} "
          f"(sum {sum(profit(r) for r in loss_rows):,.0f} USD).", "",
          "## Profit by service and channel (USD per quarter)", "", "| Service | Channel | Profit |", "|---|---|---|"]
    for k, v in sorted(by.items(), key=lambda kv: kv[1]):
        md.append(f"| {k[1]} | {k[2]} | {v:,.0f} |")
    md += ["", "## Weakest lane/service combinations", "", "| Lane | Service | Profit |", "|---|---|---|"]
    for k, v in sorted(lane.items(), key=lambda kv: kv[1])[:6]:
        md.append(f"| {k[0]} | {k[1]} | {v:,.0f} |")
    md += ["", "## Actions worth the most (profit uplift per quarter)", ""]
    md += [f"{i}. {t}: about {v:,.0f} USD." for i, (v, t) in enumerate(actions, 1)]
    (EXPECTED / "cowork-lane-margin.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    return total, uplift_disc, uplift_sea, uplift_fee


# ================================================================ 4) customer QBR
def build_qbr():
    out = folder("cowork-customer-qbr")
    wb = Workbook()
    ws = wb.active
    ws.title = "Account Health"
    write_sheet(ws, ["Account", "Account Manager", "Annual Revenue (USD)", "Contract End", "Health Score", "NPS",
                     "Claims (Qtr)", "On-time Delivery %", "On-time SLA %", "Volume Change %", "Expansion Opp (USD)", "RAG"],
                [("Woodgrove Bank", "Sydney Mattos", 640000, "2028-03-31", 84, 8.6, 2, 97.1, 95, 4, 90000, "Green"),
                 ("Tailspin Toys", "Teresa Sac", 1420000, "2027-12-31", 58, 5.9, 14, 86.4, 95, -9, 300000, "Red"),
                 ("Fourth Coffee", "Teresa Sac", 380000, "2028-06-30", 77, 7.7, 5, 94.2, 94, 1, 40000, "Yellow"),
                 ("Lucerne Publishing", "Sydney Mattos", 215000, "2027-09-30", 90, 9.1, 1, 98.0, 95, 7, 55000, "Green"),
                 ("Wide World Importers", "Vance DeLeon", 910000, "2028-01-31", 73, 7.2, 6, 93.8, 95, -2, 120000, "Yellow"),
                 ("Margie's Travel", "Vance DeLeon", 98000, "2027-08-31", 46, 5.1, 11, 88.0, 95, -15, 0, "Red")],
                widths={1: 24, 2: 18})
    ws = wb.create_sheet("QBR Agenda Template")
    write_sheet(ws, ["Minutes", "Topic", "Owner", "Goal", "Output"],
                [("0-5", "Welcome and goals of the review", "Account manager", "Agree on what we decide today", "Shared goals"),
                 ("5-15", "Delivery performance last quarter", "Logistics Fargo", "Show on-time rate, claims, root causes", "Facts and causes"),
                 ("15-30", "What we change", "Operations", "Present the recovery plan", "Plan with dates"),
                 ("30-45", "Outlook and growth", "Account manager", "Discuss volumes and new lanes", "Opportunity list"),
                 ("45-55", "Decisions and next steps", "All", "Agree on owners and dates", "Action list"),
                 ("55-60", "Wrap-up", "Account manager", "Confirm follow-up", "Next review date")],
                widths={2: 36, 4: 40, 5: 22})
    ws = wb.create_sheet("Success Plan")
    write_sheet(ws, ["Account", "Objective", "Measure", "Target", "Status", "Owner", "Due"],
                [("Tailspin Toys", "Restore on-time delivery", "On-time %", "95%", "Behind", "Teresa Sac", "2027-06-30"),
                 ("Tailspin Toys", "Reduce damage claims", "Claims per quarter", "6", "Behind", "Sonia Rees", "2027-07-31"),
                 ("Tailspin Toys", "Prepare cross-border lane", "Pilot start", "Q4", "On track", "Teresa Sac", "2027-10-01"),
                 ("Woodgrove Bank", "Expand document logistics", "Volume", "+8%", "On track", "Sydney Mattos", "2027-12-31"),
                 ("Fourth Coffee", "Stabilize peak delivery", "On-time %", "95%", "At risk", "Teresa Sac", "2027-09-30"),
                 ("Margie's Travel", "Retain contract", "Renewal signed", "Yes", "At risk", "Vance DeLeon", "2027-08-31")],
                widths={1: 20, 2: 34})
    ws = wb.create_sheet("Claims and Tickets")
    write_sheet(ws, ["Ticket", "Account", "Date", "Category", "Description", "Status", "Root Cause"],
                [("T-3101", "Tailspin Toys", "2027-01-12", "Late delivery", "Pallets for store opening arrived one day late", "Closed", "Carrier arrived before dock window"),
                 ("T-3115", "Tailspin Toys", "2027-01-19", "Damage", "Crushed cartons, 6 pallets", "Closed", "Mixed loading at Dock 3"),
                 ("T-3140", "Tailspin Toys", "2027-02-02", "Late delivery", "Weekly replenishment missed cutoff", "Closed", "Gate scanner outage"),
                 ("T-3162", "Tailspin Toys", "2027-02-16", "Late delivery", "Two trucks waited 3 hours at the gate", "Closed", "Gate scanner outage"),
                 ("T-3177", "Tailspin Toys", "2027-02-23", "Damage", "Water-damaged cartons", "Open", "Dock 3 roof repair"),
                 ("T-3190", "Tailspin Toys", "2027-03-02", "Late delivery", "Replenishment 1 day late", "Open", "Carrier slot conflict"),
                 ("T-3204", "Tailspin Toys", "2027-03-09", "Wrong item", "SKU swap on 40 cartons", "Closed", "Picking error"),
                 ("T-3219", "Tailspin Toys", "2027-03-16", "Late delivery", "Promotional shipment late", "Open", "Carrier slot conflict"),
                 ("T-3231", "Margie's Travel", "2027-02-10", "Late delivery", "Print materials late for a trade show", "Closed", "Carrier delay"),
                 ("T-3248", "Margie's Travel", "2027-03-05", "Damage", "Damaged display units", "Open", "Packaging"),
                 ("T-3260", "Fourth Coffee", "2027-03-12", "Late delivery", "Seasonal order late", "Closed", "Carrier delay"),
                 ("T-3275", "Wide World Importers", "2027-03-18", "Wrong item", "Mixed pallet", "Closed", "Picking error")],
                widths={5: 44, 7: 34})
    ws = wb.create_sheet("Pre-Read Email Draft")
    write_sheet(ws, ["Field", "Content"],
                [("To", "[customer contact]"),
                 ("Subject", "Quarterly review [date]: where we stand and what we propose"),
                 ("Opening", "[Short acknowledgement of the service problems, no excuses]"),
                 ("Facts", "[On-time rate, claims, root causes with numbers]"),
                 ("Proposal", "[Concrete changes with dates]"),
                 ("Ask", "[What we need from the customer]"),
                 ("Close", "[Agenda attached, time and place]")], widths={1: 14, 2: 70})
    save_wb(wb, out / "customer_qbr_prep.xlsx")
    exp = ["# Expected result: cowork-customer-qbr", "",
           "Red accounts: Tailspin Toys (health 58, NPS 5.9, 14 claims, on-time 86.4% vs SLA 95%, volume -9%, revenue 1,420,000 USD, "
           "contract end 2027-12-31, expansion 300,000 USD) and Margie's Travel (health 46, on-time 88.0%, volume -15%).",
           "Root causes in the ticket sheet for Tailspin Toys: gate scanner outage (2), carrier slot conflict (2), Dock 3 handling and roof "
           "repair (2 damage), picking error (1), carrier arrived before dock window (1). 8 tickets in total, 3 open.",
           "Proposal should open on the risk, close on concrete steps: new carrier slots, scanner stability, Dock 3 handling plan, cross-border lane pilot."]
    (EXPECTED / "cowork-customer-qbr.md").write_text("\n".join(exp) + "\n", encoding="utf-8")
    return out


# ================================================================ 5) disruption response
def build_disruption():
    out = folder("cowork-disruption-response")
    d = doc_with(T.RUNBOOK_TITLE, T.RUNBOOK_SUBTITLE, T.RUNBOOK_NOTE, T.RUNBOOK)
    save_doc(d, out / "fargo_outbound_disruption_playbook.docx")
    ships = [("S-91001", "Tailspin Toys", "Fabrikam Logistics", "06:30", "Gate 4", "08:30", "Yes", "Loaded, waiting at gate"),
             ("S-91002", "Tailspin Toys", "Northwind Freight", "06:45", "Gate 4", "09:00", "Yes", "Queued"),
             ("S-91003", "Woodgrove Bank", "Fabrikam Logistics", "07:00", "Gate 4", "09:30", "Yes", "Queued"),
             ("S-91004", "Fourth Coffee", "Northwind Freight", "07:15", "Gate 4", "11:00", "No", "Expected"),
             ("S-91005", "Lucerne Publishing", "Adatum Cargo", "07:30", "Gate 4", "12:00", "No", "Expected"),
             ("S-91006", "Wide World Importers", "Proseware Courier", "07:45", "Gate 4", "10:00", "Yes", "Expected"),
             ("S-91007", "Margie's Travel", "Northwind Freight", "08:00", "Gate 2", "12:00", "No", "Expected"),
             ("S-91008", "Tailspin Toys", "Fabrikam Logistics", "08:30", "Gate 4", "10:30", "Yes", "Expected"),
             ("S-91009", "Adventure Works", "Adatum Cargo", "08:45", "Gate 2", "13:00", "No", "Expected"),
             ("S-91010", "Woodgrove Bank", "Proseware Courier", "09:00", "Gate 4", "11:30", "Yes", "Expected"),
             ("S-91011", "Coho Vineyard", "Northwind Freight", "09:15", "Gate 2", "14:00", "No", "Expected"),
             ("S-91012", "Trey Research", "Fabrikam Logistics", "09:30", "Gate 4", "15:00", "No", "Expected")]
    wb = Workbook()
    ws = wb.active
    ws.title = "Outbound Morning Wave"
    write_sheet(ws, ["Shipment", "Customer", "Carrier", "Truck Arrival", "Gate", "Customer Cutoff", "Service Penalty", "Status"],
                ships, widths={2: 24, 3: 22, 8: 26})
    save_wb(wb, out / "affected_shipments.xlsx")
    (out / "trigger-message-en.txt").write_text(T.TRIGGER_EN + "\n", encoding="utf-8")
    (out / "trigger-message-de.txt").write_text(T.TRIGGER_DE + "\n", encoding="utf-8")
    exp = ["# Expected result: cowork-disruption-response", "",
           "Brief should name: gate scanners at Gate 4 down since 06:40, trucks queuing, no safety impact, next update 07:30.",
           "Gate 4 shipments: 9 of 12. Escalate first (cutoff within two hours of about 07:00, service penalty, Gate 4): "
           "S-91001 (08:30, Tailspin Toys) and S-91002 (09:00, Tailspin Toys). Next in line: S-91003 (09:30, Woodgrove Bank), S-91006 (10:00, Wide World Importers), S-91008 (10:30, Tailspin Toys), S-91010 (11:30, Woodgrove Bank).",
           "Owners from the playbook: Teresa Sac (incident lead), Billie Vester (IT on call), Sonia Rees (safety, trucks on site). External customers only after approval."]
    (EXPECTED / "cowork-disruption-response.md").write_text("\n".join(exp) + "\n", encoding="utf-8")
    return out


# ================================================================ 6) steering to board
def build_steering():
    out = folder("cowork-steering-to-board")
    d = Document()
    d.styles["Normal"].font.name = "Calibri"
    d.styles["Normal"].font.size = Pt(11)
    d.add_heading(T.TRANSCRIPT_TITLE, 0)
    for m in T.TRANSCRIPT_META:
        d.add_paragraph(m)
    for who, text in T.TRANSCRIPT:
        p = d.add_paragraph()
        p.add_run(who + ": ").bold = True
        p.add_run(text)
    save_doc(d, out / "fargo_steering_workshop_transcript.docx")
    exp = ["# Expected result: cowork-steering-to-board", "",
           "Decisions that hold: (1) Fabrikam Dock 2 slot Apr 27-29, Gate 4 from Apr 30; (2) Dock 3 closure stays Apr 27; "
           "(3) Northwind four-week trial from Apr 30 with floodlights and marshal, review May 28 by Sonia Rees; "
           "(4) wave 1 = A and B items, no mezzanine (replaces the earlier A items only); (5) go/no-go for Jun 22 on Jun 8, chaired by Sydney Mattos, "
           "Vance DeLeon brings the sprinkler test result; (6) scanner go-live moves from May 15 to Jun 1; "
           "(7) safety training week of May 17, all drivers by May 31 (Sonia Rees).",
           "Not decisions: Sunday shifts (scrapped), moving the closure by two weeks (rejected).",
           "Open: 4 forklift drivers for May 9 (Teresa Sac, no date); Wi-Fi survey owner and date (unassigned); temporary staff budget "
           "(120,000 USD not confirmed by finance, Sydney Mattos clarifies); who informs Tailspin Toys about Dock 3 and when; "
           "sprinkler test date (inspector not confirmed)."]
    (EXPECTED / "cowork-steering-to-board.md").write_text("\n".join(exp) + "\n", encoding="utf-8")
    return out


if __name__ == "__main__":
    build_freight()
    totals = freight_expected()
    build_contracts()
    m = build_margin()
    build_qbr()
    build_disruption()
    build_steering()
    print("freight totals:", {k: {s: round(v, 2) for s, v in t.items()} for k, t in totals.items()})
    print("margin total / uplifts:", [round(x) for x in m])
    print("Expected results written to", EXPECTED)
