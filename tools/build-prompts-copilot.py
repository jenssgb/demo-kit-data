from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor
from zipfile import ZipFile, ZIP_DEFLATED

OUT = Path(__file__).resolve().parent.parent / "prompts-copilot"
OUT.mkdir(exist_ok=True)

TITLE_EN = "Contoso Fargo prompt pack – Copilot Chat"
TITLE_DE = "Contoso Fargo Prompt Pack – Copilot Chat"

PROMPTS = [
    {
        "title_en": "Monday change brief",
        "title_de": "Montags-Update: Was hat sich geändert?",
        "where_en": "Copilot Chat > prompt box, grounded in Work",
        "where_de": "Copilot Chat > Eingabefeld, mit Arbeitsdaten",
        "en": "I’m pulling together the Fargo go-live stand-up for 8:30. What changed last week across emails, chats, meetings, and files about Dock 3, Gate 4, safety training, IT/Wi-Fi, and the June 22 go-live? Keep it to the ten things I should mention first, and show the source for anything that sounds risky.",
        "de": "Ich bereite gerade das Fargo-Go-live-Stand-up um 8:30 vor. Was hat sich letzte Woche in Mails, Chats, Meetings und Dateien zu Dock 3, Tor 4, Sicherheitsschulung, IT/WLAN und dem Go-live am 22. Juni geändert? Bitte nur die zehn Punkte, die ich zuerst nennen sollte, und mit Quelle bei allem, was riskant klingt.",
        "tip_en": "Goal + context + source + expectations; good weekly scheduled prompt.",
        "tip_de": "Ziel + Kontext + Quelle + Erwartung; gut als wöchentlicher geplanter Prompt.",
    },
    {
        "title_en": "Carrier impact check",
        "title_de": "Auswirkung auf Speditionen prüfen",
        "where_en": "Copilot Chat after selecting the Lisa Taylor email or adding it as context",
        "where_de": "Copilot Chat nach Auswahl der Lisa-Taylor-Mail oder mit der Mail als Kontext",
        "en": "I’m about to call Lisa about the carrier plan. Based on her latest email and the readiness tracker, what are the two decisions we need before Dock 3 closes on Apr 27? Please separate what we know from what still needs confirmation.",
        "de": "Ich telefoniere gleich mit Lisa zum Speditionsplan. Was sind auf Basis ihrer letzten Mail und des Readiness-Trackers die zwei Entscheidungen, die wir vor der Schließung von Dock 3 am 27. April brauchen? Bitte trenne, was sicher ist, von dem, was noch bestätigt werden muss.",
        "tip_en": "Narrow source and practical output for a real conversation.",
        "tip_de": "Klare Quelle und praktisches Ergebnis für ein echtes Gespräch.",
    },
    {
        "title_en": "Safety agenda for Sonia",
        "title_de": "Sicherheitsagenda für Sonia",
        "where_en": "Copilot Chat > New chat, then copy the answer into Outlook or Teams",
        "where_de": "Copilot Chat > Neuer Chat, Antwort danach in Outlook oder Teams übernehmen",
        "en": "I have a 20-minute sync with Sonia tomorrow about the Fargo safety work. Can you draft a simple agenda that covers shift lead training, the evacuation drill, Dock 2 PPE, and the Northwind 5:30 a.m. exception? I want it to sound like me, direct but not pushy.",
        "de": "Ich habe morgen einen 20-Minuten-Termin mit Sonia zur Sicherheit in Fargo. Kannst du mir eine einfache Agenda entwerfen, die Schichtleiterschulung, Räumungsübung, PSA an Dock 2 und die Northwind-Ausnahme um 5:30 Uhr abdeckt? Es soll nach mir klingen: direkt, aber nicht drängend.",
        "tip_en": "Adds tone and why the answer is needed.",
        "tip_de": "Gibt Ton und Anlass mit.",
    },
    {
        "title_en": "Kai construction risk summary",
        "title_de": "Baurisiko von Kai zusammenfassen",
        "where_en": "Copilot Chat with Kai’s weekly construction document attached",
        "where_de": "Copilot Chat mit Kais wöchentlichem Bau-Dokument als Anhang",
        "en": "I’m updating the go/no-go pack and need the construction risks in normal business language. From Kai’s latest construction weekly, which items could realistically threaten June 22, what is the current owner, and what would I ask for in the Thursday leadership sync?",
        "de": "Ich aktualisiere gerade das Go/No-go-Paket und brauche die Baurisiken in normaler Business-Sprache. Welche Punkte aus Kais letztem Bau-Weekly können den 22. Juni realistisch gefährden, wer ist aktuell Owner, und welche Entscheidung würde ich im Leadership-Sync am Donnerstag anfordern?",
        "tip_en": "Turns source material into a decision conversation.",
        "tip_de": "Macht aus Quellenmaterial eine Entscheidungsunterlage.",
    },
    {
        "title_en": "Tailspin customer answer",
        "title_de": "Antwort an Tailspin vorbereiten",
        "where_en": "Outlook > open Jordan Mitchell’s email > Copilot > Draft reply",
        "where_de": "Outlook > Mail von Jordan Mitchell öffnen > Copilot > Antwort entwerfen",
        "en": "Draft a reply to Jordan at Tailspin Toys. Say we can support the split delivery plan, with 600 T300 units by May 8 and the remaining 600 by May 22, and that we’ll send the quote PDF after internal approval. Keep it warm and specific; don’t mention internal construction problems unless needed.",
        "de": "Entwirf eine Antwort an Jordan von Tailspin Toys. Sag, dass wir die Teillieferung unterstützen können: 600 T300 bis 8. Mai und die restlichen 600 bis 22. Mai, und dass wir das Angebots-PDF nach interner Freigabe senden. Freundlich und konkret, aber interne Bauthemen bitte nicht unnötig erwähnen.",
        "tip_en": "Good prompt for Outlook: audience, facts, boundary.",
        "tip_de": "Guter Outlook-Prompt: Zielgruppe, Fakten, Grenze.",
    },
    {
        "title_en": "Stakeholder update in my style",
        "title_de": "Stakeholder-Update in meinem Stil",
        "where_en": "Copilot Chat after custom instructions are set",
        "where_de": "Copilot Chat, nachdem benutzerdefinierte Anweisungen gesetzt sind",
        "en": "I need a quick Fargo update for Lisa, Sonia, Kai, and Alex. Use my usual style: short lead sentence, then three sections — what changed, what needs a decision, and what I’ll do next. Pull only from the last week’s work data and keep it under 180 words.",
        "de": "Ich brauche ein kurzes Fargo-Update für Lisa, Sonia, Kai und Alex. Nutze meinen üblichen Stil: ein kurzer Einstiegssatz, dann drei Abschnitte – was hat sich geändert, welche Entscheidung brauchen wir, was mache ich als Nächstes. Bitte nur aus Arbeitsdaten der letzten Woche und unter 180 Wörtern.",
        "tip_en": "Shows custom instructions + grounded source + length limit.",
        "tip_de": "Zeigt benutzerdefinierte Anweisungen + Quelle + Längenlimit.",
    },
    {
        "title_en": "Find the blind spot",
        "title_de": "Blinden Fleck finden",
        "where_en": "Copilot Chat > use after a summary answer",
        "where_de": "Copilot Chat > nach einer Zusammenfassung verwenden",
        "en": "Before I share this, challenge it. What is missing from the Fargo picture that an operations lead would normally check: labor, safety, carrier capacity, IT, customer impact, or approvals? If there is no evidence, say that clearly instead of guessing.",
        "de": "Bevor ich das teile: Hinterfrage es bitte. Was fehlt im Fargo-Bild, das eine Operations-Leitung normalerweise prüfen würde: Personal, Sicherheit, Speditionskapazität, IT, Kundenauswirkung oder Freigaben? Wenn es keinen Beleg gibt, sag das bitte klar statt zu raten.",
        "tip_en": "A safe follow-up prompt: asks Copilot to identify missing evidence.",
        "tip_de": "Sicherer Follow-up-Prompt: Copilot soll fehlende Belege benennen.",
    },
    {
        "title_en": "Meeting prep for go/no-go",
        "title_de": "Vorbereitung Go/No-go-Termin",
        "where_en": "Copilot Chat > New chat, reference the June 15 go/no-go meeting",
        "where_de": "Copilot Chat > Neuer Chat, den Go/No-go-Termin am 15. Juni referenzieren",
        "en": "I have the Fargo go/no-go meeting on June 15. Please prep me with the decision points, likely objections from Logistics, Safety, Construction, and IT, and three questions I should ask first. Use the meeting invite, related emails, and the readiness tracker as sources.",
        "de": "Ich habe am 15. Juni den Fargo-Go/No-go-Termin. Bereite mich bitte vor: Entscheidungspunkte, wahrscheinliche Einwände von Logistik, Safety, Bau und IT, und drei Fragen, die ich zuerst stellen sollte. Nutze die Termineinladung, zugehörige Mails und den Readiness-Tracker als Quellen.",
        "tip_en": "Great for Work IQ: person/time context plus explicit sources.",
        "tip_de": "Gut für Work IQ: Personen-/Terminkontext plus explizite Quellen.",
    },
    {
        "title_en": "Turn answer into a Copilot Page",
        "title_de": "Antwort in eine Copilot Page überführen",
        "where_en": "Copilot Chat response > Edit in Pages / open as Page",
        "where_de": "Copilot-Chat-Antwort > In Pages bearbeiten / als Page öffnen",
        "en": "Turn this into a working page for the Fargo team. Keep the decision log at the top, then risks, open questions, and next actions. Make it easy for Lisa, Sonia, Kai, and Alex to add comments without rewriting the whole thing.",
        "de": "Mach daraus bitte eine Arbeitsseite für das Fargo-Team. Oben soll das Entscheidungslog stehen, danach Risiken, offene Fragen und nächste Aktionen. Lisa, Sonia, Kai und Alex sollen leicht Kommentare ergänzen können, ohne alles umzuschreiben.",
        "tip_en": "Use when moving from private chat to a shared artifact.",
        "tip_de": "Nützlich, wenn aus privatem Chat ein gemeinsames Arbeitsartefakt wird.",
    },
    {
        "title_en": "Prompt rewrite coach",
        "title_de": "Prompt verbessern lassen",
        "where_en": "Copilot Chat > paste a weak prompt and ask for improvement",
        "where_de": "Copilot Chat > schwachen Prompt einfügen und verbessern lassen",
        "en": "I want to teach the team why this prompt is too vague: ‘Summarize Fargo.’ Rewrite it into a better prompt using goal, context, source, and expectations. Keep the rewrite natural, like something an operations lead would really type on a Monday morning.",
        "de": "Ich möchte dem Team zeigen, warum dieser Prompt zu vage ist: ‚Fass Fargo zusammen.‘ Schreib ihn mit Ziel, Kontext, Quelle und Erwartung besser um. Bitte natürlich formuliert, so wie es eine Operations-Leitung an einem Montagmorgen wirklich eintippen würde.",
        "tip_en": "Use as a live coaching moment for GCSE.",
        "tip_de": "Gut als Live-Coaching-Moment für GCSE.",
    },
]

def add_para(doc, text, style=None, bold=False, color=None):
    p = doc.add_paragraph(style=style) if style else doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    return p

def build_docx():
    doc = Document()
    styles = doc.styles
    styles['Normal'].font.name = 'Aptos'
    styles['Normal'].font.size = Pt(10.5)
    for st in ['Heading 1','Heading 2','Heading 3']:
        styles[st].font.name = 'Aptos Display'
    add_para(doc, TITLE_EN, 'Title', color='0F6CBD')
    add_para(doc, TITLE_DE, None, bold=True)
    add_para(doc, 'Scenario: Contoso operations lead prepares the Fargo distribution center go-live. Dock 3 closes Apr 27; go/no-go is Jun 15; go-live is Jun 22. Use these prompts as reusable Copilot Chat examples, scheduled prompts, or team gallery entries.')
    add_para(doc, 'Szenario: Die Operations-Leitung von Contoso bereitet das Go-live des Distribution Centers Fargo vor. Dock 3 schließt am 27. April; Go/No-go ist am 15. Juni; Go-live ist am 22. Juni. Diese Prompts eignen sich für Copilot Chat, geplante Prompts oder die Team-Galerie.')
    add_para(doc, 'Prompt formula / Prompt-Formel: goal + context + source + expectations', None, bold=True, color='107C10')
    for i,p in enumerate(PROMPTS,1):
        doc.add_heading(f"{i}. {p['title_en']} / {p['title_de']}", level=2)
        add_para(doc, f"Where / Wo: {p['where_en']} | {p['where_de']}", bold=True)
        add_para(doc, 'EN', bold=True, color='0F6CBD')
        add_para(doc, p['en'])
        add_para(doc, 'DE', bold=True, color='0F6CBD')
        add_para(doc, p['de'])
        add_para(doc, f"Coach note / Hinweis: {p['tip_en']} | {p['tip_de']}")
    doc.add_heading('Champion handout notes / Hinweise für Champions', level=2)
    for line in [
        'Do not save prompts that contain confidential customer names unless your organization has approved that practice.',
        'When sharing prompts with a Microsoft Teams team, check that referenced files and meetings are accessible to the same audience.',
        'For scheduled prompts, start with weekly cadence and turn on email notification only for results that need action.',
        'Prompts should sound like real work, not training examples: include why you need the result and what you will do with it.',
    ]:
        doc.add_paragraph(line, style='List Bullet')
    path = OUT / 'Contoso_Fargo_Prompt_Pack.docx'
    doc.save(path)
    patch_docx_for_validation(path)



def patch_docx_for_validation(path):
    fixed = path.with_name(path.stem + '_fixed.docx')
    font = '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n<w:fonts xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" xmlns:w14="http://schemas.microsoft.com/office/word/2010/wordml" mc:Ignorable="w14">\n  <w:font w:name="Aptos"><w:charset w:val="00"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>\n  <w:font w:name="Aptos Display"><w:charset w:val="00"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>\n  <w:font w:name="Calibri"><w:charset w:val="00"/><w:family w:val="swiss"/><w:pitch w:val="variable"/></w:font>\n</w:fonts>'.encode('utf-8')
    with ZipFile(path, 'r') as zin, ZipFile(fixed, 'w', ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == 'word/settings.xml':
                text = data.decode('utf-8')
                text = text.replace('<w:zoom w:val="bestFit"/>', '<w:zoom w:val="bestFit" w:percent="100"/>')
                data = text.encode('utf-8')
            elif item.filename == 'word/fontTable.xml':
                data = font
            zout.writestr(item, data)
    fixed.replace(path)

def build_md():
    lines=[f"# {TITLE_EN}", "", f"**{TITLE_DE}**", "", "Scenario: Contoso operations lead prepares the Fargo distribution center go-live. Dock 3 closes Apr 27; go/no-go is Jun 15; go-live is Jun 22.", "", "Prompt formula: **goal + context + source + expectations**", ""]
    for i,p in enumerate(PROMPTS,1):
        lines += [f"## {i}. {p['title_en']} / {p['title_de']}", "", f"**Where / Wo:** {p['where_en']} | {p['where_de']}", "", "**EN**", "", p['en'], "", "**DE**", "", p['de'], "", f"_Coach note / Hinweis:_ {p['tip_en']} | {p['tip_de']}", ""]
    lines += ["## Champion notes", "", "- Do not save prompts that contain confidential customer names unless your organization has approved that practice.", "- When sharing prompts with a Microsoft Teams team, check that referenced files and meetings are accessible to the same audience.", "- For scheduled prompts, start with weekly cadence and turn on email notification only for results that need action.", "- Prompts should sound like real work, not training examples: include why you need the result and what you will do with it.", ""]
    (OUT / 'Contoso_Fargo_Prompt_Pack.md').write_text('\n'.join(lines), encoding='utf-8')

def build_manifest():
    import json
    manifest = {
        "title": {"en": "Copilot Chat prompt pack – Contoso Fargo", "de": "Copilot Chat Prompt Pack – Contoso Fargo"},
        "next": {
            "en": [
                "Upload or keep Contoso_Fargo_Prompt_Pack.docx in OneDrive so champions can open it during the prompt gallery demo.",
                "Use the Markdown file as the source for organizational prompts or for a Teams post to the AI champions team.",
                "Before the live demo, create or identify a Microsoft Teams team named Contoso Fargo Go-Live Champions for prompt sharing.",
                "Prepare a Monday 7:30 scheduled prompt in Copilot Chat if your tenant UI is available; otherwise demonstrate the scheduling dialog without saving."
            ],
            "de": [
                "Contoso_Fargo_Prompt_Pack.docx in OneDrive ablegen, damit Champions es während der Prompt-Gallery-Demo öffnen können.",
                "Die Markdown-Datei als Quelle für Organisationsprompts oder für einen Teams-Beitrag an das AI-Champions-Team nutzen.",
                "Vor der Live-Demo ein Microsoft-Team Contoso Fargo Go-Live Champions anlegen oder auswählen, um Prompt-Sharing zu zeigen.",
                "Wenn die Tenant-UI verfügbar ist, den Montag-7:30-Prompt in Copilot Chat vorbereiten; sonst nur den Planungsdialog zeigen und nicht speichern."
            ]
        }
    }
    (OUT / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2)+"\n", encoding='utf-8')

if __name__ == '__main__':
    build_docx(); build_md(); build_manifest(); print(f'Wrote {OUT}')
