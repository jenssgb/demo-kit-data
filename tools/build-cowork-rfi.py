"""Build only cowork-rfi-rfp, then run Update-Manifest.ps1 -All.

Use --prepare to rebuild four masters/KB/preparation. Recalculate masters
with tools\\Recalculate-Rfi.ps1 (Excel) or LibreOffice, then --copies to
clone cached formulas into runs.
Neither mode touches existing packages or a tenant.
"""
import argparse
import importlib.util
import json
import re
import sys
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

from cowork_rfi_text import COMPACT_IDS, KB, TOPICS, VERSION, bi, questions

ROOT = Path(__file__).resolve().parent.parent
PACKAGE = ROOT / "cowork-rfi-rfp"
spec = importlib.util.spec_from_file_location("cowork_base", ROOT / "tools" / "build-cowork-usecases.py")
BASE = importlib.util.module_from_spec(spec)
sys.path.insert(0, str(ROOT / "tools"))
try:
    spec.loader.exec_module(BASE)
finally:
    sys.path.pop(0)
BASE.FIXED = datetime(2026, 10, 13, 8, 0, 0)
NS = {"s": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
HEADERS = [
    bi("ID", "ID"), bi("Parent", "Hauptfrage"), bi("Priority", "Priorität"),
    bi("Requirement", "Anforderung"), bi("Format", "Format"),
    bi("Max characters", "Max. Zeichen"), bi("Applies when", "Gilt wenn"),
    bi("Evidence required", "Nachweis nötig"), bi("Answer", "Antwort"),
    bi("Status", "Status"), bi("Source ID", "Quellen-ID"), bi("Version", "Version"),
    bi("Section", "Abschnitt"), bi("Cloud link", "Cloud-Link"),
    bi("Open point", "Offener Punkt"), bi("Review", "Prüfung"),
    bi("Characters", "Zeichen"), bi("Limit check", "Limitprüfung"),
]
PRICES = [
    ("CMS-48P-POE", 24, 1, "One-time"),
    ("CAP-6E", 48, 1, "One-time"),
    ("CSFP-10G", 48, 1, "One-time"),
    ("CGW-500", 12, 1, "One-time"),
    ("CARE-NET-3Y-NBD", 24, 1, "One-time"),
    ("SVC-NET-ASSESS", 1, 1, "One-time"),
    ("SVC-ONSITE-DAY", 24, 1, "One-time"),
    ("KB-managed-service", 12, 5, "Monthly"),
    ("Care renewal years 4-5", 24, 2, "Annual"),
    ("Freight", 1, 1, "One-time"),
]


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text, encoding="utf-8", newline="\n")


def styled(ws, headers, rows, widths):
    BASE.write_sheet(ws, headers, rows, widths=widths)
    for cell in ws[1]:
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    ws.row_dimensions[1].height = 32
    ws.sheet_view.zoomScale = 85
    for row in ws.iter_rows(min_row=2):
        ws.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.font = Font(name="Arial", size=10)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.auto_filter.ref = ws.dimensions
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A3
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.print_title_rows = "1:1"


def workbook(lang, variant):
    qs = [q for q in questions() if variant == "full" or q["id"] in COMPACT_IDS]
    wb = Workbook()
    run = wb.active
    run.title = "Run"
    fields = [
        (bi("Run ID", "Lauf-ID")[lang], "MASTER"),
        (bi("Language", "Sprache")[lang], lang),
        (bi("Variant", "Variante")[lang], variant),
        (bi("Requirement count", "Anforderungszahl")[lang],
         "=" + "+".join(f"COUNTA('{t}'!A2:A12)" for t in TOPICS)),
        (bi("Output folder", "Ergebnisordner")[lang], "MASTER - DO NOT EDIT"),
        (bi("Contract version", "Vertragsversion")[lang], VERSION),
        (bi("State", "Zustand")[lang], "NEW"),
        (bi("Customer to verify live", "Live zu prüfender Kunde")[lang], "Woodgrove Bank"),
        (bi("Warning", "Hinweis")[lang], bi(
            "FICTIONAL DEMO. Only use an untouched run; never edit this master. No Save quote, sends or orders.",
            "FIKTIVE DEMO. Nur frischen Lauf verwenden; Master nie bearbeiten. Kein Save quote, Versand oder Auftrag.")[lang]),
        (bi("Editing contract", "Bearbeitungsvertrag")[lang], bi(
            "Edit I:P on question sheets, designated pricing inputs and Run B8 only. Keep IDs, questions, formats and formulas.",
            "Nur I:P der Fragenblätter, vorgesehene Preiseingaben und Run B8 bearbeiten. IDs, Fragen, Formate und Formeln erhalten.")[lang]),
        (bi("Completion", "Abschluss")[lang], bi(
            "Every question needs an evidenced answer or explained OPEN/BLOCKED/GAP/REVIEW. N/A requires a proven condition. DRAFT only.",
            "Jede Frage braucht belegte Antwort oder erklärtes OPEN/BLOCKED/GAP/REVIEW. N/A benötigt belegte Bedingung. Nur DRAFT.")[lang]),
    ]
    styled(run, [bi("Field", "Feld")[lang], bi("Value", "Wert")[lang]], fields, {1: 34, 2: 115})
    for topic in TOPICS:
        ws = wb.create_sheet(topic)
        selected = [q for q in qs if q["topic"] == topic]
        rows = []
        for r, q in enumerate(selected, 2):
            rows.append([
                q["id"], q["parent"], q["priority"], q["text"][lang],
                q["format"][lang], q["max_chars"], q["condition"][lang],
                q["evidence"][lang], *([None] * 8), f"=LEN(I{r})",
                f'=IF(I{r}="","",IF(Q{r}<=F{r},"OK","OVER"))',
            ])
        styled(ws, [h[lang] for h in HEADERS], rows,
               {1: 17, 2: 17, 3: 12, 4: 80, 5: 28, 6: 12, 7: 38, 8: 44,
                9: 80, 10: 15, 11: 19, 12: 18, 13: 18, 14: 45, 15: 45, 16: 32, 17: 12, 18: 15})
        ws.freeze_panes = "D2"
        ws.column_dimensions.group("E", "H", hidden=True)
        table = Table(displayName=f"Questions_{topic}", ref=f"A1:R{ws.max_row}")
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)
        ws.auto_filter.ref = None
        statuses = DataValidation(type="list", formula1='"ANSWERED,PARTIAL,GAP,OPEN,BLOCKED,REVIEW,N/A"',
                                  allow_blank=True)
        statuses.errorTitle = "Invalid status"
        statuses.error = "Use a listed status and explain unresolved requirements."
        statuses.showErrorMessage = True
        ws.add_data_validation(statuses)
        statuses.add(f"J2:J{ws.max_row}")
        length = DataValidation(type="custom", formula1="LEN(I2)<=F2", allow_blank=True)
        length.showErrorMessage = True
        length.errorTitle = "Answer too long"
        length.error = "Rephrase; never truncate a mandatory fact."
        ws.add_data_validation(length)
        length.add(f"I2:I{ws.max_row}")
        for row in ws.iter_rows(min_row=2):
            ws.row_dimensions[row[0].row].height = 84
            for cell in row[8:16]:
                cell.fill = PatternFill("solid", fgColor="EAF3FF")
                cell.font = Font(name="Arial", size=10, color="0000FF")
    assumptions = wb.create_sheet("Assumptions")
    styled(assumptions, [bi("Assumption", "Annahme")[lang], bi("Value", "Wert")[lang], bi("Source", "Quelle")[lang]], [
        (bi("Years", "Jahre")[lang], 5, "Buyer RFI / COMM-01"),
        (bi("Months per year", "Monate je Jahr")[lang], 12, "KB-COMM v2 / section 2"),
        (bi("Required reporting currency", "Geforderte Berichtswährung")[lang], "USD", "Buyer RFI; verify live account currency. Do not invent FX."),
        (bi("Sites", "Standorte")[lang], 12, "Buyer RFI / IMPL-01"),
        (bi("Monthly rate per site (USD)", "Monatssatz je Standort (USD)")[lang], 240, "KB-COMM v2 / section 2"),
        (bi("Currency rule", "Währungsregel")[lang], bi("No FX rate approved. Any mismatch stays open.", "Kein Wechselkurs freigegeben. Abweichung bleibt offen.")[lang], "KB-COMM v2 / section 3"),
    ], {1: 42, 2: 70, 3: 80})
    for cell in ("B2", "B3", "B5", "B6"):
        assumptions[cell].font = Font(name="Arial", size=10, color="0000FF")
    prices = wb.create_sheet("Pricing")
    headers = bi(
        ["SKU / cost", "Qty", "Unit net", "Currency", "Years / events", "Basis",
         "Net (currency in D)", "Source / timestamp", "Tax rate", "Tax", "Gross", "Review"],
        ["SKU / Kosten", "Menge", "Einzelnetto", "Währung", "Jahre / Ereignisse", "Basis",
         "Netto (Währung in D)", "Quelle / Zeitpunkt", "Steuersatz", "Steuer", "Brutto", "Prüfung"])[lang]
    rows = []
    for r, (sku, qty, periods, basis) in enumerate(PRICES, 2):
        monthly = sku == "KB-managed-service"
        rows.append([
            sku, "=Assumptions!B5" if monthly else qty,
            "=Assumptions!B6" if monthly else None, "USD" if monthly else None,
            "=Assumptions!B2" if monthly else periods, basis,
            f'=IF(COUNT(B{r},C{r},E{r})=3,ROUND(B{r}*ROUND(C{r},2)*E{r}*IF(F{r}="Monthly",Assumptions!$B$3,1),2),"")',
            "KB-COMM v2 / section 2; tax unconfirmed" if monthly else None, None,
            f'=IF(COUNT(G{r},I{r})=2,ROUND(G{r}*I{r},2),"")',
            f'=IF(COUNT(G{r},J{r})=2,G{r}+J{r},"")', None,
        ])
    styled(prices, headers, rows, {1: 32, 2: 12, 3: 18, 4: 14, 5: 18, 6: 18, 7: 23, 8: 65, 9: 14, 10: 18, 11: 18, 12: 70})
    for row in prices.iter_rows(min_row=2):
        for cell in row:
            if cell.data_type == "f":
                cell.font = Font(name="Arial", size=10, color="008000" if "Assumptions!" in cell.value else "000000")
            else:
                cell.font = Font(name="Arial", size=10, color="0000FF")
        for col in (3, 4, 8, 9, 12):
            row[col - 1].fill = PatternFill("solid", fgColor="FFF4CE")
        for col in (3, 7, 10, 11):
            row[col - 1].number_format = '#,##0.00;(#,##0.00);"-"'
        row[8].number_format = '0.0%;(0.0%);"-"'
    summary = wb.create_sheet("TCO")
    last = len(PRICES) + 1
    currency_match = f'COUNTIF(Pricing!D2:D{last},Assumptions!B4)={len(PRICES)}'
    styled(summary, bi(["Measure", "Value", "Interpretation"], ["Kennzahl", "Wert", "Einordnung"])[lang], [
        (bi("Reporting currency", "Berichtswährung")[lang], "=Assumptions!B4", bi("USD is a buyer requirement, not a claim about live Atlas.", "USD ist Käuferanforderung, keine Aussage über Live-Atlas.")[lang]),
        (bi("Known net subtotal", "Bekannte Nettoteilsumme")[lang], f'=SUMIF(Pricing!D2:D{last},Assumptions!B4,Pricing!G2:G{last})', bi("PARTIAL: only populated amounts in reporting currency; never call this complete TCO.", "PARTIAL: Nur gefüllte Beträge in Berichtswährung; keine vollständigen TCO.")[lang]),
        (bi("Complete required net TCO", "Vollständige Netto-TCO")[lang], f'=IF(AND(COUNT(Pricing!G2:G{last})={len(PRICES)},{currency_match}),SUM(Pricing!G2:G{last}),"")', bi("Blank means missing price/period/currency; renewals and freight are mandatory.", "Leer heißt fehlender Preis/Zeitraum/Währung; Verlängerung und Fracht sind Pflicht.")[lang]),
        (bi("Complete gross TCO", "Vollständige Brutto-TCO")[lang], f'=IF(AND(COUNT(Pricing!K2:K{last})={len(PRICES)},{currency_match}),SUM(Pricing!K2:K{last}),"")', bi("Blank means incomplete taxes or other required amounts.", "Leer heißt unvollständige Steuer oder andere Pflichtbeträge.")[lang]),
        (bi("Net coverage", "Netto-Abdeckung")[lang], f'=IF(AND(COUNT(Pricing!G2:G{last})={len(PRICES)},{currency_match}),"COMPLETE","PARTIAL")', bi("Keep missing costs OPEN; never replace missing values with zero.", "Fehlende Kosten OPEN lassen; fehlende Werte nicht durch Null ersetzen.")[lang]),
        (bi("Gross coverage", "Brutto-Abdeckung")[lang], f'=IF(AND(COUNT(Pricing!K2:K{last})={len(PRICES)},{currency_match}),"COMPLETE","PARTIAL")', bi("This model is a draft, not approval.", "Dieses Modell ist ein Entwurf, keine Freigabe.")[lang]),
    ], {1: 36, 2: 25, 3: 110})
    for r in (3, 4, 5):
        summary[f"B{r}"].number_format = '#,##0.00;(#,##0.00);"-"'
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    path = PACKAGE / "templates" / lang / f"RFI_{variant.title()}_Master.xlsx"
    path.parent.mkdir(parents=True, exist_ok=True)
    BASE.save_wb(wb, path)


def knowledge(lang):
    folder = PACKAGE / "knowledge" / lang
    folder.mkdir(parents=True, exist_ok=True)
    index = [
        bi("# Contoso Source Index", "# Contoso-Quellenindex")[lang], "",
        f"Version: {VERSION}. " + bi("Fictional demo; not evidence of actual vendor capabilities.", "Fiktive Demo; kein Beleg tatsächlicher Herstellerfähigkeiten.")[lang],
        bi("Scenario evidence date: 2026-10-13. Approved v2 effective 2026-10-01. No real certificates supplied.", "Nachweisstand des Szenarios: 13.10.2026. Freigegebenes v2 gilt ab 01.10.2026. Keine echten Zertifikate enthalten.")[lang],
        bi("Sources are only the PDFs in this language folder. Resolve their actual OneDrive links.", "Quellen sind nur die PDFs in diesem Sprachordner. Tatsächliche OneDrive-Links auflösen.")[lang],
        bi("| Source ID | File | Version | Status | Sections |", "| Quellen-ID | Datei | Version | Status | Abschnitte |")[lang],
        "|---|---|---|---|---|",
    ]
    for topic, paras in KB.items():
        file = f"KB_{topic}_v2.pdf"
        lines = [
            ("h1", TOPICS[topic][lang]),
            ("body", bi("CONTOSO FICTIONAL DEMO - approved content for this scenario only.",
                        "FIKTIVE CONTOSO-DEMO - nur für dieses Szenario freigegebener Inhalt.")[lang]),
            ("body", f"Source-ID: KB-{topic} | Version: v2 | APPROVED | Effective: 2026-10-01 | Scenario: 2026-10-13"),
        ]
        for n, text in enumerate(paras, 1):
            lines.extend([("h2", f"Section {n}"), ("body", text[lang])])
        BASE.pdf_lines(folder / file, lines, TOPICS[topic][lang])
        index.append(f"| KB-{topic} | {file} | v2 | APPROVED | 1-{len(paras)} |")
    BASE.pdf_lines(folder / "KB_OPS_v1_RETIRED.pdf", [
        ("h1", bi("Retired SLA", "Zurückgezogenes SLA")[lang]),
        ("body", "Source-ID: KB-OPS | Version: v1 | RETIRED 2026-10-01 | Superseded by KB_OPS_v2.pdf"),
        ("h2", "Section 1"),
        ("body", bi("OLD: 99.99% 24/7; P1 response 15 minutes. Withdrawn, not a current promise.",
                    "ALT: 99,99 % rund um die Uhr; P1-Reaktion 15 Minuten. Zurückgezogen, keine aktuelle Zusage.")[lang]),
    ], "Contoso retired SLA v1")
    index.append("| KB-OPS | KB_OPS_v1_RETIRED.pdf | v1 | RETIRED | 1; never commit |")
    write(folder / "Source_Index.md", "\n".join(index) + "\n")


def copies():
    ET.register_namespace("", NS["s"])
    ET.register_namespace("cp", "http://schemas.openxmlformats.org/package/2006/metadata/core-properties")
    ET.register_namespace("dc", "http://purl.org/dc/elements/1.1/")
    ET.register_namespace("dcterms", "http://purl.org/dc/terms/")
    ET.register_namespace("xsi", "http://www.w3.org/2001/XMLSchema-instance")
    for lang in ("en", "de"):
        for variant in ("full", "compact"):
            master = PACKAGE / "templates" / lang / f"RFI_{variant.title()}_Master.xlsx"
            BASE.freeze(master)
            with zipfile.ZipFile(master) as src:
                parts = {name: src.read(name) for name in src.namelist()}
            # Excel's session ID, local path and calculation order are not demo data.
            parts["xl/workbook.xml"] = re.sub(
                rb"<xr:revisionPtr\b[^>]*/>", b"", parts["xl/workbook.xml"])
            parts["xl/workbook.xml"] = re.sub(
                rb"<mc:AlternateContent>.*?</mc:AlternateContent>",
                lambda m: b"" if b"absPath" in m[0] else m[0],
                parts["xl/workbook.xml"], flags=re.DOTALL)
            parts.pop("xl/calcChain.xml", None)
            parts["[Content_Types].xml"] = re.sub(
                rb'<Override\b(?=[^>]*PartName="/xl/calcChain.xml")[^>]*/>',
                b"", parts["[Content_Types].xml"])
            parts["xl/_rels/workbook.xml.rels"] = re.sub(
                rb'<Relationship\b(?=[^>]*Type="[^"]*/calcChain")[^>]*/>',
                b"", parts["xl/_rels/workbook.xml.rels"])
            core = ET.fromstring(parts["docProps/core.xml"])
            for tag in ("{http://purl.org/dc/elements/1.1/}creator",
                        "{http://schemas.openxmlformats.org/package/2006/metadata/core-properties}lastModifiedBy"):
                node = core.find(tag)
                if node is not None:
                    node.text = "Contoso Demo"
            parts["docProps/core.xml"] = ET.tostring(core, encoding="utf-8", xml_declaration=True)
            with zipfile.ZipFile(master, "w", zipfile.ZIP_DEFLATED) as dst:
                for entry, data in parts.items():
                    info = zipfile.ZipInfo(entry, BASE.FIXED.timetuple()[:6])
                    info.compress_type = zipfile.ZIP_DEFLATED
                    dst.writestr(info, data)
            from openpyxl import load_workbook
            cached = load_workbook(master, data_only=True)
            try:
                if cached["TCO"]["B3"].value != 172800 or cached["TCO"]["B6"].value != "PARTIAL":
                    raise ValueError(f"Recalculate {master} before generating copies")
            finally:
                cached.close()
            # Locate Run by workbook relationship, not by assuming sheet1 forever.
            workbook_xml = ET.fromstring(parts["xl/workbook.xml"])
            sheet = next(s for s in workbook_xml.find("s:sheets", NS) if s.attrib["name"] == "Run")
            rid = sheet.attrib["{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"]
            rels = ET.fromstring(parts["xl/_rels/workbook.xml.rels"])
            target = next(r.attrib["Target"] for r in rels if r.attrib["Id"] == rid)
            name = target.lstrip("/") if target.startswith("/") else "xl/" + target
            for n in range(1, 9):
                run_id = f"{lang.upper()}-{variant.upper()}-{n:02}"
                relative = f"runs/{lang}/{variant}/{n:02}"
                out = PACKAGE / "runs" / lang / variant / f"{n:02}"
                out.mkdir(parents=True, exist_ok=True)
                # Preserve Excel's namespace declarations referenced by mc:Ignorable.
                xml = parts[name].decode("utf-8")
                for address, text in (
                    ("B2", run_id),
                    ("B6", f"Demo-Kit/Cowork/cowork-rfi-rfp/{relative}/results"),
                ):
                    pattern = rf'(<c\b[^>]*\br="{address}"[^>]*>).*?(</c>)'
                    def replace(match):
                        opening = re.sub(r'\s+t="[^"]*"', "", match[1])
                        return opening[:-1] + ' t="inlineStr"><is><t>' + escape(text) + "</t></is>" + match[2]
                    xml, changed = re.subn(pattern, replace, xml, flags=re.DOTALL)
                    if changed != 1:
                        raise ValueError(f"{master}: expected exactly one {address} cell")
                path = out / f"RFI_{variant.title()}_Run_{n:02}.xlsx"
                with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
                    for entry, data in parts.items():
                        info = zipfile.ZipInfo(entry, BASE.FIXED.timetuple()[:6])
                        info.compress_type = zipfile.ZIP_DEFLATED
                        dst.writestr(info, xml.encode("utf-8") if entry == name else data)
                write(out / "results" / "README.md", bi(
                    f"# Results for {run_id}\nPlace this run's DRAFT Customer_Response_{run_id}.docx and Internal_Review_{run_id}.docx here. This is not business evidence. Actual placement must be verified; names alone do not prove placement. Do not overwrite another run.",
                    f"# Ergebnisse für {run_id}\nDRAFT Customer_Response_{run_id}.docx und Internal_Review_{run_id}.docx hier ablegen. Dies ist kein Geschäftsbeleg. Tatsächliche Ablage prüfen; Dateinamen allein beweisen sie nicht. Keinen anderen Lauf überschreiben.")[lang] + "\n")


def preparation():
    for lang in ("en", "de"):
        text = bi(
            """# Contoso RFI preparation
Files are in OneDrive > Demo-Kit > Cowork > cowork-rfi-rfp.
1. Use the fixed global install.ps1 -Tenant command once, ideally the day before. Do not rerun it as a rehearsal reset: changed cloud workbooks can be overwritten.
2. In Edge with the demo account, enable Cowork browser access and sign in to Contoso Atlas before presenting. Verify the actual Woodgrove account/currency, visible product records and unsaved New quote > Live preview > Check availability. Never Save quote.
3. Upload skills/contoso-rfi-rfp-response/SKILL.md through Customize > Skills > arrow next to Add > Upload skill. Update an existing skill rather than creating a duplicate. Select it in Sources in the new conversation.
4. Choose runs/en/compact/01/RFI_Compact_Run_01.xlsx for the compact case or runs/en/full/01/RFI_Full_Run_01.xlsx for the large case. Read Run: NEW; all answer cells must be empty. Use 02-08 on subsequent runs. Never edit templates or reuse an answered run.
5. Add attachments > Attach cloud files and folders: reference just the chosen existing cloud workbook, not a local upload. The prompt names knowledge/en; Cowork should resolve permitted cloud PDFs there. If search is missing/ambiguous, select only that specific existing cloud source. Other runs, masters, results and skills are not answer evidence.
6. Rehearse actual same-file edits and both Word files in the Run B6 results folder. If Cowork cannot put them there, explicitly report placement incomplete; do not claim success from filenames. Prices/stock are live, so the expected reference contains no fixed commercial totals.
7. After eight runs, copy the unchanged matching master into a new run folder, set unique Run B2/B6 and prepare a results folder. No deletion/reset or global installer rerun.
The RFI is fictional and non-binding. The six difficult cases are retired SLA v1 vs v2, missing ISO/SOC2/signed DPA, 10G and CAP-6E alternatives with live supply checks, missing year-4/5 care rates, EU-only/unlimited-liability exceptions, and conditional/limited fields.
The live VM preflight must prove account access, browser controls, source indexing, file permissions, all rows, output location and two isolated runs. Package validation is not a live Cowork guarantee.
""",
            """# Vorbereitung Contoso-RFI
Dateien liegen in OneDrive > Demo-Kit > Cowork > cowork-rfi-rfp.
1. Den festen globalen install.ps1-Befehl mit -Tenant einmal ausführen, möglichst am Vortag. Nicht als Proben-Reset wiederholen: Bearbeitete Cloud-Dateien können überschrieben werden.
2. Edge mit Demo-Konto nutzen, Cowork-Browserzugriff erlauben und Contoso Atlas vorher anmelden. Tatsächliches Woodgrove-Konto/Währung, Produkte und ungespeichertes New quote > Live preview > Check availability prüfen. Niemals Save quote.
3. skills/contoso-rfi-rfp-response/SKILL.md über Customize > Skills > Pfeil neben Add > Upload skill hochladen. Bestehenden Skill aktualisieren statt duplizieren. Im neuen Verlauf unter Sources auswählen.
4. Für kompakt runs/de/compact/01/RFI_Compact_Run_01.xlsx, für groß runs/de/full/01/RFI_Full_Run_01.xlsx wählen. Run muss NEW zeigen; Antwortzellen müssen leer sein. Danach 02-08 nutzen. Nie templates bearbeiten oder beantworteten Lauf wiederverwenden.
5. Add attachments > Attach cloud files and folders: Nur die gewählte bestehende Cloud-Excel referenzieren, nicht lokal hochladen. Der Prompt nennt knowledge/de; Cowork soll die berechtigten Cloud-PDFs dort finden. Bei fehlender/mehrdeutiger Suche nur die konkret fehlende bestehende Cloud-Quelle auswählen. Andere Läufe, Master, Ergebnisse und Skills sind keine Antwortbelege.
6. Echte Änderung derselben Datei und beide Word-Dateien im Ergebnisordner aus Run B6 proben. Kann Cowork sie dort nicht ablegen, Ablage ausdrücklich als unvollständig melden; Dateinamen sind kein Erfolgsnachweis. Preise/Bestände sind live; erwartete Ergebnisse enthalten keine fixen kaufmännischen Summen.
7. Nach acht Läufen passenden unveränderten Master in neuen Laufordner kopieren, eindeutige Run B2/B6 setzen und Ergebnisordner vorbereiten. Kein Löschen/Reset und kein globaler Installer-Rerun.
Die RFI ist fiktiv und unverbindlich. Sechs schwierige Fälle: Zurückgezogenes SLA v1 gegen v2, fehlendes ISO/SOC2/unterzeichnete AVV, 10G-/CAP-6E-Alternativen mit Live-Lieferprüfung, fehlende Betreuungspreise Jahre 4/5, EU-only/unbegrenzte Haftung sowie bedingte/zeichenbegrenzte Felder.
VM-Vorprüfung muss Kontozugriff, Browser-Schaltflächen, Indexierung, Dateirechte, alle Fragen, Ablage und zwei getrennte Läufe belegen. Paketprüfung garantiert keinen erfolgreichen Cowork-Lauf.
""")[lang]
        write(PACKAGE / f"Preparation-{lang}.md", text)
    tenant = {
        "links": [
            {"name": bi("Contoso Atlas - RFI read-only preview", "Contoso Atlas - RFI nur Vorschau"),
             "url": "https://contoso-atlas.swedencentral.cloudapp.azure.com/quotes"},
        ],
        "manual": {lang: [
            bi("RFI: read Preparation-en.md, activate contoso-rfi-rfp-response, select a fresh compact/full cloud workbook and verify live Atlas, KB search and result-folder placement. Never Save quote. Do not rerun installer as a reset.",
               "RFI: Preparation-de.md lesen, contoso-rfi-rfp-response aktivieren, frische kompakte/große Cloud-Excel wählen und Atlas, KB-Suche sowie Ergebnisablage live prüfen. Nie Save quote. Installer nicht als Reset wiederholen.")[lang]
        ] for lang in ("en", "de")},
    }
    write(PACKAGE / "tenant.json", json.dumps(tenant, ensure_ascii=False, indent=2) + "\n")
    manifest = PACKAGE / "manifest.json"
    if not manifest.exists():
        write(manifest, json.dumps({"title": bi("Complex Contoso RFI", "Komplexe Contoso-RFI"),
                                   "next": {lang: [f"Read Preparation-{lang}.md"] for lang in ("en", "de")},
                                   "files": []}, ensure_ascii=False, indent=2) + "\n")
    expected = {
        "version": VERSION, "full_count": len(questions()), "compact_ids": sorted(COMPACT_IDS),
        "questions": questions(), "live_values": "Never freeze Atlas stock, account terms, approval policy or prices.",
        "critical_checks": {
            "OPS-01": "Use approved v2 (99.5% covered hours); v1 (99.99% 24/7) cannot support a commitment.",
            "EVID-02": "Parent unverified -> BLOCKED/open dependency, not N/A; never fabricate a certificate.",
            "SEC-02": "SSO supported but SCIM unsupported -> GAP, not irrelevant.",
            "DATA-01": "US/UK processing cannot satisfy EU-only; privacy/legal review.",
            "LEGAL-01": "Unlimited liability conflicts with caps; counterposition is not customer acceptance.",
            "COMM-02": "Missing required renewals, freight/tax -> incomplete TCO, not confirmed zero.",
            "SUPPLY-01": "Verify actual lifecycle, disclose CAP-7 and switch alternatives; do not save/reserve.",
        },
    }
    write(ROOT / "tools" / "expected" / "cowork-rfi-rfp.json", json.dumps(expected, ensure_ascii=False, indent=2) + "\n")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    mode = parser.add_mutually_exclusive_group(required=True)
    mode.add_argument("--prepare", action="store_true")
    mode.add_argument("--copies", action="store_true")
    args = parser.parse_args()
    if args.prepare:
        for language in ("en", "de"):
            knowledge(language)
            for size in ("full", "compact"):
                workbook(language, size)
        preparation()
    else:
        copies()
    print("RFI package updated. Recalculate masters before --copies; then run Update-Manifest.ps1 -All.")
