"""Check RFI artifacts, compact mapping, cached formulas and run isolation."""
import hashlib
import json
import tempfile
from pathlib import Path
from zipfile import ZipFile

from openpyxl import load_workbook

from cowork_rfi_text import COMPACT_IDS, CONDITIONS, TOPICS, questions

ROOT = Path(__file__).resolve().parent.parent
PACKAGE = ROOT / "cowork-rfi-rfp"
all_questions = {q["id"]: q for q in questions()}
assert len(all_questions) == 132 and len(COMPACT_IDS) == 28
assert COMPACT_IDS <= all_questions.keys()
assert all(parent in COMPACT_IDS for qid, (parent, _) in CONDITIONS.items() if qid in COMPACT_IDS)
assert {q["tag"] for q in all_questions.values() if q["id"] in COMPACT_IDS} >= {
    "conflict", "missing-evidence", "alternative", "availability", "tco", "privacy", "contract", "conditional",
}


def inventory(path, lang, variant):
    wb = load_workbook(path, data_only=False)
    cached = load_workbook(path, data_only=True)
    try:
        expected = all_questions.keys() if variant == "full" else COMPACT_IDS
        seen = set()
        for topic in TOPICS:
            ws = wb[topic]
            assert len(ws.tables) == 1 and ws.auto_filter.ref is None
            assert len(ws.data_validations.dataValidation) == 2
            for r in range(2, ws.max_row + 1):
                qid = ws.cell(r, 1).value
                assert qid not in seen
                seen.add(qid)
                q = all_questions[qid]
                assert [ws.cell(r, c).value for c in range(1, 9)] == [
                    qid, q["parent"] or None, q["priority"], q["text"][lang],
                    q["format"][lang], q["max_chars"], q["condition"][lang], q["evidence"][lang],
                ]
                assert all(ws.cell(r, c).value is None for c in range(9, 17))
                assert ws.cell(r, 17).data_type == "f" and ws.cell(r, 18).data_type == "f"
        assert seen == set(expected)
        assert cached["Run"]["B5"].value == len(expected)
        assert cached["TCO"]["B3"].value == 172800
        assert cached["TCO"]["B4"].value in (None, "")
        assert cached["TCO"]["B5"].value in (None, "")
        assert cached["TCO"]["B6"].value == cached["TCO"]["B7"].value == "PARTIAL"
        assert not wb._external_links
        assert all(cell.data_type != "e" for ws in cached for row in ws for cell in row)
        return {(ws.title, c.coordinate): c.value for ws in wb for row in ws for c in row
                if (ws.title, c.coordinate) not in {("Run", "B2"), ("Run", "B6")}}
    finally:
        wb.close()
        cached.close()


count = 0
for lang in ("en", "de"):
    for variant in ("full", "compact"):
        master = PACKAGE / "templates" / lang / f"RFI_{variant.title()}_Master.xlsx"
        baseline = inventory(master, lang, variant)
        for n in range(1, 9):
            path = PACKAGE / "runs" / lang / variant / f"{n:02}" / f"RFI_{variant.title()}_Run_{n:02}.xlsx"
            assert inventory(path, lang, variant) == baseline
            w = load_workbook(path, data_only=True)
            assert w["Run"]["B2"].value == f"{lang.upper()}-{variant.upper()}-{n:02}"
            assert w["Run"]["B6"].value.endswith(f"/{lang}/{variant}/{n:02}/results")
            assert w["Run"]["B8"].value == "NEW"
            w.close()
            assert (path.parent / "results" / "README.md").exists()
            with ZipFile(path) as z:
                assert not any(name.endswith("vbaProject.bin") for name in z.namelist())
                assert b"Contoso Demo" in z.read("docProps/core.xml")
            count += 1
assert count == 32

paths = list((PACKAGE / "templates").rglob("*.xlsx")) + list((PACKAGE / "runs").rglob("*.xlsx"))
hashes = {p: hashlib.sha256(p.read_bytes()).digest() for p in paths}
with tempfile.TemporaryDirectory(prefix="contoso-rfi-isolation-") as temp:
    target = Path(temp) / "Run_01.xlsx"
    first = PACKAGE / "runs" / "en" / "compact" / "01" / "RFI_Compact_Run_01.xlsx"
    target.write_bytes(first.read_bytes())
    edited = load_workbook(target)
    edited["QUAL"]["I2"] = "Isolated test answer"
    edited.save(target)
    edited.close()
    assert all(hashlib.sha256(p.read_bytes()).digest() == h for p, h in hashes.items())

manifest = json.loads((PACKAGE / "manifest.json").read_text(encoding="utf-8"))
listed = {f["path"] for f in manifest["files"]}
for f in manifest["files"]:
    assert hashlib.sha256((PACKAGE / f["path"]).read_bytes()).hexdigest() == f["sha256"]
assert not any("expected" in name or name == "tenant.json" for name in listed)
catalog = json.loads((ROOT / "catalog.json").read_text(encoding="utf-8"))
assert "cowork-rfi-rfp" in json.dumps(catalog)
print("RFI: 132/28 requirements, four recalculated masters, 32 isolated blank runs, formulas and manifest verified.")
