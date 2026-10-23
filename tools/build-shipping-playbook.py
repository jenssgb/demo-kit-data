"""Build the original Contoso shipping rehearsal corpus; no tenant access.

Uses the repository's existing python-docx/openpyxl Office builders.
Run tools\\Update-Manifest.ps1 -All after rebuilding.
"""
import csv
import json
from email.message import EmailMessage
from html import escape
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.worksheet.table import Table, TableStyleInfo
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches as SlideInches, Pt as SlidePt

ROOT = Path(__file__).resolve().parents[1]
BASE = "word-shipping-brief"
DOMAIN = "M365CPI98544940.OnMicrosoft.com"
DISCLAIMER = "CONTOSO | SYNTHETIC TRAINING DATA | Not operational, legal or safety advice."
CUSTOMERS = ["Fabrikam", "Northwind", "Tailspin", "Wide World", "Adatum", "Contoso Retail"]
CONTRACTED = [12400, 8200, 15600, 6900, 10300, 9800]
BILLED = [13950, 8350, 17100, 6950, 10920, 9900]
PACKET = [
    ("Scope", "Fictional administrative document-completeness exercise only. "
     "This is not a dangerous-goods shipping procedure. Never approve cargo or infer safety compliance."),
    ("Training packet", "For this exercise the administrative packet contains booking reference, consignee "
     "company, consignee contact phone, a carrier confirmation reference, and a signed internal review request. "
     "Check whether each is present. Do not validate technical classification, packaging, labels or transport law."),
    ("Missing information", "Ask for missing administrative fields. A missing contact phone blocks completion "
     "of this training packet. Do not invent a phone number or look one up."),
    ("Specialist handoff", "Sonia Rees coordinates the internal review request. Teresa Sac coordinates the "
     "carrier confirmation. A qualified authorized specialist decides actual requirements and shipment approval."),
    ("Evidence", "Cite the named document and section used. An internal checklist is not regulatory authority. "
     "If a source conflicts or the question needs professional judgment, say what evidence is missing and escalate."),
]
FILES = {
    "chat-servicenow-tickets": ["ServiceNow_incidents.csv"],
    "chat-shipping-workiq": ["Carrier_notice_ContosoStar.docx", "Booking_list.xlsx",
                             "Customer_notice_template.docx", "Ops_sync_rehearsal.docx"],
    "chat-shipping-search": ["Revised_ETA_ContosoStar.pdf", "Affected_bookings.xlsx",
                            "Customer_notice_template.docx", "Ops_sync_rehearsal.docx"],
    "chat-model-comparison": ["Model_comparison_scorecard.docx"],
    "outlook-shipping-triage": ["Revised_ETA_ContosoStar.pdf", "Affected_bookings.xlsx"],
    "outlook-shipping-reply": ["Carrier_notice_ContosoStar.docx"],
    "outlook-shipping-calendar": ["QBR_notes_Fabrikam.docx"],
    "teams-shipping-recap": ["Ops_sync_rehearsal.docx"],
    "teams-shipping-thread": ["Channel_thread_seed.docx"],
    "teams-shipping-interpreter": ["Interpreter_rehearsal.docx"],
    "word-shipping-legal": ["Supplier_Agreement_Draft.docx", "Procurement_Playbook.docx",
                            "Supplier_Agreement_expected.docx"],
    "word-shipping-finalize": ["Customer_brief_draft.docx"],
    "excel-shipping-volumes": ["Volumes_2026.xlsx"],
    "excel-shipping-changes": ["DnD_reconciliation.xlsx", "DnD_expected.csv"],
    "excel-shipping-reconcile": ["Carrier_invoices_Q3.xlsx", "Carrier_reconciliation_expected.csv"],
    "powerpoint-shipping-qbr": ["QBR_notes_Fabrikam.docx", "DnD_reconciliation.xlsx",
                               "QBR_storyboard.docx", "QBR_authored_reference.pptx"],
    "powerpoint-shipping-story": ["QBR_notes_Fabrikam.docx", "DnD_reconciliation.xlsx",
                                 "QBR_storyboard.docx", "QBR_authored_reference.pptx"],
    "powerpoint-shipping-interactive": ["DnD_reconciliation.xlsx", "DnD_expected.csv"],
    "sharepoint-dg-knowledge": ["DG_procedures.docx", "DG_precheck_checklist.docx",
                              "DG_contact_and_handoff.docx", "DG_source_register.docx",
                              "DG_lithium_booking.docx"],
    "onedrive-shipping-quotes": ["Quote_A.pdf", "Quote_B.pdf", "Quote_C.pdf"],
    "chat-dg-notebook": ["DG_precheck_checklist.docx", "DG_procedures.docx",
                        "Ops_sync_rehearsal.docx", "Survey_results.xlsx"],
    "chat-booking-vision": ["booking_4711.pdf"],
    "chat-shipping-research": ["Customer_FAQ_emissions.docx"],
    "chat-invoice-analyst": ["invoice_disputes_Q3.csv", "invoices_Q3.csv", "Analyst_expected.csv"],
    "agents-dg-skill": ["DG_precheck_checklist.docx", "booking_4711.pdf",
                        "skills/dg-pre-check/SKILL.md", "dg-pre-check.zip"],
    "agents-dg-builder": ["DG_procedures.docx", "DG_precheck_checklist.docx",
                          "DG_contact_and_handoff.docx", "DG_source_register.docx",
                          "DG_lithium_booking.docx"],
}


def write_json(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def document(name, title, sections, page_breaks=False):
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    doc.styles["Normal"].font.name = "Arial"
    doc.styles["Normal"].font.size = Pt(11)
    doc.add_heading(title, 0)
    doc.add_paragraph(DISCLAIMER)
    for index, (heading, text) in enumerate(sections):
        if page_breaks and index:
            doc.add_page_break()
        doc.add_heading(heading, 1)
        doc.add_paragraph(text)
    doc.core_properties.author = "Contoso"
    doc.core_properties.title = title
    doc.save(ROOT / BASE / name)


def workbook(name, sheets):
    wb = Workbook()
    wb.remove(wb.active)
    for title, headers, rows in sheets:
        ws = wb.create_sheet(title)
        ws.append(headers)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        if title != ".Rules":
            table = Table(displayName="Data" + str(len(wb.worksheets)), ref=ws.dimensions)
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
            ws.add_table(table)
        for column in ws.columns:
            ws.column_dimensions[column[0].column_letter].width = min(
                55, max(18, max(len(str(c.value or "")) for c in column) + 2))
        for row in ws:
            for cell in row:
                cell.font = Font(name="Arial", size=11, bold=cell.row == 1)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "#,##0;(#,##0);-"
    notes = wb.create_sheet("Source notes")
    notes.append([DISCLAIMER])
    notes.append(["All inputs are original synthetic fixtures. Copilot creates calculated results during rehearsal."])
    notes.column_dimensions["A"].width = 100
    wb.calculation.calcMode = "auto"
    wb.save(ROOT / BASE / name)


def write_csv(name, headers, rows):
    with (ROOT / BASE / name).open("w", encoding="utf-8", newline="") as stream:
        writer = csv.writer(stream)
        writer.writerow(headers)
        writer.writerows(rows)


def pdf(name, title, lines):
    output = canvas.Canvas(str(ROOT / BASE / name), pagesize=A4)
    output.setTitle(title)
    output.setAuthor("Contoso")
    output.setFont("Helvetica-Bold", 17)
    output.drawString(45, 790, title)
    output.setFont("Helvetica", 10)
    y = 755
    for line in [DISCLAIMER, *lines]:
        output.drawString(45, y, line)
        y -= 24
    output.save()


def backup_deck():
    presentation = Presentation()
    presentation.slide_width, presentation.slide_height = SlideInches(13.333), SlideInches(7.5)
    presentation.core_properties.author = "Contoso"
    presentation.core_properties.title = "Fabrikam QBR - authored rehearsal reference"
    ink, brand, soft = "242424", "0F6CBD", "F5F5F5"

    def text(slide, value, x, y, w, h, size=20, bold=False, color=ink):
        box = slide.shapes.add_textbox(*[SlideInches(n) for n in [x, y, w, h]])
        box.text_frame.word_wrap = True
        box.text_frame.margin_left = box.text_frame.margin_right = 0
        box.text_frame.margin_top = box.text_frame.margin_bottom = 0
        for i, line in enumerate(value.split("\n")):
            paragraph = box.text_frame.paragraphs[0] if not i else box.text_frame.add_paragraph()
            paragraph.text = line
            paragraph.font.name = "Aptos"
            paragraph.font.size = SlidePt(size)
            paragraph.font.bold = bold
            paragraph.font.color.rgb = RGBColor.from_string(color)
            paragraph.space_after = SlidePt(12)
        return box

    def slide(title, number, notes):
        s = presentation.slides.add_slide(presentation.slide_layouts[6])
        text(s, "CONTOSO  /  FICTIONAL SHIPPING REVIEW", .65, .5, 12, .35, 12, True, brand)
        text(s, title, .65, 1.0, 12, .65, 36, True)
        text(s, "AUTHORED REHEARSAL REFERENCE  |  Not a Copilot result or an approved Brand Kit", .65, 6.65, 11.5, .3, 10)
        text(s, str(number), 12.0, 6.65, .6, .3, 10)
        s.notes_slide.notes_text_frame.text = notes + "\n" + DISCLAIMER
        return s

    def card(s, number, title, detail, x, y, w=3.8, h=3.6):
        shape = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, *[SlideInches(n) for n in [x, y, w, h]])
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(soft)
        shape.line.fill.background()
        text(s, number, x+.25, y+.2, w-.5, .6, 32, True, brand)
        text(s, title, x+.25, y+1.0, w-.5, .55, 22, True)
        text(s, detail, x+.25, y+1.75, w-.5, h-1.9, 18)

    s = slide("Fabrikam QBR: evidence before promises", 1,
              "Authored backup for PP-02/PP-03. Rebuild with Copilot during rehearsal. "
              "Sources: QBR_notes_Fabrikam.docx and DnD_reconciliation.xlsx. Never present this as a generated result.")
    card(s, "01", "Review the quarter", "Highlights and the current\nContoso Star exception.", .65, 2.0)
    card(s, "02", "Check the variance", "Separate billed amounts\nfrom approved corrections.", 4.75, 2.0)
    card(s, "03", "Agree next evidence", "Open questions, owners\nand a review checkpoint.", 8.85, 2.0)

    s = slide("An arrival estimate is not a delivery promise", 2,
              "Source: QBR_notes_Fabrikam.docx. The carrier ETA is fictional scenario evidence. "
              "No delivery or release has been approved. Ask what evidence is needed next.")
    text(s, "4711", .65, 2.15, 5.5, 1.3, 68, True, brand)
    text(s, "Fabrikam booking\n4 containers on Contoso Star", .65, 3.85, 5.5, 1.1, 24)
    card(s, "ETA", "30 Oct 2026", "14:00 UTC\nEstimated port arrival only.\nDelivery is unconfirmed.", 7.0, 2.0, 5.65)

    s = slide("Three customers exceed 5% variance", 3,
              "Source: DnD_reconciliation.xlsx. Formula: (Billed - Contracted) / Contracted. "
              "Chart is editable and contains all six customer values. Exactly three exceed 5%; only Fabrikam exceeds 10%. No credit is approved.")
    data = CategoryChartData()
    data.categories = CUSTOMERS
    data.add_series("Variance", [(b-a)/a for a, b in zip(CONTRACTED, BILLED)])
    chart = s.shapes.add_chart(XL_CHART_TYPE.BAR_CLUSTERED, SlideInches(.65), SlideInches(1.9),
                               SlideInches(8.35), SlideInches(4.2), data).chart
    chart.has_legend = False
    chart.value_axis.minimum_scale, chart.value_axis.maximum_scale = 0, .15
    chart.value_axis.tick_labels.number_format = "0%"
    chart.value_axis.tick_labels.font.size = SlidePt(12)
    chart.category_axis.tick_labels.font.size = SlidePt(14)
    plot = chart.plots[0]
    plot.has_data_labels = True
    plot.data_labels.number_format = "0.0%"
    plot.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
    plot.data_labels.font.size = SlidePt(14)
    chart.series[0].format.fill.solid()
    chart.series[0].format.fill.fore_color.rgb = RGBColor.from_string(brand)
    text(s, "12.5%", 9.4, 2.3, 3.2, 1, 48, True, brand)
    text(s, "Fabrikam variance\nEUR 12,400 contracted\nEUR 13,950 billed", 9.4, 3.5, 3.2, 1.7, 19)
    text(s, "No credit approved", 9.4, 5.45, 3.2, .4, 16, True)

    s = slide("Three open issues still need evidence", 4,
              "Source: QBR_notes_Fabrikam.docx. This intentionally text-based content slide is suitable for "
              "the Visualize this slide demonstration. Never use that image transformation on the editable chart.")
    for y, title, detail in [
        (2.0, "Missing contact", "Booking 4711 has no consignee contact phone."),
        (3.25, "Unconfirmed delivery", "The carrier ETA is not a confirmed delivery date."),
        (4.5, "Unvalidated invoice cause", "A billing difference does not prove an approved credit."),
    ]:
        text(s, title, .85, y, 4.2, .55, 24, True, brand)
        text(s, detail, 5.4, y, 7.0, .8, 22)

    s = slide("Owners are clear; approvals are still required", 5,
              "Source: QBR_notes_Fabrikam.docx. Teresa drafts for review; Billie checks portal evidence; "
              "Sonia coordinates specialist handoff. These are synthetic source notes, not proof a meeting assigned tasks.")
    card(s, "01", "Teresa Sac", "Draft the customer notice.\nKeep it unsent for review.", .65, 2.0)
    card(s, "02", "Billie Vester", "Check the portal against\nwritten carrier evidence.", 4.75, 2.0)
    card(s, "03", "Sonia Rees", "Coordinate the qualified\nspecialist's handoff.", 8.85, 2.0)

    s = slide("What evidence would change the decision?", 6,
              "Use these discussion questions, not shipping or legal advice. "
              "Review the actual carrier source before making any commitment. Preserve all unanswered questions.")
    text(s, "?", .65, 2.1, 2.0, 2.1, 120, True, brand)
    text(s, "Who will obtain the missing contact phone?\n"
         "What confirms delivery, rather than arrival?\n"
         "Which invoice evidence explains each difference?", 3.0, 2.15, 9.5, 2.9, 26)
    text(s, "Next carrier update in the fictional scenario: 29 Oct 2026, 16:00 UTC", 3.0, 5.5, 9.5, .7, 18)
    presentation.save(ROOT / BASE / "QBR_authored_reference.pptx")


def build_sources():
    (ROOT / BASE).mkdir(exist_ok=True)
    document("Carrier_notice_ContosoStar.docx", "Carrier notice - Contoso Star", [
        ("Confirmed update", "Fictional scenario dated 29 October 2026, 09:00 UTC. The Contoso Star ETA "
         "has changed from 29 October 08:00 UTC to 30 October 14:00 UTC. This is an estimated port arrival, "
         "not delivery or release approval. Next carrier update: 29 October 16:00 UTC."),
        ("Impact", "Bookings 4711 to 4716 are affected. Teresa Sac coordinates the customer notice. "
         "Billie Vester checks the portal. Sonia Rees coordinates specialist review. No delivery guarantee."),
    ])
    bookings = [[str(4711 + i), customer, 4 + i * 2, "Contoso Star", "30 Oct 2026 14:00 UTC",
                 "Next update 29 Oct 16:00 UTC"] for i, customer in enumerate(CUSTOMERS)]
    for name in ["Booking_list.xlsx", "Affected_bookings.xlsx"]:
        workbook(name, [("Bookings", ["Booking", "Customer", "Containers", "Vessel", "ETA", "Promise"], bookings)])
    document("Customer_notice_template.docx", "Customer notice - template", [
        ("Subject", "Contoso Star - booking [reference] update"),
        ("Body", "Our carrier's latest estimated arrival is [ETA]. Your booking is affected as follows: "
         "[impact from booking evidence]. Arrival is not delivery confirmation. We will update you at [time]. "
         "Please contact Teresa Sac with questions. Review the facts before sending."),
    ])
    meeting = [
        ("Rehearsal, not a transcript", "Use the installed scheduled Teams meeting 'Ops sync - vessel delays'. "
         "With participant consent, record and transcribe 10-15 minutes in English with Facilitator on. "
         "This source is an agenda/script, not a real meeting recap or evidence of a decision."),
        ("Status", "Teresa states the carrier ETA: 30 October 2026 14:00 UTC. Bookings 4711-4716 are affected. "
         "Billie says the portal is being checked. An ETA is not a delivery promise."),
        ("Customer communication", "Agree that Teresa drafts the notice for review, and that the next update "
         "is 29 October 16:00 UTC. Say the owner and time aloud. Do not claim the customer notice was sent."),
        ("Next steps", "Sonia owns the specialist handoff. Keep booking 4711's missing contact phone open. "
         "Ask who will obtain it. The approved specialist, not Copilot, decides shipping compliance."),
        ("Saved result", "After the meeting ask @Facilitator to create a summary document based on this meeting. "
         "Open the actual returned Word link and save it during rehearsal. Never relabel this script as that result."),
    ]
    document("Ops_sync_rehearsal.docx", "Ops sync - vessel delays: rehearsal script", meeting)
    document("QBR_notes_Fabrikam.docx", "Fabrikam QBR - source notes", [
        ("Agenda", "Review service exceptions, D&D invoice variance, open issues and next-quarter actions."),
        ("Highlights", "Fabrikam has one Contoso Star booking (4711, four containers). ETA 30 October "
         "2026 14:00 UTC; delivery unconfirmed. Teresa drafts the notice; Billie checks the portal."),
        ("D&D variance", "Fabrikam contracted EUR 12,400, billed EUR 13,950. Use the workbook for the "
         "six-customer comparison. The difference is not an approved credit."),
        ("Open issues", "Missing contact phone on booking 4711; carrier delivery confirmation pending; "
         "invoice causes not yet validated. Sonia coordinates specialist review."),
        ("Next steps", "Review customer notice before sending. Validate billing evidence before proposing "
         "a correction. Keep owners and next evidence visible. No live customer or billing system changes."),
    ])
    draft_sections = [
        ("Executive summary", "Contoso is preparing a customer brief about the fictional Contoso Star delay. "
         "The carrier estimates arrival at 14:00 UTC on 30 October 2026. This is not a confirmed delivery date. "
         "The operations team must keep confirmed facts separate from open assumptions. The intended audience "
         "is the customer-service lead reviewing a draft, not a recipient of an automatic message."),
        ("Customer impact", "Bookings 4711 through 4716 affect Fabrikam, Northwind, Tailspin, Wide World, "
         "Adatum and Contoso Retail. The source workbook lists 4, 6, 8, 10, 12 and 14 containers respectively. "
         "Teresa Sac prepares communications. The next carrier update is due at 16:00 UTC on 29 October. "
         "Operations must confirm any new delivery date before the notice is shared."),
        ("Options and terms", "Two discussion options are to wait for confirmed carrier details or ask "
         "Operations to assess an alternative plan. Neither is a booked alternative or a transport instruction. "
         "The wording 'estimated port arrival' must remain unchanged. Do not introduce compensation, an "
         "approved credit or a delivery commitment. Customer-service review is required."),
        ("Review and next evidence", "Billie Vester checks portal information. Sonia Rees coordinates "
         "specialist handoff. Booking 4711 lacks the consignee contact phone. The next draft should preserve "
         "that open issue. The reviewer must confirm all dates against the carrier source before circulation. "
         "This four-page document intentionally includes headings for a table of contents demonstration."),
    ]
    document("Customer_brief_draft.docx", "Contoso Star - customer brief draft", draft_sections, True)
    document("Supplier_Agreement_Draft.docx", "Contoso supplier agreement - draft", [
        ("7.1 Payment", "The buyer shall pay invoices within 30 days of receipt."),
        ("9.2 Liability", "The buyer's liability under this agreement is unlimited."),
        ("12.4 Termination", "Either party may terminate this agreement on 14 days' written notice."),
    ])
    document("Procurement_Playbook.docx", "Contoso procurement playbook", [
        ("4.1 Payment rule", "Require a payment term of 60 days after receipt of a valid invoice. "
         "Example clause: The buyer shall pay valid invoices within 60 days of receipt."),
        ("6.3 Liability rule", "Require the buyer's aggregate liability to be capped at 12 months of fees. "
         "Example clause: The buyer's aggregate liability shall not exceed the fees paid or payable "
         "under this agreement in the preceding 12 months."),
        ("9 Termination rule", "Require at least 90 days' written notice for termination for convenience. "
         "Example clause: Either party may terminate for convenience on at least 90 days' written notice."),
    ])
    document("Supplier_Agreement_expected.docx", "Expected training comparison - not a Copilot redline", [
        ("7.1", "Compare the draft's 30 days with playbook section 4.1: 60 days."),
        ("9.2", "Compare unlimited liability with section 6.3: cap at 12 months of fees."),
        ("12.4", "Compare 14 days with section 9: at least 90 days."),
        ("Review", "This is an authored expected-result aid, not tracked-change evidence and not legal advice."),
    ])
    volume_rows = [[f"2026-{month:02}", lane, segment, 100 + month * 5 + i * 10,
                    112 + month * 6 + i * 9]
                   for month in range(1, 13)
                   for i, lane in enumerate(["North Europe", "Mediterranean", "Transatlantic"])
                   for segment in ["Dry", "Reefer", "Special", "Short sea", "Inland", "Other"]]
    workbook("Volumes_2026.xlsx", [("Volumes", ["Month", "Trade lane", "Segment", "TEU 2025", "TEU 2026"], volume_rows)])
    workbook("DnD_reconciliation.xlsx", [
        ("Charges", ["Customer", "Contracted EUR", "Billed EUR"], list(zip(CUSTOMERS, CONTRACTED, BILLED))),
        (".Rules", ["Show amounts in EUR, no decimals"], [
            ["Show percentages with one decimal place."],
            ["Keep original source amounts unchanged."],
            ["Flag variance strictly greater than 5 percent, not greater than or equal."],
            ["Put new analysis on a separate sheet. Do not approve credits."],
        ]),
    ])
    write_csv("DnD_expected.csv", ["Customer", "Contracted EUR", "Billed EUR", "Variance percent", "Above 5 percent"],
              [[c, a, b, f"{(b-a)/a*100:.6f}", str((b-a)/a > .05)] for c, a, b in zip(CUSTOMERS, CONTRACTED, BILLED)])
    accruals = [[str(5001 + i), CUSTOMERS[i % 6], 1000 + 25 * i, "EUR"] for i in range(60)]
    invoices = [row.copy() for row in accruals]
    expected = []
    for i in [7, 18, 39]:
        invoices[i][2] += 150
        expected.append([invoices[i][0], "Amount mismatch", accruals[i][2], invoices[i][2]])
    invoices.pop(54)
    expected.append([accruals[54][0], "Accrual only", accruals[54][2], ""])
    invoices.append(["5999", "Fabrikam", 1250, "EUR"])
    expected.append(["5999", "Invoice only", "", 1250])
    workbook("Carrier_invoices_Q3.xlsx", [
        ("Invoices", ["Booking", "Customer", "Amount EUR", "Currency"], invoices),
        ("Accruals", ["Booking", "Customer", "Amount EUR", "Currency"], accruals),
    ])
    write_csv("Carrier_reconciliation_expected.csv", ["Booking", "Exception", "Accrual EUR", "Invoice EUR"], expected)
    document("QBR_storyboard.docx", "Fabrikam QBR - six-slide rehearsal storyboard", [
        ("1 Agenda", "Purpose: review the quarter and agree next evidence. Not a saved Copilot presentation."),
        ("2 Highlights", "One source-backed update; preserve estimated arrival versus delivery."),
        ("3 D&D variance", "Editable bar chart from DnD_reconciliation.xlsx; flag three customers over 5%."),
        ("4 Open issues", "Text-only slide: missing contact phone, unconfirmed delivery, unvalidated invoice causes. "
         "Use Visualize this slide here, not on the editable chart."),
        ("5 Next steps", "Teresa: customer draft. Billie: portal evidence. Sonia: specialist handoff."),
        ("6 Questions", "Ask for the evidence needed before a customer promise or invoice correction."),
    ])
    for name in ["DG_procedures.docx", "DG_precheck_checklist.docx", "DG_lithium_booking.docx",
                 "DG_contact_and_handoff.docx", "DG_source_register.docx"]:
        document(name, "Contoso DG training - " + name.removesuffix(".docx").replace("_", " "), PACKET)
    workbook("Survey_results.xlsx", [("Survey", ["Team", "Responses", "Open question"], [
        ["Operations", 12, "Who obtains the missing contact phone?"],
        ["Customer service", 9, "Where is the written carrier confirmation?"],
        ["Review coordination", 6, "Who owns specialist sign-off?"],
    ])])
    pdf("booking_4711.pdf", "Contoso booking 4711 - training form", [
        "Booking reference: 4711", "Consignee company: Fabrikam", "Consignee contact phone: __________________",
        "Vessel: Contoso Star", "Carrier confirmation reference: TRN-C4711",
        "Signed internal review request: Sonia Rees (training fixture)",
        "Administrative completeness only. No safety classification or release approval.",
    ])
    pdf("Revised_ETA_ContosoStar.pdf", "Contoso Star - revised estimated arrival", [
        "Scenario update: 29 October 2026, 09:00 UTC", "Previous ETA: 29 October 2026, 08:00 UTC",
        "Revised ETA: 30 October 2026, 14:00 UTC", "Next update: 29 October 2026, 16:00 UTC",
        "Affected bookings: 4711-4716. ETA is not delivery or release confirmation.",
    ])
    for letter, price, validity, terms in [
        ("A", 12500, "30 days", "Payment 30 days; handling included"),
        ("B", 11900, "14 days", "Payment 15 days; handling EUR 800 extra"),
        ("C", 13100, "45 days", "Payment 60 days; handling included"),
    ]:
        pdf(f"Quote_{letter}.pdf", f"Contoso training - quote {letter}", [
            f"Quoted price: EUR {price:,}", f"Validity: {validity}", terms,
            "Same fictional scope: 20 standard containers, administrative comparison only.",
            "Not a real offer. Compare total scope, not just the headline price.",
        ])
    document("Customer_FAQ_emissions.docx", "Contoso customer FAQ - EU ETS review draft", [
        ("Purpose", "This is a deliberately incomplete internal draft for a research exercise. "
         "It makes no authoritative claim about current law, costs or customer obligations."),
        ("Current customer wording", "We are reviewing how EU ETS affects our shipping communications. "
         "Our draft does not distinguish reporting year, surrender year or implementation phases. "
         "Do not quote a price or contractual surcharge from this file."),
        ("Questions for Researcher", "What applies to maritime transport in 2026? Which changes should "
         "customer service explain? Check current European Commission sources and distinguish law, "
         "our draft and open questions. Escalate any advice to the qualified policy owner."),
    ])
    reasons = ["Duplicate charge", "Wrong rate", "Free-time dispute", "Data mismatch", "Other"]
    invoice_rows, dispute_rows = [], []
    for i in range(300):
        identifier = f"INV26-{i+1:04}"
        customer, month = CUSTOMERS[i % 6], f"2026-{7+i%3:02}"
        invoice_rows.append([identifier, customer, month, 1000 + 17 * i, "EUR"])
        dispute_rows.append([f"D-{i+1:04}", identifier, reasons[i % 5], [200, 120, 80, 40, 20][i % 5]])
    write_csv("invoices_Q3.csv", ["invoice_id", "customer", "month", "billed_eur", "currency"], invoice_rows)
    write_csv("invoice_disputes_Q3.csv", ["dispute_id", "invoice_id", "dispute_reason", "corrected_amount_eur"], dispute_rows)
    write_csv("Analyst_expected.csv", ["Reason", "Rows", "Corrections EUR"],
              [[reason, 60, amount * 60] for reason, amount in zip(reasons, [200, 120, 80, 40, 20])])
    write_csv("ServiceNow_incidents.csv", ["Number", "Short description", "Priority", "State", "Assigned to"], [
        ["INC0012345", "EDI upload to booking portal fails", "2", "In Progress", "Billie Vester"],
        ["INC0012346", "EDI upload retries remain queued", "2", "New", "Billie Vester"],
        ["INC0012347", "EDI upload confirmation missing", "2", "In Progress", "Billie Vester"],
        ["INC0012348", "Training printer unavailable", "3", "Resolved", "Billie Vester"],
    ])
    document("Model_comparison_scorecard.docx", "Model comparison - rehearsal scorecard", [
        ("Prompt", "Compare three options to reduce detention charges at a congested port: pros, cons and effort."),
        ("Compare", "Judge explicit assumptions, evidence, clarity and reviewability; not which brand always wins. "
         "No rate quotes, live system changes or automatic customer commitments."),
    ])
    document("Interpreter_rehearsal.docx", "Interpreter - bilingual rehearsal", [
        ("English", "We are reviewing the Contoso Star customer notice. The arrival is estimated, "
         "not a delivery promise. Teresa drafts the update and Sonia coordinates the review."),
        ("Deutsch", "Wir besprechen die Kundeninformation zur Contoso Star. Die Ankunft ist geschätzt, "
         "keine Lieferzusage. Teresa entwirft das Update und Sonia koordiniert die Prüfung."),
    ])
    thread = [
        ("Teresa Sac", "Contoso Star update: estimated arrival 30 October 14:00 UTC, not a delivery promise. "
         "Please keep the next carrier update at 29 October 16:00 UTC visible."),
        ("Billie Vester", "I will check the portal and compare it with the written carrier notice. "
         "A portal status alone does not authorize shipment release."),
        ("Sonia Rees", "Booking 4711 still lacks a consignee contact phone. I coordinate specialist review; "
         "do not treat this training checklist as dangerous-goods approval."),
    ]
    replies = []
    for i in range(12):
        person, text = thread[i % 3]
        replies.append((f"Reply {i+1} - {person}", text + [
            " Decision: Teresa drafts the customer notice for review; it is not sent yet.",
            " Open question: who obtains the missing phone number before the review?",
            " Next evidence: written carrier confirmation, not an inferred delivery time.",
            " Decision: keep all unverified assumptions in the open-issues list.",
        ][i // 3]))
    document("Channel_thread_seed.docx", "Contoso Shipping Ops - channel rehearsal", [
        ("Setup", "In an approved Contoso Shipping Ops team, channel Ops - Vessel delays, publish one "
         "main post and the twelve replies below using the actual listed internal CDX users. "
         "This document is seed content, not a Teams conversation. Do not attribute a post to a user who did not post it."),
        ("Main post - Teresa Sac", "Contoso Star delay: please agree customer communication and next evidence."),
        *replies,
        ("Other post 1", "Training rota review: unrelated to the vessel-delay thread."),
        ("Other post 2", "Office equipment inventory: unrelated to the vessel-delay thread."),
        ("Other post 3", "Quarterly review agenda: use the QBR sources separately."),
    ])
    skill = """---
name: dg-pre-check
description: Check administrative completeness of the fictional Contoso booking training packet.
---
Use the selected booking and the supplied Contoso training checklist only.
Check booking reference, consignee company, consignee contact phone, carrier confirmation reference,
and signed internal review request. Quote the evidence for each field.
Return present, missing, or unclear; never invent a phone number or a confirmation.
Ask who will supply missing information and state that the training packet is incomplete if needed.
This is not dangerous-goods advice, legal advice, technical validation or shipment authorization.
Never assess classification, packaging, labels or transport safety. Escalate those to a qualified specialist.
"""
    path = ROOT / BASE / "skills" / "dg-pre-check" / "SKILL.md"
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(skill, encoding="utf-8", newline="\n")
    with ZipFile(ROOT / BASE / "dg-pre-check.zip", "w", ZIP_DEFLATED) as archive:
        archive.writestr("SKILL.md", skill)
    backup_deck()


def mail(folder, filename, sender, alias, subject, text):
    path = folder / "tenant" / filename
    path.parent.mkdir(parents=True, exist_ok=True)
    message = EmailMessage()
    message["From"] = f"{sender} <{alias}@{DOMAIN}>"
    message["To"] = f"MOD Administrator <admin@{DOMAIN}>"
    message["Subject"] = subject
    message.set_content(DISCLAIMER + "\n\n" + text)
    path.write_bytes(message.as_bytes())
    return "tenant/" + filename


def build_folders():
    for folder_id, files in {BASE: [], **FILES}.items():
        folder = ROOT / folder_id
        folder.mkdir(exist_ok=True)
        title = folder_id.replace("-", " ").title()
        write_json(folder / "manifest.json", {
            "title": {"en": title, "de": title},
            "next": {
                "en": ["Run the global -Tenant command; rehearse this app's page and check the manual prerequisites."],
                "de": ["Das globale -Tenant-Setup ausführen; diese App-Seite proben und die manuellen Voraussetzungen prüfen."],
            },
        })
        if files:
            write_json(folder / "shared.json", [f"{BASE}/{name}" for name in files])
        app_name, app_url = (
            ("Outlook", "https://outlook.cloud.microsoft/") if folder_id.startswith("outlook-") else
            ("Teams", "https://teams.microsoft.com/") if folder_id.startswith("teams-") else
            ("Word & Excel", "https://www.microsoft365.com/") if folder_id.startswith(("word-", "excel-")) else
            ("PowerPoint", "https://www.microsoft365.com/") if folder_id.startswith("powerpoint-") else
            ("OneDrive", "https://www.microsoft365.com/") if folder_id.startswith("onedrive-") else
            ("SharePoint", "https://www.microsoft365.com/") if folder_id.startswith("sharepoint-") else
            ("Microsoft Copilot", "https://copilot.cloud.microsoft/")
        )
        config = {"links": [{"name": {"en": app_name, "de": app_name}, "url": app_url}]}
        if folder_id == "powerpoint-shipping-qbr":
            config["links"].append({"name": {"en": "Brand Kits", "de": "Brand Kits"},
                                    "url": "https://office.com/create/brandkit"})
        if folder_id in {"agents-dg-skill", "word-shipping-finalize"}:
            config["links"].append({
                "name": {"en": "Message Center (admin verification)", "de": "Message Center (Admin-Prüfung)"},
                "url": "https://admin.microsoft.com/#/MessageCenter",
            })
        write_json(folder / "tenant.json", config)

    folder = ROOT / "chat-shipping-workiq"
    config = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    config["mails"] = []
    for i in range(6):
        config["mails"].append(mail(folder, f"{i+1}_Teresa_Delay_notice.eml", "Teresa Sac", "TeresaS",
                                    f"Delay notice - Contoso Star - update {i+1}",
                                    "Contoso Star delay training update. The latest carrier ETA is "
                                    "30 October 2026 14:00 UTC, not delivery confirmation. Bookings 4711-4716 "
                                    "are affected. I will draft the customer notice for review. "
                                    "Next update: 29 October 16:00 UTC. Teresa."))
    for i, text in enumerate(["Equipment inventory", "Training rota", "Office move"]):
        config["mails"].append(mail(folder, f"{i+7}_Billie_Other.eml", "Billie Vester", "BillieV",
                                    "Contoso shipping - " + text, "Unrelated training mail: " + text))
    config["mails"].append(mail(folder, "10_Teresa_Attachments.eml", "Teresa Sac", "TeresaS",
                                "Contoso Star delay - carrier attachments",
                                "The attached revised ETA notice and affected-bookings workbook are "
                                "the synthetic source evidence. ETA 30 October 2026 14:00 UTC is not delivery."))
    config["mailOptions"] = [{
        "subject": "Contoso Star delay - carrier attachments",
        "attachments": ["Revised_ETA_ContosoStar.pdf", "Affected_bookings.xlsx"],
    }]
    write_json(folder / "shared.json", [f"{BASE}/{name}" for name in
               [*FILES[folder.name], "Revised_ETA_ContosoStar.pdf", "Affected_bookings.xlsx"]])
    config["requiredUsers"] = [f"{a}@{DOMAIN}" for a in ["TeresaS", "BillieV"]]
    config["manual"] = {
        "en": ["Work IQ readiness: as the presenter, verify real mailbox and document citations without attachments. "
               "File upload and mail submission are not indexing or delivery proofs. Teams evidence requires the rehearsed meeting/thread."],
        "de": ["Work-IQ-Bereitschaft: als Präsentator echte Mail- und Dokumentzitate ohne Anhänge prüfen. "
               "Upload und Mailübergabe sind keine Indizierungs- oder Zustellnachweise. Teams-Nachweise brauchen die geprobte Besprechung/Unterhaltung."],
    }
    write_json(folder / "tenant.json", config)

    folder = ROOT / "outlook-shipping-reply"
    config = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    config["mails"] = [mail(folder, "1_Teresa_Booking_4711.eml", "Teresa Sac", "TeresaS",
        "Booking 4711 - new ETA?",
        "Please draft a reply for Fabrikam and keep it unsent until review. Can we confirm the new ETA?\n\n"
        "Forwarded carrier notice (synthetic source copied into this message):\n"
        "From: Teresa Sac, carrier coordination\n"
        "Contoso Star ETA: 30 October 2026 14:00 UTC. This is estimated port arrival, not delivery. "
        "Next update 29 October 16:00 UTC.\n\n"
        "Earlier request from Fabrikam (fictional company): What changed for booking 4711? "
        "Please offer two discussion options and a call on Thursday. No alternative has been booked.")]
    config["requiredUsers"] = [f"TeresaS@{DOMAIN}"]
    config["mailOptions"] = [{"subject": "Booking 4711 - new ETA?", "html": True}]
    write_json(folder / "tenant.json", config)

    folder = ROOT / "outlook-shipping-calendar"
    config = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    config["events"] = [{
        "subject": {"en": "Fabrikam QBR - shipping rehearsal", "de": "Fabrikam QBR - Versandgeneralprobe"},
        "weekday": "Thursday", "start": "10:00", "minutes": 45, "onlineMeeting": True,
        "attendees": ["Teresa Sac", "Sonia Rees"],
        "body": {"en": "Synthetic QBR. Review QBR_notes_Fabrikam.docx in Demo-Kit > Outlook > "
                       "outlook-shipping-calendar. Discuss booking 4711, the missing phone and unvalidated D&D variance.",
                 "de": "Synthetischer QBR. QBR_notes_Fabrikam.docx unter Demo-Kit > Outlook > "
                       "outlook-shipping-calendar prüfen. Buchung 4711, fehlende Telefonnummer und ungeklärte D&D-Abweichung besprechen."},
    }]
    config["mails"] = [mail(folder, f"{i+1}_Teresa_QBR.eml", "Teresa Sac", "TeresaS",
        "Fabrikam QBR - " + topic,
        "For the synthetic Fabrikam QBR: " + content) for i, (topic, content) in enumerate([
            ("open issues", "Booking 4711 lacks the contact phone. Delivery remains unconfirmed."),
            ("invoice evidence", "Contracted EUR 12,400; billed EUR 13,950. No credit has been approved."),
        ])]
    config["requiredUsers"] = [f"{a}@{DOMAIN}" for a in ["TeresaS", "SoniaR"]]
    write_json(folder / "tenant.json", config)

    folder = ROOT / "teams-shipping-recap"
    config = json.loads((folder / "tenant.json").read_text(encoding="utf-8"))
    config["events"] = [{
        "subject": {"en": "Ops sync - vessel delays", "de": "Ops sync - vessel delays"},
        "weekday": "Tuesday", "start": "10:00", "minutes": 15, "onlineMeeting": True,
        "attendees": ["Teresa Sac", "Sonia Rees", "Billie Vester"],
        "body": {"en": "Synthetic rehearsal: delay status, customer communication and next steps. "
                       "Use Ops_sync_rehearsal.docx. Record/transcribe only with consent and policy approval. "
                       "Say owners, times and unresolved questions aloud; enable Facilitator.",
                 "de": "Synthetische Generalprobe: Verspätung, Kundeninformation und nächste Schritte. "
                       "Ops_sync_rehearsal.docx verwenden. Aufnahme/Transkription nur mit Zustimmung und Richtlinienfreigabe. "
                       "Verantwortliche, Zeiten und offene Fragen laut nennen; Facilitator aktivieren."},
    }]
    config["requiredUsers"] = [f"{a}@{DOMAIN}" for a in ["TeresaS", "SoniaR", "BillieV"]]
    config["manual"] = {
        "en": ["Rehearse Ops sync - vessel delays with recording, English transcription and Facilitator. "
               "Verify Recap, Notes, Tasks and Chapters as available. Ask @Facilitator to create a summary document; "
               "save the actual Word result link. The installed script is not a transcript."],
        "de": ["Ops sync - vessel delays mit Aufnahme, englischer Transkription und Facilitator proben. "
               "Recap, Notizen, Aufgaben und Kapitel nach Verfügbarkeit prüfen. @Facilitator ein Zusammenfassungsdokument "
               "erstellen lassen; echten Word-Ergebnislink speichern. Das installierte Skript ist kein Transkript."],
    }
    write_json(folder / "tenant.json", config)
    manual = {
        "teams-shipping-thread": (
            "Use Channel_thread_seed.docx in an approved Contoso Shipping Ops / Ops - Vessel delays channel. "
            "Post the main message and 12 replies with the named internal CDX accounts; do not impersonate users. "
            "Verify >1,000 characters. Channel Agent requires Public Preview, app permission and Loop; pre-generate its report.",
            "Channel_thread_seed.docx im genehmigten Kanal Contoso Shipping Ops / Ops - Vessel delays verwenden. "
            "Hauptnachricht und 12 Antworten mit den genannten internen CDX-Konten posten; keine Identitäten vortäuschen. "
            ">1.000 Zeichen prüfen. Channel Agent braucht Public Preview, App-Freigabe und Loop; Bericht vorab erzeugen."),
        "teams-shipping-interpreter": (
            "Use the installed scheduled Ops sync meeting with two licensed internal users and headsets. "
            "Check Interpreter policy; rehearse German/English. No Polish or consecutive-mode promise.",
            "Die installierte geplante Ops-sync-Besprechung mit zwei lizenzierten internen Benutzern und Headsets verwenden. "
            "Interpreter-Richtlinie prüfen; Deutsch/Englisch proben. Kein Polnisch und kein Versprechen für Konsekutivmodus."),
        "word-shipping-finalize": (
            "Word Windows rehearsal: enable Review > Track Changes manually. Verify comments, table of contents, "
            "header/footer and page numbers; save the actual result. If unsupported, perform those operations manually.",
            "Word-Windows-Probe: Überprüfen > Änderungen nachverfolgen von Hand einschalten. Kommentare, Inhaltsverzeichnis, "
            "Kopf-/Fußzeile und Seitenzahlen prüfen; echtes Ergebnis speichern. Nicht unterstützte Schritte manuell ausführen."),
        "excel-shipping-changes": (
            "Let Copilot edit the cloud DnD_reconciliation.xlsx directly during rehearsal. Preserve the same file "
            "for Show Changes; an uploaded result or copied workbook cannot supply that edit history.",
            "Copilot während der Probe die Cloud-Datei DnD_reconciliation.xlsx direkt bearbeiten lassen. "
            "Für Show Changes dieselbe Datei behalten; hochgeladene Ergebnisse oder Kopien liefern diesen Verlauf nicht."),
        "powerpoint-shipping-qbr": (
            "Before rehearsal, an approved Brand Manager publishes and shares a synthetic Contoso brand kit. "
            "Permissions can take 24 hours. Generate and save the real six-slide QBR; the storyboard is not a Copilot result.",
            "Vor der Probe veröffentlicht und teilt ein genehmigter Brand Manager ein synthetisches Contoso-Brand-Kit. "
            "Rechte können 24 Stunden brauchen. Echten QBR mit sechs Folien erzeugen und speichern; das Storyboard ist kein Copilot-Ergebnis."),
        "powerpoint-shipping-story": (
            "Open the real saved six-slide QBR from the preceding PowerPoint exercise. Keep the D&D chart editable; "
            "Visualize this slide belongs on the Open issues text slide.",
            "Den echten gespeicherten QBR mit sechs Folien aus der vorherigen PowerPoint-Übung öffnen. "
            "D&D-Diagramm bearbeitbar lassen; Visualize this slide auf der Textfolie Open issues nutzen."),
        "powerpoint-shipping-interactive": (
            "Verify Home/Insert > Interactive Slide or the skill before the visit. Rehearse and save the actual "
            "interactive slide; if absent, explain the rollout instead. Never ship a static mockup as an interactive result.",
            "Vor dem Besuch Home/Insert > Interactive Slide oder den Skill prüfen. Echte interaktive Folie proben "
            "und speichern; fehlt der Einstieg, Rollout erklären. Kein statisches Mockup als interaktives Ergebnis ausgeben."),
        "sharepoint-dg-knowledge": (
            "In an approved existing Dangerous Goods - Procedures demo site, publish the five DG training "
            "documents and 4-6 draft pages using them. No shipping/safety advice. Obtain owner approval for site access, "
            "Site AI settings and Restricted Content Discovery. No site or tenant policy is changed by this installer.",
            "In einer genehmigten bestehenden Demo-Site Dangerous Goods - Procedures die fünf DG-Übungsdokumente "
            "und 4-6 darauf beruhende Entwurfsseiten bereitstellen. Keine Transport-/Sicherheitsberatung. "
            "Sitezugriff, Site-AI-Einstellungen und Restricted Content Discovery vom Eigentümer freigeben lassen. "
            "Der Installer ändert weder Site noch Tenant-Richtlinien."),
        "chat-dg-notebook": (
            "Create DG process review manually in Copilot Notebooks. Add the checklist, approved procedures page, "
            "real saved meeting recap, QBR email and Survey_results.xlsx explicitly. If email reference selection is unavailable, "
            "omit it and state the reduced scope. Generate Audio overview before the session; daily limits may apply.",
            "DG process review in Copilot Notebooks manuell anlegen. Checkliste, genehmigte Verfahrensseite, "
            "echtes gespeichertes Besprechungsergebnis, QBR-Mail und Survey_results.xlsx ausdrücklich hinzufügen. "
            "Ist die Mail-Auswahl nicht verfügbar, weglassen und geringeren Umfang nennen. Audioüberblick vorher erzeugen; Tageslimits möglich."),
        "chat-booking-vision": (
            "Print booking_4711.pdf or show it on a second screen. Test mobile voice, camera/microphone permissions "
            "and screen mirroring; do not expose personal or production content.",
            "booking_4711.pdf ausdrucken oder auf zweitem Bildschirm zeigen. Mobile Spracheingabe, Kamera-/Mikrofonrechte "
            "und Bildschirmspiegelung prüfen; keine persönlichen oder produktiven Inhalte zeigen."),
        "chat-servicenow-tickets": (
            "In an approved existing ServiceNow developer instance create the synthetic CSV incidents. "
            "Set the presenter's actual mapped identity in Opened by. Configure the synchronized Tickets connector, "
            "restricted user access and a full crawl. No connector, external system or identity mapping is created by this installer.",
            "In einer genehmigten bestehenden ServiceNow-Entwicklungsinstanz die synthetischen CSV-Incidents anlegen. "
            "Tatsächlich zugeordnete Präsentatoridentität in Opened by setzen. Synchronisierten Tickets-Connector, "
            "eingeschränkten Zugriff und vollen Crawl konfigurieren. Der Installer legt weder Connector noch externes System oder Identitätsmapping an."),
        "chat-model-comparison": (
            "Verify model availability with the presenter. Only an authorized admin may approve Anthropic "
            "subprocessor access; EU/EFTA/UK defaults and EU Data Boundary exclusions must be reviewed. Do not change policy for this demo.",
            "Modellverfügbarkeit mit dem Vorführkonto prüfen. Nur berechtigte Admins dürfen Anthropic freigeben; "
            "EU/EFTA/UK-Standards und EU-Data-Boundary-Ausnahme prüfen. Für diese Demo keine Richtlinie ändern."),
        "chat-shipping-research": (
            "Start the actual Researcher report at least 45 minutes before the session; review citations and "
            "save its history. Researcher: 25 queries/user/month. Do not portray the draft FAQ as legal authority.",
            "Echten Researcher-Bericht mindestens 45 Minuten vor dem Termin starten; Zitate prüfen und Verlauf behalten. "
            "Researcher: 25 Anfragen pro Benutzer/Monat. FAQ-Entwurf nicht als Rechtsquelle darstellen."),
        "agents-dg-skill": (
            "Check MC1493223 and the presenter's Agents and Skills UI before rehearsal. Only demonstrate direct "
            "Chat skill creation/use after tenant verification. Agent Builder custom skills are a separate Frontier preview. "
            "Keep the dg-pre-check ZIP ready; do not claim automatic installation or create a substitute agent silently.",
            "Vor der Probe MC1493223 und Agents and Skills im Vorführkonto prüfen. Direkte Chat-Skill-Erstellung/"
            "Nutzung erst nach Tenant-Prüfung vorführen. Agent-Builder-Skills sind eine getrennte Frontier-Preview. "
            "dg-pre-check-ZIP bereithalten; keine automatische Installation oder stillen Ersatzagenten behaupten."),
        "agents-dg-builder": (
            "Use the approved existing Dangerous Goods - Procedures site, publish the DG DOCX sources, "
            "and grant the approved DG Validation group site access separately. Build/test the agent and share "
            "Can chat only after approval. Paste real site/agent links into the private Desktop link page; no invented URL.",
            "Genehmigte bestehende Site Dangerous Goods - Procedures verwenden, DG-DOCX-Quellen veröffentlichen "
            "und genehmigter DG-Validation-Gruppe gesondert Sitezugriff geben. Agent bauen/testen und erst nach Freigabe "
            "mit Can chat teilen. Echte Site-/Agentlinks privat in die Desktop-Linkseite eintragen; keine erfundene URL."),
    }
    for folder_id, (en, de) in manual.items():
        path = ROOT / folder_id / "tenant.json"
        config = json.loads(path.read_text(encoding="utf-8"))
        config["manual"] = {"en": [en], "de": [de]}
        write_json(path, config)


if __name__ == "__main__":
    build_sources()
    build_folders()
    print(f"Built shipping sources and {len(FILES) + 1} self-contained demo folders. "
          "Run tools\\Update-Manifest.ps1 -All.")
