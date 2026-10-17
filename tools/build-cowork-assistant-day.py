"""Generate independent assistant-day inputs, skill, presenter guide and JSON page.

Local only. Uses the repository's deterministic Office helpers, not Cowork.
Run tools\\Update-Manifest.ps1 -All after generation.
"""
import argparse
import hashlib
import importlib.util
import json
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
ID = "cowork-assistant-day"
PACKAGE = ROOT / ID
DOMAIN = "M365CPI98544940.OnMicrosoft.com"
ACCOUNT = f"admin@{DOMAIN}"
DATE = "2026-10-17"
SKILL = "contoso-assistant-day"
CAST = {"Teresa Sac": "TeresaS", "Vance DeLeon": "VanceD", "Sonia Rees": "SoniaR",
        "Billie Vester": "BillieV", "Sydney Mattos": "SydneyM"}
OVERVIEW = "https://learn.microsoft.com/microsoft-365/copilot/cowork/"
USE = OVERVIEW + "use-cowork"
FAQ = OVERVIEW + "cowork-faq"
CUSTOMIZE = OVERVIEW + "cowork-customize"
SOURCES = [
    ("Cowork capabilities", "Cowork-Funktionen", OVERVIEW),
    ("Sources, saved work and one-time approvals", "Quellen, gespeicherte Arbeit und einzelne Freigaben", USE),
    ("Permissions and limitations", "Berechtigungen und Grenzen", FAQ),
    ("License, access and billing prerequisites", "Lizenz-, Zugriffs- und Abrechnungsvoraussetzungen", OVERVIEW + "get-started"),
    ("Upload and select a custom skill", "Eigenen Skill hochladen und auswählen", CUSTOMIZE),
    ("Executive scenario inspiration, not Cowork execution evidence", "Executive-Szenarioideen, kein Cowork-Ausführungsnachweis",
     "https://learn.microsoft.com/training/modules/empower-workforce-copilot-executives/"),
    ("Retirement check", "Abkündigungsprüfung",
     "https://support.microsoft.com/en-us/microsoft-365-copilot/frequently-asked-questions-about-retired-copilot-features"),
]
spec = importlib.util.spec_from_file_location("assistant_base", ROOT / "tools" / "build-cowork-usecases.py")
BASE = importlib.util.module_from_spec(spec)
spec.loader.exec_module(BASE)
BASE.FIXED = datetime(2026, 10, 17, 8)


def bi(en, de):
    return {"en": en, "de": de}


def write(path, value):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(value.rstrip() + "\n", encoding="utf-8", newline="\n")


def write_json(path, value):
    write(path, json.dumps(value, ensure_ascii=False, indent=2))


def case(code, title, story, prompt, files, outputs, procedure, human, check, asof="2026-10-19T08:30:00+02:00"):
    return dict(code=code, title=title, story=story, prompt=prompt, files=files, outputs=outputs,
                procedure=procedure, human=human, check=check, asOf=asof)


CASES = [
    case("CO-30", bi("Inbox triage", "Posteingang sortieren"),
         bi("The executive office has six unrelated messages. Two concern the same request; one needs no response. Prepare decisions without making them.",
            "Im Executive Office liegen sechs unterschiedliche Nachrichten. Zwei betreffen dieselbe Anfrage; eine braucht keine Antwort. Entscheidungen vorbereiten, nicht treffen."),
         bi("Please sort these inbox requests for me and prepare the two replies that need my attention. Save the overview so I can review it with Teresa.",
            "Sortiere bitte diese Anfragen im Posteingang und bereite die zwei Antworten vor, um die ich mich kümmern muss. Speichere die Übersicht für meine Abstimmung mit Teresa."),
         ["Inbox_Rules.docx"], ["Inbox_Review.docx"],
         bi("Read only the six CO-30 inbox messages listed in Case.md. Save a compact table with all six IDs, five categories (decision required, meeting request, information, delegable, no action), reason, proposed next step and source. Flag 30-05 as the duplicate of 30-01, not a second approval. Save exactly two actual unsent Outlook drafts: Sydney (30-01) asking for the missing purchase justification without approving EUR 2,200; Teresa (30-02) asking which existing commitment takes priority, without promising a slot. Do not file, flag, delete or move mail.",
            "Nur die sechs in Case.md genannten CO-30-Mails lesen. Kompakte Tabelle mit allen sechs IDs, fünf Kategorien (Entscheidung nötig, Terminwunsch, Information, delegierbar, keine Aktion), Begründung, nächstem Schritt und Quelle speichern. 30-05 als Duplikat von 30-01 markieren, nicht als zweite Freigabe. Genau zwei echte ungesendete Outlook-Entwürfe speichern: an Sydney (30-01), fehlende Begründung erfragen ohne 2.200 EUR freizugeben; an Teresa (30-02), Priorität bestehender Verpflichtungen erfragen ohne einen Termin zuzusagen. Keine Mail verschieben, markieren oder löschen."),
         bi("Choose the escalation and tone personally.", "Eskalation und Ton persönlich entscheiden."),
         bi("Six rows; five categories; duplicate linked; two saved unsent drafts; no approval of EUR 2,200.",
            "Sechs Zeilen; fünf Kategorien; Duplikat zugeordnet; zwei gespeicherte ungesendete Entwürfe; keine Freigabe von 2.200 EUR.")),
    case("CO-31", bi("Morning brief", "Morning Brief"),
         bi("Before the day starts, a supplier decision has moved forward, IT has an update and the office calendar contains a clash.",
            "Vor Tagesbeginn wurde eine Lieferantenentscheidung vorgezogen, IT meldet einen neuen Stand und im eigenen Kalender kollidieren Termine."),
         bi("Please prepare my morning brief for the scenario day. Bring together the calendar, Teresa's update and the current papers, with the decisions I should tackle first.",
            "Bereite bitte meinen Morning Brief für unseren Szenariotag vor. Bring Kalender, Teresas Update und die aktuellen Unterlagen zusammen und zeig mir, welche Entscheidungen zuerst anstehen."),
         ["Morning_Papers.docx", "Calendar_Snapshot.docx"], ["Morning_Brief.docx"],
         bi("Use only CO-31 mail, Morning_Papers and the listed EAD26 own-calendar entries. Treat the supplied snapshot as 19 October 2026 08:30 Europe/Berlin, irrespective of the actual clock. Verify the actual calendar against the selected snapshot; disclose discrepancies. Save a readable one-page Word briefing with top three priorities, changed facts, decision deadlines, source links, missing evidence and proposed work order. Replace EUR 72,000 / five days with EUR 88,000 / eight days from the dated current paper. Teresa's decision deadline is 09:30 today, not tomorrow. Billie's test is reported complete in synthetic evidence, not verified by you. Exclude CO-35 afternoon notes and CO-37 day-close sources.",
            "Nur CO-31-Mails, Morning_Papers und die genannten eigenen EAD26-Kalendereinträge nutzen. Stand ist 19. Oktober 2026, 08:30 Europe/Berlin, unabhängig von der echten Uhrzeit. Tatsächlichen Kalender mit der ausgewählten Momentaufnahme abgleichen; Abweichungen benennen. Lesbaren einseitigen Word-Brief mit Top 3, Änderungen, Entscheidungsfristen, Quellenlinks, Lücken und Arbeitsreihenfolge speichern. 72.000 EUR / fünf Tage durch 88.000 EUR / acht Tage aus dem aktuellen Papier ersetzen. Teresas Entscheidungsfrist ist heute 09:30, nicht morgen. Billies Test ist laut synthetischem Beleg erledigt, nicht von dir geprüft. CO-35-Nachmittagsnotizen und CO-37-Abschlussbelege ausschließen."),
         bi("Decide what to raise with Teresa first.", "Entscheiden, was zuerst persönlich zu Teresa gehört."),
         bi("One page; 09:30 deadline; EUR 88,000/eight days; calendar overlap; no afternoon information.",
            "Eine Seite; Frist 09:30; 88.000 EUR/acht Tage; Kalenderkollision; keine Nachmittagsinformationen.")),
    case("CO-32", bi("Executive meeting brief", "Meeting Brief"),
         bi("A supplier service review needs a decision. The old dashboard looks positive; a newer quality note changes the picture.",
            "Ein Lieferantenreview braucht eine Entscheidung. Das alte Dashboard sieht positiv aus; eine neuere Qualitätsnotiz verändert das Bild."),
         bi("Please get my supplier review briefing ready. Compare the old dashboard with the latest quality update and give me the questions Teresa needs to decide.",
            "Bereite bitte mein Briefing fürs Lieferantenreview vor. Vergleiche das alte Dashboard mit dem aktuellen Qualitätsstand und zeig mir die Fragen, die Teresa entscheiden muss."),
         ["Review_Previous.docx", "Review_Current.docx"], ["Meeting_Brief.docx"],
         bi("Use the two CO-32 messages and the two selected review papers. Save a one-page Word briefing with meeting goal, earlier/current comparison, decisions, three talking points, unknowns and real source links. Keep the periods separate: September 96% on-time and six issues; October-to-date 88% and fourteen issues. Sonia's training capacity remains unconfirmed; do not designate her as approved owner. The supplier review is 20 October at 11:00, not the CO-35 internal process review. No access to anyone else's Teams chats is assumed.",
            "Die zwei CO-32-Mails und beide ausgewählten Review-Papiere nutzen. Einseitigen Word-Brief mit Terminziel, Alt-/Neu-Vergleich, Entscheidungen, drei Talking Points, offenen Fragen und echten Quellenlinks speichern. Zeiträume trennen: September 96 Prozent pünktlich/sechs Fälle; Oktober bisher 88 Prozent/vierzehn Fälle. Sonias Trainingskapazität ist unbestätigt, sie ist keine freigegebene Verantwortliche. Lieferantenreview am 20. Oktober um 11:00, nicht das interne CO-35-Prozessreview. Kein Zugriff auf fremde persönliche Teams-Chats unterstellen."),
         bi("Choose the negotiating position, not just the numbers.", "Die Gesprächsposition bestimmen, nicht nur Zahlen übernehmen."),
         bi("One page; distinct periods; 96 to 88%, six to fourteen; training decision open.",
            "Eine Seite; getrennte Zeiträume; 96 auf 88 Prozent, sechs auf vierzehn; Trainingsentscheidung offen.")),
    case("CO-33", bi("Commitments radar", "Wiedervorlage"),
         bi("Three promises have different evidence: one is closed, one has no completion evidence, and one response is ambiguous.",
            "Drei Zusagen mit unterschiedlicher Beleglage: eine abgeschlossen, eine ohne Abschlussbeleg und eine mehrdeutige Rückmeldung."),
         bi("Please check what became of these three commitments and prepare the follow-ups. Keep uncertain cases separate so I don't chase someone who has already delivered.",
            "Schau bitte nach, was aus diesen drei Zusagen geworden ist, und bereite das Nachfassen vor. Trenne unklare Fälle, damit ich niemanden unnötig erinnere."),
         ["Handover_List.docx"], ["Commitments.docx"],
         bi("Read the five CO-33 messages chronologically using their explicit synthetic timestamps, not delivery time. Save a three-row evidence register: Vance's handover list confirmed delivered by 33-02 and selected Handover_List; Sonia's listening summary has no completion evidence; Billie's 'looks fine' is not explicit accessibility-test completion. Missing evidence is not proof of failure. Save two unsent drafts: a gentle request to Sonia for the summary/status and a clarification to Billie about the test evidence. Do not remind Vance, mark promises complete or invent deadlines.",
            "Die fünf CO-33-Mails anhand ihrer ausdrücklichen synthetischen Zeitstempel lesen, nicht nach Zustellzeit. Drei Zeilen mit Belegen speichern: Vances Übergabeliste durch 33-02 und Handover_List als geliefert bestätigt; für Sonias Gesprächsauswertung fehlt ein Abschlussbeleg; Billies 'sieht gut aus' bestätigt keinen abgeschlossenen Barrierefreiheitstest. Fehlende Belege sind kein Beweis für Versäumnis. Zwei ungesendete Entwürfe: Sonia freundlich nach Auswertung/Stand fragen, Billie um Testbeleg bitten. Vance nicht erinnern, keine Zusage als erledigt markieren und keine Frist erfinden."),
         bi("Choose whom to call rather than send a reminder.", "Entscheiden, wen man lieber persönlich anspricht."),
         bi("Closed / no completion evidence / uncertain separated; no Vance reminder; two unsent drafts.",
            "Abgeschlossen / kein Abschlussbeleg / unklar getrennt; keine Erinnerung an Vance; zwei ungesendete Entwürfe.")),
    case("CO-34", bi("Calendar week", "Kalenderwoche"),
         bi("The office week has an overlap, a building change without travel time and a workshop without an agenda. Other people's availability is unknown.",
            "Die eigene Woche enthält eine Kollision, einen Ortswechsel ohne Wegezeit und einen Workshop ohne Agenda. Fremde Verfügbarkeit ist unbekannt."),
         bi("Please review my scenario week and prepare sensible options for the calendar problems. Draft the coordination, but don't change any appointments.",
            "Prüfe bitte meine Szenariowoche und bereite sinnvolle Optionen für die Kalenderprobleme vor. Entwirf die Abstimmung, ändere aber noch keine Termine."),
         ["Calendar_Snapshot.docx", "Calendar_Rules.docx"], ["Calendar_Options.docx"],
         bi("Inspect only the EAD26 events for 19-23 October in the presenter's own calendar; compare with the selected snapshot. Save a decision paper separating hard overlap, missing travel/preparation buffer and missing agenda. Monday site review 10:00-10:45 overlaps budget 10:30-11:00. Building B visit at 11:00 follows Building A budget without the required 30-minute journey. Monday workshop 13:00 has no agenda. Tuesday supplier review 11:00 follows a protected focus block ending 11:00 with no 15-minute preparation. Offer alternatives as proposals only; do not assert other attendees are available. Save two unsent drafts: Sydney about the budget conflict and Sonia requesting the workshop agenda. Keep calendar IDs, times and attendees unchanged.",
            "Nur EAD26-Termine vom 19. bis 23. Oktober im eigenen Presenter-Kalender prüfen und mit der ausgewählten Momentaufnahme vergleichen. Entscheidungsvorlage mit echter Überschneidung, fehlender Wege-/Vorbereitungszeit und fehlender Agenda speichern. Montag Standortreview 10:00-10:45 kollidiert mit Budget 10:30-11:00. Besuch in Gebäude B um 11:00 folgt auf Budget in Gebäude A ohne nötige 30 Minuten Weg. Montagsworkshop um 13:00 ohne Agenda. Lieferantenreview Dienstag 11:00 folgt auf geschützten Fokusblock bis 11:00 ohne 15 Minuten Vorbereitung. Alternativen nur vorschlagen, keine fremde Verfügbarkeit behaupten. Zwei ungesendete Entwürfe: Sydney zur Budgetkollision, Sonia zur Workshopagenda. Kalender-IDs, Zeiten und Teilnehmende unverändert lassen."),
         bi("Protect the commitments that matter politically.", "Politisch und persönlich wichtige Termine schützen."),
         bi("Four problem types; two actual drafts; all original calendar entries unchanged.",
            "Vier Problemarten; zwei echte Entwürfe; alle Originaltermine unverändert.")),
    case("CO-35", bi("Meeting to action", "Vom Meeting zur Umsetzung"),
         bi("A separate internal process review has released notes, one existing workbook row, and new actions with one missing deadline.",
            "Ein unabhängiges internes Prozessreview hat freigegebene Notizen, eine vorhandene Aufgabenzeile und neue Aufgaben mit einer fehlenden Frist."),
         bi("Please turn these released review notes into decisions and next steps. Update our action register and prepare the individual follow-ups.",
            "Mach bitte aus diesen freigegebenen Review-Notizen die Beschlüsse und nächsten Schritte. Aktualisiere unser Aufgabenregister und bereite das individuelle Nachfassen vor."),
         ["Released_Notes.docx", "@Action_Register.xlsx"], ["Decisions.docx", "@Action_Register.xlsx"],
         bi("Use only the selected released CO-35 notes and active-run original workbook. Save Decisions.docx and edit that exact cloud workbook in place. Preserve sheets, column order, original P01 cell values/types/formats, Readme and table name. Append A01 Sydney Mattos, reconcile handover exceptions, 2026-10-22; A02 Billie Vester, accessibility test evidence, deadline not decided (blank); A03 Sonia Rees, onboarding checklist draft, 2026-10-23. Expand Actions table to A1:F5. D01 is a two-week pilot, D02 no software purchase; neither is approval to send. Save three individualized unsent drafts to the named owners. Billie's draft asks for a due date, not an invented commitment. Do not overwrite the baseline or make a replacement workbook.",
            "Nur ausgewählte freigegebene CO-35-Notizen und Originalmappe des aktiven Laufs nutzen. Decisions.docx speichern und genau diese Cloud-Mappe am bestehenden Ort bearbeiten. Blätter, Spaltenreihenfolge, Werte/Typen/Formate von P01, Readme und Tabellennamen erhalten. Ergänzen: A01 Sydney Mattos, Übergabeausnahmen abgleichen, 2026-10-22; A02 Billie Vester, Barrierefreiheitstest belegen, Frist nicht beschlossen (leer); A03 Sonia Rees, Einarbeitungscheckliste entwerfen, 2026-10-23. Actions-Tabelle auf A1:F5 erweitern. D01 zweiwöchiger Pilot, D02 kein Softwarekauf; keine Sendefreigabe. Drei individuelle ungesendete Entwürfe an die Verantwortlichen speichern. Billie nach Frist fragen statt Zusage zu erfinden. Keine Baseline überschreiben oder Ersatzmappe anlegen."),
         bi("Confirm the missing due date and have the difficult conversation.", "Fehlende Frist klären und das schwierige Gespräch führen."),
         bi("Original workbook identity; exact P01/Readme preservation; A01-A03 and blank A02 date; three drafts.",
            "Identität der Originalmappe; P01/Readme exakt erhalten; A01-A03, A02 ohne Datum; drei Entwürfe."),
         "2026-10-19T14:00:00+02:00"),
    case("CO-36", bi("Short-notice executive visit", "Kurzfristiger Executive-Besuch"),
         bi("Teresa plans an internal site visit on Thursday, with a tight arrival/departure window, a safety constraint and uncertain technical availability.",
            "Teresa plant kurzfristig einen internen Standortbesuch am Donnerstag: enges Zeitfenster, Sicherheitsvorgabe und unbestätigte IT-Verfügbarkeit."),
         bi("Please get Teresa's Thursday visit ready using these messages and the site notes. Save a workable visit pack and draft the coordination for my review.",
            "Bereite bitte Teresas Besuch am Donnerstag mit diesen Nachrichten und Standortnotizen vor. Speichere ein stimmiges Besuchspaket und entwirf die Abstimmung für meine Prüfung."),
         ["Visit_Brief.docx"], ["Visit_Agenda.docx", "Visit_People.docx", "Visit_Decisions.docx"],
         bi("Use only CO-36 evidence. Produce three matching saved documents: one-page agenda for Thursday 22 October 10:00-12:00; one-page participant briefing grounded in synthetic responsibilities; decision list separating approved constraints from open choices. Include Sonia's mandatory safety introduction before the 30-minute walk, 10-minute transitions, and at least 20 minutes of personal discussion. Teresa leaves at 12:00; Billie is available only from 11:15 and his technical demo is not confirmed. Never invent attendee biographies, a room booking or acceptance. Save three unsent drafts: Teresa agenda proposal, Sonia safety/logistics questions, Billie feasibility of the 11:15 demo. No invitations or calendar writes.",
            "Nur CO-36-Belege nutzen. Drei passende gespeicherte Dokumente erstellen: einseitiger Ablauf Donnerstag 22. Oktober 10:00-12:00; einseitiges Teilnehmerbriefing aus synthetischen Zuständigkeiten; Entscheidungsliste mit festen Vorgaben und offenen Punkten. Sonias verpflichtende Sicherheitseinweisung vor dem 30-minütigen Rundgang, zehnminütige Übergänge und mindestens 20 Minuten persönliches Gespräch einplanen. Teresa geht um 12:00; Billie erst ab 11:15 verfügbar, seine Technikdemo ist unbestätigt. Keine Biografien, Raumbuchung oder Zusage erfinden. Drei ungesendete Entwürfe: Teresa Agendavorschlag, Sonia Sicherheits-/Logistikfragen, Billie Machbarkeit der Demo ab 11:15. Keine Einladungen oder Kalendereingriffe."),
         bi("Decide the visit's real purpose and speak to the hosts.", "Eigentliches Besuchsziel entscheiden und Gastgeber persönlich einbinden."),
         bi("Three consistent saved documents; safety before walk; 10:00-12:00, 20 conversation minutes; three drafts.",
            "Drei stimmige gespeicherte Dokumente; Sicherheit vor Rundgang; 10:00-12:00, 20 Gesprächsminuten; drei Entwürfe."),
         "2026-10-19T14:30:00+02:00"),
    case("CO-37", bi("Day closure", "Tagesabschluss"),
         bi("A deliberately synthetic end-of-day checkpoint records one completed decision, an unsent draft and an unresolved test. It is not a record of earlier Cowork runs.",
            "Ein ausdrücklich synthetischer Tagesstand belegt eine erledigte Entscheidung, einen ungesendeten Entwurf und einen offenen Test. Kein Protokoll vorheriger Cowork-Läufe."),
         bi("Please prepare my day closure from this checkpoint and get tomorrow's morning brief started. Separate completed work, open approvals and anything the evidence doesn't settle.",
            "Bereite bitte meinen Tagesabschluss aus diesem Zwischenstand vor und fang den Morning Brief für morgen an. Trenne Erledigtes, offene Freigaben und alles, was die Belege nicht klären."),
         ["Day_Checkpoint.docx", "Tomorrow_Calendar.docx"], ["Day_Closure.docx", "Tomorrow_Brief.docx"],
         bi("Standalone route: use only CO-37 synthetic checkpoint and tomorrow calendar, as of 19 October 17:00. These are staged source records, not Cowork output evidence. Save a one-page closure and one-page draft briefing for 20 October. C01 purchase declined by Teresa in the supplied source; C02 draft prepared but NOT sent/approved; C03 accessibility test completion unknown. The 11:00 review needs 15 minutes preparation; financial reconciliation remains due 22 October. Give top three proposed priorities and source links; no claimed real task completion. Optional real-results route only when the user explicitly selects actual saved results AND separate human approval records: label it as a different route, do not merge staged and real statuses or infer that a draft was sent. No automation.",
            "Eigenständiger Start: nur synthetischen CO-37-Checkpoint und morgigen Kalender mit Stand 19. Oktober 17:00 nutzen. Das sind vorbereitete Quellenbelege, keine Cowork-Ergebnisse. Einseitigen Abschluss und einseitigen Briefentwurf für 20. Oktober speichern. C01 laut Quelle Kauf von Teresa abgelehnt; C02 Entwurf vorbereitet, NICHT gesendet/freigegeben; C03 Abschluss Barrierefreiheitstest unbekannt. Review 11:00 braucht 15 Minuten Vorbereitung; Finanzabgleich bleibt zum 22. Oktober fällig. Drei vorgeschlagene Prioritäten und Quellenlinks nennen; keine tatsächliche Cowork-Erledigung behaupten. Optionaler echter Folgepfad nur bei ausdrücklich ausgewählten tatsächlichen Ergebnissen UND separaten menschlichen Freigabebelegen: getrennt kennzeichnen, synthetische/echte Zustände nicht mischen und Versand nicht aus Entwurf ableiten. Keine Automation."),
         bi("Set tomorrow's priorities and protect relationship time.", "Prioritäten für morgen setzen und Zeit für Beziehungen schützen."),
         bi("Two saved one-page documents; declined / draft only / unknown separated; no inferred sends.",
            "Zwei gespeicherte Einseiter; abgelehnt / nur Entwurf / unklar getrennt; kein behaupteter Versand."),
         "2026-10-19T17:00:00+02:00"),
]

# Original, synthetic Contoso sources; no training binaries or real business data.
DOCS = {
    "Inbox_Rules.docx": bi(
        "Executive Office, 19 October 08:30. Synthetic responsibilities only.\n\nSydney owns cost analysis, not final spending approval. Teresa decides disputed priorities. Sonia coordinates onboarding material. No purchase, time slot or delegation is approved by these rules. Six separate inbox messages are evidence, not a verified Outlook conversation thread.",
        "Executive Office, 19. Oktober 08:30. Nur synthetische Zuständigkeiten.\n\nSydney prüft Kosten, erteilt keine endgültige Ausgabenfreigabe. Teresa entscheidet strittige Prioritäten. Sonia koordiniert Einarbeitungsmaterial. Diese Regeln genehmigen keinen Kauf, Termin oder Delegation. Sechs einzelne Posteingangsmails sind Belege, kein nachgewiesener Outlook-Thread."),
    "Morning_Papers.docx": bi(
        "MORNING sources only. 19 October 2026, 08:20 Europe/Berlin.\n\n17 October planning estimate: Fabrikam replacement capacity EUR 72,000, five-day lead time; decision planned for 20 October.\n\n19 October 08:15 revised quotation: EUR 88,000, eight-day lead time. No order placed. Teresa needs an option decision by 09:30 on 19 October before supplier allocation closes. The quoted amount is not approval.\n\nBillie's 08:20 internal note: scanner connectivity test completed in the synthetic scenario; accessibility testing is a different work item with no completion evidence here. Vance needs a personal call about operational impact. No afternoon decisions available.",
        "Nur MORGEN-Quellen. 19. Oktober 2026, 08:20 Europe/Berlin.\n\nPlanannahme 17. Oktober: Fabrikam-Ersatzkapazität 72.000 EUR, fünf Tage Vorlauf; Entscheidung für 20. Oktober geplant.\n\nAktualisiertes Angebot 19. Oktober 08:15: 88.000 EUR, acht Tage Vorlauf. Kein Auftrag erteilt. Teresa braucht die Optionsentscheidung bis 09:30 am 19. Oktober, bevor die Kapazitätszuordnung schließt. Betrag ist keine Freigabe.\n\nBillies interne Notiz 08:20: Scannerverbindung im synthetischen Szenario getestet; Barrierefreiheit ist ein anderer Vorgang ohne Abschlussbeleg hier. Vance braucht ein persönliches Gespräch über operative Folgen. Keine Nachmittagsentscheidungen verfügbar."),
    "Review_Previous.docx": bi(
        "Supplier service review baseline, 1 October 2026. September reporting period: 96% on-time delivery; six unresolved quality issues at period close. Proposal: renew unchanged. This is the earlier assessment, not today's approval.",
        "Ausgangsstand Lieferantenreview, 1. Oktober 2026. Berichtszeitraum September: 96 Prozent pünktlich; sechs ungelöste Qualitätsfälle zum Monatsende. Vorschlag: unverändert verlängern. Frühere Bewertung, keine heutige Freigabe."),
    "Review_Current.docx": bi(
        "Quality update, 19 October 2026 08:00. October-to-date: 88% on-time delivery; fourteen unresolved quality issues. Periods differ; no claim of an identical monthly sample. Teresa must choose unchanged renewal, a conditional extension or further negotiation at the 20 October 11:00 review. No option is approved.\n\nSonia's capacity for supplier training is not confirmed. Vance has reported two operational bottlenecks, handover timing and missing exception ownership. These are synthetic roles and facts, not personal profiles.",
        "Qualitätsupdate, 19. Oktober 2026 08:00. Oktober bisher: 88 Prozent pünktlich; vierzehn ungelöste Qualitätsfälle. Unterschiedliche Zeiträume, keine Behauptung gleicher Monatsstichprobe. Teresa entscheidet beim Review am 20. Oktober 11:00 über unveränderte Verlängerung, bedingte Verlängerung oder Nachverhandlung. Keine Option freigegeben.\n\nSonias Kapazität für Lieferantentraining ist unbestätigt. Vance meldet zwei operative Engpässe: Übergabezeit und fehlende Zuständigkeit für Ausnahmen. Synthetische Rollen und Fakten, keine persönlichen Profile."),
    "Handover_List.docx": bi(
        "Handover list, version 1. Delivered by Vance, 19 October 2026 08:10, linked to evidence 33-02.\n\nReceiving: Vance DeLeon; safety questions: Sonia Rees; scanner connectivity: Billie Vester. This delivered list closes Vance's promise 33-01 only. It does not establish completion of Sonia's listening summary or Billie's accessibility test.",
        "Übergabeliste, Version 1. Von Vance am 19. Oktober 2026 08:10 geliefert, Beleg 33-02.\n\nWarenannahme: Vance DeLeon; Sicherheitsfragen: Sonia Rees; Scannerverbindung: Billie Vester. Diese gelieferte Liste schließt nur Vances Zusage 33-01. Sie belegt weder Sonias Gesprächsauswertung noch Billies Barrierefreiheitstest."),
    "Calendar_Rules.docx": bi(
        "Own executive-office calendar, 19-23 October 2026; all events are synthetic appointments with no invited attendees. This is not an executive's delegated calendar.\n\nBuilding A to Building B requires 30 minutes. Prepare for the supplier review for 15 minutes. Tuesday focus 10:00-11:00 is protected unless Teresa decides otherwise. A workshop needs a named purpose and agenda; none is supplied for Monday 13:00. No other person's availability or meeting acceptance is established. Never move a calendar event as part of this demo.",
        "Eigener Executive-Office-Kalender, 19.-23. Oktober 2026; alle Termine synthetisch ohne eingeladene Teilnehmende. Kein delegierter Chef-Kalender.\n\nGebäude A nach Gebäude B braucht 30 Minuten. Lieferantenreview braucht 15 Minuten Vorbereitung. Fokus Dienstag 10:00-11:00 geschützt, sofern Teresa nichts anderes entscheidet. Workshop braucht Zweck und Agenda; Montag 13:00 fehlt beides. Keine fremde Verfügbarkeit oder Terminzusage belegt. In dieser Demo niemals einen Termin ändern."),
    "Released_Notes.docx": bi(
        "RELEASED NOTES, internal process review, 19 October 2026 13:00-13:45. Deliberate afternoon checkpoint; not an actual transcript. The own-calendar workshop is a separate earlier source with no published agenda at 08:30. These notes become available only at 14:00.\n\nD01: Teresa approved a two-week handover pilot. D02: No software purchase approved.\n\nA01 Sydney Mattos: reconcile handover exceptions by 22 October 2026. A02 Billie Vester: provide accessibility test evidence; deadline not decided. A03 Sonia Rees: draft onboarding checklist by 23 October 2026. These synthetic assignments are not actual task creation or message approvals.\n\nExisting P01 was completed before this review and must be preserved. Vance should be spoken to personally about adoption friction. Sending the follow-ups is not approved.",
        "FREIGEGEBENE NOTIZEN, internes Prozessreview, 19. Oktober 2026 13:00-13:45. Bewusster Nachmittagsstand, kein tatsächliches Transkript. Der eigene Kalenderworkshop ist eine frühere Quelle ohne veröffentlichte Agenda um 08:30. Diese Notizen sind erst um 14:00 verfügbar.\n\nD01: Teresa genehmigt einen zweiwöchigen Übergabepiloten. D02: Kein Softwarekauf genehmigt.\n\nA01 Sydney Mattos: Übergabeausnahmen bis 22. Oktober 2026 abgleichen. A02 Billie Vester: Barrierefreiheitstest belegen; Frist nicht beschlossen. A03 Sonia Rees: Einarbeitungscheckliste bis 23. Oktober 2026 entwerfen. Synthetische Zuweisungen, keine tatsächliche Aufgabenerstellung oder Nachrichtenfreigabe.\n\nP01 war vor diesem Review erledigt und muss erhalten bleiben. Vance persönlich zu Akzeptanzproblemen ansprechen. Versand des Nachfassens ist nicht freigegeben."),
    "Visit_Brief.docx": bi(
        "Internal site visit, Thursday 22 October 2026, Building B. No external guest. Teresa arrives 10:00 and leaves 12:00. Vance hosts. Sonia requires a ten-minute safety introduction before the thirty-minute walk. Allow ten minutes between separate activities and at least twenty minutes for personal discussion.\n\nBillie can join only from 11:15. A live technical demo is merely an option, not confirmed. A room and any equipment must still be confirmed by Vance. No reservations, invitations, transport or catering orders exist. Participant roles are scenario responsibilities only.",
        "Interner Standortbesuch, Donnerstag 22. Oktober 2026, Gebäude B. Kein externer Gast. Teresa kommt 10:00 und geht 12:00. Vance ist Gastgeber. Sonia verlangt zehn Minuten Sicherheitseinweisung vor dem dreißigminütigen Rundgang. Zehn Minuten zwischen getrennten Aktivitäten und mindestens zwanzig Minuten persönliches Gespräch einplanen.\n\nBillie kann erst ab 11:15. Live-Technikdemo ist eine unbestätigte Option. Raum und Ausstattung muss Vance noch bestätigen. Keine Reservierung, Einladung, Transport- oder Cateringbestellung vorhanden. Teilnehmerrollen sind nur Aufgaben im Szenario."),
    "Day_Checkpoint.docx": bi(
        "STAGED SYNTHETIC SOURCE CHECKPOINT, 19 October 2026 17:00 Europe/Berlin. This is authored input, not genuine Cowork output or proof that earlier demos ran.\n\nC01 15:00: Teresa's supplied decision record declines the EUR 2,200 equipment purchase. Scope: this purchase only, not the separate EUR 88,000 supplier capacity decision. C02 16:00: an onboarding follow-up draft was prepared in the fictional source record; it is not sent and has no human approval. C03 16:30: accessibility test completion remains unknown. A request is not evidence of completion.\n\nFinancial reconciliation remains due 22 October. Tomorrow's supplier review at 11:00 needs a fifteen-minute preparation slot. No automatic morning task has been created. These staged states must never be merged with actual Cowork results without explicit user selection and separate approval evidence.",
        "VORBEREITETER SYNTHETISCHER QUELLENSTAND, 19. Oktober 2026 17:00 Europe/Berlin. Verfasstes Eingangsmaterial, kein echtes Cowork-Ergebnis oder Beweis vorheriger Demos.\n\nC01 15:00: Teresas vorgegebener Entscheidungsvermerk lehnt den 2.200-EUR-Gerätekauf ab. Nur dieser Kauf, nicht die separate 88.000-EUR-Kapazitätsentscheidung. C02 16:00: Im fiktiven Quellenprotokoll ist ein Einarbeitungs-Nachfassentwurf vorbereitet; ungesendet und ohne menschliche Freigabe. C03 16:30: Abschluss Barrierefreiheitstest bleibt unbekannt. Eine Anfrage ist kein Abschlussbeleg.\n\nFinanzabgleich bleibt zum 22. Oktober fällig. Morgen Lieferantenreview 11:00 mit 15 Minuten Vorbereitungsbedarf. Keine automatische Morgenaufgabe angelegt. Synthetische Zustände niemals ohne ausdrückliche Auswahl und getrennte Freigabebelege mit echten Cowork-Ergebnissen vermischen."),
    "Tomorrow_Calendar.docx": bi(
        "Snapshot for Tuesday 20 October 2026, Europe/Berlin. Own EAD26 appointments: protected focus 10:00-11:00; supplier review 11:00-11:45. No preparation buffer supplied. The assistant must decide how to protect fifteen minutes; no time has actually been moved. Other attendees' availability unknown.",
        "Momentaufnahme Dienstag 20. Oktober 2026, Europe/Berlin. Eigene EAD26-Termine: geschützter Fokus 10:00-11:00; Lieferantenreview 11:00-11:45. Keine Vorbereitungszeit vorhanden. Assistenz entscheidet über 15 Minuten Schutzzeit; kein Termin tatsächlich verschoben. Fremde Verfügbarkeit unbekannt."),
}

# Every item is a separate delivered email. Dates in the body are authoritative
# scenario timestamps: sendMail does not backdate actual delivery.
MAILS = [
    ("30-01", "Sydney Mattos", bi("Equipment purchase decision", "Entscheidung Gerätekauf"),
     bi("19 October 08:00. Please decide by 15:00 whether we may buy the EUR 2,200 screen bundle. I have not supplied a business justification or approved spending.",
        "19. Oktober 08:00. Bitte bis 15:00 entscheiden, ob wir das Bildschirmset für 2.200 EUR kaufen dürfen. Begründung noch nicht geliefert; keine Ausgabenfreigabe.")),
    ("30-02", "Teresa Sac", bi("Thirty-minute discussion request", "Wunsch nach 30 Minuten Gespräch"),
     bi("19 October 08:05. I need thirty minutes on Tuesday between 10:00 and 11:00 to discuss office priorities. Please check your commitments first; nothing is agreed.",
        "19. Oktober 08:05. Ich brauche Dienstag zwischen 10:00 und 11:00 dreißig Minuten für Office-Prioritäten. Bitte erst deine Verpflichtungen prüfen; nichts vereinbart.")),
    ("30-03", "Billie Vester", bi("Maintenance completed", "Wartung abgeschlossen"),
     bi("19 October 08:10. The synthetic overnight printer maintenance is complete. Information only; no response or acceptance needed.",
        "19. Oktober 08:10. Synthetische nächtliche Druckerwartung abgeschlossen. Nur Information; keine Antwort oder Abnahme nötig.")),
    ("30-04", "Sonia Rees", bi("Onboarding checklist formatting", "Formatierung Einarbeitungscheckliste"),
     bi("19 October 08:12. Please have the onboarding checklist formatted. I coordinate this material and can handle it if you hand it over. No decision is required today.",
        "19. Oktober 08:12. Bitte die Einarbeitungscheckliste formatieren lassen. Ich koordiniere das Material und kann es nach Übergabe übernehmen. Heute keine Entscheidung nötig.")),
    ("30-05", "Vance DeLeon", bi("Same equipment request, not a second order", "Dieselbe Geräteanfrage, kein zweiter Auftrag"),
     bi("19 October 08:15. Passing on Sydney's EUR 2,200 screen-bundle request [EAD26 30-01]. Same purchase, not another order; please avoid two approvals.",
        "19. Oktober 08:15. Weitergabe von Sydneys Anfrage zum Bildschirmset für 2.200 EUR [EAD26 30-01]. Derselbe Kauf, kein weiterer Auftrag; bitte nicht doppelt freigeben.")),
    ("30-06", "Sonia Rees", bi("Thank you, closed", "Danke, abgeschlossen"),
     bi("19 October 08:18. Thanks for yesterday's directions. Everything is clear; no action or reply needed.",
        "19. Oktober 08:18. Danke für die gestrige Wegbeschreibung. Alles klar; keine Aktion oder Antwort nötig.")),
    ("31-01", "Teresa Sac", bi("Supplier decision moved to 09:30", "Lieferantenentscheidung jetzt 09:30"),
     bi("19 October 08:15. Use today's Morning_Papers: EUR 88,000 and eight-day lead time replace the older estimate. I need a decision recommendation by 09:30 today, not tomorrow. No purchase approved.",
        "19. Oktober 08:15. Heutige Morning_Papers nutzen: 88.000 EUR und acht Tage Vorlauf ersetzen die alte Schätzung. Entscheidungsempfehlung heute bis 09:30, nicht morgen. Kein Kauf genehmigt.")),
    ("31-02", "Billie Vester", bi("Scanner test update", "Stand Scannertest"),
     bi("19 October 08:20. The scanner connectivity test is complete in this synthetic scenario. Accessibility testing is separate and not covered by this message.",
        "19. Oktober 08:20. Scannerverbindung im synthetischen Szenario getestet. Barrierefreiheit ist separat und nicht durch diese Nachricht belegt.")),
    ("32-01", "Teresa Sac", bi("Review decision questions", "Entscheidungsfragen fürs Review"),
     bi("19 October 08:00. For the supplier service review on 20 October at 11:00, compare the previous and current reports. I need options, not an automatic renewal decision.",
        "19. Oktober 08:00. Fürs Lieferantenreview am 20. Oktober 11:00 alten und aktuellen Bericht vergleichen. Ich brauche Optionen, keine automatische Verlängerung.")),
    ("32-02", "Sonia Rees", bi("Training capacity not confirmed", "Trainingskapazität unbestätigt"),
     bi("19 October 08:10. My capacity for supplier training has not been confirmed. Please leave that as a question, not an assigned commitment.",
        "19. Oktober 08:10. Meine Kapazität fürs Lieferantentraining ist unbestätigt. Bitte als Frage offenlassen, nicht als Zusage zuordnen.")),
    ("33-01", "Vance DeLeon", bi("Handover list promise", "Zusage Übergabeliste"),
     bi("16 October 14:00. I will provide the handover list by 19 October 09:00.",
        "16. Oktober 14:00. Ich liefere die Übergabeliste bis 19. Oktober 09:00.")),
    ("33-02", "Vance DeLeon", bi("Handover list delivered", "Übergabeliste geliefert"),
     bi("19 October 08:10. My promised handover list is delivered: selected source Handover_List.docx. This closes my 33-01 promise.",
        "19. Oktober 08:10. Meine zugesagte Übergabeliste liegt vor: ausgewählte Quelle Handover_List.docx. Meine Zusage 33-01 ist damit abgeschlossen.")),
    ("33-03", "Sonia Rees", bi("Listening summary promise", "Zusage Gesprächsauswertung"),
     bi("16 October 15:00. I'll provide the listening summary by 19 October 08:00. No later completion evidence is supplied in this source set.",
        "16. Oktober 15:00. Ich liefere die Gesprächsauswertung bis 19. Oktober 08:00. In diesem Quellensatz kein späterer Abschlussbeleg vorhanden.")),
    ("33-04", "Billie Vester", bi("Accessibility evidence promise", "Zusage Barrierefreiheitsbeleg"),
     bi("16 October 15:30. I will check accessibility and get back to you. We have not agreed a due date.",
        "16. Oktober 15:30. Ich prüfe Barrierefreiheit und melde mich. Wir haben keine Frist vereinbart.")),
    ("33-05", "Billie Vester", bi("Looks fine, scope unspecified", "Sieht gut aus, Umfang unklar"),
     bi("19 October 08:15. Looks fine to me. This message does not identify which check, results or date; please do not interpret it as a signed test record.",
        "19. Oktober 08:15. Sieht für mich gut aus. Nachricht nennt weder Prüfung noch Ergebnisse oder Datum; bitte nicht als Testabnahme auslegen.")),
    ("36-01", "Teresa Sac", bi("Thursday internal visit window", "Zeitfenster interner Besuch Donnerstag"),
     bi("19 October 14:00. I can visit Building B on Thursday 22 October from 10:00 to 12:00, then must leave. I want at least twenty minutes of direct conversation, not only presentations.",
        "19. Oktober 14:00. Ich kann Donnerstag 22. Oktober 10:00-12:00 Gebäude B besuchen und muss dann gehen. Mindestens zwanzig Minuten direktes Gespräch, nicht nur Präsentationen.")),
    ("36-02", "Sonia Rees", bi("Visit safety requirement", "Sicherheitsvorgabe Besuch"),
     bi("19 October 14:10. The thirty-minute walk requires a ten-minute safety introduction beforehand. Plan ten-minute transitions. Room and equipment are still to be confirmed.",
        "19. Oktober 14:10. Vor dem dreißigminütigen Rundgang sind zehn Minuten Sicherheitseinweisung nötig. Zehn Minuten Übergänge planen. Raum und Technik noch zu bestätigen.")),
    ("36-03", "Billie Vester", bi("Only available from 11:15", "Erst ab 11:15 verfügbar"),
     bi("19 October 14:20. I can join only from 11:15 on Thursday. A technical demo might be possible; I have not confirmed it.",
        "19. Oktober 14:20. Donnerstag kann ich erst ab 11:15 dazukommen. Technikdemo vielleicht möglich, noch nicht bestätigt.")),
]

# No attendees: setup creates private synthetic appointments, never invitations.
EVENTS = [
    ("site", "2026-10-19", "10:00", 45, bi("Site review", "Standortreview"), "Building A", bi("Review handover risks.", "Übergaberisiken prüfen.")),
    ("budget", "2026-10-19", "10:30", 30, bi("Budget discussion", "Budgetgespräch"), "Building A", bi("Review spending options with Sydney; no approval.", "Ausgabenoptionen mit Sydney prüfen; keine Freigabe.")),
    ("visit-b", "2026-10-19", "11:00", 30, bi("Building B walk", "Rundgang Gebäude B"), "Building B", bi("Vance hosts the walk.", "Vance begleitet den Rundgang.")),
    ("workshop", "2026-10-19", "13:00", 60, bi("Internal workshop", "Interner Workshop"), "Building A", bi("Agenda not yet supplied at the morning snapshot.", "Zum Morgenstand noch keine Agenda geliefert.")),
    ("focus", "2026-10-20", "10:00", 60, bi("Protected focus", "Geschützter Fokus"), "Building A", bi("Protected unless Teresa decides otherwise.", "Geschützt, sofern Teresa nichts anderes entscheidet.")),
    ("supplier", "2026-10-20", "11:00", 45, bi("Supplier service review", "Lieferantenreview"), "Building A", bi("Renewal options; prepare for 15 minutes.", "Verlängerungsoptionen; 15 Minuten Vorbereitung.")),
    ("writing", "2026-10-21", "09:00", 60, bi("Office writing block", "Office-Schreibzeit"), "Building A", bi("Flexible proposal, not permission to move.", "Flexibler Vorschlag, keine Verschiebefreigabe.")),
    ("thursday", "2026-10-22", "10:00", 120, bi("Internal visit planning hold", "Planungsblock interner Besuch"), "Building B", bi("Own planning hold, no attendance acceptance.", "Eigener Planungsblock, keine Teilnehmerzusage.")),
    ("friday", "2026-10-23", "10:00", 30, bi("Weekly office review", "Office-Wochenreview"), "Building A", bi("Review open decisions.", "Offene Entscheidungen prüfen.")),
]


def subject(mail, lang):
    return f"[EAD26 {mail[0]}] {mail[2][lang]}"


def case_mails(c):
    return [m for m in MAILS if m[0].startswith(c["code"][3:] + "-")]


def document(path, title, text, lang):
    d = BASE.doc_with(title, "Contoso | EAD26",
                      bi("Synthetic business source, not a Cowork result.", "Synthetische Geschäftsquelle, kein Cowork-Ergebnis.")[lang],
                      [(bi("Source record", "Quellenvermerk")[lang], text.split("\n\n"))])
    for style in ("Normal", "Title", "Heading 1"):
        d.styles[style].font.name = "Arial"
    d.styles["Normal"].font.size = Pt(11)
    zoom = d.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    section = d.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(1.8)
    section.left_margin = section.right_margin = Cm(2)
    path.parent.mkdir(parents=True, exist_ok=True)
    BASE.save_doc(d, path)


def workbook(path, lang):
    wb = Workbook()
    ws = wb.active
    ws.title = "Actions"
    ws.append(bi(["ID", "Task", "Owner", "Due date", "Status", "Source"],
                 ["ID", "Aufgabe", "Verantwortlich", "Frist", "Status", "Quelle"])[lang])
    ws.append(["P01", bi("Confirm review inputs", "Review-Eingaben bestätigen")[lang], "Vance DeLeon",
               "2026-10-16", bi("Completed", "Erledigt")[lang], "EAD26-P01"])
    table = Table(displayName="OfficeActions", ref="A1:F2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    for n, width in enumerate([12, 54, 26, 18, 18, 35], 1):
        ws.column_dimensions[chr(64 + n)].width = width
    for row in ws:
        for cell in row:
            cell.font = Font(name="Arial", size=11, bold=cell.row == 1)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="D9E8F5")
    ws.row_dimensions[2].height = 36
    info = wb.create_sheet("Readme")
    info.append(["Scenario", "EAD26"])
    info.append(["Source", "Released_Notes.docx | CO-35 | 2026-10-19T14:00:00+02:00"])
    info.append(["Invariant", bi("Preserve P01, sheets, table name and this Readme. Append A01-A03 only.",
                                "P01, Blätter, Tabellenname und Readme erhalten. Nur A01-A03 ergänzen.")[lang]])
    info.append(["Safety", bi("Synthetic source; no approval to send or create appointments.",
                             "Synthetische Quelle; keine Versand- oder Terminefreigabe.")[lang]])
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 95
    for row in info:
        for cell in row:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(wrap_text=True)
    path.parent.mkdir(parents=True, exist_ok=True)
    BASE.save_wb(wb, path)


def run_context(lang, run):
    return bi(
        f"# EAD26 active run {run}\n\nScenario: Contoso Executive Office. Language: {lang}. "
        f"Read this selected cloud original's actual location: its parent is runs/{lang}/{run}. "
        "Save this run's new files under its outputs/<exercise-code>/ folder. Create the case subfolder if needed; "
        "never replace an existing result. Include this run name and exercise code in unsent draft subjects. "
        "The only editable business input is this run's Action_Register.xlsx for CO-35. "
        "Read the selected case's Case.md for its snapshot and exact source allowlist. "
        "A new conversation starts each case; earlier generated files are not required. "
        "Do not read sibling cases or previous runs. Baseline is immutable. No send/post/invite/book/pay/automation.",
        f"# EAD26 aktiver Lauf {run}\n\nSzenario: Contoso Executive Office. Sprache: {lang}. "
        f"Tatsächliche Cloud-Ablage dieses ausgewählten Originals lesen: sein Elternordner ist runs/{lang}/{run}. "
        "Neue Dateien unter outputs/<Übungscode>/ dieses Laufs speichern. Fall-Unterordner bei Bedarf anlegen; "
        "vorhandene Ergebnisse niemals ersetzen. Laufname und Übungscode in Betreffzeilen ungesendeter Entwürfe aufnehmen. "
        "Einzige bearbeitbare Geschäftseingabe: Action_Register.xlsx dieses Laufs für CO-35. "
        "Stand und exakte Quellenliste stehen in der ausgewählten Case.md. "
        "Jeder Fall startet im neuen Verlauf; frühere Ergebnisse sind keine Voraussetzung. "
        "Keine Nachbarfälle oder älteren Läufe lesen. Baseline unveränderlich. Nicht senden/posten/einladen/buchen/bezahlen/automatisieren.")[lang]


def case_context(c, lang):
    lines = [f"# {c['code']} — {c['title'][lang]}", c["story"][lang],
             f"Scenario snapshot: {c['asOf']} | Europe/Berlin | EAD26",
             bi("Only this case's selected files and listed inbox subjects are allowed. Exact file list:",
                "Nur ausgewählte Dateien dieses Falls und genannte Posteingangsbetreffzeilen nutzen. Exakte Dateiliste:")[lang]]
    lines.extend(f"- {f.removeprefix('@')}" for f in c["files"])
    lines.append("- Run_Context.md (selected active cloud run)")
    lines.append(bi("Inbox evidence (separate messages, not a verified conversation thread):",
                    "Posteingangsbelege (einzelne Nachrichten, kein nachgewiesener Gesprächsverlauf):")[lang])
    lines.extend(f"- {subject(m, lang)}" for m in case_mails(c))
    if c["code"] in ("CO-31", "CO-34"):
        lines.append(bi("Own calendar: only [EAD26 site], [EAD26 budget], [EAD26 visit-b], [EAD26 workshop], "
                        "[EAD26 focus], [EAD26 supplier], [EAD26 writing], [EAD26 thursday], [EAD26 friday], 19-23 October 2026.",
                        "Eigener Kalender: nur [EAD26 site], [EAD26 budget], [EAD26 visit-b], [EAD26 workshop], "
                        "[EAD26 focus], [EAD26 supplier], [EAD26 writing], [EAD26 thursday], [EAD26 friday], 19.-23. Oktober 2026.")[lang])
    lines.append(bi(
        "The dated source content determines chronology; actual email delivery happens at setup. "
        "Later dated explicit updates supersede earlier facts only within the same business scope. "
        "If a listed live source is missing, stop and identify it. Never fill gaps from another case, public web research, "
        "private executive chats or unselected outputs. These inputs are not action approvals.",
        "Chronologie nach datiertem Quelleninhalt; echte Mailzustellung erfolgt beim Setup. "
        "Spätere ausdrückliche Updates ersetzen frühere Fakten nur zum selben Geschäftsvorgang. "
        "Fehlt eine genannte Live-Quelle, stoppen und benennen. Keine Lücken aus anderen Fällen, Webrecherche, "
        "privaten Chef-Chats oder nicht ausgewählten Ergebnissen füllen. Eingaben sind keine Aktionsfreigaben.")[lang])
    return "\n\n".join(lines)


def skill_text():
    parts = [f"---\nname: {SKILL}\ndescription: Prepare executive-office inbox triage, morning and meeting briefs, commitments, calendar options, meeting actions, internal visits or day closure from selected EAD26 case sources. Save real work and unsent drafts; run only the requested case.\n---",
             "# Contoso assistant day",
             "## Shared procedure / Gemeinsames Verfahren\n\n"
             "- Use the request language. Execute exactly ONE matching case CO-30 to CO-37, never the entire day.\n"
             "- For EAD26 use this custom skill, not the separate LEF26 event skills; built-in Office, search and email skills remain available.\n"
             "- Require the selected Case.md and cloud-original Run_Context.md; read their actual locations and allowlists. "
             "Use the explicit scenario snapshot, never the computer clock as 'today'. No sibling/later-phase contamination.\n"
             "- Read actual inbox messages with the exact listed subjects. Verify all required sources before writing. "
             "Separate message subjects are not proof of a real Outlook conversation. Mail delivery timestamps are not scenario chronology.\n"
             "- Save actual Word files under the selected run's outputs/<code>/; create the case folder only if needed. "
             "No overwrites. If prior results exist, require a fresh run. Report verified saved cloud links, not previews or hypothetical paths. "
             "Use Arial 11 and readable A4 layouts; keep one-page outputs genuinely one page. Include real source links; never invent URLs.\n"
             "- For drafts, use only the relevant named approved demo mailboxes below. Include [EAD26 <code> <run>] in subjects; "
             "no CC/BCC. Save real unsent Outlook drafts and verify them. A chat text, approval preview or file containing email text is not a saved draft.\n"
             "- Never send, post, invite, move/delete mail, alter calendars, book, pay or create scheduled/event-driven automations. "
             "Reject such approval cards. Review file edits individually one time; never Always allow. Read an original again after any partial failure before retrying.\n"
             "- Preserve baselines and other runs. Only CO-35 may edit the selected active-run original Action_Register.xlsx, "
             "without replacing its cloud identity. Stop and report errors, missing permission, skill or source rather than claiming success.\n"
             "- Missing evidence is not evidence of noncompletion. Do not invent owners, deadlines, availability, purchases, "
             "approvals, delivered messages or completed tests. Distinguish quoted synthetic business records from actual actions performed by you.\n"
             "- Explain the human decision left open. No personal biographies, confidential profiles, inferred job titles or private Teams access.\n\n"
             "Approved draft recipients (scenario roles only):\n"
             + "\n".join(f"- {name}: {alias}@{DOMAIN}" for name, alias in CAST.items())]
    for lang, label in (("en", "English procedures"), ("de", "Deutsche Verfahren")):
        parts.append(f"## {label}")
        for c in CASES:
            parts.extend([f"### {c['code']} {c['title'][lang]}", c["procedure"][lang],
                          "Files: " + ", ".join(f.removeprefix("@") for f in c["outputs"]), c["human"][lang]])
    return "\n\n".join(parts) + "\n"


def preparation(lang):
    return bi(
        "# Assistant day: start here\n\n"
        "Demo Kit: Cowork end to end > Executive assistant: A day with Cowork. The separate leadership-event page remains available.\n\n"
        "Use the standard global tenant setup, first with -WhatIf. It installs every demo; do not run a per-case installer. "
        "Confirm the signed-in approved demo admin and current Exchange licenses. Actual mail appears in this admin's Inbox with [EAD26 nn-nn]; "
        "appointments appear in this admin's own calendar on 19-23 October 2026, under [EAD26 key]. Files: Demo-Kit > Cowork > cowork-assistant-day.\n\n"
        "One-time skill setup: download skills/contoso-assistant-day.zip. In Cowork select Customize > Skills > arrow next to Add > Upload skill. "
        "Choose the ZIP and keep Only you. If that name already exists, inspect it instead of creating a numbered duplicate. "
        "Start a new conversation and select the skill in Sources; verify it is active in the side panel. UI translations/rollout require your tenant check.\n\n"
        "For each case, attach runs/en/01/Run_Context.md, inputs/en/<code>/Case.md and exactly the files named there through "
        "Add attachments > Attach cloud files. For CO-35 also select that run's original Action_Register.xlsx. "
        "Installing data does not select it in Cowork. Check the listed inbox messages/calendar before starting. "
        "Use the short request in Demo Kit; inspect persistent files and Outlook Drafts outside chat.\n\n"
        "Runs 01-03 are independent clean starting slots. Use a new slot after any edits. "
        "If used up, the local Reset-CoworkAssistantDay.ps1 -Language en -WhatIf previews a non-destructive fresh-run reset; "
        "remove -WhatIf only after review. It uses the existing cached authorized Graph context, never starts sign-in, "
        "downloads and SHA-256 checks the original baseline, and creates a unique run. "
        "No calendar or mail reset is needed because live tasks may not change them. If they did, stop; do not silently reinstall.\n\n"
        "The installer updates baselines/fixtures and is NOT a reset; existing run workbooks and contexts are preserved even when edited. "
        "New reset runs are not manifest-managed and earlier results are retained. "
        "For a reset, the ZIP-uploaded skill is unchanged; re-register only an intentional skill revision, avoiding duplicates.\n\n"
        "All eight live checks remain OPEN. No new Cowork backups are supplied. "
        "Prerequisites: active Copilot license, Cowork enabled, usage billing, reachable cloud files, inbox discovery and your own calendar. "
        "No delegation, private executive Teams chats, browser-control add-on or Frontier App skill is required. "
        "Stop on missing sources or send/post/invite/move/automation approvals. Presenter expectations are outside this folder in tools/expected.",
        "# Assistenzalltag: hier starten\n\n"
        "Demo Kit: Cowork Ende zu Ende > Assistenzalltag: Ein Tag mit Cowork. Die separate Seite Executive Office: Leadership-Event bleibt erhalten.\n\n"
        "Globales Tenant-Setup verwenden, zuerst mit -WhatIf. Es installiert alle Demos; kein Installer pro Fall. "
        "Angemeldeten freigegebenen Demo-Admin und aktuelle Exchange-Lizenzen prüfen. Mails liegen im Posteingang dieses Admins mit [EAD26 nn-nn]; "
        "Termine im eigenen Kalender dieses Admins am 19.-23. Oktober 2026 unter [EAD26 key]. Dateien: Demo-Kit > Cowork > cowork-assistant-day.\n\n"
        "Skill einmalig: skills/contoso-assistant-day.zip herunterladen. In Cowork Anpassen > Skills > Pfeil neben Hinzufügen > Skill hochladen wählen. "
        "ZIP auswählen, Nur du beibehalten. Existiert der Name schon, vorhandenen Skill prüfen statt nummeriertes Duplikat anzulegen. "
        "Neuen Verlauf starten und Skill in Quellen auswählen; Aktivierung rechts prüfen. UI-Übersetzungen und Rollout im eigenen Tenant prüfen.\n\n"
        "Pro Fall über Anlagen hinzufügen > Clouddateien anfügen runs/de/01/Run_Context.md, inputs/de/<code>/Case.md und genau deren Dateien auswählen. "
        "Für CO-35 zusätzlich Action_Register.xlsx desselben Laufs als Cloud-Original auswählen. "
        "Installation wählt keine Quellen in Cowork aus. Genannte Mails/Kalender vor dem Start prüfen. "
        "Kurzen Auftrag aus dem Demo Kit verwenden, gespeicherte Dateien und Outlook-Entwürfe außerhalb des Chats öffnen.\n\n"
        "Läufe 01-03 sind unabhängige saubere Startplätze. Nach Änderungen neuen Platz nutzen. "
        "Sind alle benutzt, zeigt Reset-CoworkAssistantDay.ps1 -Language de -WhatIf lokal einen nicht löschenden Fresh-Run-Reset; "
        "-WhatIf erst nach Prüfung weglassen. Bestehende autorisierte Graph-Sitzung wird wiederverwendet, keine Anmeldung gestartet. "
        "Originalbaseline wird heruntergeladen und per SHA-256 geprüft; eindeutiger Lauf wird neu angelegt. "
        "Mail-/Kalenderreset unnötig, weil Live-Aufträge nichts daran ändern dürfen. Falls doch: stoppen, nicht still neu installieren.\n\n"
        "Installer aktualisiert Baselines/Fixtures und ist KEIN Reset; vorhandene Laufmappen und Kontexte bleiben auch nach Bearbeitung erhalten. "
        "Neue Reset-Läufe gehören nicht zum Manifest; ältere Ergebnisse bleiben erhalten. "
        "Skill beim Reset nicht erneut registrieren; nur bewusste Skill-Revisionen austauschen, keine Duplikate.\n\n"
        "Alle acht Live-Abnahmen OFFEN. Keine neuen Cowork-Backups mitgeliefert. "
        "Voraussetzungen: aktive Copilot-Lizenz, Cowork freigeschaltet, Nutzungsabrechnung, erreichbare Clouddateien, Mailfindung und eigener Kalender. "
        "Keine Delegation, privaten Chef-Chats, Browser-Erweiterung oder Frontier-App-Skill nötig. "
        "Bei fehlenden Quellen oder Senden/Posten/Einladen/Verschieben/Automation-Freigaben stoppen. Presenter-Erwartungen außerhalb dieses Ordners unter tools/expected.")[lang]


def runbook(lang):
    parts = [bi("# CSMAD presenter selection", "# CSMAD Presenter-Auswahl")[lang],
             preparation(lang), bi(
                 "## Select, do not race through sixteen jobs\n\n"
                 "Pool: eight new independent everyday cases plus the unchanged eight event moments. "
                 "For the demo portion choose CO-30, CO-31, CO-34, CO-35 and CO-36 as the core, "
                 "with CO-24 presentation/change as a contrasting saved artifact and CO-37 as the closing thought. "
                 "CO-32 and CO-33 are reserves, not mandatory extra steps. Start longer jobs before explaining another saved result. "
                 "Suggested 40-minute speaking allocation (not Cowork runtime): 4 minutes framing, 6 inbox, 5 morning, "
                 "6 calendar, 8 meeting-to-action, 6 visit or event-pack contrast, 5 closure/judgment. "
                 "Do not fill the slot with generation waits. Use only your own genuinely saved and verified backups; "
                 "if no accepted backup exists, omit that live case instead of presenting a fabricated result.\n\n"
                 "Existing event pool, unchanged: CO-26 reply; CO-27 calendar; CO-28 conversations; CO-29 travel; "
                 "CO-23 executive briefing/venue comparison; CO-24 event pack; CO-24 speaker change; CO-25 follow-up. "
                 "These remain a separate story with their separate three skills and private historical evidence.",
                 "## Auswählen, nicht durch sechzehn Jobs hetzen\n\n"
                 "Pool: acht neue unabhängige Alltagsfälle plus acht unveränderte Event-Momente. "
                 "Für den Demoanteil CO-30, CO-31, CO-34, CO-35 und CO-36 als Kern auswählen, "
                 "CO-24 Präsentation/Änderung als kontrastierendes gespeichertes Ergebnis und CO-37 als Schlussgedanken. "
                 "CO-32 und CO-33 sind Reserven, keine Pflichtzusatzschritte. Lange Jobs starten, bevor ein anderes gespeichertes Ergebnis erklärt wird. "
                 "Vorschlag für 40 Minuten Redeanteil (keine Cowork-Laufzeit): 4 Minuten Einordnung, 6 Posteingang, 5 Morgen, "
                 "6 Kalender, 8 Meeting-to-Action, 6 Besuch oder Eventpaket-Kontrast, 5 Abschluss/Urteil. "
                 "Keine Generierungswartezeiten als Inhalt einplanen. Nur eigene tatsächlich gespeicherte und geprüfte Backups zeigen. "
                 "Ohne abgenommenes Backup den Live-Fall weglassen statt ein Ergebnis vorzutäuschen.\n\n"
                 "Unveränderter Event-Pool: CO-26 Antwort; CO-27 Kalender; CO-28 Gespräche; CO-29 Reise; "
                 "CO-23 Executive-Briefing/Ortsvergleich; CO-24 Eventpaket; CO-24 Speaker-Änderung; CO-25 Nachbereitung. "
                 "Separate Geschichte mit ihren drei Skills und privaten historischen Belegen.")[lang]]
    for c in CASES:
        parts.extend([f"## {c['code']} — {c['title'][lang]}", c["story"][lang],
                      f"Skill: `{SKILL}` | Snapshot: {c['asOf']}",
                      f"Sources: `inputs/{lang}/{c['code']}/Case.md`, " + ", ".join(
                          f"`runs/{lang}/<active>/{f[1:]}`" if f.startswith("@") else f"`inputs/{lang}/{c['code']}/{f}`" for f in c["files"])
                      + f", `runs/{lang}/<active>/Run_Context.md`.",
                      c["prompt"][lang], c["procedure"][lang],
                      bi("Human judgment: ", "Menschliches Urteil: ")[lang] + c["human"][lang],
                      bi("Approval boundary: create documents/drafts; only CO-35 edits its run workbook. One-time approval only. "
                         "Stop at missing sources, unexpected recipients, send/post/invite/calendar/move/automation actions or inconsistent saved files. "
                         "Record a limitation and switch to an accepted result; no repair marathon.",
                         "Freigabegrenze: Dateien/Entwürfe anlegen; nur CO-35 bearbeitet seine Laufmappe. Nur einzelne Freigaben. "
                         "Bei fehlenden Quellen, unerwarteten Empfängern, Senden/Posten/Einladen/Kalender/Verschieben/Automation oder widersprüchlichen Dateien stoppen. "
                         "Grenze festhalten und zu abgenommenem Ergebnis wechseln; kein Reparaturmarathon.")[lang]])
    parts.append(bi("## Your live acceptance — all OPEN", "## Deine Live-Abnahme — alles OFFEN")[lang])
    parts.append("| Code | Check outside chat / Außerhalb des Chats prüfen | Status |\n|---|---|---|")
    parts.extend(f"| {c['code']} | {c['check'][lang]} | {'OPEN' if lang == 'en' else 'OFFEN'} |" for c in CASES)
    parts.append(bi(
        "Keep actual task links, cloud file identities, draft IDs, screenshots and downloaded accepted outputs privately after your rehearsal. "
        "Check saved documents in their native apps, not only chat previews. No agent UI rehearsal was performed. "
        "Separate three states: documented capability, locally validated inputs, presenter-accepted live output. "
        "Use the explicit synthetic date on stage; 'today' is not a claim that these appointments are current.",
        "Nach deiner Probe tatsächliche Verläufe, Cloud-Dateiidentitäten, Entwurf-IDs, Screenshots und heruntergeladene abgenommene Ergebnisse privat sichern. "
        "Gespeicherte Dokumente in ihren Apps öffnen, nicht nur Chatvorschau. Keine Agenten-UI-Probe durchgeführt. "
        "Drei Zustände trennen: dokumentierte Funktion, lokal geprüfte Eingaben, vom Presenter abgenommenes Live-Ergebnis. "
        "Synthetisches Datum auf der Bühne nennen; 'heute' behauptet keine aktuellen Termine.")[lang])
    parts.append("## Official references\n\n" + "\n".join(f"- {s[2]}" for s in SOURCES))
    return "\n\n".join(parts)


def page():
    steps = [
        {"kind": "callout", "tone": "info", "title": bi("Less screen time. More human interaction.", "Weniger Bildschirmzeit. Mehr menschliche Begegnung."),
         "text": bi("Eight independent office cases, from inbox triage to a coordinated internal visit. They complement the separate leadership-event story. Each starts independently and produces saved work; no delegation or automation is assumed.",
                    "Acht unabhängige Office-Fälle vom Posteingang bis zum koordinierten internen Besuch. Ergänzung zur separaten Leadership-Eventstory. Jeder Fall startet eigenständig und erzeugt gespeicherte Arbeit; keine Delegation oder Automation vorausgesetzt.")},
        {"kind": "checklist", "title": bi("Presenter preflight", "Presenter-Vorbereitung"),
         "checks": [
             {"status": "warn", "label": bi("Local package readiness is not live acceptance", "Lokales Paket ist keine Live-Abnahme"),
              "detail": bi("Install through the global Tenant setup. Check Copilot license, Cowork, usage billing, actual inbox/calendar sources, skill and saved results yourself. All eight live checks are open.",
                           "Über globales Tenant-Setup installieren. Copilot-Lizenz, Cowork, Nutzungsabrechnung, tatsächliche Mail-/Kalenderquellen, Skill und gespeicherte Ergebnisse selbst prüfen. Alle acht Live-Abnahmen offen.")},
             {"status": "warn", "label": bi("Own account, synthetic dates, drafts only", "Eigenes Konto, synthetische Daten, nur Entwürfe"),
              "detail": bi("Scenario week is 19-23 October 2026, not the current clock. Never send/post/invite/move/book/pay or activate automation. No delegated mailbox or private executive chats.",
                           "Szenariowoche 19.-23. Oktober 2026, nicht aktuelle Uhrzeit. Nicht senden/posten/einladen/verschieben/buchen/bezahlen oder automatisieren. Kein delegiertes Postfach, keine privaten Chef-Chats.")}]},
        {"kind": "action", "title": bi("Add one reusable skill", "Einen wiederverwendbaren Skill hinzufügen"),
         "items": bi(
             ["Download **skills › contoso-assistant-day.zip** from **Demo-Kit › Cowork › cowork-assistant-day**.",
              "In Cowork select **Customize › Skills**, then the arrow next to **Add › Upload skill**. Choose the ZIP. Keep **Only you**; inspect an existing same-named skill instead of adding duplicates.",
              "Start a new conversation. Select the skill in **Sources** and verify it is active in the side panel. Check exact labels in your tenant. Read **Preparation-en.md** for files, inbox, calendar and fresh-run reset."],
             ["**skills › contoso-assistant-day.zip** unter **Demo-Kit › Cowork › cowork-assistant-day** herunterladen.",
              "In Cowork **Anpassen › Skills**, dann den Pfeil neben **Hinzufügen › Skill hochladen** wählen. ZIP auswählen. **Nur du** beibehalten; vorhandenen gleichnamigen Skill prüfen statt Duplikate anzulegen.",
              "Neuen Verlauf starten. Skill in **Quellen** auswählen und rechts Aktivierung prüfen. Genaue Bezeichnungen im Tenant prüfen. **Preparation-de.md** beschreibt Dateien, Posteingang, Kalender und Fresh-Run-Reset."])} ,
        {"kind": "screen", "title": bi("Where to add the skill (schematic)", "Wo der Skill hinzugefügt wird (Schema)"),
         "screen": {"app": bi("Cowork — Customize › Skills", "Cowork — Anpassen › Skills"), "platform": "Web",
                    "url": "copilot.cloud.microsoft", "tabs": [], "button": bi("Add", "Hinzufügen"),
                    "contextMenu": [bi("Create new", "Neu erstellen"), bi("Upload skill", "Skill hochladen")],
                    "menuHighlight": bi("Upload skill", "Skill hochladen"),
                    "markers": [{"target": "button", "text": bi("Open the arrow next to **Add** on **Skills**.", "Auf **Skills** den Pfeil neben **Hinzufügen** öffnen.")},
                                {"target": "context-menu", "text": bi("Select **Upload skill** and choose the ZIP.", "**Skill hochladen** und ZIP wählen.")}]}},
        {"kind": "screen", "title": bi("Select the real cloud sources", "Die echten Cloudquellen auswählen"),
         "screen": {"layout": "chat", "app": "Microsoft Copilot", "platform": "Web", "url": "copilot.cloud.microsoft",
                    "mode": "cowork", "markers": [
                        {"target": "mode-toggle", "text": bi("Select **Cowork**.", "**Cowork** auswählen.")},
                        {"target": "input", "text": bi("Use **Add attachments › Attach cloud files**. Installation does not select a file.", "**Anlagen hinzufügen › Clouddateien anfügen** verwenden. Installation wählt keine Datei aus.")}]}},
    ]
    for c in CASES:
        prep = bi([], [])
        for lang in ("en", "de"):
            prep[lang] = [
                bi(f"New conversation; select **{SKILL}**. Under **Demo-Kit › Cowork › {ID}**, choose **runs › {lang} › 01 › Run_Context.md** (or a fresh unused run).",
                   f"Neuer Verlauf; **{SKILL}** auswählen. Unter **Demo-Kit › Cowork › {ID}** die Datei **runs › {lang} › 01 › Run_Context.md** auswählen (oder einen frischen unbenutzten Lauf).")[lang],
                bi(f"Attach **inputs › {lang} › {c['code']} › Case.md** and these files from that folder: ",
                   f"**inputs › {lang} › {c['code']} › Case.md** und diese Dateien aus demselben Ordner anfügen: ")[lang]
                + ", ".join(f"**{f}**" for f in c["files"] if not f.startswith("@")),
            ]
            if "@Action_Register.xlsx" in c["files"]:
                prep[lang].append(bi("Also select **Action_Register.xlsx** from the SAME active run as a cloud original, not the baseline.",
                                     "Zusätzlich **Action_Register.xlsx** aus DEMSELBEN aktiven Lauf als Cloud-Original auswählen, nicht die Baseline.")[lang])
            if case_mails(c):
                prep[lang].append(bi("Verify the following actual inbox subjects: ", "Diese tatsächlichen Posteingangsbetreffzeilen prüfen: ")[lang]
                                  + "; ".join(f"**{subject(m, lang)}**" for m in case_mails(c)))
            if c["code"] in ("CO-31", "CO-34"):
                prep[lang].append(bi("Check your own **[EAD26 key]** calendar appointments on **19-23 October 2026**. Do not select someone else's calendar.",
                                     "Eigene **[EAD26 key]**-Kalendertermine am **19.-23. Oktober 2026** prüfen. Keinen fremden Kalender auswählen.")[lang])
            prep[lang].append(bi(f"Use scenario snapshot **{c['asOf']}**. No previous live result is needed; do not attach other case folders.",
                                 f"Szenariostand **{c['asOf']}** verwenden. Kein vorheriges Live-Ergebnis nötig; keine anderen Fallordner anfügen.")[lang])
        steps.extend([
            {"kind": "section", "title": c["title"], "exercise": {"code": c["code"], "chatTitle": c["title"], "story": c["story"]}},
            {"kind": "prompt", "title": bi("Delegate this work", "Diese Arbeit übergeben"),
             "text": c["prompt"], "preparation": prep,
             "where": bi("Cowork, with the selected skill and exact cloud originals.", "Cowork mit ausgewähltem Skill und genauen Cloud-Originalen."),
             "expect": c["check"]},
            {"kind": "compare", "before": {"label": bi("Cowork's legwork", "Coworks Vorarbeit"),
                                          "text": c["check"]},
             "after": {"label": bi("The assistant's judgment", "Urteil der Assistenz"), "text": c["human"]}},
        ])
    for m in MAILS:
        steps.append({"kind": "content",
                      "title": bi(f"Email from {m[1]} (in the inbox)", f"Mail von {m[1]} (im Posteingang)"),
                      "text": {lang: f"Subject: {subject(m, lang)}\n\n{m[3][lang]}" for lang in ("en", "de")}})
    steps.append({"kind": "content", "title": bi("Own calendar (scenario reference)", "Eigener Kalender (Szenarioreferenz)"),
                  "text": {lang: "\n\n".join(
                      f"[EAD26 {key}] {title[lang]} | {day} {start} | {minutes} min | {place}\n{body[lang]}"
                      for key, day, start, minutes, title, place, body in EVENTS) for lang in ("en", "de")}})
    return {"id": ID, "band": "usecases", "accent": "cowork", "data": ID,
            "title": bi("Executive assistant: A day with Cowork", "Assistenzalltag: Ein Tag mit Cowork"),
            "kicker": bi("Eight independent jobs, simple to advanced", "Acht unabhängige Aufgaben, einfach bis erweitert"),
            "why": bi("Prepare real work while keeping decisions, trust and relationships human.",
                      "Echte Arbeit vorbereiten, Entscheidungen, Vertrauen und Beziehungen menschlich halten."),
            "steps": steps, "sources": [{"label": bi(en, de), "url": url} for en, de, url in SOURCES]}


def build_data():
    for lang in ("en", "de"):
        calendar = "\n\n".join(f"[EAD26 {key}] {title[lang]} | {day} {start} | {minutes} min | {place}\n{body[lang]}"
                              for key, day, start, minutes, title, place, body in EVENTS)
        for c in CASES:
            folder = PACKAGE / "inputs" / lang / c["code"]
            write(folder / "Case.md", case_context(c, lang))
            for filename in c["files"]:
                if filename.startswith("@"):
                    continue
                content = calendar if filename == "Calendar_Snapshot.docx" else DOCS[filename][lang]
                document(folder / filename, filename.removesuffix(".docx").replace("_", " "), content, lang)
        for m in MAILS:
            msg = EmailMessage(policy=SMTP)
            msg["From"] = f"{m[1]} <{CAST[m[1]]}@{DOMAIN}>"
            msg["To"] = f"MOD Administrator <{ACCOUNT}>"
            msg["Subject"] = subject(m, lang)
            msg["Date"] = "Mon, 19 Oct 2026 08:00:00 +0200"
            msg["Message-ID"] = f"<ead26-{lang}-{m[0]}@{DOMAIN}>"
            msg.set_content(bi("Synthetic source; delivery time is setup time, not the scenario timestamp below.\n\n",
                               "Synthetische Quelle; Zustellzeit ist Setup-Zeit, nicht der unten genannte Szenariostand.\n\n")[lang] + m[3][lang])
            path = PACKAGE / f"mails-{lang}" / f"{m[0]}_{m[1].split()[0]}.eml"
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(msg.as_bytes())
        workbook(PACKAGE / "baseline" / lang / "Action_Register.xlsx", lang)
        write(PACKAGE / "baseline" / lang / "Run_Context.md", run_context(lang, "01"))
        for run in ("01", "02", "03"):
            folder = PACKAGE / "runs" / lang / run
            workbook(folder / "Action_Register.xlsx", lang)
            write(folder / "Run_Context.md", run_context(lang, run))
            write(folder / "outputs" / "README.md", bi("Empty result destination; no Cowork outputs supplied.",
                                                       "Leeres Ergebnisziel; keine Cowork-Ausgaben mitgeliefert.")[lang])
        write(PACKAGE / f"Preparation-{lang}.md", preparation(lang))
        write(ROOT / "tools" / "expected" / f"{ID}-runbook-{lang}.md", runbook(lang))
    hashes = {lang: hashlib.sha256((PACKAGE / "baseline" / lang / "Action_Register.xlsx").read_bytes()).hexdigest()
              for lang in ("en", "de")}
    hashes["contexts"] = {lang: hashlib.sha256((PACKAGE / "baseline" / lang / "Run_Context.md").read_bytes()).hexdigest()
                          for lang in ("en", "de")}
    write_json(PACKAGE / "baseline" / "hashes.json", hashes)
    write(PACKAGE / "skills" / SKILL / "SKILL.md", skill_text())
    info = ZipInfo("SKILL.md", (2026, 10, 17, 8, 0, 0))
    info.compress_type = ZIP_DEFLATED
    with ZipFile(PACKAGE / "skills" / f"{SKILL}.zip", "w") as archive:
        archive.writestr(info, skill_text().encode())
    write_json(PACKAGE / "tenant.json", {
        "presenterAccount": ACCOUNT,
        "preserveExistingFiles": [f"runs/{lang}/{run}/{name}" for lang in ("en", "de")
                                  for run in ("01", "02", "03") for name in ("Action_Register.xlsx", "Run_Context.md")],
        "requiredUsers": [f"{alias}@{DOMAIN}" for alias in CAST.values()],
        "mails": [f"mails-{{lang}}/{m[0]}_{m[1].split()[0]}.eml" for m in MAILS],
        "events": [
            {"fixedDate": day, "start": start, "minutes": minutes, "timeZone": "W. Europe Standard Time",
             "scenarioKey": f"EAD26-{key}", "attendees": [],
             "subject": {lang: f"[EAD26 {key}] {title[lang]}" for lang in ("en", "de")},
             "location": place,
             "body": {lang: f"Synthetic EAD26 | {body[lang]}" for lang in ("en", "de")}}
            for key, day, start, minutes, title, place, body in EVENTS],
        "links": [{"name": bi("Microsoft Copilot Cowork", "Microsoft Copilot Cowork"), "url": "https://copilot.cloud.microsoft/"}],
        "manual": {lang: [
            bi("Read Preparation-en.md in the assistant-day folder. Verify the EAD26 Inbox messages and own calendar appointments dated 19-23 October 2026.",
               "Preparation-de.md im Assistenzalltag-Ordner lesen. EAD26-Mails im Posteingang und eigene Kalendertermine vom 19.-23. Oktober 2026 prüfen.")[lang],
            bi("Register contoso-assistant-day.zip once through Customize > Skills > Add menu > Upload skill. Keep Only you. Start a new conversation and select the skill in Sources.",
               "contoso-assistant-day.zip einmalig über Anpassen > Skills > Menü Hinzufügen > Skill hochladen registrieren. Nur du beibehalten. Neuen Verlauf starten, Skill in Quellen auswählen.")[lang],
            bi("Presenter performs all eight live checks. No new Cowork backup results supplied; locally generated fixtures are not execution evidence.",
               "Presenter führt alle acht Live-Prüfungen durch. Keine neuen Cowork-Backupergebnisse; lokal erzeugte Fixtures sind kein Ausführungsnachweis.")[lang]]
                   for lang in ("en", "de")},
    })
    manifest_path = PACKAGE / "manifest.json"
    if not manifest_path.exists():
        write_json(manifest_path, {"title": page()["title"], "next": bi(["Read Preparation-en.md."], ["Preparation-de.md lesen."])})
    write_json(ROOT / "tools" / "expected" / f"{ID}.json", {
        "id": ID, "marker": "EAD26", "readiness": "Locally prepared inputs; all eight live acceptance checks OPEN",
        "researchDate": DATE, "caseContracts": CASES,
        "sourceMatrix": [{"code": c["code"], "documented": [OVERVIEW, USE, FAQ],
                          "local": "Fixture, short request, skill and output contract parity",
                          "tenant": "Presenter checks actual sources, skill activation, saved originals and draft persistence",
                          "liveAcceptance": "OPEN"} for c in CASES],
        "materialResearch": "Official executive training currently uses Northwind Traders in Word. No matching ready-to-install Contoso assistant-day fixture found in the reviewed material. Original synthetic Contoso sources authored; no training binaries redistributed.",
        "workbook": {"preserve": ["P01", "Readme", "OfficeActions"], "append": [
            ["A01", "Sydney Mattos", "2026-10-22"], ["A02", "Billie Vester", None], ["A03", "Sonia Rees", "2026-10-23"]]},
    })
    for name in ("Reset-CoworkAssistantDay.ps1", "DemoCalendar.ps1"):
        source = ROOT / "tools" / name
        if source.exists():
            (PACKAGE / name).write_bytes(source.read_bytes())


def build_deck(decks):
    path = decks / "cowork.json"
    deck = json.loads(path.read_text(encoding="utf-8"))
    for file in decks.glob("*.json"):
        other = json.loads(file.read_text(encoding="utf-8"))
        for demo in other["demos"]:
            if demo["id"] == ID:
                continue
            for e in [demo.get("exercise"), *(s.get("exercise") for s in demo["steps"])]:
                if e and e["code"] in {c["code"] for c in CASES}:
                    raise ValueError(f"Exercise collision {file.name}/{demo['id']}/{e['code']}")
    view = page()
    deck["demos"] = [d for d in deck["demos"] if d["id"] != ID]
    deck["demos"].insert(0, view)
    deck["agenda"]["items"] = [i for i in deck["agenda"]["items"] if i["title"] != view["title"]]
    deck["agenda"]["items"].insert(0, {"title": view["title"], "detail": view["kicker"]})
    addition = bi("Also includes eight independent assistant-day cases CO-30 to CO-37, from inbox triage to day closure.",
                  "Außerdem acht unabhängige Assistenzalltagsfälle CO-30 bis CO-37, vom Posteingang bis zum Tagesabschluss.")
    for lang in ("en", "de"):
        if addition[lang] not in deck["meta"]["scope"][lang]:
            deck["meta"]["scope"][lang] += " " + addition[lang]
    deck["meta"].update(version="1.9", date=DATE)
    write_json(path, deck)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--decks", type=Path, default=ROOT.parent / "demo-kit-new2026" / "decks")
    args = parser.parse_args()
    build_data()
    build_deck(args.decks)
    print("Built eight independent EAD26 cases CO-30 to CO-37. No tenant actions or Cowork outputs.")
