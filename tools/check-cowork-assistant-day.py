"""Offline input contracts; never a claim of live Cowork acceptance."""
import hashlib
import importlib.util
import json
import re
import subprocess
from datetime import datetime, timedelta
from email import policy
from email.parser import BytesParser
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
spec = importlib.util.spec_from_file_location("assistant_builder", ROOT / "tools" / "build-cowork-assistant-day.py")
B = importlib.util.module_from_spec(spec)
spec.loader.exec_module(B)
P = B.PACKAGE


def read(path):
    return json.loads(path.read_text(encoding="utf-8"))


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


deck_path = ROOT.parent / "demo-kit-new2026" / "decks" / "cowork.json"
deck = read(deck_path)
page = next(d for d in deck["demos"] if d["id"] == B.ID)
assert page == B.page(), "Canonical page mismatch"
codes = [f"CO-{n}" for n in range(30, 38)]
sections = [s for s in page["steps"] if s["kind"] == "section"]
prompts = [s for s in page["steps"] if s["kind"] == "prompt"]
assert [s["exercise"]["code"] for s in sections] == codes
assert len(prompts) == 8
assert all(s["kind"] in {"section", "prompt"} and "screen" not in s for s in page["steps"])
for prompt in prompts:
    for lang in ("en", "de"):
        assert all(f["kind"] and f["origin"] for f in prompt["files"][lang])
expected = read(ROOT / "tools" / "expected" / f"{B.ID}.json")
assert expected["caseContracts"] == B.CASES
assert len(expected["sourceMatrix"]) == 8
assert all(row["liveAcceptance"] == "OPEN" for row in expected["sourceMatrix"])

# JSON references are framework paths, not filesystem paths.
for path in deck_path.parent.glob("*.json"):
    for d in read(path)["demos"]:
        if d["id"] == B.ID:
            continue
        exercises = [d.get("exercise"), *(s.get("exercise") for s in d["steps"])]
        assert not any(e and e["code"] in codes for e in exercises), f"Collision: {path}/{d['id']}"


def bilingual(value):
    if isinstance(value, dict):
        if "en" in value or "de" in value:
            assert set(value) == {"en", "de"} and value["en"] and value["de"], value
        for child in value.values():
            bilingual(child)
    elif isinstance(value, list):
        for child in value:
            bilingual(child)


bilingual(page)
hashes = read(P / "baseline" / "hashes.json")
tenant = read(P / "tenant.json")
assert tenant["presenterAccount"] == B.ACCOUNT
assert set(tenant["requiredUsers"]) == {f"{alias}@{B.DOMAIN}" for alias in B.CAST.values()}
assert len(tenant["mails"]) == len(B.MAILS) == 18
assert len(tenant["events"]) == len(B.EVENTS) == 9
assert not tenant.get("executiveEvent") and not tenant.get("onenote")
assert len(tenant["preserveExistingFiles"]) == 12
assert all((P / f).is_file() for f in tenant["preserveExistingFiles"])
for lang in ("en", "de"):
    guide = ROOT / "tools" / "expected" / f"{B.ID}-runbook-{lang}.md"
    assert guide.read_text(encoding="utf-8") == B.runbook(lang).rstrip() + "\n"
    assert (P / f"Preparation-{lang}.md").read_text(encoding="utf-8") == B.preparation(lang).rstrip() + "\n"
    assert guide.read_text(encoding="utf-8").count("| OPEN |" if lang == "en" else "| OFFEN |") == 8
    assert sha(P / "baseline" / lang / "Action_Register.xlsx") == hashes[lang]
    assert sha(P / "baseline" / lang / "Run_Context.md") == hashes["contexts"][lang]
    for c, prompt in zip(B.CASES, prompts):
        text = prompt["text"][lang]
        assert text == c["prompt"][lang] and len(text) <= 330
        assert not re.search(r"\n|Act as|#{1,6} |^\s*[-*] ", text)
        assert len(prompt["preparation"][lang]) >= 3
        folder = P / "inputs" / lang / c["code"]
        context = (folder / "Case.md").read_text(encoding="utf-8")
        assert context == B.case_context(c, lang).rstrip() + "\n"
        assert c["asOf"] in context
        assert c["files"] and c["outputs"] and c["procedure"][lang]
        for f in c["files"]:
            if not f.startswith("@"):
                assert (folder / f).is_file(), f
        for m in B.case_mails(c):
            assert B.subject(m, lang) in context
    for rel in tenant["mails"]:
        path = P / rel.replace("{lang}", lang)
        msg = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        assert msg["To"].addresses[0].addr_spec == B.ACCOUNT
        assert not msg["Cc"] and not msg["Bcc"] and not msg["In-Reply-To"] and not msg["References"]
        assert msg["From"].addresses[0].addr_spec in tenant["requiredUsers"]
        assert msg["Subject"].startswith("[EAD26 ")
        assert "setup" in msg.get_content() if lang == "en" else "Setup" in msg.get_content()
    for run in ("01", "02", "03"):
        folder = P / "runs" / lang / run
        assert sha(folder / "Action_Register.xlsx") == hashes[lang]
        assert (folder / "Run_Context.md").read_text(encoding="utf-8") == B.run_context(lang, run).rstrip() + "\n"
        assert {p.name for p in (folder / "outputs").iterdir()} == {"README.md"}

for path in P.rglob("*.docx"):
    doc = Document(path)
    assert len(doc.paragraphs) >= 4 and any(p.text.strip() for p in doc.paragraphs)
    assert doc.styles["Normal"].font.name == "Arial"
    assert 20.9 <= doc.sections[0].page_width.cm <= 21.1
    with ZipFile(path) as z:
        assert not z.testzip()
        assert all(i.date_time == (2026, 10, 17, 8, 0, 0) for i in z.infolist())
    lang = path.parts[path.parts.index("inputs") + 1]
    expected_text = B.DOCS.get(path.name)
    if expected_text:
        assert all(p in "\n\n".join(x.text for x in doc.paragraphs) for p in expected_text[lang].split("\n\n"))

for path in P.rglob("*.xlsx"):
    wb = load_workbook(path)
    assert wb.sheetnames == ["Actions", "Readme"]
    ws = wb["Actions"]
    assert ws.max_row == 2 and ws.max_column == 6
    assert ws["A2"].value == "P01" and ws["C2"].value == "Vance DeLeon"
    assert ws["D2"].value == "2026-10-16" and ws.tables["OfficeActions"].ref == "A1:F2"
    for sheet in wb:
        for row in sheet:
            for cell in row:
                assert cell.data_type != "e" and cell.data_type != "f"
                assert cell.font.name == "Arial"

event_map = {e["scenarioKey"]: e for e in tenant["events"]}
for event in tenant["events"]:
    assert not event["attendees"] and "weekday" not in event
    assert event["timeZone"] == "W. Europe Standard Time"
    assert datetime.fromisoformat(event["fixedDate"]).year == 2026
    assert 0 < event["minutes"] <= 1440
site, budget = event_map["EAD26-site"], event_map["EAD26-budget"]
a = datetime.fromisoformat(site["fixedDate"] + "T" + site["start"])
b = datetime.fromisoformat(budget["fixedDate"] + "T" + budget["start"])
assert max(a, b) < min(a + timedelta(minutes=site["minutes"]), b + timedelta(minutes=budget["minutes"]))
assert event_map["EAD26-visit-b"]["start"] == "11:00"
assert event_map["EAD26-focus"]["start"] == "10:00"
assert event_map["EAD26-supplier"]["start"] == "11:00"

skill = P / "skills" / B.SKILL / "SKILL.md"
assert skill.read_text(encoding="utf-8") == B.skill_text()
with ZipFile(P / "skills" / f"{B.SKILL}.zip") as z:
    assert z.namelist() == ["SKILL.md"] and z.read("SKILL.md") == skill.read_bytes()
    assert len(z.read("SKILL.md")) < 1_000_000
for name in ("Reset-CoworkAssistantDay.ps1", "DemoCalendar.ps1"):
    assert (P / name).read_bytes() == (ROOT / "tools" / name).read_bytes()
manifest = read(P / "manifest.json")
assert {f["path"] for f in manifest["files"]} == {
    p.relative_to(P).as_posix() for p in P.rglob("*") if p.is_file() and p.name not in ("manifest.json", "tenant.json")
}
for entry in manifest["files"]:
    assert sha(P / entry["path"]) == entry["sha256"], entry["path"]
catalog = read(ROOT / "catalog.json")
assert B.ID in json.dumps(catalog)
assert "tools/expected" not in json.dumps(manifest)

# The approved global compact migration preserves existing prompt/exercise identities.
old = subprocess.check_output(["git", "show", "HEAD:decks/cowork.json"], cwd=deck_path.parent.parent)
previous = json.loads(old)
assert [d["id"] for d in deck["demos"]] == [d["id"] for d in previous["demos"]]
for current, original in zip(deck["demos"], previous["demos"]):
    assert [s["text"] for s in current["steps"] if s["kind"] == "prompt"] == [
        s["text"] for s in original["steps"] if s["kind"] == "prompt"], current["id"]
    identity = lambda d: [(e["code"], e["chatTitle"]) for e in
                         [d.get("exercise"), *(s.get("exercise") for s in d["steps"])] if e]
    assert identity(current) == identity(original), current["id"]
for path in (ROOT / "cowork-executive-event" / "skills").rglob("*"):
    if path.is_file():
        committed = subprocess.check_output(["git", "show", f"HEAD:{path.relative_to(ROOT).as_posix()}"], cwd=ROOT)
        assert committed == path.read_bytes(), f"Existing event skill changed: {path}"
print("PASS: eight compact bilingual independent cases, immutable requests/identities, explicit sources, Office inputs, skill, reset copies, manifest/catalog and unchanged event skills. Live acceptance remains OPEN.")
