"""Check installed starter inputs, replay isolation and JSON-view parity."""
import hashlib
import importlib.util
import json
from collections import defaultdict
from copy import deepcopy
from email import policy
from email.parser import BytesParser
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
PACKAGE = ROOT / "cowork-starters"
DECKS = ROOT.parent / "demo-kit-new2026" / "decks"
EXPECTED = json.loads((ROOT / "tools" / "expected" / "cowork-starters.json").read_text(encoding="utf-8"))


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def presentation_content(page):
    return {k: v for k, v in page.items() if k not in {"band", "data"}}


cowork = read_json(DECKS / "cowork.json")
bonn = read_json(DECKS / "cowork-bonn.json")
rfi = read_json(DECKS / "rfi-rfp.json")
page = next(d for d in cowork["demos"] if d["id"] == "cowork-starters")
spec = importlib.util.spec_from_file_location("starter_builder", ROOT / "tools" / "build-cowork-starters.py")
builder = importlib.util.module_from_spec(spec)
spec.loader.exec_module(builder)
assert page == builder.demo_page(), "Compact generator and canonical page differ"
mirror = next(d for d in bonn["demos"] if d["id"] == page["id"])
assert presentation_content(page) == presentation_content(mirror)
assert "data" not in mirror and bonn["files"]["demoData"] == {}
assert presentation_content(rfi["demos"][0]) == presentation_content(
    next(d for d in cowork["demos"] if d["id"] == "cowork-rfi-rfp"))
assert rfi["files"]["demoData"] == {} and "data" not in rfi["demos"][0]
assert not rfi.get("hidden")
rfi_page = next(d for d in cowork["demos"] if d["id"] == "cowork-rfi-rfp")
assert [s["kind"] for s in rfi_page["steps"]] == ["section", "prompt"]
assert set([rfi_page["steps"][0]["anchor"], *rfi_page["steps"][0]["aliases"]]) == {
    "abschnitt-0", "abschnitt-4", "abschnitt-7", "abschnitt-10",
}
for lang in ("en", "de"):
    inputs = rfi_page["steps"][1]["files"][lang]
    assert len(inputs) == 4, "Only the workbook, grouped live KB, active skill and Atlas"
    knowledge = inputs[1]
    expected_knowledge = {
        f"cowork-rfi-rfp/knowledge/{lang}/{p.name}"
        for p in (ROOT / "cowork-rfi-rfp" / "knowledge" / lang).iterdir()
        if p.suffix == ".pdf" or p.name == "Source_Index.md"
    }
    assert len(expected_knowledge) == 14
    assert set(knowledge["materialIds"]) == expected_knowledge
    assert all(Path(material).name in knowledge["location"] for material in expected_knowledge)
    assert inputs[0]["materialIds"] == [
        f"cowork-rfi-rfp/runs/{lang}/compact/01/RFI_Compact_Run_01.xlsx"
    ]
    assert inputs[2]["kind"] == "skill" and inputs[2]["origin"] == "manual"
    assert inputs[3]["kind"] == "site" and inputs[3]["origin"] == "manual"
tutorial = deepcopy(rfi_page)
tutorial["steps"][0]["aliases"].remove("abschnitt-0")
tutorial["steps"][:0] = [
    {"kind": "section", "title": {"en": "Choose a fresh run", "de": "Frischen Lauf wählen"},
     "anchor": "abschnitt-0"},
    {"kind": "action", "items": {"en": ["Read Preparation-en.md."], "de": ["Preparation-de.md lesen."]}},
]
assert builder.rfi_page(tutorial) == rfi_page, "RFI producer must remove reintroduced tutorials"
assert builder.rfi_page(rfi["demos"][0], "rfi-rfp") == rfi["demos"][0]
assert [s["exercise"]["code"] for s in page["steps"] if s["kind"] == "section"] == [
    "CO-19", "CO-20", "CO-21", "CO-22",
]
assert sum(s["kind"] == "prompt" for s in page["steps"]) == 8
for s in page["steps"]:
    assert s["kind"] in {"section", "prompt"} and "screen" not in s
    if s["kind"] == "prompt":
        assert s["preparation"]["en"] and s["preparation"]["de"]
        assert s["text"]["en"] and s["text"]["de"]
        for lang in ("en", "de"):
            assert all(f["kind"] and f["origin"] for f in s["files"][lang])
for f in DECKS.glob("*.json"):
    d = read_json(f)
    assert all(p["band"] in {b["id"] for b in d["bands"]} for p in d["demos"]), f
    assert len({p["id"] for p in d["demos"]}) == len(d["demos"]), f

tenant = read_json(PACKAGE / "tenant.json")
assert len(tenant["mails"]) == 3 and len(tenant["events"]) == 1
event = tenant["events"][0]
assert event["subject"] == EXPECTED["meeting"]["subject"] and event["weekday"] == 4
assert event["start"] == "10:00" and event["minutes"] == 30
assert set(event["attendees"]) == {"Teresa Sac", "Vance DeLeon", "Sonia Rees"}
assert not any("[CO-20]" in str(e["subject"]) for e in tenant["events"])
cast = {"Teresa Sac": "TeresaS", "Vance DeLeon": "VanceD", "Sonia Rees": "SoniaR"}
for lang in ("en", "de"):
    for relative in tenant["mails"]:
        msg = BytesParser(policy=policy.default).parsebytes((PACKAGE / relative.replace("{lang}", lang)).read_bytes())
        sender = msg["From"].addresses[0]
        assert sender.addr_spec == f"{cast[sender.display_name]}@M365CPI98544940.OnMicrosoft.com"
        assert str(msg["Subject"]).startswith("[CO-19]")
        assert msg["To"].addresses[0].username == "admin"
        assert msg.get_body(preferencelist=("plain",)).get_content().strip()
        assert any(str(msg["Subject"]) in f["name"] for s in page["steps"] if s["kind"] == "prompt"
                   for f in s["files"][lang] if f["kind"] == "email")
    assert any(event["subject"][lang] in f["name"] for s in page["steps"] if s["kind"] == "prompt"
               for f in s["files"][lang] if f["kind"] == "event")

baseline = None
paths = []
for lang in ("en", "de"):
    for run in range(1, 5):
        path = PACKAGE / "runs" / lang / f"{run:02}" / "Revenue_Q3.xlsx"
        paths.append(path)
        wb = load_workbook(path, data_only=False)
        try:
            assert wb.sheetnames == ["Revenue", "Readme"]
            rows = list(wb["Revenue"].values)
            assert len(rows) - 1 == EXPECTED["revenue"]["records"]
            assert baseline is None or rows == baseline
            baseline = rows
            assert len(wb["Revenue"].tables) == 1
            assert not wb._external_links
            assert all(c.data_type not in {"e", "f"} for ws in wb for row in ws for c in row)
            assert wb["Readme"]["B5"].value == f"{lang.upper()}-{run:02}"
            assert wb["Readme"]["B6"].value == f"Demo-Kit/Cowork/cowork-starters/runs/{lang}/{run:02}/results"
            assert (path.parent / "results" / "README.md").exists()
            assert all(c.font.name == "Arial" for ws in wb for row in ws for c in row if c.value is not None)
        finally:
            wb.close()
monthly = defaultdict(int)
regions = defaultdict(lambda: {"revenue": 0, "budget": 0})
for month, region, revenue, budget in baseline[1:]:
    monthly[month] += revenue
    regions[region]["revenue"] += revenue
    regions[region]["budget"] += budget
total = sum(monthly.values())
budget = sum(v["budget"] for v in regions.values())
assert total == EXPECTED["revenue"]["total"] and budget == EXPECTED["revenue"]["budget"]
assert total - budget == EXPECTED["revenue"]["variance"]
assert round((total - budget) / budget * 100, 6) == EXPECTED["revenue"]["variance_percent"]
assert monthly == EXPECTED["revenue"]["monthly"]
for values in regions.values():
    values["variance"] = values["revenue"] - values["budget"]
assert regions == EXPECTED["revenue"]["regions"]
assert len({str(p.parent) for p in paths}) == 8

manifest = read_json(PACKAGE / "manifest.json")
for f in manifest["files"]:
    assert hashlib.sha256((PACKAGE / f["path"]).read_bytes()).hexdigest() == f["sha256"], f["path"]
    assert "expected" not in f["path"] and f["path"] != "tenant.json"
assert manifest["shared"] == [{
    "from": "cowork-copilot",
    "path": "Contoso_Fargo_GoLive_Readiness_Tracker.xlsx",
    "sha256": hashlib.sha256((ROOT / "cowork-copilot" / "Contoso_Fargo_GoLive_Readiness_Tracker.xlsx").read_bytes()).hexdigest(),
}]
catalog = read_json(ROOT / "catalog.json")
owners = defaultdict(list)
for app in catalog["apps"]:
    for demo in app["demos"]:
        owners[demo].append(app["id"])
assert owners["cowork-starters"] == owners["cowork-rfi-rfp"] == ["cowork"]
assert all(len(v) == 1 for v in owners.values())
assert all(name in read_json(ROOT / "people.json") for name in cast)
print("Starters: 4 bilingual exercises, 8 isolated source workbooks, mail/calendar references, EUR totals, manifests and view parity verified.")
