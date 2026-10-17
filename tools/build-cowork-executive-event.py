"""Build the bilingual executive-event inputs and their generic Cowork page.

Run this script, then tools\\Update-Manifest.ps1 -All.
Never creates tenant data or fabricates Cowork-generated backup results.
"""
import argparse
import hashlib
import importlib.util
import json
from datetime import datetime
from email.message import EmailMessage
from email.policy import SMTP
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx.oxml.ns import qn
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
ID = "cowork-executive-event"
PACKAGE = ROOT / ID
DOMAIN = "M365CPI98544940.OnMicrosoft.com"
DATE = "2026-10-17"
CAST = {"Teresa Sac": "TeresaS", "Vance DeLeon": "VanceD",
        "Sonia Rees": "SoniaR", "Billie Vester": "BillieV", "Sydney Mattos": "SydneyM"}
USE = "https://learn.microsoft.com/microsoft-365/copilot/cowork/use-cowork"
SOURCES = [
    ("Use Cowork: research, files and saved drafts", "Cowork: Recherche, Dateien und gespeicherte Entwürfe", USE),
    ("Cowork prerequisites", "Cowork-Voraussetzungen",
     "https://learn.microsoft.com/microsoft-365/copilot/cowork/get-started"),
    ("Cowork approval and limitations", "Cowork-Freigaben und Grenzen",
     "https://learn.microsoft.com/microsoft-365/copilot/cowork/cowork-faq"),
    ("Create, upload and select Cowork skills", "Cowork-Skills erstellen, hochladen und auswählen",
     "https://learn.microsoft.com/microsoft-365/copilot/cowork/cowork-customize"),
    ("Official executive scenario inspiration", "Offizielle Executive-Szenarien als Anregung",
     "https://learn.microsoft.com/training/modules/empower-workforce-copilot-executives/"),
    ("Retired Copilot features", "Abgekündigte Copilot-Funktionen",
     "https://support.microsoft.com/en-us/microsoft-365-copilot/frequently-asked-questions-about-retired-copilot-features"),
    ("Official DB station information", "Offizielle DB-Bahnhofsinformationen",
     "https://www.bahnhof.de/frankfurt-main-hbf"),
    ("Le Méridien: operator event information", "Le Méridien: Veranstaltungsinformationen des Betreibers",
     "https://www.marriott.com/de/hotels/framd-le-meridien-frankfurt/events/"),
    ("Frankfurter Hof: operator event information", "Frankfurter Hof: Veranstaltungsinformationen des Betreibers",
     "https://hrewards.com/de/steigenberger-icon-frankfurter-hof/meeting-event"),
]
spec = importlib.util.spec_from_file_location("event_base", ROOT / "tools" / "build-cowork-usecases.py")
BASE = importlib.util.module_from_spec(spec)
spec.loader.exec_module(BASE)
BASE.FIXED = datetime(2026, 10, 17, 8)


def bi(en, de):
    return {"en": en, "de": de}


def write(path, text):
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(text.rstrip() + "\n", encoding="utf-8", newline="\n")


def write_json(path, value):
    write(path, json.dumps(value, ensure_ascii=False, indent=2))


FACTS = {
    "id": ID, "event": "Contoso Leadership Forum Frankfurt",
    "date": "2026-10-22", "timeZone": "W. Europe Standard Time",
    "start": "09:30", "end": "14:00", "headcount": 42, "budgetEUR": 9000,
    "preEventAsOf": "2026-10-19T16:00:00+02:00",
    "postEventAsOf": "2026-10-22T14:30:00+02:00",
    "change": {"speaker": "Teresa Sac", "availableFrom": "11:30", "availableUntil": "12:15", "mode": "remote"},
    "travel": {"person": "Billie Vester", "origin": "Berlin Hbf", "arrivalDate": "2026-10-21",
               "arrivalBy": "18:00", "returnDate": "2026-10-22", "returnNotBefore": "15:00",
               "planningCeilingEUR": 650, "insideEventBudget": True},
}

BRIEF = bi(
    "As of 19 October 2026, 16:00 Europe/Berlin. Synthetic rehearsal, not a real booked event.\n\n"
    "Contoso Leadership Forum Frankfurt takes place on 22 October 2026, 09:30-14:00 Europe/Berlin. "
    "The executive office (MOD Administrator) coordinates 42 participants and a EUR 9,000 total ceiling. "
    "Travel is inside this ceiling. The older 36-person / EUR 8,000 figures are superseded by Teresa's "
    "[LEF26-BEFORE-02] email. Five named demo users represent workstream leads; the remaining 37 participants "
    "are a synthetic aggregate, not people to invite.\n\n"
    "Purpose: agree three priorities for the next quarter and leave with accountable decisions, while "
    "making room for conversations between leaders. There is no confirmed venue, supplier order or travel booking. "
    "Find two real Frankfurt venue options on operator websites and assess suitability for 42 people, "
    "rail arrival, presentation equipment and conversation space. Separate published facts from prices, "
    "availability and details requiring confirmation. Use official DB information for Frankfurt Hauptbahnhof. "
    "Do not contact anyone, submit a form, book or pay.\n\n"
    "Known calendar conflicts: preparation overlaps the executive office's site review on 20 October; the event overlaps "
    "the executive office's investment call on 22 October, 10:00-10:30. Identify the conflict and propose "
    "a decision; do not move appointments or assume delegation. Only the calendar and phase BEFORE sources "
    "belong in the first briefing. Do not use CHANGE, AFTER, presenter expectations or backup answers.",
    "Stand 19. Oktober 2026, 16:00 Europe/Berlin. Synthetische Probe, keine real gebuchte Veranstaltung.\n\n"
    "Das Contoso Leadership Forum Frankfurt findet am 22. Oktober 2026 von 09:30 bis 14:00 Europe/Berlin statt. "
    "Das Executive Office (MOD Administrator) koordiniert 42 Teilnehmende bei einer Gesamtobergrenze von "
    "9.000 EUR einschließlich Anreise. Die älteren Angaben 36 Personen / 8.000 EUR werden durch Teresas "
    "Mail [LEF26-BEFORE-02] ersetzt. Fünf namentliche Demobenutzer vertreten die Arbeitsbereiche; "
    "die weiteren 37 Personen sind eine synthetische Gesamtzahl und keine einzuladenden Personen.\n\n"
    "Ziel: drei Prioritäten für das nächste Quartal und verbindliche Entscheidungen vereinbaren, mit Raum "
    "für persönliche Gespräche zwischen den Führungskräften. Es gibt weder einen bestätigten Veranstaltungsort "
    "noch einen Lieferantenauftrag oder eine Reisebuchung. Zwei reale Veranstaltungsorte in Frankfurt auf "
    "Betreiberseiten recherchieren und auf Eignung für 42 Personen, Bahnanreise, Präsentationstechnik und "
    "Gesprächsflächen prüfen. Veröffentlichte Fakten von Preisen, Verfügbarkeit und klärungsbedürftigen Angaben "
    "trennen. Für Frankfurt Hauptbahnhof offizielle DB-Informationen nutzen. Niemanden kontaktieren, "
    "kein Formular absenden, nichts buchen oder bezahlen.\n\n"
    "Bekannte Kalenderkonflikte: Die Vorbereitung überschneidet sich am 20. Oktober mit dem Standortreview des Executive Office; "
    "das Forum überschneidet sich am 22. Oktober von 10:00 bis 10:30 mit dem Investment-Call des Executive Office. "
    "Konflikt benennen und eine Entscheidung vorschlagen; keine Termine verschieben und keine Vertretung "
    "unterstellen. Ins erste Briefing gehören nur Kalender und Quellen der Phase BEFORE. "
    "CHANGE, AFTER, Presenter-Erwartungen und Reserveantworten nicht verwenden.")

ROLES = bi(
    "Synthetic responsibilities, not the users' real job titles.\n\n"
    "Teresa Sac: leadership sponsor; owns quarterly priorities and the final agenda decision.\n"
    "Vance DeLeon: venue/logistics workstream; confirms feasibility, not a booking authority.\n"
    "Sonia Rees: facilitation and participant communications; protects personal conversation time.\n"
    "Billie Vester: presentation technology and remote-speaker test.\n"
    "Sydney Mattos: budget review; a recommendation is not spending approval.\n"
    "MOD Administrator: executive office and demo presenter; own calendar, no delegated executive mailbox.\n\n"
    "All work is synthetic. No external attendee addresses, personal biographies or sensitive participant data.",
    "Synthetische Aufgaben im Szenario, nicht die tatsächlichen Stellen der Benutzer.\n\n"
    "Teresa Sac: Sponsorin; verantwortet Quartalsprioritäten und die endgültige Agendaentscheidung.\n"
    "Vance DeLeon: Veranstaltungsort und Logistik; prüft Machbarkeit, darf nicht buchen.\n"
    "Sonia Rees: Moderation und Teilnehmerkommunikation; schützt Zeit für persönliche Gespräche.\n"
    "Billie Vester: Präsentationstechnik und Test der Remote-Zuschaltung.\n"
    "Sydney Mattos: Budgetprüfung; eine Empfehlung ist keine Ausgabenfreigabe.\n"
    "MOD Administrator: Executive Office und Presenter; eigener Kalender, kein delegiertes Chef-Postfach.\n\n"
    "Alle Geschäftsinhalte sind synthetisch. Keine externen Teilnehmeradressen, persönlichen Biografien "
    "oder sensiblen Teilnehmerdaten.")

AGENDA = [
    ("09:30", "10:00", bi("Arrival and conversation", "Ankommen und Gespräche")),
    ("10:00", "10:15", bi("Welcome and purpose", "Begrüßung und Ziel")),
    ("10:15", "11:00", bi("Teresa: quarterly priorities", "Teresa: Quartalsprioritäten")),
    ("11:00", "11:15", bi("Break", "Pause")),
    ("11:15", "12:00", bi("Decisions and trade-offs", "Entscheidungen und Zielkonflikte")),
    ("12:00", "12:45", bi("Lunch and conversation", "Mittagessen und Gespräche")),
    ("12:45", "13:30", bi("Leadership exchange", "Austausch der Führungskräfte")),
    ("13:30", "14:00", bi("Owners and next steps", "Verantwortliche und nächste Schritte")),
]

CHANGE = bi(
    "PHASE CHANGE. Reveal only after the first preparation pack has been saved. "
    "Synthetic update for the same event on 22 October 2026.\n\n"
    "Teresa can now join only remotely, 11:30-12:15 Europe/Berlin. Her 45-minute priorities session "
    "must move into that slot. Keep the 09:30 start, 14:00 finish, 42 participants, EUR 9,000 ceiling "
    "and at least 60 minutes of informal conversation. Billie must test the remote connection. "
    "Do not invent a Teams meeting URL or imply that a venue has been booked. Update only the current "
    "run's agenda, presentation and saved drafts; explain what changed and what still needs a human decision.",
    "PHASE CHANGE. Erst zeigen, wenn das erste Vorbereitungspaket gespeichert ist. "
    "Synthetische Änderung für dieselbe Veranstaltung am 22. Oktober 2026.\n\n"
    "Teresa kann jetzt nur remote von 11:30 bis 12:15 Europe/Berlin teilnehmen. Ihre 45-minütige "
    "Prioritätensession muss in dieses Zeitfenster verschoben werden. Beginn 09:30, Ende 14:00, "
    "42 Personen, die Obergrenze von 9.000 EUR und mindestens 60 Minuten informelle Gespräche bleiben bestehen. "
    "Billie muss die Zuschaltung testen. Keinen Teams-Link erfinden und keinen Veranstaltungsort als gebucht "
    "darstellen. Nur Agenda, Präsentation und gespeicherte Entwürfe des aktuellen Laufs aktualisieren; "
    "Änderungen und noch notwendige menschliche Entscheidungen benennen.")

AFTER = bi(
    "PHASE AFTER. Synthetic post-event snapshot, 22 October 2026, 14:30 Europe/Berlin. "
    "This is a deliberate time jump in the rehearsal, not an actual meeting transcript.\n\n"
    "D01: Adopt three priorities: shorten handovers, make ownership visible, preserve direct conversation. "
    "Teresa approved the wording during the fictional forum.\n"
    "D02: Start a four-week handover pilot; no new software purchase was approved.\n"
    "A01: Vance sends the pilot scope to the executive office by 26 October 2026.\n"
    "A02: Sonia prepares a listening session outline by 27 October 2026. No meeting is yet scheduled.\n"
    "A03: Billie documents the remote-speaker check by 23 October 2026.\n"
    "A04: Sydney reconciles actual event costs by 28 October 2026. Invoices are not available; "
    "do not report an actual spend or assert that we were under budget.\n"
    "A05: Executive office asks Teresa to clarify the pilot success measure by 23 October 2026; "
    "a success threshold was not decided.\n\n"
    "Sonia heard that the lunch conversations surfaced more issues than the slide discussion. "
    "The assistant should talk to Vance personally about handover friction rather than delegating that "
    "relationship to a generic follow-up. Prepare messages, do not send, post or create appointments.",
    "PHASE AFTER. Synthetischer Stand nach der Veranstaltung, 22. Oktober 2026, 14:30 Europe/Berlin. "
    "Bewusster Zeitsprung in der Probe, kein tatsächliches Meetingtranskript.\n\n"
    "D01: Drei Prioritäten beschlossen: Übergaben verkürzen, Zuständigkeiten sichtbar machen, "
    "direkte Gespräche schützen. Teresa hat die Formulierung im fiktiven Forum freigegeben.\n"
    "D02: Vierwöchigen Übergabe-Piloten starten; kein neuer Softwarekauf freigegeben.\n"
    "A01: Vance sendet den Pilotumfang bis 26. Oktober 2026 an das Executive Office.\n"
    "A02: Sonia bereitet bis 27. Oktober 2026 ein Konzept für eine Zuhörrunde vor. Noch kein Termin angelegt.\n"
    "A03: Billie dokumentiert bis 23. Oktober 2026 den Test der Remote-Zuschaltung.\n"
    "A04: Sydney gleicht bis 28. Oktober 2026 die tatsächlichen Veranstaltungskosten ab. Rechnungen fehlen; "
    "keine Ist-Ausgaben nennen und keine Budgeteinhaltung behaupten.\n"
    "A05: Executive Office bittet Teresa bis 23. Oktober 2026, das Erfolgskriterium für den Piloten zu klären; "
    "ein Zielwert wurde nicht beschlossen.\n\n"
    "Sonia berichtet, dass die Gespräche beim Mittagessen mehr Themen sichtbar machten als die Folienrunde. "
    "Die Assistenz sollte mit Vance persönlich über Reibungen bei Übergaben sprechen, statt diese Beziehung "
    "an eine allgemeine Nachfassmail zu delegieren. Nachrichten vorbereiten, nichts senden, posten oder terminieren.")

TRAVEL = bi(
    "PHASE BEFORE. Synthetic business-travel request for the same Contoso forum, planning snapshot "
    "19 October 2026, 16:00 Europe/Berlin. These are fictional business requirements, not facts about "
    "Billie's real location or an existing reservation.\n\n"
    "Billie Vester starts this synthetic business trip at Berlin Hbf. Travel by rail and arrive at "
    "Frankfurt(Main) Hbf by 18:00 on 21 October. One hotel night, 21-22 October. On 22 October, plan "
    "to reach the forum venue by 09:15; the event is 09:30-14:00. The venue is still undecided, so the "
    "final local transfer remains open. The return train must leave Frankfurt(Main) Hbf no earlier "
    "than 15:00 on 22 October. Leave sensible transfer buffers; do not assume transport is booked.\n\n"
    "The provisional ceiling for this one trip is EUR 650, already inside the event's EUR 9,000 "
    "ceiling, not extra money. No travel policy, fare, room rate, availability, booking, invoice or "
    "spending approval is supplied. Do not claim policy compliance, an actual cost or remaining event "
    "budget from this ceiling alone. Research public official DB and hotel-operator information, "
    "present two hotel options and the remaining booking questions. If dated train details or rates "
    "cannot be verified, name the outstanding check instead of inventing them. No enquiries, form "
    "submissions, bookings, payments or external contact.",
    "PHASE BEFORE. Synthetischer Dienstreiseauftrag für dasselbe Contoso-Forum, Planungsstand "
    "19. Oktober 2026, 16:00 Europe/Berlin. Fiktive geschäftliche Vorgaben, keine Angaben über "
    "Billies tatsächlichen Standort und keine bestehende Reservierung.\n\n"
    "Billie Vester startet diese synthetische Dienstreise in Berlin Hbf. Bahnanreise mit Ankunft "
    "in Frankfurt(Main) Hbf am 21. Oktober bis 18:00. Eine Hotelnacht vom 21. auf den 22. Oktober. "
    "Am 22. Oktober um 09:15 am Veranstaltungsort sein; das Forum dauert 09:30-14:00. Der Ort ist "
    "noch offen, deshalb bleibt der letzte Transfer zu klären. Rückfahrt ab Frankfurt(Main) Hbf "
    "am 22. Oktober frühestens um 15:00. Sinnvolle Umstiegspuffer vorsehen; keine Buchung unterstellen.\n\n"
    "Vorläufiger Planungsrahmen für diese eine Reise: 650 EUR, bereits in der Gesamtobergrenze "
    "von 9.000 EUR enthalten, kein zusätzliches Budget. Reiserichtlinie, Tarife, Zimmerpreise, "
    "Verfügbarkeit, Buchungen, Rechnungen und Ausgabenfreigaben liegen nicht vor. Weder Richtlinien- "
    "noch Budgeteinhaltung oder Ist-Kosten behaupten; aus dieser Obergrenze keinen Rest des gesamten "
    "Eventbudgets errechnen. Öffentliche offizielle DB- und Hotelbetreiberinformationen recherchieren, "
    "zwei Hoteloptionen und offene Buchungsfragen darstellen. Nicht belegbare datierte Zugdetails "
    "oder Preise als offenen Prüfschritt benennen, statt sie zu erfinden. Keine Anfragen, Formulare, "
    "Buchungen, Zahlungen oder externe Kontaktaufnahme.")

MAILS = [
    ("01_Teresa_Objectives.eml", "Teresa Sac", "01", bi("Objectives and first assumptions", "Ziele und erste Annahmen"),
     bi("Our Frankfurt forum is planned for 22 October, 09:30-14:00. Initial assumption: 36 people, EUR 8,000 including travel. These are provisional; I will confirm the final figures. We need decisions, not a slide marathon.",
        "Unser Frankfurt-Forum ist für den 22. Oktober, 09:30-14:00 geplant. Erste Annahme: 36 Personen, 8.000 EUR einschließlich Anreise. Noch vorläufig; die endgültigen Zahlen bestätige ich. Wir brauchen Entscheidungen, keinen Folienmarathon.")),
    ("02_Teresa_Confirmed_Scope.eml", "Teresa Sac", "02", bi("Confirmed planning scope", "Bestätigter Planungsrahmen"),
     bi("Update to my earlier note: use 42 participants and a EUR 9,000 total ceiling including travel. This replaces 36 / EUR 8,000. The 22 October 09:30-14:00 window stays. Venue selection is still open; please prepare options, not bookings.",
        "Update zu meiner vorigen Nachricht: Bitte mit 42 Teilnehmenden und einer Gesamtobergrenze von 9.000 EUR einschließlich Anreise arbeiten. Das ersetzt 36 / 8.000 EUR. Der 22. Oktober 09:30-14:00 bleibt. Der Ort ist noch offen; bitte Optionen vorbereiten, keine Buchungen.")),
    ("03_Vance_Venue.eml", "Vance DeLeon", "03", bi("Venue decision still open", "Veranstaltungsort noch offen"),
     bi("We have not contacted or booked any venue. Please compare two real Frankfurt options using operator information. Capacity for 42, rail access and room for conversations matter. Published web information is not a quote or proof of availability.",
        "Wir haben noch keinen Veranstaltungsort kontaktiert oder gebucht. Bitte zwei reale Frankfurt-Optionen anhand von Betreiberinformationen vergleichen. Platz für 42, Bahnanreise und Gesprächsflächen sind wichtig. Webangaben sind weder ein Angebot noch ein Verfügbarkeitsnachweis.")),
    ("04_Sonia_Conversation.eml", "Sonia Rees", "04", bi("Protect conversation time", "Gesprächszeit schützen"),
     bi("Please protect at least 60 minutes of informal conversation across arrival, breaks and lunch. Keep the deck to six slides. The purpose is three shared priorities and named decision owners, not a presentation contest.",
        "Bitte mindestens 60 Minuten informelle Gespräche über Ankommen, Pausen und Mittagessen schützen. Höchstens sechs Folien. Ziel sind drei gemeinsame Prioritäten und benannte Entscheidungsverantwortliche, kein Präsentationswettbewerb.")),
    ("05_Billie_Technology.eml", "Billie Vester", "05", bi("Technology is not confirmed", "Technik ist noch nicht bestätigt"),
     bi("No venue or remote-speaker equipment is confirmed. I can validate screen, audio and remote participation after a shortlist decision. Do not promise a working hybrid setup yet.",
        "Ort und Technik für Remote-Beiträge sind nicht bestätigt. Nach einer Shortlist-Entscheidung kann ich Bild, Ton und Zuschaltung prüfen. Bitte noch keine funktionierende Hybridtechnik zusagen.")),
    ("06_Sydney_Budget.eml", "Sydney Mattos", "06", bi("Budget decision and calendar clash", "Budgetentscheidung und Kalenderkonflikt"),
     bi("The EUR 9,000 ceiling includes travel. Venue quotes and actual travel costs are missing. The executive office has an investment call on 22 October 10:00-10:30; please show the conflict rather than silently moving it. A recommendation is not spending approval.",
        "Die Obergrenze von 9.000 EUR schließt Anreise ein. Raumangebote und tatsächliche Reisekosten fehlen. Das Executive Office hat am 22. Oktober 10:00-10:30 einen Investment-Call; bitte den Konflikt zeigen, nicht still verschieben. Eine Empfehlung ist keine Ausgabenfreigabe.")),
]

RUN_PREP = bi(
    "Choose the fresh run from the reset report and attach its **Action_Register.xlsx** through "
    "**Add attachments › Attach cloud files and folders**. It only identifies the run; do not edit it. "
    "The selected skill derives the run name from that workbook's actual folder; do not type technical run identifiers into the live prompt.",
    "Frischen Lauf aus dem Reset-Bericht wählen und dessen **Action_Register.xlsx** über "
    "**Anlagen hinzufügen › Clouddateien und -ordner anfügen** auswählen. Sie dient nur als Laufkennung; "
    "nicht bearbeiten. Der ausgewählte Skill liest den Laufnamen aus dem tatsächlichen Ordner der Mappe; keine technischen Laufkennungen in den Live-Prompt schreiben.")

STARTERS = [
    {
        "code": "CO-26", "title": bi("Get the reply ready", "Die Antwort ist vorbereitet"),
        "chatTitle": bi("Frankfurt planning reply", "Frankfurt Planungsantwort"),
        "story": bi("Teresa has corrected the headcount and budget. You want a short, dependable reply, not another summary to copy by hand.",
                    "Teresa hat Personenzahl und Budget korrigiert. Du brauchst eine kurze, verlässliche Antwort, keine weitere Zusammenfassung zum Abtippen."),
        "prep": bi(
            [RUN_PREP["en"], "Under **Add work context › Emails**, select **[LEF26-BEFORE-02] Confirmed planning scope**. The conversation may appear under MOD Administrator after the setup reply. Use only the BEFORE phase."],
            [RUN_PREP["de"], "Unter **Arbeitskontext hinzufügen › E-Mails** **[LEF26-BEFORE-02] Bestätigter Planungsrahmen** auswählen. Nach der Setup-Antwort kann die Unterhaltung unter MOD Administrator erscheinen. Nur Phase BEFORE verwenden."]),
        "instructions": bi(
            "Please prepare a brief reply to Teresa Sac based on her selected planning update. Confirm the current participant count, total budget including travel and event window, without repeating the superseded figures as current. Ask her which agenda decision she still needs from us. Save exactly one English Outlook draft to Teresa, with [LEF26], this run's name and Planning reply in the subject. Keep the tone personal and concise. A venue and spending have not been approved. Show the saved draft; don't send, post, book anything or change the selected workbook.",
            "Bereite bitte eine kurze Antwort an Teresa Sac auf ihre ausgewählte Planungsnachricht vor. Bestätige die aktuelle Personenzahl, das Gesamtbudget einschließlich Anreise und das Veranstaltungsfenster, ohne die ersetzten Zahlen als aktuell zu übernehmen. Frage, welche Agendaentscheidung sie noch von uns braucht. Speichere genau einen deutschen Outlook-Entwurf an Teresa, mit [LEF26], dem Laufnamen und Planungsantwort im Betreff. Persönlich und knapp, bitte. Ort und Ausgaben sind nicht freigegeben. Zeig den gespeicherten Entwurf; nichts senden, posten oder buchen und die ausgewählte Mappe nicht bearbeiten."),
        "expect": bi("One actual unsent Outlook draft to Teresa; 42 participants, EUR 9,000 including travel, 22 October 09:30-14:00; no venue or spending approval.",
                     "Ein tatsächlich gespeicherter ungesendeter Outlook-Entwurf an Teresa; 42 Personen, 9.000 EUR einschließlich Anreise, 22. Oktober 09:30-14:00; keine Orts- oder Ausgabenfreigabe."),
    },
    {
        "code": "CO-27", "title": bi("Untangle the calendar", "Terminkollisionen entschärfen"),
        "chatTitle": bi("Frankfurt calendar decisions", "Frankfurt Terminabstimmung"),
        "story": bi("The executive office is double-booked during preparation and the forum. Make the decision visible without silently moving anyone's meeting.",
                    "Das Executive Office ist bei der Vorbereitung und während des Forums doppelt verplant. Die Entscheidung soll sichtbar werden, ohne heimlich Termine zu verschieben."),
        "prep": bi(
            [RUN_PREP["en"], "Use a new conversation. Check the presenter's own calendar on **20 and 22 October 2026**, only subjects starting with **[LEF26]**. These four synthetic entries are already installed; no delegated executive calendar is used."],
            [RUN_PREP["de"], "Neuen Verlauf verwenden. Eigenen Kalender des Presenters am **20. und 22. Oktober 2026** prüfen, nur Betreffzeilen mit **[LEF26]**. Die vier synthetischen Einträge sind bereits vorhanden; kein delegierter Führungskräftekalender."]),
        "instructions": bi(
            "Please check my calendar on 20 and 22 October 2026 for the Frankfurt forum, using only entries whose subjects start with [LEF26]. Identify the overlaps with their exact local times in Europe/Berlin. Prepare a concise English Outlook draft to Sonia Rees that explains both conflicts and proposes options for us to agree, without assuming anyone is free or authorized to stand in. Put [LEF26], this run's name and Calendar decisions in the subject. Save one draft and show it. Don't use unrelated appointments, change events, create invitations, send or post anything, or edit the workbook.",
            "Schau bitte für das Frankfurt-Forum in meinen Kalender am 20. und 22. Oktober 2026, ausschließlich auf Einträge, deren Betreff mit [LEF26] beginnt. Benenne die Überschneidungen mit genauen lokalen Zeiten in Europe/Berlin. Bereite einen kurzen deutschen Outlook-Entwurf an Sonia Rees vor, der beide Konflikte und mögliche gemeinsame Lösungen erklärt, ohne freie Zeiten oder eine Vertretungsbefugnis zu unterstellen. In den Betreff gehören [LEF26], der Laufname und Terminabstimmung. Einen Entwurf speichern und zeigen. Keine fremden Termine verwenden, nichts verschieben, einladen, senden oder posten und die Mappe nicht bearbeiten."),
        "expect": bi("One unsent Sonia draft showing 20 October 10:15-10:45 and 22 October 10:00-10:30 overlaps, proposed decisions rather than rescheduling; all four original events unchanged.",
                     "Ein ungesendeter Sonia-Entwurf mit den Überschneidungen 20. Oktober 10:15-10:45 und 22. Oktober 10:00-10:30, Entscheidungsvorschläge statt Terminänderungen; alle vier ursprünglichen Termine unverändert."),
    },
    {
        "code": "CO-28", "title": bi("Prepare for the conversation", "Begegnungen vorbereiten"),
        "chatTitle": bi("Frankfurt conversation preparation", "Frankfurt Gesprächsvorbereitung"),
        "story": bi("You want useful personal conversations at the forum, not a stack of generic biographies or automated connection requests.",
                    "Du möchtest beim Forum gute persönliche Gespräche führen, keine allgemeinen Biografien sammeln oder automatisch Kontaktanfragen verschicken."),
        "prep": bi(
            [RUN_PREP["en"], "Attach the verified SharePoint **Participants.docx** from the setup report's shared BEFORE library, not an assumed OneDrive copy. It contains five named workstream leads and the executive-office role, not LinkedIn profiles or actual job titles. Do not select AFTER notes."],
            [RUN_PREP["de"], "Die geprüfte SharePoint-Datei **Participants.docx** aus der gemeinsamen BEFORE-Bibliothek laut Setup-Bericht auswählen, keine angenommene OneDrive-Kopie. Im Clouddialog bei Bedarf **Weitere Orte...** öffnen und die freigegebene Site und Bibliothek wählen. Enthält fünf namentliche Arbeitsbereichsverantwortliche und die Executive-Office-Rolle, keine LinkedIn-Profile oder tatsächlichen Stellenangaben. Keine AFTER-Notizen auswählen."]),
        "instructions": bi(
            "I'm preparing for personal conversations at our Frankfurt forum. Use only the selected scenario responsibilities to create Conversation_Brief.docx in the outputs folder next to this run's Action_Register.xlsx. In one readable English page, suggest three useful pairings among Teresa, Vance, Sonia, Billie and Sydney, explain the shared business question and give me one natural opening question for each. Make clear these are fictional scenario roles, not researched biographies. Don't invent relationships or post-event decisions, search LinkedIn, send connection requests or contact anyone. Show the actual saved file and leave the workbook unchanged.",
            "Ich bereite mich auf persönliche Gespräche bei unserem Frankfurt-Forum vor. Erstelle bitte nur aus den ausgewählten Aufgaben im Szenario Conversation_Brief.docx im Ordner outputs neben der Action_Register.xlsx dieses Laufs. Schlage auf einer gut lesbaren deutschen Seite drei sinnvolle Gesprächspaarungen zwischen Teresa, Vance, Sonia, Billie und Sydney vor, erkläre die gemeinsame geschäftliche Frage und gib mir jeweils eine natürlich formulierte Einstiegsfrage. Mach deutlich, dass es fiktive Szenariorollen sind, keine recherchierten Biografien. Keine Beziehungen oder späteren Beschlüsse erfinden, nicht auf LinkedIn recherchieren, keine Kontaktanfragen oder Nachrichten versenden. Zeig die tatsächlich gespeicherte Datei und lass die Mappe unverändert."),
        "expect": bi("One persisted, readable one-page conversation brief; three source-grounded pairings/questions, no invented biography, LinkedIn access, contacts or AFTER decisions.",
                     "Ein gespeichertes, gut lesbares einseitiges Gesprächsbriefing; drei aus den Rollen abgeleitete Paarungen und Fragen, keine erfundene Biografie, LinkedIn-Zugriffe, Kontaktaktionen oder AFTER-Beschlüsse."),
    },
    {
        "code": "CO-29", "title": bi("Prepare the business trip", "Die Dienstreise vorbereiten"),
        "chatTitle": bi("Frankfurt business travel proposal", "Frankfurt Reisevorschlag"),
        "story": bi("Billie needs a practical rail-and-hotel plan with buffers and an approval boundary. The assistant keeps the booking decision, rather than trusting invented prices.",
                    "Billie braucht einen brauchbaren Bahn-/Hotelplan mit Puffern und einer klaren Freigabegrenze. Die Assistenz entscheidet über die Buchung, statt erfundenen Preisen zu vertrauen."),
        "prep": bi(
            [RUN_PREP["en"], "Attach **inputs/en/before/Travel_Request.md**. It defines a synthetic Berlin business-trip origin, the 21-22 October journey and a EUR 650 planning ceiling inside EUR 9,000. No actual policy, fare or booking is supplied. Use public official DB and hotel-operator sources only."],
            [RUN_PREP["de"], "**inputs/de/before/Travel_Request.md** auswählen. Enthält den fiktiven geschäftlichen Startpunkt Berlin, die Reise vom 21.-22. Oktober und 650 EUR Planungsrahmen innerhalb der 9.000 EUR. Keine tatsächliche Richtlinie, Tarife oder Buchungen vorhanden. Nur öffentliche offizielle DB- und Hotelbetreiberquellen nutzen."]),
        "instructions": bi(
            "Please turn Billie's selected travel request into a practical English travel proposal. Save Travel_Plan.docx, at most two readable pages, in the outputs folder next to this run's Action_Register.xlsx. Include the outward and return travel windows, the one-night stay, sensible buffers, two real Frankfurt hotel options and clickable official DB and hotel-operator sources. Distinguish verified information from a proposed itinerary and from the dated fares, availability and policy checks still needed. The EUR 650 ceiling is inside the EUR 9,000 event budget, not extra; don't invent quotes, train numbers or proof that the trip fits the ceiling. Leave the final venue transfer open until the venue is chosen. Show the saved file. No external contact, form submissions, bookings, payments, sends, calendar changes or workbook edits.",
            "Mach bitte aus Billies ausgewähltem Reiseauftrag einen brauchbaren deutschen Reisevorschlag. Speichere Travel_Plan.docx mit höchstens zwei gut lesbaren Seiten im Ordner outputs neben der Action_Register.xlsx dieses Laufs. Darin sollen Hin- und Rückreisefenster, die eine Übernachtung, sinnvolle Puffer, zwei reale Frankfurter Hoteloptionen und anklickbare offizielle DB- und Hotelbetreiberquellen stehen. Trenne belegte Informationen vom vorgeschlagenen Ablauf und den noch nötigen Prüfungen datierter Tarife, Verfügbarkeit und Reiserichtlinie. Die 650 EUR sind in den 9.000 EUR Eventbudget enthalten, nicht zusätzlich; keine Angebote, Zugnummern oder angebliche Einhaltung der Obergrenze erfinden. Den letzten Transfer bis zur Ortsentscheidung offenlassen. Zeig die gespeicherte Datei. Niemanden kontaktieren, keine Formulare absenden, buchen, bezahlen, senden, Termine oder die Mappe ändern."),
        "expect": bi("A persisted travel proposal with two operator-backed hotel options and actual source hyperlinks; 21 October arrival by 18:00, one night, return on 22 October not before 15:00; EUR 650 inside EUR 9,000; no fabricated fare, availability, policy approval or booking.",
                     "Ein gespeicherter Reisevorschlag mit zwei betreiberbelegten Hoteloptionen und tatsächlichen Quellenlinks; Ankunft 21. Oktober bis 18:00, eine Nacht, Rückfahrt 22. Oktober frühestens 15:00; 650 EUR innerhalb 9.000 EUR; keine erfundenen Tarife, Verfügbarkeit, Richtlinienfreigaben oder Buchungen."),
    },
]

CHANGE_TITLE = bi("Handle the last-minute change", "Die kurzfristige Änderung auffangen")
CHANGE_PREP = bi(
    ["After reviewing the initial pack, reattach **Leadership_Forum.pptx** and **Agenda.docx** from this run's actual **outputs** folder, plus **inputs/en/change/Speaker_Change.md**. Stay in the same conversation and use the existing Teresa draft. Do not select default Cowork working copies or create replacements."],
    ["Nach Prüfung des ersten Pakets **Leadership_Forum.pptx** und **Agenda.docx** aus dem tatsächlichen Ordner **outputs** dieses Laufs erneut auswählen, dazu **inputs/de/change/Speaker_Change.md**. Im selben Verlauf bleiben und den vorhandenen Teresa-Entwurf verwenden. Keine Cowork-Arbeitskopien auswählen oder Ersatzdateien erzeugen."])

ACTS = [
    {
        "code": "CO-23", "title": bi("From chaos to clarity", "Vom Chaos zur Klarheit"),
        "chatTitle": bi("Frankfurt executive briefing", "Frankfurt Executive-Briefing"),
        "story": bi("The executive office has conflicting updates and an unresolved venue decision. You need a dependable briefing before talking to the sponsor.",
                    "Im Executive Office liegen widersprüchliche Updates und eine offene Ortsentscheidung. Du brauchst ein verlässliches Briefing für das Gespräch mit der Sponsorin."),
        "prep": bi(
            ["Open **Add attachments › Attach cloud files and folders**. Choose an unused **runs/en/01–04** or the fresh run from the reset report, and select its **Action_Register.xlsx** so Cowork can identify the actual run folder. Select the verified SharePoint **Event_Brief.docx**, **Participants.docx** and **Agenda_Baseline.docx**, all phase **BEFORE**.",
             "Through **Add attachments › Add work context › Emails**, select the six **[LEF26-BEFORE-01]** to **[LEF26-BEFORE-06]** conversations. A conversation with a setup reply can appear under the presenter's name.",
             "Under **Meetings**, select the forum and preparation meeting. Check only calendar entries whose subjects begin with **[LEF26]** on 20 and 22 October; unrelated calendar entries are outside this synthetic demo. Not every appointment appears in the meeting picker.",
             "Paste the exact **BEFORE Teams discussion** link from the setup report into the same input. Do not attach the whole channel, which also contains the later AFTER discussion.",
             "Use Cowork on the web in signed-in Edge. This workflow uses built-in web research, not local browser control. Do not select CHANGE, AFTER, backup outputs or presenter expectations."],
            ["**Anlagen hinzufügen › Clouddateien und -ordner anfügen** öffnen. Einen unbenutzten Lauf **runs/de/01–04** oder den frischen Lauf aus dem Reset-Bericht wählen und dessen **Action_Register.xlsx** auswählen, damit Cowork den tatsächlichen Laufordner kennt. Die geprüften SharePoint-Dateien **Event_Brief.docx**, **Participants.docx** und **Agenda_Baseline.docx**, alle Phase **BEFORE**, auswählen.",
             "Über **Anlagen hinzufügen › Arbeitskontext hinzufügen › E-Mails** die sechs Unterhaltungen **[LEF26-BEFORE-01]** bis **[LEF26-BEFORE-06]** auswählen. Unterhaltungen mit einer Setup-Antwort können unter dem Namen des Presenters erscheinen.",
             "Unter **Besprechungen** das Forum und die Vorbereitung auswählen. Am 20. und 22. Oktober nur Kalendereinträge prüfen, deren Betreff mit **[LEF26]** beginnt; andere Termine gehören nicht zu dieser synthetischen Demo. Nicht jeder Kalendereintrag erscheint in der Besprechungsauswahl.",
             "Den genauen Link der **BEFORE-Teams-Diskussion** aus dem Setup-Bericht in dieselbe Eingabe einfügen. Nicht den ganzen Kanal auswählen: Er enthält auch die spätere AFTER-Diskussion.",
             "Cowork im Web in angemeldetem Edge nutzen. Dieser Ablauf verwendet die integrierte Webrecherche, keine lokale Browsersteuerung. CHANGE, AFTER, Reserveergebnisse und Presenter-Erwartungen nicht auswählen."]),
        "instructions": bi(
            "I'm preparing the Contoso Leadership Forum Frankfurt on 22 October. For this rehearsal, use the synthetic planning snapshot of 19 October at 16:00. Please turn the selected BEFORE documents and email conversations and the linked BEFORE Teams discussion into an executive briefing I can use with Teresa. Read only that discussion and its replies, not later scenario phases. Check the calendar on 20 and 22 October, but use only entries whose subjects start with [LEF26]; don't use or mention unrelated appointments. Reconcile the older and newer figures, check those calendar clashes and make the decisions still needed obvious. Research two real Frankfurt venue options using only their operators' websites and rail arrival using only official DB information, not news reports. Use the selected Action_Register.xlsx's existing OneDrive folder to identify this run; don't edit that workbook yet. Save Briefing.docx and Venue_Comparison.docx in its outputs subfolder and show me the actual links. Keep each document to two readable pages with clear tables and clickable source links. Focus on what we need to decide; about 350 words per file is enough. Distinguish published facts from unconfirmed price or availability; leave unsupported claims out. Don't contact anyone, send messages, move appointments, book or pay. Compare Le Méridien Frankfurt at https://www.marriott.com/de/hotels/framd-le-meridien-frankfurt/events/ and Steigenberger Icon Frankfurter Hof at https://hrewards.com/de/steigenberger-icon-frankfurter-hof/meeting-event. Link short source names rather than printing long addresses. Reply and create all file contents in English. Suggested decision owners are proposals, not proof of booking or spending authority.",
            "Ich bereite das Contoso Leadership Forum Frankfurt am 22. Oktober vor. Wir proben mit dem synthetischen Planungsstand vom 19. Oktober um 16:00. Mach bitte aus den ausgewählten BEFORE-Dokumenten und Mail-Unterhaltungen und der verlinkten BEFORE-Teams-Diskussion ein Executive-Briefing für mein Gespräch mit Teresa. Lies nur diese Diskussion mit ihren Antworten, keine späteren Szenariophasen. Prüfe den Kalender am 20. und 22. Oktober, aber nutze ausschließlich Einträge, deren Betreff mit [LEF26] beginnt; andere Termine weder verwenden noch erwähnen. Kläre die alten und neuen Zahlen, prüfe diese Kalenderkonflikte und mach deutlich, welche Entscheidungen noch fehlen. Recherchiere zwei reale Veranstaltungsorte nur auf deren Betreiberseiten und die Bahnanreise nur anhand offizieller DB-Informationen, nicht aus Presseberichten. Der vorhandene OneDrive-Ordner der ausgewählten Action_Register.xlsx bestimmt diesen Lauf; die Mappe noch nicht bearbeiten. Speichere Briefing.docx und Venue_Comparison.docx in dessen Unterordner outputs und zeig mir die tatsächlichen Links. Mach beide Dokumente übersichtlich, mit jeweils höchstens zwei gut lesbaren Seiten, klaren Tabellen und anklickbaren Quellenlinks. Konzentriere dich auf die entscheidungsrelevanten Punkte; ungefähr 350 Wörter pro Datei reichen. Trenne veröffentlichte Fakten von unbestätigten Preisen und Verfügbarkeit; unbelegte Aussagen weglassen. Niemanden kontaktieren, nichts senden, verschieben, buchen oder bezahlen. Vergleiche Le Méridien Frankfurt auf https://www.marriott.com/de/hotels/framd-le-meridien-frankfurt/events/ und Steigenberger Icon Frankfurter Hof auf https://hrewards.com/de/steigenberger-icon-frankfurter-hof/meeting-event. Verlinke kurze Quellennamen, statt lange Adressen abzudrucken. Antworte und erstelle sämtliche Datei-Inhalte auf Deutsch. Vorgeschlagene Entscheidungsverantwortliche sind Vorschläge, kein Nachweis für eine Buchungs- oder Ausgabenbefugnis."),
        "expect": bi("Two persisted Word files, at most two readable pages each; 42 participants, EUR 9,000, both [LEF26] calendar clashes, official-only venue/rail sources, no unrelated appointments and explicitly open commitments.",
                     "Zwei gespeicherte Word-Dateien mit jeweils höchstens zwei gut lesbaren Seiten; 42 Personen, 9.000 EUR, beide [LEF26]-Kalenderkonflikte, nur offizielle Orts-/Bahnquellen, keine fremden Termine und klar offene Zusagen."),
    },
    {
        "code": "CO-24", "title": bi("Let Cowork do the work", "Cowork erledigt die Arbeit"),
        "chatTitle": bi("Frankfurt event preparation pack", "Frankfurt Vorbereitungspaket"),
        "story": bi("The briefing is reviewed. You need presentation, agenda and communication drafts that stay consistent when a speaker's availability changes.",
                    "Das Briefing ist geprüft. Präsentation, Agenda und Kommunikationsentwürfe sollen auch nach einer geänderten Sprecherverfügbarkeit zusammenpassen."),
        "prep": bi(
            ["Use the reviewed **Briefing.docx** and **Venue_Comparison.docx** from CO-23 plus the selected BEFORE sources. If the preceding live act failed, select only genuine Cowork checkpoint files from a successful rehearsal, clearly identified as such.",
             "Select this run's **Action_Register.xlsx**. The skill derives the run name from the actual folder; do not edit baseline source files. No venue or expenditure has been approved."],
            ["Das geprüfte **Briefing.docx** und **Venue_Comparison.docx** aus CO-23 mit den ausgewählten BEFORE-Quellen nutzen. Ist der vorherige Live-Lauf fehlgeschlagen, nur echte Cowork-Checkpointdateien einer erfolgreichen Probe wählen und das offen sagen.",
             "**Action_Register.xlsx** dieses Laufs auswählen. Der Skill liest den Laufnamen aus dem tatsächlichen Ordner; keine Ausgangsquellen bearbeiten. Weder Veranstaltungsort noch Ausgabe sind freigegeben."]),
        "instructions": bi(
            "Please turn the reviewed Frankfurt briefing into an English preparation pack for the same event. Use only the reviewed BEFORE LEF26 sources, not unrelated appointments, later decisions or projects from the conversation. Create Leadership_Forum.pptx with no more than six useful, visual slides and Agenda.docx for 09:30-14:00 Europe/Berlin, protecting at least an hour for informal conversation. Give the slides a clear, light layout with large readable type, a timeline and useful number cards rather than walls of text. Keep each timeline label short, in one paragraph without manual line breaks, so it is easy to update later. The three quarterly priorities are still to be agreed; don't invent decisions. Save both in this run's output folder. Save separate English Outlook drafts to Teresa for the agenda decision and to Vance for the venue questions, with this run's name in each subject. Keep prices, venue and spending approval open where evidence is missing; any proposed decision owners are suggestions, not verified booking authority. Show the files and saved drafts. Don't send or post anything, invite anyone or make bookings.",
            "Mach bitte aus dem geprüften Frankfurt-Briefing ein deutsches Vorbereitungspaket für dieselbe Veranstaltung. Nutze nur die geprüften BEFORE-Quellen von LEF26, keine anderen Termine, späteren Beschlüsse oder Projekte aus dem Verlauf. Erstelle Leadership_Forum.pptx mit höchstens sechs hilfreichen, visuellen Folien und Agenda.docx für 09:30 bis 14:00 Europe/Berlin, mit mindestens einer Stunde für informelle Gespräche. Für die Folien bitte ein klares, helles Layout mit großer, gut lesbarer Schrift, einer Zeitachse und sinnvollen Zahlenkacheln statt Textwänden. Halte die Beschriftungen der Zeitachse kurz, jeweils in einem Absatz ohne händische Zeilenumbrüche, damit sie sich später leicht ändern lassen. Die drei Quartalsprioritäten sind noch zu vereinbaren; keine Beschlüsse erfinden. Speichere beide im Ergebnisordner dieses Laufs. Speichere getrennte deutsche Outlook-Entwürfe an Teresa für die Agendaentscheidung und an Vance für die offenen Ortsfragen, jeweils mit dem Laufnamen im Betreff. Lass Preise, Veranstaltungsort und Ausgabenfreigaben offen, wenn Belege fehlen; vorgeschlagene Entscheidungsverantwortliche sind keine nachgewiesene Buchungsbefugnis. Zeig die Dateien und gespeicherten Entwürfe. Nichts senden, posten, einladen oder buchen."),
        "expect": bi("An actual presentation, agenda and two saved Outlook drafts; no invented venue, price or approval.",
                     "Tatsächliche Präsentation, Agenda und zwei gespeicherte Outlook-Entwürfe; keine erfundenen Räume, Preise oder Freigaben."),
        "changeInstructions": bi(
            "Teresa's availability has changed, as described in the selected Speaker_Change.md. Please update this run's actual saved presentation, agenda and existing agenda-decision draft to Teresa rather than making duplicates. Keep everything in English, the event window, headcount and budget unchanged, protect conversation time and show what moved. Keep the existing layout and formatting; replace the full timeline labels so no old lines remain. The maximum of six slides applies to the whole forum, not just Teresa's session; make that clear in the agenda. Add a saved English draft to Billie asking for the remote-connection check. The test is still pending, not commissioned or completed. Don't invent a meeting link, send anything or change calendar events.",
            "Teresas Verfügbarkeit hat sich geändert, wie in der ausgewählten Speaker_Change.md beschrieben. Aktualisiere bitte die tatsächlich gespeicherte Präsentation, Agenda und den vorhandenen Entwurf zur Agendaentscheidung an Teresa dieses Laufs, statt Duplikate anzulegen. Lass alle Inhalte auf Deutsch und Veranstaltungszeit, Personenzahl und Budget gleich; schütze die Gesprächszeit und zeig, was verschoben wurde. Behalte Layout und Formatierung bei und ersetze die Beschriftungen der Zeitachse vollständig, damit keine alten Zeilen zurückbleiben. Höchstens sechs Folien gelten für das ganze Forum, nicht nur für Teresas Block; mach das in der Agenda deutlich. Speichere zusätzlich einen deutschen Entwurf an Billie mit der Bitte um den Test der Zuschaltung. Der Test bleibt offen, nicht beauftragt oder abgeschlossen. Keinen Meeting-Link erfinden, nichts senden und keine Kalendertermine ändern."),
    },
    {
        "code": "CO-25", "title": bi("Back to being human", "Wieder Zeit für Menschen"),
        "chatTitle": bi("Frankfurt follow-up and commitments", "Frankfurt Nachbereitung"),
        "story": bi("We now jump deliberately to the event's synthetic post-event notes. Cowork records commitments while you concentrate on a personal conversation with Vance.",
                    "Wir springen bewusst zu den synthetischen Notizen nach der Veranstaltung. Cowork hält Zusagen fest, während du dich auf das persönliche Gespräch mit Vance konzentrierst."),
        "prep": bi(
            ["Say explicitly: **We now move to the prepared post-event snapshot of the same event.** Select only **inputs/en/after/Post_Event_Notes.docx** and this run's **Action_Register.xlsx**. Paste the exact AFTER Teams discussion link from the setup report into the input; do not attach the whole channel.",
             "Include the final reviewed agenda if available. The skill derives the run name from the selected original workbook. Do not treat pre-event assumptions or rehearsal backup answers as post-event decisions."],
            ["Ausdrücklich sagen: **Wir wechseln jetzt zum vorbereiteten Stand nach derselben Veranstaltung.** Nur **inputs/de/after/Post_Event_Notes.docx** und **Action_Register.xlsx** dieses Laufs auswählen. Den genauen Link der AFTER-Teams-Diskussion aus dem Setup-Bericht in die Eingabe einfügen, nicht den ganzen Kanal auswählen.",
             "Die abschließend geprüfte Agenda bei Bedarf hinzufügen. Der Skill liest den Laufnamen aus der ausgewählten Originalmappe. Vorherige Annahmen und Reserveantworten nicht als Beschlüsse nach der Veranstaltung behandeln."]),
        "instructions": bi(
            "I'm going to talk to Vance personally about the handover issues. Meanwhile, please use the selected post-event notes and linked AFTER discussion to save an English Decisions.docx in the outputs subfolder next to the selected Action_Register.xlsx and update that original workbook with every recorded task, owner, due date and source. Keep the decision document to two readable pages. Preserve P01 as historical pre-event work, the Readme and the template's formatting; use Arial 11 consistently, wrap long source references and fit the row heights so nothing is cut off. The table must include all new rows and columns. Separate decisions, proposals and unanswered questions. Save individual English Outlook follow-up drafts to Vance, Sonia, Billie, Sydney and Teresa, using this run's name and the corresponding A01-A05 task code in each subject. Ask Teresa about the missing success measure and don't claim we know actual spending. Show the saved files and drafts and flag anything still requiring my judgment. Don't send, post, schedule anything or invent commitments.",
            "Ich spreche jetzt persönlich mit Vance über die Probleme bei den Übergaben. Lege bitte währenddessen aus den ausgewählten Notizen nach der Veranstaltung und der verlinkten AFTER-Diskussion eine deutsche Decisions.docx im Unterordner outputs neben der ausgewählten Action_Register.xlsx ab und aktualisiere diese ursprüngliche Mappe mit allen dokumentierten Aufgaben, Verantwortlichen, Fristen und Quellen. Für das Beschlussdokument reichen zwei gut lesbare Seiten. Erhalte P01 als historische Vorbereitungsaufgabe, das Readme und die Vorlagenformatierung; verwende einheitlich Arial 11, brich lange Quellennachweise um und passe die Zeilenhöhen an, damit nichts abgeschnitten wird. Die Tabelle muss alle neuen Zeilen und Spalten umfassen. Trenne Beschlüsse, Vorschläge und offene Fragen. Speichere einzelne deutsche Outlook-Nachfassentwürfe an Vance, Sonia, Billie, Sydney und Teresa, jeweils mit dem Laufnamen und dem zugehörigen Aufgabencode A01 bis A05 im Betreff. Frage Teresa nach dem fehlenden Erfolgskriterium und behaupte keine bekannten Ist-Kosten. Zeig die gespeicherten Dateien und Entwürfe und markiere, wo mein Urteil noch gebraucht wird. Nichts senden, posten oder terminieren und keine Zusagen erfinden."),
        "expect": bi("Saved decisions, an edited action workbook and five unsent drafts; A01–A05 have the recorded owners/dates; success measure and actual costs remain open.",
                     "Gespeicherte Beschlüsse, bearbeitete Aufgabenmappe und fünf ungesendete Entwürfe; A01–A05 mit dokumentierten Verantwortlichen/Fristen; Erfolgskriterium und Ist-Kosten bleiben offen."),
    },
]


pack_quality = bi(
    " Keep Agenda.docx to one readable A4 page, including source notes; shorten repeated explanations rather than shrinking the text. "
    "Use equally readable time labels of at least 18 pt, with explicit start and end times and an unambiguous association to each agenda item. "
    "Render and inspect the finished presentation and agenda before saving: no tiny labels, clipped text or a source-note-only extra page.",
    " Halte Agenda.docx einschließlich Quellenhinweis auf einer gut lesbaren A4-Seite; kürze Wiederholungen statt die Schrift zu verkleinern. "
    "Verwende gleich gut lesbare Zeitangaben mit mindestens 18 Punkt, ausdrücklichen Anfangs- und Endzeiten und eindeutiger Zuordnung zum Programmpunkt. "
    "Rendere und prüfe die fertige Präsentation und Agenda vor dem Speichern: keine winzigen Beschriftungen, abgeschnittenen Texte oder Zusatzseite nur für Quellen.")
for language in ("en", "de"):
    ACTS[1]["instructions"][language] += pack_quality[language]
    ACTS[1]["changeInstructions"][language] += pack_quality[language]

ACTS[2]["instructions"]["en"] += (
    " Read the historical P01 row before editing. Append A01-A05 starting at row 3 on Actions. "
    "Leave Actions!A2:F2 values, cell types and number formats unchanged; in particular D2 must remain "
    "the text 2026-10-20 with General number format, not an Excel date. Do not rewrite that historical row when adding the new tasks. "
    "Arial 11 applies to the worksheet cells; the decision document may retain its readable professional Word body font. Do not repeatedly reformat saved Word fonts. "
    "Treat the personal conversation with Vance as planned or ongoing, not completed, unless the selected AFTER evidence explicitly confirms completion. Do not thank him for a conversation already held without that evidence.")
ACTS[2]["instructions"]["de"] += (
    " Lies die historische P01-Zeile vor der Bearbeitung. Hänge A01-A05 ab Zeile 3 im Blatt Actions an. "
    "Lass Werte, Zelltypen und Zahlenformate von Actions!A2:F2 unverändert; insbesondere bleibt D2 "
    "der Text 2026-10-20 mit dem Zahlenformat Standard (General), kein Excel-Datum. Schreibe die historische Zeile beim Ergänzen nicht neu. "
    "Arial 11 gilt für die Zellen der Aufgabenmappe; das Beschlussdokument darf seine gut lesbare professionelle Word-Textschrift behalten. Keine wiederholten Schriftänderungen an gespeicherten Word-Dateien. "
    "Das persönliche Gespräch mit Vance ist geplant oder läuft gerade, aber noch nicht abgeschlossen, solange die ausgewählten AFTER-Belege den Abschluss nicht ausdrücklich bestätigen. Ohne diesen Beleg nicht für ein bereits geführtes Gespräch danken.")

WORKFLOWS = STARTERS + ACTS

SKILLS = {
    "contoso-executive-briefing": {
        "name": "contoso-executive-briefing",
        "description": "Prepare Contoso Leadership Forum Frankfurt / LEF26 replies, calendar options, conversation preparation, travel proposals and executive briefings. Für Planungsantwort, Terminabstimmung, Gesprächsvorbereitung, Reisevorschlag und Executive-Briefing zum Frankfurt-Forum.",
        "codes": ["CO-26", "CO-27", "CO-28", "CO-29", "CO-23"],
    },
    "contoso-event-pack": {
        "name": "contoso-event-pack",
        "description": "Create the Contoso Frankfurt forum presentation, agenda and unsent coordination drafts, or apply a speaker change to the saved originals. Für Eventpaket und kurzfristige Änderungen am Frankfurter Forum.",
        "codes": ["CO-24"],
    },
    "contoso-event-follow-up": {
        "name": "contoso-event-follow-up",
        "description": "Turn explicitly staged post-event Contoso Frankfurt notes into saved decisions, the original action workbook and individual unsent follow-ups. Für Beschlüsse, Aufgaben und persönliche Nachbereitung nach dem Frankfurt-Forum.",
        "codes": ["CO-25"],
    },
}

research_quality = bi(
    " Use two short entries, one per date, for the calendar conflicts rather than a narrow five-column table. Keep dates and words intact and keep each heading with its first content row. Verify that the saved venue document contains actual clickable operator hyperlinks, not just source labels; a workspace preview is not proof that those links survived saving. Mark accessibility of the selected room configuration as an open check unless the operator explicitly confirms it.",
    " Beschreibe die Kalenderkonflikte in zwei kurzen Einträgen, einem pro Datum, statt in einer schmalen fünfspaltigen Tabelle. Datumsangaben und Wörter dürfen nicht auseinanderbrechen; Überschriften bleiben mit der ersten Inhaltszeile zusammen. Prüfe im gespeicherten Ortsvergleich tatsächliche anklickbare Betreiberlinks, nicht nur Quellennamen; eine Arbeitsbereichsvorschau belegt nicht, dass die Links beim Speichern erhalten blieben. Die Barrierefreiheit der gewählten Raumkonfiguration bleibt offen, solange der Betreiber sie nicht ausdrücklich bestätigt.")
travel_quality = bi(
    " General route pages and sample-day first or last departures are orientation only, not a dated fallback connection. Do not use their last-train times as a fallback for 21 or 22 October. If no date-specific connection is verified, leave the fallback connection as an explicit dated timetable check.",
    " Allgemeine Streckenseiten und erste oder letzte Abfahrten eines Beispieltags dienen nur der Orientierung, nicht als datierte Rückfallverbindung. Verwende deren letzte Zugabfahrt nicht als Rückfallebene für den 21. oder 22. Oktober. Ohne geprüfte datierte Verbindung bleibt die Rückfallverbindung ausdrücklich eine noch nötige Fahrplanprüfung.")
for language in ("en", "de"):
    ACTS[0]["instructions"][language] += research_quality[language]
    next(act for act in STARTERS if act["code"] == "CO-29")["instructions"][language] += travel_quality[language]

SHORT_REQUESTS = {
    "CO-26": bi(
        "Please draft a short reply to Teresa's planning update. Confirm the current arrangements and ask what she still needs from us; save it in Outlook without sending.",
        "Antworte Teresa bitte kurz auf ihre Planungsmail. Bestätige die aktuellen Eckdaten und frage, was sie noch von uns braucht; speichere die Antwort als Outlook-Entwurf."),
    "CO-27": bi(
        "Where do our forum appointments clash on 20 and 22 October? Suggest options and save a coordination email to Sonia as a draft.",
        "Wo kollidieren unsere Forum-Termine am 20. und 22. Oktober? Schlage Lösungen vor und speichere eine Abstimmungsmail an Sonia als Entwurf."),
    "CO-28": bi(
        "Help me prepare for good conversations at the Frankfurt forum. From the selected roles, suggest three useful pairings and a natural opening question for each.",
        "Ich möchte beim Frankfurt-Forum gute Gespräche führen. Bereite mir aus den ausgewählten Rollen drei sinnvolle Gesprächspaarungen mit jeweils einer passenden Einstiegsfrage vor."),
    "CO-29": bi(
        "Please prepare a rail-and-hotel proposal for Billie's trip to our Frankfurt forum, using the travel request. Include practical buffers and the checks we still need; don't book anything.",
        "Bereite bitte Billies Bahnreise mit Übernachtung zu unserem Frankfurt-Forum anhand des Reiseauftrags vor. Ich brauche einen Vorschlag mit Puffern und offenen Prüfungen, noch keine Buchung."),
    "CO-23": bi(
        "I'm preparing my conversation with Teresa about the Frankfurt forum. Please bring the selected planning updates together into a briefing and compare two suitable venues.",
        "Ich bereite mein Gespräch mit Teresa zum Frankfurt-Forum vor. Führe bitte die ausgewählten Planungsstände in einem Briefing zusammen und vergleiche zwei passende Veranstaltungsorte."),
    "CO-24": bi(
        "Please turn the reviewed briefing into our event pack: the presentation, agenda and coordination drafts. Keep everything consistent and save the results.",
        "Mach bitte aus dem geprüften Briefing unser Eventpaket mit Präsentation, Agenda und Abstimmungsentwürfen. Die Unterlagen sollen zusammenpassen und gespeichert vorliegen."),
    "CO-25": bi(
        "I'll speak to Vance personally about the handovers. Meanwhile, please prepare the decisions, action register and individual follow-up drafts from our post-event notes.",
        "Ich spreche persönlich mit Vance über die Übergaben. Bereite währenddessen bitte aus unseren Nachbereitungsnotizen die Beschlüsse, Aufgabenmappe und persönlichen Nachfassentwürfe vor."),
}

for act in WORKFLOWS:
    act["prompt"] = SHORT_REQUESTS[act["code"]]
    act["skill"] = next(key for key, skill in SKILLS.items() if act["code"] in skill["codes"])
    skill_name = SKILLS[act["skill"]]["name"]
    for lang in ("en", "de"):
        act["prep"][lang].insert(0, bi(
            f"Start a new task after installing **{skill_name}**. Select the listed sources and type the short request. Once the task starts, check **Workspace › Plugins** for **{skill_name}** before accepting its results.",
            f"Nach Installation von **{skill_name}** eine neue Aufgabe starten. Genannte Quellen auswählen und den kurzen Auftrag eingeben. Nach dem Start rechts unter **Arbeitsbereich › Plug-Ins** prüfen, dass **{skill_name}** aufgeführt ist, bevor du Ergebnisse übernimmst.")[lang])

ACTS[1]["follow"] = bi(
    "Teresa can now join only remotely from 11:30 to 12:15. Please update our saved event pack and prepare the necessary coordination.",
    "Teresa kann jetzt nur noch von 11:30 bis 12:15 remote teilnehmen. Passe bitte unser gespeichertes Eventpaket an und bereite die nötige Abstimmung vor.")


def skill_text(key):
    skill = SKILLS[key]
    parts = [
        f"---\nname: {json.dumps(skill['name'])}\ndescription: {json.dumps(skill['description'], ensure_ascii=False)}\n---",
        f"# {skill['name']}\n\nVersion: {DATE}. Reusable instructions, not evidence or pre-generated results.",
        "## Operating rules\n\n"
        "- Execute only the procedure matching the user's requested work; never run the entire collection automatically.\n"
        "- Use the user's language for all responses, documents and drafts. Use only that language's procedure below; English is the default only when the request does not establish a language.\n"
        "- Derive the actual run name and output folder from the selected Action_Register.xlsx metadata. Never ask the user to paste a technical run ID when its original cloud location is available. If its location cannot be verified, ask one short question before writing.\n"
        "- Use only the selected scenario evidence. BEFORE is preparation, CHANGE is a selected availability update, AFTER is an explicitly staged post-event snapshot. Skill instructions are not evidence that a business decision occurred.\n"
        "- Save real files and actual unsent Outlook drafts, then report their verified cloud links. Never pass a chat response, approval preview or workspace-only copy off as a saved result.\n"
        "- Verify the exact saved cloud originals, not only a workspace or live-render copy: required page counts, readable tables, requested fonts and real source hyperlinks must survive saving. If verification fails, report that boundary rather than claiming completion.\n"
        "- To correct a saved Office document, reattach that exact OneDrive or SharePoint file and edit it in place. Editing does not require deleting or replacing it. Do not create _Arial or _final alternatives, or ask the presenter to fix formatting manually.\n"
        "- During BEFORE or CHANGE work, the action workbook's metadata identifies the run. Task rows from a later AFTER execution are not planning evidence and must not be imported into the preparation pack.\n"
        "- No external contact, sending, posting, invitations, calendar changes, booking, payment or automation. Draft recipients must be the named approved demo-tenant mailboxes in the selected evidence; no other recipients, CC or BCC.\n"
        "- Do not infer booking authority, spending approval, availability, completed tests, prices, actual costs or success criteria. Missing evidence stays an explicit question.\n"
        "- File edits require individually reviewed one-time approvals, never Always allow. An error may leave partial edits: reread the actual original before retrying. Stop and report a failed operation instead of repeating the same failed batch blindly.\n"
        "- Preserve immutable inputs. Only the post-event procedure edits the selected original action workbook. Never overwrite the baseline or another run.",
    ]
    for lang, label in (("en", "English requests"), ("de", "Deutsche Aufträge")):
        parts.append(f"## {label}")
        for act in WORKFLOWS:
            if act["code"] in skill["codes"]:
                parts.extend([f"### {act['code']}: {act['title'][lang]}", act["instructions"][lang]])
                if "changeInstructions" in act:
                    parts.extend([f"### {CHANGE_TITLE[lang]}", act["changeInstructions"][lang]])
    return "\n\n".join(parts) + "\n"


def build_skills():
    for key in SKILLS:
        content = skill_text(key)
        folder = PACKAGE / "skills" / key
        write(folder / "SKILL.md", content)
        info = ZipInfo("SKILL.md", (2026, 10, 17, 8, 0, 0))
        info.compress_type = ZIP_DEFLATED
        with ZipFile(PACKAGE / "skills" / f"{key}.zip", "w") as archive:
            archive.writestr(info, content.encode("utf-8"))


def doc(path, title, text, lang):
    document = BASE.doc_with(title, bi("Contoso | LEF26 | Synthetic rehearsal", "Contoso | LEF26 | Synthetische Probe")[lang],
                             bi("BEFORE / CHANGE / AFTER are separate evidence phases.",
                                "BEFORE / CHANGE / AFTER sind getrennte Quellenphasen.")[lang],
                             [(title, text.split("\n\n"))])
    for style in ("Normal", "Title", "Heading 1"):
        document.styles[style].font.name = "Arial"
    zoom = document.settings.element.find(qn("w:zoom"))
    if zoom is not None:
        zoom.set(qn("w:percent"), "100")
    BASE.save_doc(document, path)


def workbook(path, lang, run):
    wb = Workbook()
    ws = wb.active
    ws.title = "Actions"
    headers = bi(["ID", "Task", "Owner", "Due date", "Status", "Source"],
                 ["ID", "Aufgabe", "Verantwortlich", "Frist", "Status", "Quelle"])[lang]
    ws.append(headers)
    ws.append(["P01", bi("Verify venue feasibility; no booking", "Machbarkeit des Orts prüfen; keine Buchung")[lang],
               "Vance DeLeon", "2026-10-20", bi("Open", "Offen")[lang], "LEF26-BEFORE-03"])
    table = Table(displayName="EventActions", ref="A1:F2")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(table)
    ws.freeze_panes = "A2"
    for index, width in enumerate([12, 56, 26, 18, 18, 36], 1):
        ws.column_dimensions[chr(64 + index)].width = width
    for row in ws:
        for cell in row:
            cell.font = Font(name="Arial", size=11, bold=cell.row == 1, color="FFFFFF" if cell.row == 1 else "000000")
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.fill = PatternFill("solid", fgColor="0B3D91")
    info = wb.create_sheet("Readme")
    for row in bi(
        [("Scenario", "LEF26"), ("Baseline slot", f"{lang.upper()}-{run:02}"),
         ("Phase", "P01 is pre-event; append A01-A05 only in CO-25."),
         ("Output", "Use the selected active run's actual output folder from the setup/reset report."),
         ("Safety", "Synthetic data. No send/post/invite/booking. Inputs are not approvals.")],
        [("Szenario", "LEF26"), ("Ausgangskopie", f"{lang.upper()}-{run:02}"),
         ("Phase", "P01 liegt vor dem Event; A01-A05 erst in CO-25 ergänzen."),
         ("Ausgabe", "Tatsächlichen Ergebnisordner des aktiven Laufs aus dem Setup-/Reset-Bericht nutzen."),
         ("Sicherheit", "Synthetische Daten. Nicht senden/posten/einladen/buchen. Eingaben sind keine Freigaben.")]
    )[lang]:
        info.append(row)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 100
    for row in info:
        for cell in row:
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(wrap_text=True)
    BASE.save_wb(wb, path)


def build_data():
    build_skills()
    for lang in ("en", "de"):
        before = PACKAGE / "inputs" / lang / "before"
        before.mkdir(parents=True, exist_ok=True)
        doc(before / "Event_Brief.docx", "Contoso Leadership Forum Frankfurt", BRIEF[lang], lang)
        doc(before / "Participants.docx", bi("Workstream roles", "Aufgaben im Szenario")[lang], ROLES[lang], lang)
        agenda = "\n\n".join(f"{start}-{end}: {title[lang]}" for start, end, title in AGENDA)
        doc(before / "Agenda_Baseline.docx", bi("Proposed agenda, not approved", "Agendavorschlag, nicht freigegeben")[lang], agenda, lang)
        write(before / "Travel_Request.md", TRAVEL[lang])
        write(PACKAGE / "inputs" / lang / "change" / "Speaker_Change.md", CHANGE[lang])
        after = PACKAGE / "inputs" / lang / "after"
        after.mkdir(parents=True, exist_ok=True)
        doc(after / "Post_Event_Notes.docx", bi("Post-event snapshot", "Stand nach der Veranstaltung")[lang], AFTER[lang], lang)
        for filename, sender, number, subject, body in MAILS:
            message = EmailMessage(policy=SMTP)
            message["From"] = f"{sender} <{CAST[sender]}@{DOMAIN}>"
            message["To"] = f"MOD Administrator <admin@{DOMAIN}>"
            message["Subject"] = f"[LEF26-BEFORE-{number}] {subject[lang]}"
            message["Date"] = f"Mon, 19 Oct 2026 {9 + int(number):02}:00:00 +0200"
            message["Message-ID"] = f"<lef26-{lang}-{number}@{DOMAIN}>"
            message.set_content(
                bi("Synthetic BEFORE planning evidence, as of 19 October 2026. Delivery time is the setup time.\n\n",
                   "Synthetische BEFORE-Planungsquelle, Stand 19. Oktober 2026. Zustellung erfolgt beim Setup.\n\n")[lang]
                + body[lang])
            path = PACKAGE / f"mails-{lang}" / filename
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(message.as_bytes())
        for run in range(1, 5):
            folder = PACKAGE / "runs" / lang / f"{run:02}"
            folder.mkdir(parents=True, exist_ok=True)
            workbook(folder / "Action_Register.xlsx", lang, run)
            write(folder / "outputs" / "README.md", bi(
                "Output destination for this run. No generated results are supplied. Verify the real Cowork files and their actual cloud locations.",
                "Ergebnisziel dieses Laufs. Keine erzeugten Ergebnisse mitgeliefert. Echte Cowork-Dateien und ihre tatsächliche Cloud-Ablage prüfen.")[lang])
        baseline = PACKAGE / "baseline" / lang
        baseline.mkdir(parents=True, exist_ok=True)
        workbook(baseline / "Action_Register.xlsx", lang, 0)
        write(PACKAGE / f"Preparation-{lang}.md", preparation(lang))
        write(ROOT / "tools" / "expected" / f"{ID}-runbook-{lang}.md", runbook(lang))
    events = [
        ("forum", "2026-10-22", "09:30", 270,
         bi("Leadership Forum Frankfurt - venue open", "Leadership Forum Frankfurt - Ort offen"),
         list(CAST)),
        ("prep", "2026-10-20", "10:00", 45,
         bi("Forum preparation", "Forum-Vorbereitung"), ["Vance DeLeon", "Sonia Rees"]),
        ("prep-clash", "2026-10-20", "10:15", 30,
         bi("Site review - preparation conflict", "Standortreview - Vorbereitungskonflikt"), []),
        ("forum-clash", "2026-10-22", "10:00", 30,
         bi("Investment call - forum conflict", "Investment-Call - Forum-Konflikt"), []),
    ]
    write_json(PACKAGE / "baseline" / "hashes.json", {
        lang: hashlib.sha256((PACKAGE / "baseline" / lang / "Action_Register.xlsx").read_bytes()).hexdigest()
        for lang in ("en", "de")
    })
    write_json(PACKAGE / "tenant.json", {
        "mails": [f"mails-{{lang}}/{m[0]}" for m in MAILS],
        "requiredUsers": [f"{alias}@{DOMAIN}" for alias in CAST.values()],
        "executiveEvent": {
            "id": "LEF26", "account": f"admin@{DOMAIN}", "domain": DOMAIN,
            "people": {name: f"{alias}@{DOMAIN}" for name, alias in CAST.items()},
            "mailReplies": [
                {"key": "scope-ack", "parentSender": "Teresa Sac",
                 "parentSubject": bi("[LEF26-BEFORE-02] Confirmed planning scope", "[LEF26-BEFORE-02] Bestätigter Planungsrahmen"),
                 "text": bi("Thanks Teresa. The executive office will use 42 people and EUR 9,000 including travel. The earlier figures are superseded. We will prepare an agenda decision; no venue or spend is approved.",
                            "Danke Teresa. Das Executive Office arbeitet mit 42 Personen und 9.000 EUR einschließlich Anreise. Die früheren Zahlen sind ersetzt. Wir bereiten die Agendaentscheidung vor; Ort und Ausgaben sind nicht freigegeben.")},
                {"key": "venue-ack", "parentSender": "Vance DeLeon",
                 "parentSubject": bi("[LEF26-BEFORE-03] Venue decision still open", "[LEF26-BEFORE-03] Veranstaltungsort noch offen"),
                 "text": bi("Thanks Vance. We will compare operator information and record open questions. No enquiry, reservation or booking is authorized. We need your judgment on practical feasibility after reviewing the options.",
                            "Danke Vance. Wir vergleichen Betreiberinformationen und halten offene Fragen fest. Keine Anfrage, Reservierung oder Buchung ist freigegeben. Nach Prüfung der Optionen brauchen wir dein Urteil zur praktischen Machbarkeit.")},
            ],
            "sharedFiles": ["inputs/{lang}/before/Event_Brief.docx", "inputs/{lang}/before/Participants.docx",
                            "inputs/{lang}/before/Agenda_Baseline.docx"],
            "events": [{"key": key, "date": day, "start": start, "minutes": minutes,
                        "timeZone": FACTS["timeZone"], "attendees": attendees,
                        "subject": bi(f"[LEF26] {subject['en']}", f"[LEF26] {subject['de']}"),
                        "body": bi("Synthetic rehearsal. BEFORE phase. No venue booked. LEF26",
                                   "Synthetische Probe. Phase BEFORE. Kein Ort gebucht. LEF26")}
                       for key, day, start, minutes, subject, attendees in events],
            "threads": [
                {"key": "before", "phase": "BEFORE",
                 "text": bi("Planning discussion: use 42 people and EUR 9,000 including travel. Venue not booked. Review the linked shared event brief. LEF26-BEFORE",
                            "Planungsdiskussion: 42 Personen und 9.000 EUR einschließlich Anreise. Ort nicht gebucht. Verlinkten gemeinsamen Veranstaltungsbrief prüfen. LEF26-BEFORE"),
                 "replies": [
                     bi("Executive office quotes Sonia's synthetic update: protect at least 60 minutes of informal conversation. No final agenda approval.",
                        "Executive Office zitiert Sonias synthetisches Update: mindestens 60 Minuten informelle Gespräche schützen. Keine endgültige Agendafreigabe."),
                     bi("Executive office quotes Billie's synthetic update: remote-speaker technology still needs a test. Do not promise it works.",
                        "Executive Office zitiert Billies synthetisches Update: Zuschaltung muss noch getestet werden. Keine Funktionszusage.")],
                 },
                {"key": "after", "phase": "AFTER",
                 "text": bi("STAGED POST-EVENT SNAPSHOT, 22 October 2026 14:30. D01 three priorities agreed; D02 handover pilot, no software purchase. Use only in CO-25. LEF26-AFTER",
                            "VORBEREITETER STAND NACH DEM EVENT, 22. Oktober 2026 14:30. D01 drei Prioritäten beschlossen; D02 Übergabe-Pilot, kein Softwarekauf. Nur in CO-25 nutzen. LEF26-AFTER"),
                 "replies": [bi("Executive office quotes Sonia's synthetic reflection: lunch conversations surfaced issues; speak to Vance personally. Actual costs and the pilot success measure remain open.",
                                "Executive Office zitiert Sonias synthetische Reflexion: Mittagsgespräche machten Themen sichtbar; persönlich mit Vance sprechen. Ist-Kosten und Erfolgskriterium des Piloten bleiben offen.")],
                 },
            ],
        },
        "links": [{"name": bi("Microsoft Copilot Cowork", "Microsoft Copilot Cowork"),
                   "url": "https://copilot.cloud.microsoft/"},
                  {"name": bi("DB Frankfurt Hauptbahnhof", "DB Frankfurt Hauptbahnhof"),
                   "url": "https://www.bahnhof.de/frankfurt-main-hbf"}],
        "manual": bi(
            ["Verify Cowork access, usage billing, cloud discovery and built-in web research on the presentation device. Local browser control is not required for this scenario.",
             "Install the three trusted skills from this scenario's skills folder through Cowork Customize and Upload skill: contoso-executive-briefing, contoso-event-pack and contoso-event-follow-up. Keep them private. Start new tasks and verify the matching skill under Workspace > Plugins after execution starts; short live requests depend on these instructions. Do not upload duplicates.",
             "Before using this scenario, an authorized operator records the approved existing demo team/channel and tenant ID in the private executive-event settings file. No sites or teams are created.",
             "Open Preparation-en.md. Inspect actual setup links, calendar conflicts and saved drafts; rehearse all eight demonstrations. Actual mail threads, indexing, UI labels and output folders need tenant verification.",
             "Use a fresh run and conversation. The installer is not a reset. The presenter-only runbook stays in tools/expected, outside indexed inputs."],
            ["Cowork-Zugriff, Nutzungsabrechnung, Quellenfindung und integrierte Webrecherche auf dem Präsentationsgerät prüfen. Dieses Szenario benötigt keine lokale Browsersteuerung.",
             "Die drei vertrauenswürdigen Skills aus dem Ordner skills dieses Szenarios über Cowork Anpassen und Skill-Upload installieren: contoso-executive-briefing, contoso-event-pack und contoso-event-follow-up. Privat lassen. Neue Aufgaben starten und nach dem Start den passenden Skill unter Arbeitsbereich > Plug-Ins prüfen; die kurzen Live-Aufträge setzen seine Anweisungen voraus. Keine Duplikate hochladen.",
             "Vor Verwendung trägt ein berechtigter Operator das freigegebene vorhandene Demo-Team, den Kanal und die Tenant-ID in die private Executive-Event-Konfiguration ein. Keine Sites oder Teams werden angelegt.",
             "Preparation-de.md öffnen. Tatsächliche Setup-Links, Kalenderkonflikte und gespeicherte Entwürfe prüfen; alle acht Demos proben. Mail-Threads, Indexierung, UI-Bezeichnungen und Ergebnisordner im Tenant prüfen.",
             "Frischen Lauf und Verlauf nutzen. Der Installer ist kein Reset. Der Presenter-Leitfaden bleibt unter tools/expected außerhalb indexierter Eingaben."]),
    })
    write_json(PACKAGE / "manifest.json", {
        "title": bi("Executive office: Frankfurt leadership event", "Executive Office: Leadership-Event Frankfurt"),
        "next": bi(["Read Preparation-en.md and complete the tenant preflight."],
                   ["Preparation-de.md lesen und Tenant-Preflight abschließen."]),
    })
    write_json(ROOT / "tools" / "expected" / f"{ID}.json", {
        "facts": FACTS, "codes": [a["code"] for a in WORKFLOWS], "stageDemonstrations": 8,
        "actions": [
            ["A01", "Vance DeLeon", "2026-10-26"], ["A02", "Sonia Rees", "2026-10-27"],
            ["A03", "Billie Vester", "2026-10-23"], ["A04", "Sydney Mattos", "2026-10-28"],
            ["A05", "MOD Administrator", "2026-10-23"],
        ],
        "requiredOutputs": ["Conversation_Brief.docx", "Travel_Plan.docx", "Briefing.docx", "Venue_Comparison.docx",
                            "Leadership_Forum.pptx", "Agenda.docx", "Decisions.docx"],
        "skills": {key: {"name": value["name"], "codes": value["codes"]} for key, value in SKILLS.items()},
        "readiness": "Preparation package, not live acceptance. Short requests require the installed matching skills. The presenter owns the live rehearsal and presentation-device acceptance; local checks do not prove Cowork execution. Actual results and rehearsal evidence remain private.",
    })
    for name in ("ExecutiveEvent.ps1", "Reset-CoworkExecutiveEvent.ps1"):
        (PACKAGE / name).write_bytes((ROOT / "tools" / name).read_bytes())


def preparation(lang):
    return bi(
        "# Executive-event preparation\n\nSynthetic Contoso rehearsal. No venue is booked. "
        "Install with the global -Tenant command; use -Language en. The six BEFORE mails arrive in the presenter's inbox. "
        "The four [LEF26] entries appear in the presenter's Outlook calendar with fixed Frankfurt dates. "
        "The approved Teams channel contains real presenter-authored discussions with explicitly quoted synthetic updates. "
        "The shared BEFORE documents are in that channel's SharePoint files folder; use the setup report's exact links. "
        "OneDrive holds change/after inputs and runs/en/01-04.\n\n"
        "One-time skill setup: download the three ZIPs from this scenario's **skills** folder. In Cowork open **Customize › Add › Upload** "
        "(some versions show **Skills › Add › Upload skill**). Add contoso-executive-briefing, contoso-event-pack and contoso-event-follow-up, keeping them private. "
        "Do not upload a second copy of an existing skill. Start a new task with the listed sources and short request; after execution starts, verify the matching skill under **Workspace › Plugins** before accepting its results. "
        "The skill holds repeatable checks and formatting; type only the short business request. Select real sources separately, not by pasting a page of instructions.\n\n"
        "Pick a fresh run and a new conversation. Set the run output folder explicitly; verify where Cowork actually saved files. "
        "If it cannot target the folder, record the actual Output folder link and select those exact files next; don't claim relocation. "
        "Use the eight sections under Cowork in Demo Kit: reply, calendar, conversation preparation, business travel, venue/briefing, event pack, last-minute change and follow-up. "
        "Use a new conversation for each introductory job and attach only its specified sources. Verify the complete request is visible before Send and in the posted message afterward; attachment chips alone are not a request. "
        "Select BEFORE only for CO-23; reveal CHANGE only after CO-24's initial outputs, "
        "then explicitly switch to AFTER for CO-25. No external send, Teams post, invitation, booking or payment during live execution.\n\n"
        "If required tenant setup is blocked, the demo is not stage-ready. Installer reruns can overwrite workbooks; archive successful results privately first.",
        "# Vorbereitung Executive-Event\n\nSynthetische Contoso-Probe. Kein Veranstaltungsort gebucht. "
        "Mit dem globalen -Tenant-Befehl und -Language de installieren. Sechs BEFORE-Mails erscheinen im Posteingang des Presenters. "
        "Vier [LEF26]-Einträge stehen mit festen Frankfurt-Terminen im Outlook-Kalender des Presenters. "
        "Der freigegebene Teams-Kanal enthält echte Beiträge des Presenters mit ausdrücklich zitierten synthetischen Updates. "
        "Die gemeinsamen BEFORE-Dokumente liegen in den SharePoint-Dateien dieses Kanals; genaue Links im Setup-Bericht nutzen. "
        "OneDrive enthält Änderungs-/Nachbereitungsquellen und runs/de/01-04.\n\n"
        "Skills einmalig einrichten: die drei ZIPs aus dem Ordner **skills** dieses Szenarios herunterladen. In Cowork links **Anpassen**, oben **Hinzufügen** "
        "und dann **Hochladen** wählen. contoso-executive-briefing, contoso-event-pack und contoso-event-follow-up hinzufügen und privat lassen. "
        "Einen vorhandenen Skill nicht nochmals hochladen. Eine neue Aufgabe mit den genannten Quellen und dem kurzen Auftrag starten; danach rechts unter **Arbeitsbereich › Plug-Ins** den passenden Skill prüfen, bevor du Ergebnisse übernimmst. "
        "Wiederkehrende Prüfungen und Formatierung stehen im Skill; nur den kurzen Geschäftsauftrag tippen. Echte Quellen getrennt auswählen, keine Seite Anweisungen hineinkopieren.\n\n"
        "Frischen Lauf und neuen Verlauf wählen. Ergebnisordner ausdrücklich festlegen; tatsächliche Cowork-Ablage prüfen. "
        "Kann Cowork das Ziel nicht nutzen, tatsächlichen Output-folder-Link notieren und anschließend genau diese Dateien auswählen; "
        "keine Verschiebung behaupten. Die acht Abschnitte unter Cowork im Demo Kit nutzen: Antwort, Kalender, Gesprächsvorbereitung, Dienstreise, Ortsvergleich/Briefing, Eventpaket, kurzfristige Änderung und Nachbereitung. "
        "Für jeden Einstiegsauftrag einen neuen Verlauf verwenden und nur seine angegebenen Quellen auswählen. Vor Senden und danach in der gesendeten Nachricht den vollständigen Auftrag prüfen; Dateianhänge allein sind kein Auftrag. "
        "Für CO-23 nur BEFORE wählen, CHANGE erst nach den ersten CO-24-Ergebnissen zeigen und für CO-25 ausdrücklich auf AFTER wechseln. "
        "Live weder extern senden noch in Teams posten, einladen, buchen oder bezahlen.\n\n"
        "Ist notwendiges Tenant-Setup blockiert, ist die Demo nicht bühnenbereit. Installer-Reruns können Arbeitsmappen überschreiben; "
        "erfolgreiche Ergebnisse vorher privat sichern.")[lang]


def runbook(lang):
    heading = bi("# Presenter runbook: Frankfurt leadership event", "# Presenter-Leitfaden: Leadership-Event Frankfurt")[lang]
    introduction = bi(
        "HANDOFF STATUS: This is a preparation package, not acceptance of a complete live workflow. "
        "The presenter owns the live rehearsal, tenant/device checks and stage decision. Local checks do not establish stage readiness. "
        "Agent delivery does not require further autonomous UI runs or two consecutive complete rehearsals. "
        "Use the checklist below and keep actual result links and acceptance evidence privately.\n\n"
        "The user allocated 20 minutes to the already-complete pitch and 40 minutes to demos. Do not promise Cowork completion times. "
        "The schedule below allocates presentation time, not guaranteed processing time. Record observed timings during rehearsal and adjust cutovers from that evidence.\n\n"
        "Narrative: less screen time, more human interaction. Start by showing the conflicting source evidence, then the persisted work. "
        "Never claim saved files, drafts or approvals based on chat text alone.",
        "ÜBERGABESTAND: Dieses Paket bereitet die Vorführung vor; es bestätigt keinen vollständigen erfolgreichen Live-Ablauf. "
        "Die eigene Generalprobe, Tenant-/Geräteprüfung und Bühnenfreigabe übernimmt der Presenter. Lokale Prüfungen belegen keine Bühnenreife. "
        "Weitere autonome UI-Durchläufe oder zwei vollständige Proben hintereinander sind keine Bedingung für die Agenten-Übergabe. "
        "Die folgende Testliste verwenden und tatsächliche Ergebnislinks und Abnahmen privat festhalten.\n\n"
        "Vorgabe: 20 Minuten vorhandener Pitch, 40 Minuten Demos. Keine Cowork-Laufzeiten versprechen. "
        "Der folgende Plan verteilt Präsentationszeit und garantiert keine Bearbeitungsdauer. Beobachtete Zeiten in der Probe protokollieren und daraus Abbruch-/Wechselpunkte bestimmen.\n\n"
        "Roter Faden: weniger Bildschirmzeit, mehr menschliche Begegnung. Erst die widersprüchlichen Quellen, dann die gespeicherte Arbeit zeigen. "
        "Dateien, Entwürfe und Freigaben nie allein aus Chattext als erledigt bezeichnen.")[lang]
    parts = [heading, introduction, rehearsal_checklist(lang), preparation(lang)]
    parts.append(bi(
        "## 40-minute presentation budget\n\n"
        "| Demo clock | Work to show | Human takeaway |\n"
        "|---|---|---|\n"
        "| 00:00-04:00 | CO-26: open the saved planning reply | Review the answer; keep the relationship personal. |\n"
        "| 04:00-08:00 | CO-27: conflicts and saved coordination draft | A person decides which commitment moves. |\n"
        "| 08:00-12:00 | CO-28: open the conversation brief | Prepare better questions, not scraped profiles. |\n"
        "| 12:00-17:00 | CO-29: inspect the sourced travel proposal | Approve a proposal; no booking or payment. |\n"
        "| 17:00-23:00 | CO-23: briefing and venue comparison | Discuss the trade-off and missing quotations. |\n"
        "| 23:00-30:00 | CO-24: actual slides, agenda and drafts | Review completed work outside the chat. |\n"
        "| 30:00-35:00 | CO-24 continued: change the saved originals | Check consistency, then make the human decision. |\n"
        "| 35:00-40:00 | CO-25: decisions, workbook and five drafts | Use the recovered time for personal follow-up. |\n\n"
        "The demo clock starts after the separate 20-minute pitch. Source selection and approval checks count inside these budgets. "
        "Start longer research while discussing an independent everyday job only after rehearsing that tab sequence; keep its conversation and BEFORE evidence separate. "
        "At each slot boundary, move to a reviewed genuine rehearsal result if the live output is not ready. Say explicitly that it is an earlier Cowork result. "
        "Never compress required approval checks or repair formatting on stage. Before CO-25, say: 'We are now jumping to the prepared post-event snapshot from 22 October at 14:30; this event has not happened during our presentation.'",
        "## Präsentationsbudget: 40 Minuten\n\n"
        "| Demo-Uhr | Sichtbare Arbeit | Menschlicher Mehrwert |\n"
        "|---|---|---|\n"
        "| 00:00-04:00 | CO-26: gespeicherte Planungsantwort öffnen | Antwort prüfen; die Beziehung bleibt persönlich. |\n"
        "| 04:00-08:00 | CO-27: Konflikte und gespeicherten Abstimmungsentwurf zeigen | Ein Mensch entscheidet, welcher Termin weicht. |\n"
        "| 08:00-12:00 | CO-28: Gesprächsvorbereitung öffnen | Bessere Fragen statt ausgelesener Profile. |\n"
        "| 12:00-17:00 | CO-29: belegten Reisevorschlag prüfen | Vorschlag freigeben; nichts buchen oder bezahlen. |\n"
        "| 17:00-23:00 | CO-23: Briefing und Ortsvergleich zeigen | Abwägung und fehlende Angebote besprechen. |\n"
        "| 23:00-30:00 | CO-24: echte Folien, Agenda und Entwürfe öffnen | Erledigte Arbeit außerhalb des Chats prüfen. |\n"
        "| 30:00-35:00 | CO-24 fortgesetzt: gespeicherte Originale ändern | Konsistenz prüfen, dann menschlich entscheiden. |\n"
        "| 35:00-40:00 | CO-25: Beschlüsse, Aufgabenmappe und fünf Entwürfe öffnen | Gewonnene Zeit für persönliche Nachbereitung nutzen. |\n\n"
        "Die Demo-Uhr startet nach dem getrennten 20-minütigen Pitch. Quellenauswahl und Freigaben zählen in diese Zeitbudgets. "
        "Längere Recherche nur dann während einer unabhängigen Alltagsaufgabe starten, wenn genau dieser Tabwechsel geprobt ist; Verlauf und BEFORE-Belege getrennt halten. "
        "An jeder Abschnittsgrenze auf ein geprüftes echtes Ergebnis einer früheren Probe wechseln, falls der Live-Lauf noch nicht fertig ist. Den Wechsel ausdrücklich benennen. "
        "Niemals notwendige Freigabeprüfungen verkürzen oder auf der Bühne Formatierung reparieren. Vor CO-25 sagen: 'Wir springen jetzt zum vorbereiteten Nachbereitungsstand vom 22. Oktober um 14:30. Diese Veranstaltung hat nicht während unserer Präsentation stattgefunden.'"
    )[lang])
    stage_number = 0
    for act in WORKFLOWS:
        stage_number += 1
        parts.extend([f"## {stage_number}. {act['code']}: {act['title'][lang]}", act["story"][lang],
                      "\n".join(f"- {p}" for p in act["prep"][lang]),
                      "### Prompt", act["prompt"][lang], bi("### Result", "### Ergebnis")[lang], act["expect"][lang]])
        if "follow" in act:
            stage_number += 1
            parts.extend([f"## {stage_number}. {CHANGE_TITLE[lang]} ({act['code']})",
                "\n".join(f"- {p}" for p in CHANGE_PREP[lang]),
                "### Prompt", act["follow"][lang], bi("### Result", "### Ergebnis")[lang],
                bi("The same original presentation and agenda now show Teresa remote 11:30-12:15 for the full 45 minutes and 105 informal minutes. "
                   "The existing Teresa draft is updated and one Billie test-request draft added; do not invent a successful technology test.",
                   "Dieselben Originale von Präsentation und Agenda zeigen Teresa remote 11:30-12:15 für volle 45 Minuten und 105 informelle Minuten. "
                   "Der vorhandene Teresa-Entwurf ist aktualisiert und ein Billie-Testanfrageentwurf ergänzt; keinen erfolgreichen Techniktest erfinden.")[lang]])
    parts.extend([
        bi("## Approval and result check", "## Freigaben und Ergebniskontrolle")[lang],
        bi("After submitting, verify that this request has actual Cowork activity; a sent-message notice or a new task URL alone does not prove delivery. Retry an explicitly undelivered message with its own retry control once, not by creating another task. If delivery still fails, switch to the disclosed genuine checkpoint. Approve only reviewed file uploads/edits once. Check the requested file and destination, opening parameter details when available. Wait for the exact filename to appear; a DOCX badge alone is not enough to identify an edit target. For Office edit cards, choose **Approve once**, not **Always allow**; the one-time choice can be under **More options**, and several separate edit requests can appear. A document can still be marked as being edited while its final upload is waiting for approval; inspect the chat approval card rather than waiting only for the file status. Cancel Send/Post/Create invitation/booking actions. Open actual OneDrive/SharePoint files "
           "and Outlook Drafts; verify recipients, subjects, unsent state and content. Check 42 / EUR 9,000, 09:30-14:00, "
           "Teresa remote 11:30-12:15 after the change, and A01-A05 with their recorded owners and dates. "
           "A failed edit can leave partial changes in the original file. Reopen and reread its current contents before any retry; never assume rollback or approve the same failed batch blindly. "
           "Use only [LEF26] calendar entries and official operator/DB research; reject unrelated appointments, press-derived claims and later-phase leakage. Do not invent actual spend, a success threshold, a venue approval or a meeting transcript.",
           "Nach dem Absenden prüfen, ob Cowork diesen Auftrag tatsächlich bearbeitet; die Meldung zum Nachrichtenversand oder eine neue Aufgaben-URL allein bestätigt keine Zustellung. Bei **Nicht zugestellt. Nochmal versuchen?** den vorhandenen Auftrag einmal über genau diese Schaltfläche erneut versuchen, keine zusätzliche Aufgabe anlegen. Scheitert die Zustellung weiterhin, zum offen benannten echten Checkpoint wechseln. Nur geprüfte Datei-Uploads/-Bearbeitungen einmalig freigeben. Angeforderte Datei und Ziel prüfen, verfügbare Parameterdetails öffnen. Auf den genauen Dateinamen warten; allein die Kennzeichnung DOCX identifiziert das Bearbeitungsziel nicht. Bei Office-Bearbeitungskarten **Einmal genehmigen** statt **Immer zulassen** wählen; die einmalige Freigabe kann unter **Weitere Optionen** stehen, und mehrere getrennte Bearbeitungsanfragen sind möglich. Eine Datei kann noch als in Bearbeitung erscheinen, während ihr abschließender Upload auf Freigabe wartet; die Freigabekarte im Chat prüfen, nicht nur auf den Dateistatus warten. Senden/Posten/Einladen/Buchen abbrechen. Tatsächliche OneDrive-/SharePoint-Dateien "
           "und Outlook-Entwürfe öffnen; Empfänger, Betreff, ungesendeten Zustand und Inhalt prüfen. 42 / 9.000 EUR, 09:30-14:00, "
           "Teresa nach Änderung remote 11:30-12:15 und A01-A05 mit dokumentierten Verantwortlichen und Fristen prüfen. "
           "Eine fehlgeschlagene Bearbeitung kann Teiländerungen in der Originaldatei hinterlassen. Vor einem neuen Versuch den aktuellen Inhalt erneut öffnen und prüfen; keine automatische Rücknahme annehmen und denselben fehlgeschlagenen Auftrag nicht blind erneut freigeben. "
           "Nur [LEF26]-Kalendereinträge und offizielle Betreiber-/DB-Quellen verwenden; fremde Termine, Pressebehauptungen und spätere Quellenphasen zurückweisen. Keine Ist-Kosten, Zielwerte, Ortsfreigaben oder Meetingtranskripte erfinden.")[lang],
        bi("## Rehearsal, backup and reset", "## Probe, Reserve und Reset")[lang],
        bi("The presenter runs the live rehearsal using the checklist above. Further autonomous agent UI tests are not part of this handoff. "
           "Log actual session/source/output/draft links, phase leakage, elapsed observations and sent-message checks privately. "
           "Use existing real Cowork outputs and screenshots or native Office renders as disclosed backups. A screenshot or render is not a recording or proof of the full workflow. "
           "If a backup is missing, keep that moment open until your own rehearsal produces and checks it. "
           "Recordings and tenant-bearing screenshots must stay outside the public repos. "
           "Backup checklist: briefing/comparison; slide deck/agenda; decisions/action register; unsent-draft evidence; local playback. "
           "If a live act stalls, say you are switching to an actual earlier rehearsal result and open its files; never manufacture a success screenshot. Do not spend the live slot repairing lengthy document formatting or unrelated-source leakage; stop that path and use a reviewed checkpoint.\n\n"
           "Reset: in the same PowerShell window as setup, run tools\\Reset-CoworkExecutiveEvent.ps1 -WhatIf, inspect the proposed fresh run, then rerun without -WhatIf in the verified demo session. "
           "A fresh PowerShell process must first reconnect through the normal cached Graph sign-in described in README; do not rerun the installer merely to reconnect. "
           "It creates a new action workbook/output namespace; it does not delete results, mail, calendar or Teams content. "
           "Start a new Cowork conversation, choose BEFORE sources and use the new run name in draft subjects. "
           "The untouched fixture baseline remains the source of truth; live prompts never modify it. Review and cancel any accidentally created automation manually. "
           "If baseline sources were altered, stop and restore only verified scenario-owned objects; don't reinstall all demos as a reset.",
           "Der Presenter führt die eigene Generalprobe mit der obigen Testliste durch. Weitere autonome UI-Tests des Agenten gehören nicht zu dieser Übergabe. "
           "Tatsächliche Verlauf-/Quell-/Ergebnis-/Entwurfslinks, Phasenvermischung, beobachtete Laufzeiten und Versandkontrollen privat protokollieren. "
           "Vorhandene echte Cowork-Ergebnisse und Screenshots oder native Office-Renderings als offen benannte Reserven nutzen. Ein Bild oder Rendering ist kein Video und kein Beleg für den gesamten Ablauf. "
           "Fehlt ein Backup, bleibt der Moment offen, bis die eigene Probe es erzeugt und prüft. "
           "Aufzeichnungen und Tenant-Screenshots bleiben außerhalb öffentlicher Repos. "
           "Reserveprüfung: Briefing/Vergleich; Präsentation/Agenda; Beschlüsse/Aufgabenmappe; Nachweis ungesendeter Entwürfe; lokale Wiedergabe. "
           "Stockt ein Live-Akt, ausdrücklich zum echten Ergebnis einer früheren Probe wechseln; keinen Erfolgsscreenshot erzeugen. Nicht während des Live-Slots lange Formatkorrekturen oder fremde Quellinhalte reparieren; diesen Pfad stoppen und einen geprüften Checkpoint nutzen.\n\n"
           "Reset: im selben PowerShell-Fenster wie das Setup tools\\Reset-CoworkExecutiveEvent.ps1 -WhatIf starten, vorgeschlagenen frischen Lauf prüfen und im verifizierten Demokontext ohne -WhatIf wiederholen. "
           "Ein neuer PowerShell-Prozess muss zuerst den normalen Graph-Login mit Cache laut README initialisieren; den Installer nicht allein zum Wiederverbinden erneut starten. "
           "Erzeugt eine neue Aufgabenmappe und einen neuen Ergebnisnamensraum; löscht weder Ergebnisse noch Mails, Kalender oder Teams-Inhalte. "
           "Neuen Cowork-Verlauf starten, BEFORE-Quellen wählen und neuen Laufnamen in Entwurfsbetreffzeilen nutzen. "
           "Die unveränderte Ausgangsbasis bleibt maßgeblich; Live-Prompts verändern sie nicht. Versehentlich angelegte Automatisierungen manuell prüfen und abbrechen. "
           "Bei geänderten Ausgangsquellen stoppen und nur nachgewiesen eigene Szenarioobjekte wiederherstellen; nicht alle Demos als Reset neu installieren.")[lang],
    ])
    return "\n\n".join(parts)


def rehearsal_checklist(lang):
    introduction = bi(
        "## Your rehearsal checklist\n\n"
        "All rows start OPEN for your own acceptance, even when historical results exist. "
        "For each moment, record PASS / FIX / BACKUP ONLY, the actual task/result link and your observation privately. "
        "Do not infer current cloud state from an older downloaded file. No send, post, invitation, booking or payment.\n\n"
        "| Moment | Check outside the chat | Your result |\n|---|---|---|",
        "## Testliste für deine Generalprobe\n\n"
        "Für deine eigene Abnahme starten alle Zeilen OFFEN, auch wenn frühere Ergebnisse vorliegen. "
        "Je Moment OK / KORREKTUR / NUR BACKUP, tatsächlichen Verlauf-/Ergebnislink und Beobachtung privat festhalten. "
        "Eine ältere heruntergeladene Datei bestätigt nicht den aktuellen Cloud-Stand. Nicht senden, posten, einladen, buchen oder bezahlen.\n\n"
        "| Moment | Außerhalb des Chats prüfen | Dein Ergebnis |\n|---|---|---|")[lang]
    checks = bi(
        [
            ("CO-26", "Teresa draft saved and unsent; 42 people, EUR 9,000 including travel, unchanged event window."),
            ("CO-27", "Sonia draft saved; both recorded overlaps present; all four original calendar entries unchanged."),
            ("CO-28", "Conversation_Brief.docx saved and readable; three role-grounded pairings; no invented biographies."),
            ("CO-29", "Travel_Plan.docx saved; dated travel constraints and real source links; EUR 650 inside EUR 9,000; unknown fares remain unknown."),
            ("CO-23", "Briefing.docx and venue comparison saved; current figures, both conflicts, source links and open venue decision; no CHANGE/AFTER leakage."),
            ("CO-24", "Actual six-slide PowerPoint, one-page agenda and two unsent drafts; matching BEFORE facts and readable saved files."),
            ("CO-24 change", "Same original file identities; remote 11:30-12:15, 45 minutes, 105 informal minutes; existing Teresa draft reused and one Billie draft added."),
            ("CO-25", "Decisions.docx, original workbook and five unsent follow-ups; exact P01 preservation, A01-A05 owners/dates, unknown costs and success measure; no invented completed conversation."),
        ],
        [
            ("CO-26", "Teresa-Entwurf gespeichert und ungesendet; 42 Personen, 9.000 EUR einschließlich Anreise, unverändertes Veranstaltungsfenster."),
            ("CO-27", "Sonia-Entwurf gespeichert; beide Überschneidungen enthalten; alle vier Originaltermine unverändert."),
            ("CO-28", "Conversation_Brief.docx gespeichert und lesbar; drei rollenbasierte Paarungen; keine erfundenen Biografien."),
            ("CO-29", "Travel_Plan.docx gespeichert; datierte Reisevorgaben und echte Quellenlinks; 650 EUR innerhalb 9.000 EUR; unbekannte Tarife bleiben offen."),
            ("CO-23", "Briefing.docx und Ortsvergleich gespeichert; aktuelle Zahlen, beide Konflikte, Quellenlinks und offene Ortsentscheidung; keine CHANGE-/AFTER-Vermischung."),
            ("CO-24", "Echte sechs Folien, einseitige Agenda und zwei ungesendete Entwürfe; übereinstimmende BEFORE-Fakten und lesbare gespeicherte Dateien."),
            ("CO-24 Änderung", "Dieselben Originaldateien; remote 11:30-12:15, 45 Minuten, 105 informelle Minuten; Teresa-Entwurf wiederverwendet und ein Billie-Entwurf ergänzt."),
            ("CO-25", "Decisions.docx, Originalmappe und fünf ungesendete Entwürfe; P01 exakt erhalten, A01-A05 mit Verantwortlichen/Fristen, Kosten und Erfolgskriterium offen; kein erfundener Gesprächsabschluss."),
        ])[lang]
    rows = [f"| {code} | {check} | {'OPEN' if lang == 'en' else 'OFFEN'} |" for code, check in checks]
    boundary = bi(
        "\n\nBefore the first live request, verify the demo account, required access, installed skill versions, actual source links and a fresh run. "
        "Do not re-upload already registered skills or rerun the installer as a reset. "
        "Open your private backup index locally and verify the actual files on the presentation device. "
        "Backups can be repaired outputs from different runs: label their origin and do not present them as one successful chain. "
        "Historical cloud links are navigation aids, not current availability checks. "
        "Stop at a missing source, unsafe approval or incorrect result; keep the issue in this checklist rather than starting a formatting loop.",
        "\n\nVor dem ersten Live-Auftrag Demo-Konto, benötigten Zugriff, installierte Skill-Versionen, tatsächliche Quellenlinks und frischen Lauf prüfen. "
        "Bereits registrierte Skills nicht nochmals hochladen und den Installer nicht als Reset wiederholen. "
        "Den privaten Backup-Index lokal öffnen und die tatsächlichen Dateien auf dem Präsentationsgerät prüfen. "
        "Backups können korrigierte Ergebnisse verschiedener Läufe sein: Herkunft benennen und nicht als eine erfolgreiche Kette darstellen. "
        "Historische Cloud-Links helfen bei der Navigation, bestätigen aber keine aktuelle Verfügbarkeit. "
        "Bei fehlender Quelle, unsicherer Freigabe oder falschem Ergebnis stoppen; den Punkt in dieser Testliste festhalten statt eine Formatierungsschleife zu beginnen.")[lang]
    return introduction + "\n" + "\n".join(rows) + boundary


def page():
    steps = [
        {"kind": "callout", "tone": "info", "title": bi("Less screen time. More human interaction.", "Weniger Bildschirmzeit. Mehr menschliche Begegnung."),
         "text": bi("Eight assistant jobs, from a saved reply to coordinated event follow-up, in one synthetic Contoso story. A last-minute change and a clearly staged post-event snapshot show the progression. Inspect saved work, not just chat answers.",
                    "Acht Assistenzaufgaben, vom gespeicherten Antwortentwurf bis zur koordinierten Event-Nachbereitung, in einer synthetischen Contoso-Geschichte. Kurzfristige Änderung und klarer Zeitsprung zeigen die Steigerung. Gespeicherte Arbeit zeigen, nicht nur Chatantworten.")},
        {"kind": "callout", "tone": "tip", "title": bi("Short requests, reusable skills", "Kurze Aufträge, wiederverwendbare Skills"),
         "text": bi("The requests below are one or two natural sentences. Recurring source, safety and output requirements live in three reusable Cowork skills, not in a page-long prompt. Complete their one-time setup before rehearsing.",
                    "Die Aufträge unten bestehen aus ein bis zwei natürlichen Sätzen. Wiederkehrende Quellen-, Sicherheits- und Ergebnisvorgaben stehen in drei wiederverwendbaren Cowork-Skills, nicht in einem seitenlangen Prompt. Die Skills vor der Probe einmalig einrichten.")},
        {"kind": "action", "title": bi("Add the three skills once", "Die drei Skills einmalig hinzufügen"),
         "items": bi(
             ["Download the three ZIP files from **Demo-Kit › Cowork › cowork-executive-event › skills**.",
              "In Cowork open **Customize › Add › Upload**; some versions show **Skills › Add › Upload skill**. Add **contoso-executive-briefing**, **contoso-event-pack** and **contoso-event-follow-up**. Keep access private; avoid duplicate uploads.",
              "Start a new task with the listed sources and short request. After execution starts, check **Workspace › Plugins** for the expected skill name before accepting results. The short request depends on its installed skill; verify this path in your own rehearsal."],
             ["Die drei ZIP-Dateien aus **Demo-Kit › Cowork › cowork-executive-event › skills** herunterladen.",
              "In Cowork links auf **Anpassen**, oben auf **Hinzufügen** und dann auf **Hochladen** klicken. **contoso-executive-briefing**, **contoso-event-pack** und **contoso-event-follow-up** hinzufügen. Zugriff privat lassen; keine Duplikate hochladen.",
              "Neue Aufgabe mit den genannten Quellen und dem kurzen Auftrag starten. Nach dem Start rechts unter **Arbeitsbereich › Plug-Ins** den erwarteten Skillnamen prüfen, bevor du Ergebnisse übernimmst. Der kurze Auftrag benötigt seinen installierten Skill; diesen Weg in der eigenen Generalprobe prüfen."])},
        {"kind": "screen", "title": bi("Upload the ZIP on Customize (schematic)", "ZIP in Anpassen hochladen (Schema)"),
         "screen": {"app": bi("Microsoft Copilot Cowork — Customize", "Microsoft Copilot Cowork — Anpassen"), "platform": "Web / Edge",
                    "url": "copilot.cloud.microsoft/cowork",
                    "tabs": [],
                    "button": bi("Add", "Hinzufügen"),
                    "contextMenu": [bi("Upload", "Hochladen")],
                    "menuHighlight": bi("Upload", "Hochladen"),
                    "markers": [{"target": "button", "text": bi("On **Customize**, select **Add** at the top.", "In **Anpassen** oben **Hinzufügen** wählen.")},
                                {"target": "context-menu", "text": bi("Select **Upload**, then choose your skill ZIP.", "**Hochladen** wählen, dann die Skill-ZIP auswählen.")}]}},
        {"kind": "checklist", "title": bi("Before presenting", "Vor der Vorführung"),
         "checks": [{"status": "warn", "label": bi("Verify the real tenant and rehearse", "Tatsächlichen Tenant prüfen und proben"),
                     "detail": bi("Cowork access, billing, built-in web research, seeded app data, indexing, actual saved files/drafts and reset must pass on the presentation device.",
                                  "Cowork-Zugriff, Abrechnung, integrierte Webrecherche, angelegte App-Daten, Indexierung, tatsächlich gespeicherte Dateien/Entwürfe und Reset müssen auf dem Präsentationsgerät funktionieren.")},
                    {"status": "warn", "label": bi("Draft-only; no booking or external contact", "Nur Entwürfe; keine Buchung oder externe Kontaktaufnahme"),
                     "detail": bi("Cancel Send/Post/invitation/booking approval cards. Do not use Always allow. Each phase uses its own selected evidence.",
                                  "Freigaben für Senden/Posten/Einladen/Buchen abbrechen. Nicht Always allow nutzen. Jede Phase nutzt nur ihre ausgewählten Quellen.")}]},
        {"kind": "screen", "title": bi("Where do I work?", "Wo arbeite ich?"),
         "screen": {"layout": "chat", "app": "Microsoft Copilot", "platform": "Web / Edge",
                    "url": "copilot.cloud.microsoft", "mode": "cowork",
                    "markers": [{"target": "mode-toggle", "text": bi("Select **Cowork** at the top.", "Oben **Cowork** wählen.")},
                                {"target": "input", "text": bi("Choose the exact phase sources through **Add attachments**, then paste the request.",
                                                             "Über **Anlagen hinzufügen** die genauen Quellen der Phase auswählen, dann den Auftrag einfügen.")}]}},
    ]
    for act in WORKFLOWS:
        steps.extend([
            {"kind": "section", "title": act["title"],
             "exercise": {"code": act["code"], "chatTitle": act["chatTitle"], "story": act["story"]}},
            {"kind": "prompt", "title": bi("Complete this work", "Diese Arbeit erledigen"),
             "where": bi("Cowork on the web in the approved demo Edge profile.", "Cowork im Web im freigegebenen Demo-Edge-Profil."),
             "preparation": act["prep"], "text": act["prompt"], "expect": act["expect"]},
        ])
        if "follow" in act:
            steps.append({"kind": "section", "title": CHANGE_TITLE,
                          "meta": bi("Continue CO-24 in the same conversation; edit the saved originals.",
                                     "CO-24 im selben Verlauf fortsetzen; gespeicherte Originale bearbeiten.")})
            steps.append({"kind": "prompt", "title": CHANGE_TITLE,
                          "preparation": CHANGE_PREP,
                          "text": act["follow"], "expect": bi(
                              "Same files updated; Teresa remote 11:30-12:15; at least 60 minutes of conversation; existing draft updated and one Billie draft added.",
                              "Dieselben Dateien aktualisiert; Teresa remote 11:30-12:15; mindestens 60 Minuten Gesprächszeit; vorhandener Entwurf aktualisiert und ein Billie-Entwurf ergänzt.")})
        steps.append({"kind": "action", "title": bi("Show the completed work", "Die erledigte Arbeit zeigen"),
                      "items": bi(
                          ["Open the actual persistent files and Outlook drafts outside the chat. Reject unsupported claims and keep decisions requiring judgment open.",
                           "Explain the human decision or conversation this result makes possible. No Send/Post/Create invitation."],
                          ["Tatsächlich gespeicherte Dateien und Outlook-Entwürfe außerhalb des Chats öffnen. Unbelegte Aussagen zurückweisen und notwendige menschliche Entscheidungen offenhalten.",
                           "Zeigen, welche persönliche Entscheidung oder welches Gespräch dadurch möglich wird. Nicht senden, posten oder einladen."])})
    for _, sender, number, subject, body in MAILS:
        steps.append({"kind": "content", "title": bi(f"Email from {sender} (in the inbox)", f"Mail von {sender} (im Posteingang)"),
                      "text": {lang: f"Subject: [LEF26-BEFORE-{number}] {subject[lang]}\n\n{body[lang]}" for lang in ("en", "de")}})
    return {
        "id": ID, "band": "usecases", "accent": "cowork", "data": ID,
        "title": bi("Executive office: leadership event", "Executive Office: Leadership-Event"),
        "kicker": bi("Eight real assistant jobs: reply → calendar → conversations → travel → event → follow-up",
                     "Acht echte Assistenzaufgaben: Antwort → Kalender → Gespräche → Reise → Event → Nachbereitung"),
        "why": bi("Cowork completes the paperwork while the assistant keeps decisions and relationships.", "Cowork erledigt die Unterlagen; Entscheidungen und Beziehungen bleiben bei der Assistenz."),
        "steps": steps,
        "sources": [{"label": bi(en, de), "url": url} for en, de, url in SOURCES],
    }


def build_deck(decks):
    path = decks / "cowork.json"
    deck = json.loads(path.read_text(encoding="utf-8"))
    codes = {act["code"] for act in WORKFLOWS}
    for file in decks.glob("*.json"):
        other = json.loads(file.read_text(encoding="utf-8"))
        for demo in other["demos"]:
            if demo["id"] == ID:
                continue
            exercises = [demo.get("exercise"), *(s.get("exercise") for s in demo["steps"])]
            if any(e and e["code"] in codes for e in exercises):
                raise ValueError(f"Exercise collision: {file.name}/{demo['id']}")
    view = page()
    old_title = next((d["title"] for d in deck["demos"] if d["id"] == ID), None)
    deck["demos"] = [d for d in deck["demos"] if d["id"] != ID]
    deck["demos"].insert(0, view)
    deck["agenda"]["items"] = [item for item in deck["agenda"]["items"] if item["title"] not in [old_title, view["title"]]]
    deck["agenda"]["items"].insert(0, {"title": view["title"], "detail": view["kicker"]})
    deck["meta"].update(version="1.8", date=DATE)
    write_json(path, deck)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--decks", type=Path, default=ROOT.parent / "demo-kit-new2026" / "decks")
    args = parser.parse_args()
    build_data()
    build_deck(args.decks)
    print("Built LEF26 bilingual inputs and eight stage demonstrations (CO-23 through CO-29). No tenant writes or Cowork results. Run Update-Manifest.ps1 -All.")
